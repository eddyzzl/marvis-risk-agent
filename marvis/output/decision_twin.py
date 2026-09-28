"""Exports read the exact authenticated receipt shown by the API."""

from io import BytesIO
import json

from docx import Document
from openpyxl import Workbook

from marvis.spreadsheet_safety import safe_xlsx_cell


EVENT_COLUMNS = (
    "scenario",
    "record_id",
    "status",
    "error_code",
    "event_task_id",
    "event_request_id",
    "content_hash",
    "contract_hash",
    "snapshot_hash",
    "decision_at",
    "knowledge_cutoff",
    "availability_mode",
)


def _event_rows(payload):
    for scenario in payload.get("scenarios", []):
        for row in scenario["decisions"]:
            evidence = row.get("event_evidence")
            if evidence:
                yield (
                    scenario["name"],
                    row["record_id"],
                    evidence["status"],
                    row.get("error_code"),
                    evidence["task_id"],
                    evidence["request_id"],
                    evidence["content_hash"],
                    evidence["contract_hash"],
                    evidence.get("snapshot_hash"),
                    evidence["decision_at"],
                    evidence["knowledge_cutoff"],
                    evidence["availability_mode"],
                )


def _temporal_rows(payload):
    """Project signed values only; rendering never recomputes a drift metric."""
    rows = []
    for scenario in payload.get("scenarios", []):
        result = scenario["metrics"].get("stability", {})
        if result.get("schema_version") != "decision_twin.temporal_stability.v1":
            continue
        reference = result["reference"]["window"]
        for comparison in result["comparisons"]:
            window = comparison["window"]
            for check in comparison["checks"]:
                rows.append(
                    (
                        scenario["name"],
                        result["verdict"],
                        reference["name"],
                        reference["sample_count"],
                        reference["members_hash"],
                        window["name"],
                        window["start"],
                        window["end"],
                        window["sample_count"],
                        window["members_hash"],
                        check["metric"],
                        check["value"],
                        check["threshold"],
                        check["unit"],
                        check["status"],
                        check["reason"],
                    )
                )
    return rows


def render_historical_replay(receipt, format):
    payload = receipt["payload"]
    facts = [
        ("artifact_id", receipt["artifact_id"]),
        ("kind", receipt["kind"]),
        ("contract_hash", payload["contract_hash"]),
        ("authority", "proposal_only"),
        ("causal_gain_verified", False),
        ("marvis_historical_execution_verified", False),
    ]
    summary = {k: v for k, v in payload.items() if k not in {"scenarios", "records"}}
    rows = []
    for scenario in payload.get("scenarios", []):
        for row in scenario["decisions"]:
            rows.append(
                (
                    scenario["name"],
                    row["record_id"],
                    row["decision_at"],
                    row["score"],
                    row["action"]["type"],
                    row["action"].get("reason_code"),
                    row["facts_hash"],
                    row["package_hash"],
                )
            )
    columns = (
        "scenario",
        "record_id",
        "decision_at",
        "score",
        "action",
        "reason_code",
        "facts_hash",
        "package_hash",
    )
    stream = BytesIO()
    temporal_rows = _temporal_rows(payload)
    event_rows = list(_event_rows(payload))
    if format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "回放证据"
        for fact in facts:
            sheet.append([safe_xlsx_cell(value) for value in fact])
        # JSON is split by top-level sections so wide population receipts never
        # disappear into Excel's 32,767-character single-cell limit.
        sheet.append(
            [
                "source",
                safe_xlsx_cell(json.dumps(payload.get("source"), ensure_ascii=False)),
            ]
        )
        metrics = workbook.create_sheet("场景指标")
        metrics.append(["scenario", "metric", "value"])
        for scenario in payload.get("scenarios", []):
            for key in ("metrics", "constraints", "comparison"):
                encoded = json.dumps(scenario[key], ensure_ascii=False)
                for offset in range(0, len(encoded), 30_000):
                    metrics.append(
                        [
                            safe_xlsx_cell(scenario["name"]),
                            key,
                            safe_xlsx_cell(encoded[offset : offset + 30_000]),
                        ]
                    )
        if temporal_rows:
            temporal = workbook.create_sheet("时间稳定性")
            temporal.append(
                [
                    "scenario",
                    "verdict",
                    "reference_window",
                    "reference_count",
                    "reference_members_hash",
                    "comparison_window",
                    "start",
                    "end",
                    "comparison_count",
                    "comparison_members_hash",
                    "metric",
                    "value",
                    "threshold",
                    "unit",
                    "status",
                    "reason",
                ]
            )
            for row in temporal_rows:
                temporal.append([safe_xlsx_cell(value) for value in row])
        decisions = workbook.create_sheet("逐笔回放")
        decisions.append(list(columns))
        for row in rows:
            decisions.append([safe_xlsx_cell(value) for value in row])
        if event_rows:
            events = workbook.create_sheet("原生事件证据")
            events.append(list(EVENT_COLUMNS))
            for row in event_rows:
                events.append([safe_xlsx_cell(value) for value in row])
        imported = payload.get("observed_actions", {}).get("records", [])
        if imported:
            observed = workbook.create_sheet("外部历史动作")
            keys = tuple(imported[0])
            observed.append(list(keys))
            for row in imported:
                observed.append([safe_xlsx_cell(row[key]) for key in keys])
        if payload.get("records"):
            outcomes = workbook.create_sheet("现金流对账")
            keys = tuple(payload["records"][0])
            outcomes.append(list(keys))
            for row in payload["records"]:
                outcomes.append([safe_xlsx_cell(row[k]) for k in keys])
        contract = workbook.create_sheet("口径与边界")
        for key, value in summary.items():
            if key == "observed_actions":
                value = {k: v for k, v in value.items() if k != "records"}
            encoded = json.dumps(value, ensure_ascii=False, indent=2)
            for offset in range(0, len(encoded), 30_000):
                contract.append(
                    [key, safe_xlsx_cell(encoded[offset : offset + 30_000])]
                )
        workbook.save(stream)
    elif format == "docx":
        document = Document()
        document.add_heading("历史决策回放与对账", 0)
        document.add_paragraph(
            "结果来自平台冻结包的历史模拟。导入的历史动作和现金流不等于 MARVIS 线上执行证明；场景差异不等于因果增益。"
        )
        for key, value in facts:
            document.add_paragraph(f"{key}: {value}")
        document.add_heading("口径与证据", 1)
        document.add_paragraph(json.dumps(summary, ensure_ascii=False, indent=2))
        for scenario in payload.get("scenarios", []):
            document.add_heading(scenario["name"], 1)
            document.add_paragraph(
                json.dumps(
                    {k: v for k, v in scenario.items() if k != "decisions"},
                    ensure_ascii=False,
                    indent=2,
                )
            )
        if temporal_rows:
            document.add_heading("时间稳定性检查", 1)
            document.add_paragraph(
                "分箱仅由参考窗拟合；本检查描述历史分布变化，不证明历史部署、样本外效果或因果增益。缺乏支持的检查保留 unknown。"
            )
            table = document.add_table(rows=1, cols=6)
            for cell, value in zip(
                table.rows[0].cells,
                ("方案", "对照窗", "指标", "结果", "阈值", "状态 / 原因"),
                strict=True,
            ):
                cell.text = value
            for row in temporal_rows:
                values = (
                    row[0],
                    row[5],
                    row[10],
                    row[11],
                    row[12],
                    row[14] + (" / " + row[15] if row[15] else ""),
                )
                for cell, value in zip(table.add_row().cells, values, strict=True):
                    cell.text = "unknown" if value is None else str(value)
        if event_rows:
            document.add_heading("逐笔原生事件收据引用", 1)
            document.add_paragraph(
                "每笔申请绑定原生签名事件收据；unknown 表示来源覆盖不足，触发已批准的复核或拒绝策略。来源为发布者声明，不构成身份或欺诈证明。"
            )
            for row in event_rows:
                document.add_paragraph(
                    json.dumps(
                        dict(zip(EVENT_COLUMNS, row, strict=True)), ensure_ascii=False
                    )
                )
        # Complete row evidence is retained in the canonical JSON / Excel; the
        # narrative export explicitly states its summary scope.
        document.add_paragraph(
            "逐笔明细见相同 artifact_id 的 JSON 或 Excel；本 Word 展示完整场景汇总及合同。"
        )
        document.save(stream)
    else:
        raise ValueError("unsupported historical export format")
    return stream.getvalue()
