"""Render the persisted deterministic verdict without recalculating it."""

from io import BytesIO
import json

from docx import Document
from openpyxl import Workbook

from marvis.business_acceptance import acceptance_display_rows
from marvis.spreadsheet_safety import safe_xlsx_cell


def render_business_acceptance(review: dict, format: str) -> bytes:
    result = review["business_acceptance"]
    rows = acceptance_display_rows(result)
    facts = [
        ("execution_completed", review.get("execution_completed")),
        ("business_status", result["status"]),
        ("summary_ref", review.get("summary_ref")),
        ("objective_hash", result.get("objective_hash")),
        ("target", json.dumps(result.get("target"), ensure_ascii=False)),
        ("evidence", json.dumps(result.get("evidence"), ensure_ascii=False)),
        ("reasons", "; ".join(result.get("reasons") or [])),
    ]
    columns = ("metric", "value", "unit", "denominator", "status", "reason")
    stream = BytesIO()
    if format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "业务验收"
        for fact in facts:
            sheet.append([safe_xlsx_cell(value) for value in fact])
        sheet.append([])
        sheet.append(list(columns))
        for row in rows:
            sheet.append([safe_xlsx_cell(row.get(key)) for key in columns])
        contract = workbook.create_sheet("验收合同")
        for key, value in (result.get("objective") or {}).items():
            contract.append(
                [
                    safe_xlsx_cell(key),
                    safe_xlsx_cell(json.dumps(value, ensure_ascii=False)),
                ]
            )
        workbook.save(stream)
    elif format == "docx":
        document = Document()
        document.add_heading("业务验收", 0)
        document.add_paragraph(
            "执行完成与业务达标分别记录。以下结论来自同一份已保存的确定性判定。"
        )
        table = document.add_table(rows=0, cols=2)
        for key, value in facts:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = key, "" if value is None else str(value)
        table = document.add_table(rows=1, cols=len(columns))
        for cell, name in zip(table.rows[0].cells, columns, strict=True):
            cell.text = name
        for row in rows:
            for cell, name in zip(table.add_row().cells, columns, strict=True):
                cell.text = "" if row.get(name) is None else str(row[name])
        document.add_heading("验收合同", 1)
        document.add_paragraph(
            json.dumps(result.get("objective"), ensure_ascii=False, indent=2)
        )
        document.save(stream)
    else:
        raise ValueError("unsupported acceptance export format")
    return stream.getvalue()
