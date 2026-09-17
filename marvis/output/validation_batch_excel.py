from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
import math
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from marvis.artifacts import ArtifactUnitOfWork
from marvis.output.styles import (
    BORDER_COLOR,
    BRAND_HEADER_FILL,
    BRAND_HEADER_FONT_COLOR,
    FONT_NAME,
    STRESS_HIGH_FILL,
    STRESS_LOW_FILL,
    STRESS_MEDIUM_FILL,
    stress_psi_risk,
    stress_risk_cell_color,
)
from marvis.output.xlsx_safety import safe_xlsx_cell
from marvis.validation.stress_risk import stress_risk_label


SUMMARY_SHEET_NAME = "模型验证汇总"
OPS_SHEET_NAME = "运营汇总"
_ASSETS_DIR = Path(__file__).resolve().parent / "assets"
_PASS_MARK_PATH = _ASSETS_DIR / "validation_pass_mark.png"
_FAIL_MARK_PATH = _ASSETS_DIR / "validation_fail_mark.png"

_VISIBLE_HEADERS = (
    "模型版本",
    "有效性 KS（OOT，仅展示不参与通过判定）",
    "稳定性 PSI（OOT vs train）",
    "PMML打分",
    "压力测试",
    "报告完整性",
    "处理结果",
    "人工校验",
    "失败/复核原因",
)
_INTERNAL_HEADERS = ("批次项ID", "子任务ID")
_THIN_SIDE = Side(style="thin", color=BORDER_COLOR)
_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
_WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")


@dataclass(frozen=True)
class ValidationBatchSummaryRow:
    ordinal: int
    item_id: str
    child_task_id: str
    model_name: str
    model_version: str
    oot_ks: float | None
    oot_psi: float | None
    pmml_status: str
    stress_risk: str | None
    report_complete: bool
    outcome: str
    manual_review_url: str
    error_message: str


def write_validation_batch_excel(
    *,
    batch_name: str,
    created_at: str | datetime,
    rows: Iterable[ValidationBatchSummaryRow],
    output_path: Path,
) -> Path:
    """Write the batch workbook: audit sheet plus the operational KS/PSI sheet.

    OOT KS is deliberately informational. The workbook only visualizes the
    deterministic gates supplied by the batch runner; it does not recompute a
    different pass/fail verdict while rendering.
    """

    row_values = sorted(list(rows), key=lambda row: row.ordinal)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    uow = ArtifactUnitOfWork()
    artifact = uow.stage_file(output_path.parent, output_path.name)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SUMMARY_SHEET_NAME
    try:
        _write_summary_sheet(
            sheet,
            batch_name=batch_name,
            created_at=created_at,
            rows=row_values,
        )
        _write_ops_sheet(workbook.create_sheet(OPS_SHEET_NAME), created_at=created_at, rows=row_values)
        workbook.save(artifact.path)
        uow.promote_all()
        uow.commit()
        return artifact.final_path
    except Exception:
        uow.rollback()
        raise


def _write_summary_sheet(
    sheet,
    *,
    batch_name: str,
    created_at: str | datetime,
    rows: list[ValidationBatchSummaryRow],
) -> None:
    year, month = _year_month(created_at)
    title = f"模型验证：{year}年{month}月验证{len(rows)}个模型"

    sheet.merge_cells("A1:I1")
    title_cell = sheet["A1"]
    title_cell.value = title
    title_cell.font = Font(name=FONT_NAME, size=18, bold=True, color="E25555")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 36

    sheet.merge_cells("A2:I2")
    subtitle = sheet["A2"]
    subtitle.value = safe_xlsx_cell(
        f"批次：{batch_name or '-'}｜通过判定不使用 OOT KS；非通过项请点击“人工校验”查看单模型证据。"
    )
    subtitle.font = Font(name=FONT_NAME, size=9, color="666666")
    subtitle.alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 24

    headers = _VISIBLE_HEADERS + _INTERNAL_HEADERS
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor=BRAND_HEADER_FILL)
        cell.font = Font(
            name=FONT_NAME,
            size=10,
            bold=True,
            color=BRAND_HEADER_FONT_COLOR,
        )
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    sheet.row_dimensions[3].height = 42

    for output_row, summary in enumerate(rows, start=4):
        _write_model_row(sheet, output_row, summary)

    legend_row = max(6, 5 + len(rows))
    _write_risk_legend(sheet, legend_row)

    widths = {
        "A": 35,
        "B": 25,
        "C": 24,
        "D": 14,
        "E": 16,
        "F": 16,
        "G": 17,
        "H": 15,
        "I": 42,
        "J": 24,
        "K": 24,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["J"].hidden = True
    sheet.column_dimensions["K"].hidden = True

    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:I{max(3, 3 + len(rows))}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:3"
    sheet.print_area = f"A1:I{legend_row + 4}"


_OPS_HEADERS = (
    "模型版本",
    "有效性 KS",
    "稳定性 PSI",
    "一致性",
    "可复现性",
    "完整性",
)


def _write_ops_sheet(sheet, *, created_at: str | datetime, rows: list[ValidationBatchSummaryRow]) -> None:
    """Write the operator-facing sheet matching the 2026-08 nine-model workbook."""

    year, month = _year_month(created_at)
    title = f"模型验证：{year}年{month}月验证{len(rows)}个模型"

    sheet.merge_cells("A1:F1")
    title_cell = sheet["A1"]
    title_cell.value = title
    title_cell.font = Font(name=FONT_NAME, size=18, bold=True, color="E25555")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 32

    for column, header in enumerate(_OPS_HEADERS, start=1):
        cell = sheet.cell(row=2, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor=BRAND_HEADER_FILL)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color=BRAND_HEADER_FONT_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    sheet.row_dimensions[2].height = 24

    for output_row, summary in enumerate(rows, start=3):
        _write_ops_model_row(sheet, output_row, summary)

    last_row = max(2, 2 + len(rows))
    widths = {"A": 32, "B": 14, "C": 14, "D": 12, "E": 12, "F": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:2"
    sheet.print_area = f"A1:F{last_row}"


def _write_ops_model_row(sheet, output_row: int, summary: ValidationBatchSummaryRow) -> None:
    model_label = summary.model_name.strip() or f"模型 {summary.ordinal}"
    values = (
        model_label,
        _ks_percent(summary.oot_ks),
        _finite_number(summary.oot_psi),
        None,
        None,
        None,
    )
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=output_row, column=column, value=safe_xlsx_cell(value) if value is not None else None)
        cell.font = Font(name=FONT_NAME, size=11, color="333333")
        cell.alignment = Alignment(
            horizontal="left" if column == 1 else "center",
            vertical="center",
        )
        cell.border = _BORDER
        cell.fill = _WHITE_FILL
    sheet.row_dimensions[output_row].height = 28
    sheet.cell(output_row, 2).number_format = "0.0"
    sheet.cell(output_row, 3).number_format = "0.00"

    consistency, reproducibility, completeness = _ops_gate_flags(summary)
    for column, passed in enumerate((consistency, reproducibility, completeness), start=4):
        _add_ops_mark(sheet, output_row, column, passed)


def _ops_gate_flags(summary: ValidationBatchSummaryRow) -> tuple[bool, bool, bool]:
    consistency = str(summary.pmml_status).lower() == "pass"
    # V2 主路径把 Notebook 分数一致性换成 PMML 全量打分；可复现性绿灯与一致性同源。
    reproducibility = consistency and str(summary.outcome).lower() not in {"failed", "fail"}
    completeness = bool(summary.report_complete)
    return consistency, reproducibility, completeness


def _add_ops_mark(sheet, row: int, column: int, passed: bool) -> None:
    mark_path = _PASS_MARK_PATH if passed else _FAIL_MARK_PATH
    if not mark_path.is_file():
        cell = sheet.cell(row, column)
        cell.value = "✓" if passed else "!"
        cell.font = Font(name=FONT_NAME, size=14, bold=True, color="548235" if passed else "C00000")
        return
    image = XLImage(str(mark_path))
    image.anchor = f"{_column_letter(column)}{row}"
    sheet.add_image(image)


def _column_letter(column: int) -> str:
    return chr(ord("A") + column - 1)


def _write_model_row(sheet, output_row: int, summary: ValidationBatchSummaryRow) -> None:
    pmml_passed = str(summary.pmml_status).lower() == "pass"
    outcome = str(summary.outcome).lower()
    model_label = " ".join(
        part for part in (summary.model_name.strip(), summary.model_version.strip()) if part
    ) or f"模型 {summary.ordinal}"
    stress_risk = summary.stress_risk if summary.stress_risk in {"low", "medium", "high"} else None

    values = (
        model_label,
        _ks_percent(summary.oot_ks),
        _finite_number(summary.oot_psi),
        "✓" if pmml_passed else "! 人工校验",
        f"✓ {stress_risk_label(stress_risk)}" if stress_risk == "low" else f"! {stress_risk_label(stress_risk)}",
        "✓" if summary.report_complete else "! 人工校验",
        _outcome_label(outcome),
        "人工校验" if outcome != "pass" and summary.manual_review_url else "",
        summary.error_message,
        summary.item_id,
        summary.child_task_id,
    )
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=output_row, column=column, value=safe_xlsx_cell(value))
        cell.font = Font(name=FONT_NAME, size=9, bold=column in {4, 5, 6, 7, 8})
        cell.alignment = Alignment(
            horizontal="left" if column in {1, 9} else "center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = _BORDER
        cell.fill = _WHITE_FILL
    sheet.row_dimensions[output_row].height = 32

    sheet.cell(output_row, 2).number_format = "0.0"
    sheet.cell(output_row, 3).number_format = "0.0000"

    _apply_risk_fill(sheet.cell(output_row, 3), stress_psi_risk(summary.oot_psi))
    _apply_risk_fill(sheet.cell(output_row, 4), "low" if pmml_passed else "high")
    _apply_risk_fill(sheet.cell(output_row, 5), stress_risk)
    _apply_risk_fill(sheet.cell(output_row, 6), "low" if summary.report_complete else "high")
    _apply_risk_fill(
        sheet.cell(output_row, 7),
        "low" if outcome == "pass" else "high" if outcome in {"failed", "fail"} else "medium",
    )

    link_cell = sheet.cell(output_row, 8)
    if link_cell.value and summary.manual_review_url:
        link_cell.hyperlink = summary.manual_review_url
        link_cell.font = Font(name=FONT_NAME, size=9, bold=True, color="0563C1", underline="single")
        link_cell.fill = PatternFill("solid", fgColor=STRESS_MEDIUM_FILL)


def _write_risk_legend(sheet, start_row: int) -> None:
    panel_columns = ((1, 3), (4, 5), (6, 7), (8, 9))
    for start_column, end_column in panel_columns:
        sheet.merge_cells(
            start_row=start_row,
            start_column=start_column,
            end_row=start_row,
            end_column=end_column,
        )
        sheet.merge_cells(
            start_row=start_row + 1,
            start_column=start_column,
            end_row=start_row + 3,
            end_column=end_column,
        )

    intro_heading = sheet.cell(start_row, 1, "压力测试")
    intro_heading.fill = PatternFill("solid", fgColor="E25555")
    intro_heading.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
    intro_body = sheet.cell(
        start_row + 1,
        1,
        "压力测试主要基于某个信源缺失情境下模型关键指标的偏移及影响进行测试，"
        "关注有效性 KS、稳定性 PSI 及 OOT 分数偏移。",
    )

    legends = (
        (
            "✓ 低风险",
            "KS衰减范围：[0,10%)\nPSI范围：[0,0.10)\n特征缺失影响不大，可继续使用，但需常规稳定性监控。",
            STRESS_LOW_FILL,
        ),
        (
            "− 中风险",
            "KS衰减范围：[10%,20%)\nPSI范围：[0.10,0.25)\n达到预警阈值，异常时需人工核验并及时通知模型团队。",
            STRESS_MEDIUM_FILL,
        ),
        (
            "! 高风险",
            "KS衰减范围：[20%,+∞)\nPSI范围：[0.25,+∞)\n触发实时监控与熔断评估，应立即降级至备用方案或启动人工复核。",
            STRESS_HIGH_FILL,
        ),
    )
    for index, (label, description, fill) in enumerate(legends, start=1):
        start_column, _ = panel_columns[index]
        marker = sheet.cell(start_row, start_column, label)
        description_cell = sheet.cell(start_row + 1, start_column, description)
        marker.fill = PatternFill("solid", fgColor=fill)
        marker.font = Font(name=FONT_NAME, size=10, bold=True)
        marker.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        description_cell.font = Font(name=FONT_NAME, size=9)
        description_cell.alignment = Alignment(vertical="center", wrap_text=True)
    intro_body.font = Font(name=FONT_NAME, size=9)
    intro_body.alignment = Alignment(vertical="center", wrap_text=True)
    for row in range(start_row, start_row + 4):
        for column in range(1, 10):
            sheet.cell(row, column).border = _BORDER
    sheet.row_dimensions[start_row].height = 26
    for row in range(start_row + 1, start_row + 4):
        sheet.row_dimensions[row].height = 28


def _apply_risk_fill(cell, risk: str | None) -> None:
    color = stress_risk_cell_color(risk)
    if color:
        cell.fill = PatternFill("solid", fgColor=color)


def _outcome_label(outcome: str) -> str:
    if outcome == "pass":
        return "✓ 通过"
    if outcome in {"failed", "fail"}:
        return "! 验证失败"
    return "! 人工校验"


def _finite_number(value: float | None) -> float | None:
    if value is None:
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def _ks_percent(value: float | None) -> float | None:
    number = _finite_number(value)
    return None if number is None else number * 100.0


def _year_month(value: str | datetime) -> tuple[int, int]:
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value).strip().replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(text)
        except ValueError as exc:
            raise ValueError("created_at must be an ISO-8601 datetime") from exc
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone()
    return parsed.year, parsed.month


__all__ = [
    "OPS_SHEET_NAME",
    "SUMMARY_SHEET_NAME",
    "ValidationBatchSummaryRow",
    "write_validation_batch_excel",
]
