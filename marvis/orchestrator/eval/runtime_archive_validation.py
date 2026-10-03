"""Read-only domain checks of retained V2 validation evidence.

Caller-held bindings are mandatory. Archive flags are never proof. Only installed
platform algorithms and the installed PMML scorer execute; uploaded notebooks,
pickles, plugins and archived Python are never imported or executed. Temporary
stress outputs are written outside the archive and original workspace.
"""
from __future__ import annotations

from contextlib import ExitStack, closing
from dataclasses import asdict, replace
import io
import json
import math
from pathlib import Path
import re
import sqlite3
import tempfile
import zipfile

import numpy as np
import pandas as pd
from pydantic import BaseModel, ConfigDict, Field, model_validator

from .runtime_contracts import RuntimeCase, digest
from .runtime_custody import verify_case_archive


_HASH = r"^[0-9a-f]{64}$"
_MAX_BYTES = 128 * 1024**2
_MAX_ROWS = 100_000
_MAX_CELLS = 2_000_000
_MAX_SNAPSHOT_BYTES = 1024**3
_CHECKS = (
    "archive_integrity", "recomputation_snapshot", "external_bindings", "task_contract", "materials",
    "pmml_rescoring", "metrics_recomputation", "stress_recomputation",
    "confirmation_record_binding", "report_download_binding", "excel_values", "word_confirmed_text",
)


class FrozenValidationBinding(BaseModel):
    """Supply from independently held run/case records, never from this archive."""
    model_config = ConfigDict(extra="forbid", strict=True, frozen=True)
    case: RuntimeCase
    run_id: str = Field(min_length=1, max_length=160)
    task_id: str = Field(pattern=r"^[a-zA-Z0-9_-]{1,160}$")
    cases_sha256: str = Field(pattern=_HASH)
    expected_sha256: str = Field(pattern=_HASH)
    source: dict
    model_connection_sha256: str = Field(pattern=_HASH)
    input_contract_sha256: str = Field(pattern=_HASH)
    confirmed_draft_sha256: str | None
    report_revision: int | None
    report_sha256: dict[str, str | None]
    bin_count: int = Field(ge=2, le=100)

    @model_validator(mode="after")
    def validate_external_scope(self):
        if (self.case.task.task_type != "validation" or not self.case.actions
                or self.case.actions[0].kind != "start_validation_agent"):
            raise ValueError("only the standard V2 validation Agent entry is supported")
        if set(self.report_sha256) != {"word", "excel"}:
            raise ValueError("both report bindings must be explicitly provided, including absent values")
        hashes = [self.source.get("source_sha256"), self.confirmed_draft_sha256, *self.report_sha256.values()]
        if not isinstance(hashes[0], str) or not re.fullmatch(_HASH, hashes[0]):
            raise ValueError("external source identity is required")
        if any(value is not None and (not isinstance(value, str) or not re.fullmatch(_HASH, value)) for value in hashes):
            raise ValueError("external identity must be a SHA256 digest")
        if self.report_revision is not None and (type(self.report_revision) is not int or self.report_revision < 1):
            raise ValueError("external report revision must be positive")
        return self


class _Unsupported(ValueError):
    pass


def _file(archive, relative, *, maximum=_MAX_BYTES):
    relative = Path(relative)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("archive path boundary")
    path = archive / relative
    if path.resolve() != path or not path.is_file() or path.stat().st_size > maximum:
        raise ValueError("archive file boundary")
    return path


def _json(archive, relative):
    return json.loads(_file(archive, relative, maximum=16 * 1024**2).read_bytes())


def _authenticated_snapshot(archive, destination, expected_digest):
    from .runtime_custody import _copy_bounded

    raw = _file(archive, "manifest.json", maximum=16 * 1024**2).read_bytes()
    if digest(raw) != expected_digest:
        raise ValueError("manifest changed before snapshot")
    inventory = json.loads(raw)["files"]
    if sum(item["size_bytes"] for item in inventory.values()) > _MAX_SNAPSHOT_BYTES:
        raise _Unsupported("private recomputation snapshot exceeds one GiB")
    # The independently held manifest authenticates this copy before any domain
    # parser sees it. Replacing the original afterward cannot change its bytes.
    for relative, item in inventory.items():
        target = destination / relative
        for parent in reversed(target.parents):
            if parent.is_relative_to(destination):
                parent.mkdir(mode=0o700, exist_ok=True)
        _copy_bounded(_file(archive, relative, maximum=_MAX_SNAPSHOT_BYTES), target, item["size_bytes"])
    manifest = destination / "manifest.json"
    manifest.touch(mode=0o600, exist_ok=False)
    manifest.write_bytes(raw)
    verify_case_archive(destination, expected_manifest_sha256=expected_digest)


def _validate_pmml_resources(raw):
    from defusedxml.ElementTree import iterparse
    from marvis.validation.pmml_manifest import MAX_XML_DEPTH, MAX_XML_NODES

    count = depth = 0
    for event, element in iterparse(io.BytesIO(raw), events=("start", "end"),
                                   forbid_dtd=True, forbid_entities=True, forbid_external=True):
        if event == "start":
            count += 1
            depth += 1
            if depth > MAX_XML_DEPTH or count > MAX_XML_NODES:
                raise _Unsupported("PMML XML inspection limit")
            if element.tag.rsplit("}", 1)[-1] in {"TableLocator", "Extension"}:
                raise _Unsupported("PMML external or extension content")
        else:
            depth -= 1
            element.clear()


def _matches(left, right):
    if isinstance(left, dict) and isinstance(right, dict):
        return left.keys() == right.keys() and all(_matches(left[k], right[k]) for k in left)
    if isinstance(left, (list, tuple)) and isinstance(right, (list, tuple)):
        return len(left) == len(right) and all(_matches(a, b) for a, b in zip(left, right, strict=True))
    if type(left) in {int, float} and type(right) in {int, float}:
        return (left == right if not math.isfinite(left) or not math.isfinite(right)
                else math.isclose(left, right, rel_tol=1e-10, abs_tol=1e-12))
    return type(left) is type(right) and left == right


def _require_equal(actual, expected):
    if not _matches(actual, expected):
        raise ValueError("recomputed value differs")


def _public_source_identity(source):
    # A separately supplied binding may contain accidental credentials or
    # institution metadata. Compare its complete contents internally, but never
    # echo arbitrary extensions in a review result (including mismatch results).
    patterns = {"commit": r"[0-9a-f]{40}", "source_sha256": r"[0-9a-f]{64}",
                "dirty_diff_sha256": r"[0-9a-f]{64}"}
    return {key: source[key] for key, pattern in patterns.items()
            if isinstance(source.get(key), str) and re.fullmatch(pattern, source[key])}


def _original_path(archive, original_root, absolute):
    root, path = Path(original_root), Path(absolute)
    if not root.is_absolute() or not path.is_absolute() or ".." in path.parts:
        raise ValueError("original path identity invalid")
    relative = path.relative_to(root)
    # This maps immutable identities to retained bytes; it does not rewrite DB,
    # signatures or the old directory, and never opens the former absolute path.
    return _file(archive, Path("workspace") / relative)


def _read_contract(archive, binding):
    from marvis.repositories.validation_contracts import _row_to_contract_record

    path = _file(archive, "workspace/marvis.sqlite")
    with closing(sqlite3.connect(path.as_uri() + "?mode=ro&immutable=1", uri=True)) as conn:
        conn.row_factory = sqlite3.Row
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (binding.task_id,)).fetchone()
        row = conn.execute("SELECT * FROM validation_input_contracts WHERE task_id=?", (binding.task_id,)).fetchone()
        if task is None or row is None:
            raise ValueError("task or input contract absent")
        record = _row_to_contract_record(row)
        messages = [dict(r) for r in conn.execute(
            "SELECT * FROM agent_messages WHERE task_id=? ORDER BY created_at,id LIMIT 5001", (binding.task_id,))]
        audits = [dict(r) for r in conn.execute(
            "SELECT * FROM audit WHERE target_ref=? AND kind='report.agent_conclusions.confirm' LIMIT 5001", (binding.task_id,))]
    if len(messages) > 5000 or len(audits) > 5000:
        raise _Unsupported("archive message review limit")
    return dict(task), record.contract, messages, audits


def _materials(archive, manifest, binding, task, contract):
    from marvis.validation.feature_metadata import FeatureMetadataSelection, normalize_feature_metadata
    from marvis.validation.pmml_manifest import parse_pmml_input_manifest
    from marvis.validation.sample_chunks import iter_sample_chunks
    from marvis.validation.sample_schema import inspect_sample_schema

    if len({Path(m.path).name for m in binding.case.materials}) != len(binding.case.materials):
        raise ValueError("ambiguous material basenames")
    material_paths = {}
    for index, material in enumerate(binding.case.materials):
        path = _file(archive, f"inputs/{index:03d}{Path(material.path).suffix}")
        if task[material.role + "_path"] != material.path:
            raise ValueError("task material selection changed")
        uploaded = _original_path(archive, manifest["original_workspace"], Path(task["source_dir"]) / material.path)
        if (digest(path.read_bytes()) != material.sha256 or digest(uploaded.read_bytes()) != material.sha256
                or contract.material_hashes.get(material.role) != material.sha256):
            raise ValueError("frozen material mismatch")
        # Select retained bytes by the original task selection, not a basename
        # search. The original absolute schema identity is checked below.
        material_paths[material.role] = uploaded
    original_schema = contract.require_sample_schema()
    if _original_path(archive, manifest["original_workspace"], original_schema.path) != material_paths["sample"]:
        raise ValueError("sample schema original path does not identify selected material")
    reading_schema = inspect_sample_schema(material_paths["sample"])
    original_fields, reading_fields = asdict(original_schema), asdict(reading_schema)
    original_fields.pop("path")
    reading_fields.pop("path")
    _require_equal(reading_fields, original_fields)
    # Separate in-memory read context, independently inspected from retained
    # bytes. The frozen contract, original path, DB and signature are untouched.
    reading_contract = replace(contract, sample_schema=reading_schema)
    # XML-aware checks cover namespaced elements and alternate encodings before
    # handing declarative data to the installed PMML engine.
    _validate_pmml_resources(material_paths["pmml"].read_bytes())
    parsed_manifest = parse_pmml_input_manifest(material_paths["pmml"])
    _require_equal(asdict(parsed_manifest), asdict(contract.require_pmml_manifest()))
    selected = contract.confirmed
    metadata = normalize_feature_metadata(
        material_paths["dictionary"], manifest=parsed_manifest,
        selection=FeatureMetadataSelection(selected["metadata_sheet"], selected["feature_col"],
                                            selected["category_col"], selected["importance_col"]),
    )
    _require_equal(asdict(metadata), asdict(contract.require_feature_metadata()))
    count = 0
    for chunk in iter_sample_chunks(material_paths["sample"], columns=(), chunk_size=10_000, schema=reading_schema):
        count += len(chunk.row_ids)
        if count > _MAX_ROWS or count * max(1, len(parsed_manifest.raw_required_fields)) > _MAX_CELLS:
            raise _Unsupported("domain recomputation row or cell limit")
    if len(metadata.per_category_raw_fields) > 64:
        raise _Unsupported("domain recomputation stress category limit")
    return material_paths, metadata, count, reading_contract


def _rescore(archive, outputs, paths, contract, count, reading_contract):
    from marvis.validation.field_transformations import apply_confirmed_transformations, required_transformation_inputs, transformation_closure
    from marvis.validation.pmml_score_artifacts import SCORING_ENGINE, build_pmml_scoring_identity, validate_pmml_score_artifact
    from marvis.validation.pmml_scoring import load_pmml_scorer
    from marvis.validation.results import pmml_scoring_result_from_dict
    from marvis.validation.sample_chunks import iter_sample_chunks

    scoring = pmml_scoring_result_from_dict(_json(archive, outputs / "pmml_scoring_result.json"))
    identity = build_pmml_scoring_identity(contract=contract, pmml_sha256=contract.material_hashes["pmml"],
        sample_sha256=contract.material_hashes["sample"], chunk_size=scoring.chunk_size)
    if scoring.engine_version != identity.engine_version:
        raise _Unsupported("installed PMML engine differs from original")
    if (scoring.input_row_count != count or scoring.pmml_sha256 != contract.material_hashes["pmml"]
            or scoring.sample_sha256 != contract.material_hashes["sample"]
            or scoring.score_artifact_path != "pmml_scores.parquet"
            or scoring.output_field != contract.require_output_field()):
        raise ValueError("score material identity mismatch")
    score_path = _file(archive, outputs / "pmml_scores.parquet")
    validate_pmml_score_artifact(scoring, score_path, expected_cache_key=identity.cache_key)
    scores = pd.read_parquet(score_path)
    if not np.array_equal(scores["row_id"].to_numpy(), np.arange(count)):
        raise ValueError("score row ordering differs")
    fields = contract.require_pmml_manifest().raw_required_fields
    transforms = transformation_closure(fields, contract.transformations)
    projection = tuple(required_transformation_inputs(fields, transforms))
    scorer = load_pmml_scorer(paths["pmml"], contract.require_output_field())
    offset = 0
    for chunk in iter_sample_chunks(paths["sample"], columns=projection, chunk_size=10_000, schema=reading_contract.require_sample_schema()):
        inputs = apply_confirmed_transformations(chunk.frame, transforms).loc[:, list(fields)]
        computed = scorer.score_chunk(inputs).to_numpy(dtype=float)
        retained = scores["pmml_score"].iloc[offset:offset + len(computed)].to_numpy(dtype=float)
        if not np.isfinite(computed).all() or not np.allclose(computed, retained, rtol=1e-10, atol=1e-12):
            raise ValueError("recomputed PMML scores differ")
        offset += len(computed)
    if offset != count:
        raise ValueError("recomputed score population differs")
    if (scoring.status != "pass" or scoring.success_count != count or scoring.failure_count
            or scoring.null_count or scoring.non_finite_count or scoring.missing_inputs or scoring.bounded_errors
            or scoring.required_input_count != len(fields) or scoring.engine != SCORING_ENGINE):
        raise ValueError("scoring summary differs from independently recomputed finite scores")
    return scorer, scoring, score_path


def _confirmation(task, messages, audits, binding):
    from marvis.repositories.tasks import AGENT_REPORT_CONCLUSION_KEYS

    if binding.confirmed_draft_sha256 is None or binding.report_revision is None:
        raise _Unsupported("no separately retained confirmation binding")
    values = json.loads(task["report_values_json"])
    if task["report_values_revision"] != binding.report_revision:
        raise ValueError("report revision mismatch")
    matches = []
    for message in messages:
        metadata = json.loads(message["metadata_json"])
        draft = metadata.get("draft_values")
        if (message["role"] == "assistant" and message["stage"] == "word_conclusion_draft"
                and isinstance(draft, dict) and digest(draft) == binding.confirmed_draft_sha256):
            if (metadata.get("report_revision") != binding.report_revision - 1
                    or not AGENT_REPORT_CONCLUSION_KEYS <= draft.keys()
                    or any(values.get(key) != value for key, value in draft.items())):
                raise ValueError("confirmed draft contents mismatch")
            matches.append(draft)
    if not matches or not any(a["outcome"] == "succeeded" and json.loads(a["detail_json"]).get("expected_revision") == binding.report_revision - 1 for a in audits):
        raise ValueError("confirmation records absent")
    return values


def _ooxml(path, kind):
    from defusedxml.ElementTree import fromstring

    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if (len(entries) > 2048 or len({e.filename for e in entries}) != len(entries)
                or sum(e.file_size for e in entries) > 64 * 1024**2
                or any(e.file_size > 32 * 1024**2 for e in entries)):
            raise _Unsupported("report ZIP expansion limit")
        if any("vbaproject" in e.filename.lower() for e in entries):
            raise _Unsupported("active report content")
        for item in entries:
            if item.filename.endswith(".rels"):
                root = fromstring(archive.read(item))
                if any(e.attrib.get("TargetMode") == "External" for e in root.iter()):
                    raise _Unsupported("external report relationship")
        part = "word/document.xml" if kind == "word" else "xl/workbook.xml"
        return fromstring(archive.read(part))


def _excel(path, results, report_values):
    from openpyxl import Workbook, load_workbook
    from marvis.output import excel

    expected = Workbook()
    expected.remove(expected.active)
    for writer in (excel._write_overview, excel._write_basic_info, excel._write_monthly_distribution, excel._write_hyperparameters,
                   excel._write_feature_importance, excel._write_effectiveness_overall,
                   excel._write_psi_stability, excel._write_monthly_effectiveness, excel._write_stress_summary):
        writer(expected, results)
    excel._write_report_texts(expected, report_values)
    for split in ("train", "test", "oot"):
        excel._write_bins(expected, f"分箱_{split}", results.effectiveness.bin_tables[split], first_header=f"{split}(按照train分箱)")
        excel._write_bins(expected, f"分箱_独立10等分_{split}", results.effectiveness.independent_quantile_bin_tables[split], first_header=f"{split}(独立10等分)")
    for category in results.stress_test.per_category:
        excel._write_bins(expected, f"压力测试_分箱_{category.category}", category.bin_table,
                          first_header=category.category, include_bin_share=True)
    # Match the text cells of the standard chart sheet without rendering images.
    charts = expected.create_sheet("ROC_KS曲线")
    for index, split in enumerate(("train", "test", "oot")):
        charts.cell(index * 32 + 1, 1, f"{split} ROC曲线和KS曲线")
    actual = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if set(actual.sheetnames) != set(expected.sheetnames):
            raise ValueError("report worksheet set differs from standard V2")
        checked = 0
        for sheet in expected:
            other = actual[sheet.title]
            if other.max_row > 200_000 or other.max_column > 128 or other.max_row * other.max_column > _MAX_CELLS:
                raise _Unsupported("report cell inspection limit")
            rows = [tuple(None if c.value == "" else c.value for c in row) for row in sheet]
            for row_index, observed in enumerate(other.iter_rows(
                    max_row=max(len(rows), other.max_row), max_col=max(sheet.max_column, other.max_column), values_only=True)):
                expected_row = rows[row_index] if row_index < len(rows) else ()
                for column, value in enumerate(observed):
                    expected_value = expected_row[column] if column < len(expected_row) else None
                    # Only the two actual timing values are excluded. Their
                    # labels and extra cells on these rows still must match.
                    if (sheet.title == "验证总览" and column == 1 and expected_row
                            and expected_row[0] in {"打分耗时(秒)", "吞吐(行/秒)"}):
                        continue
                    _require_equal(value, expected_value)
                    checked += column < len(expected_row)
        return checked
    finally:
        actual.close()
        expected.close()


def revalidate_validation_archive(
    archive: Path, *, expected_manifest_sha256: str, frozen_binding: FrozenValidationBinding,
) -> dict:
    """Integrity, bound identities, and domain recomputation are separate claims."""
    from .runtime_runner import _source_identity
    from marvis.validation.config import ValidationConfig
    from marvis.validation.input_contracts import input_contract_to_dict
    from marvis.validation.pmml_analysis import load_pmml_analysis_frame
    from marvis.validation.platform_metrics import compute_platform_validation_results
    from marvis.validation.pmml_stress import run_pmml_stress
    from marvis.validation.results import validation_results_from_dict

    if not isinstance(frozen_binding, FrozenValidationBinding):
        raise ValueError("an explicit FrozenValidationBinding is required")
    binding = frozen_binding.model_copy(deep=True)
    archive = original_archive = Path(archive).absolute()
    temporary = ExitStack()
    verifier_source = _source_identity()
    result = {"schema": "marvis.runtime-validation-revalidation.v1", "case_id": binding.case.id,
        "algorithm": "installed_platform_v2_deterministic_recomputation.v1",
        "verifier_source": verifier_source, "original_source_binding": _public_source_identity(binding.source),
        "original_source_binding_sha256": digest(binding.source),
        "verifier_source_matches_original": verifier_source["source_sha256"] == binding.source["source_sha256"],
        "recomputation_isolation": "private_authenticated_copy; excludes_adversarial_same_uid_or_root_access",
        "source_authentication": "not_established", "acceptance_claim": "not_established",
        "checks": {key: {"status": "unverified"} for key in _CHECKS},
        "unsupported": {
            "native_signature_authentication": "original_key_not_available",
            "narrative_semantic_correctness": "requires_independent_review",
            "business_labels_and_training_validity": "not_established_by_recomputation",
            "original_runtime_cost_and_timing": "not_recomputed",
            "report_chart_pixels_and_word_numeric_layout": "not_recomputed",
        }}
    checks = result["checks"]
    current = "archive_integrity"
    try:
        verify_case_archive(archive, expected_manifest_sha256=expected_manifest_sha256)
        checks[current] = {"status": "verified", "scope": "bytes_match_caller_held_manifest_digest"}
        current = "recomputation_snapshot"
        snapshot = Path(temporary.enter_context(tempfile.TemporaryDirectory(prefix="marvis-archive-read-"))).resolve()
        _authenticated_snapshot(archive, snapshot, expected_manifest_sha256)
        archive = snapshot
        checks[current] = {"status": "verified", "scope": "all_parser_inputs_in_private_manifest_authenticated_copy"}
        current = "external_bindings"
        manifest, bindings = _json(archive, "manifest.json"), _json(archive, "bindings.json")
        for key in ("run_id", "cases_sha256", "expected_sha256", "source", "model_connection_sha256"):
            _require_equal(bindings.get(key), getattr(binding, key))
        _require_equal(_json(archive, "case.json"), binding.case.model_dump())
        if bindings["execution_before_scoring"]["execution"]["task_id"] != binding.task_id or manifest["case_id"] != binding.case.id:
            raise ValueError("run task binding mismatch")
        checks[current] = {"status": "verified", "scope": "binding_consistency_not_origin_authentication"}
        current = "task_contract"
        task, contract, messages, audits = _read_contract(archive, binding)
        if (task["task_type"] != "validation" or task["run_mode"] != "agent" or task["validation_workflow_version"] != 2
                or task["model_name"] != binding.case.task.model_name or contract.status != "ready"
                or digest(input_contract_to_dict(contract)) != binding.input_contract_sha256):
            raise ValueError("task contract binding mismatch")
        checks[current] = {"status": "verified", "scope": "frozen_contract_and_readonly_native_rows"}
        current = "materials"
        paths, metadata, count, reading_contract = _materials(archive, manifest, binding, task, contract)
        checks[current] = {"status": "verified", "rows": count, "scope": "frozen_material_hashes_reparsed_sample_pmml_and_dictionary",
                           "read_context": "archived_material_path_mapped_for_recomputation"}
        outputs = Path("workspace/tasks") / binding.task_id / "outputs"
        current = "pmml_rescoring"
        scorer, scoring, score_path = _rescore(archive, outputs, paths, contract, count, reading_contract)
        checks[current] = {"status": "verified", "rows_recomputed": count, "engine_version": scoring.engine_version}
        config = ValidationConfig(target_col="__target__", score_col="__pmml_score__", split_col="__split__",
                                  time_col="__time__", bin_count=binding.bin_count)
        frame = load_pmml_analysis_frame(sample_path=paths["sample"], score_path=score_path,
                                        contract=reading_contract, scoring_result=scoring)
        retained = validation_results_from_dict(_json(archive, outputs / "validation_results.json"))
        with tempfile.TemporaryDirectory(prefix="marvis-archive-recompute-") as tmp:
            current = "stress_recomputation"
            stress = run_pmml_stress(contract=reading_contract, config=config, sample_path=paths["sample"],
                baseline_score_path=score_path, scoring_result=scoring, scenario_dir=Path(tmp) / "stress",
                feature_categories=contract.require_feature_metadata().per_category_raw_fields, scorer=scorer, chunk_size=10_000,
                category_source_counts={"dictionary": len(metadata.rows), "unresolved": 0})
            _require_equal(asdict(stress), asdict(retained.stress_test))
            checks[current] = {"status": "verified", "categories_recomputed": len(stress.per_category)}
            current = "metrics_recomputation"
            recomputed = compute_platform_validation_results(model_name=task["model_name"], model_version=task["model_version"],
                contract=contract, sample_scored=frame, config=config, scoring_result=scoring,
                metadata_resolution=metadata, stress_test=stress)
            _require_equal(asdict(recomputed), asdict(retained))
            checks[current] = {"status": "verified", "scope": "basic_info_split_and_monthly_KS_AUC_PSI_lift_bins_curves_and_stress"}
        current = "confirmation_record_binding"
        values = _confirmation(task, messages, audits, binding)
        checks[current] = {"status": "verified", "scope": "contents_revision_and_audit_consistency_not_actor_authentication"}
        current = "report_download_binding"
        report_paths = {}
        for kind, filename in (("word", "validation_report.docx"), ("excel", "validation.xlsx")):
            if binding.report_sha256[kind] is None:
                raise _Unsupported("no separately retained report download binding")
            path = _file(archive, outputs / filename, maximum=32_000_000)
            if digest(path.read_bytes()) != binding.report_sha256[kind]:
                raise ValueError("report download binding differs")
            _ooxml(path, kind)
            report_paths[kind] = path
        checks[current] = {"status": "verified", "scope": "externally_bound_report_bytes_and_bounded_OOXML"}
        current = "excel_values"
        checks[current] = {"status": "verified", "cells_compared": _excel(report_paths["excel"], recomputed, values)}
        current = "word_confirmed_text"
        document = _ooxml(report_paths["word"], "word")
        text = re.sub(r"\s+", "", "".join(document.itertext()))
        for key in ("TEXT:pressure_test_summary", "TEXT:pressure_impact_recommendation", "TEXT:final_validation_conclusion", "TEXT:model_training_description"):
            narrative = re.sub(r"!!([^!\n]+)!!|\*\*([^*\n]+)\*\*", lambda m: m.group(1) or m.group(2), values[key])
            if not narrative.strip() or re.sub(r"\s+", "", narrative) not in text:
                raise ValueError("confirmed report text absent")
        checks[current] = {"status": "verified", "scope": "four_confirmed_narratives_present_not_semantic_correctness"}
    except _Unsupported as exc:
        checks[current] = {"status": "unsupported", "reason": str(exc)}
    except Exception as exc:
        checks[current] = {"status": "mismatch", "error_type": type(exc).__name__}
    finally:
        if checks["archive_integrity"]["status"] == "verified":
            try:
                verify_case_archive(original_archive, expected_manifest_sha256=expected_manifest_sha256)
                verify_case_archive(archive, expected_manifest_sha256=expected_manifest_sha256)
                result["archive_unchanged_during_revalidation"] = True
            except Exception:
                checks["archive_integrity"] = {"status": "mismatch", "reason": "archive_changed_during_revalidation"}
                result["archive_unchanged_during_revalidation"] = False
        temporary.close()
    result["supported_checks_verified"] = all(c["status"] == "verified" for c in checks.values())
    return result
