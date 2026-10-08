"""Independent revalidation of normal single-dataset feature metric reports.

The default IV/KS/AUC/coverage and explicitly skipped optional binning are
covered here. Other selected metric families require their own calculators.
"""
from contextlib import ExitStack, closing
import json
from pathlib import Path
import re
import sqlite3
import tempfile
import unicodedata

from pydantic import BaseModel, ConfigDict, Field, model_validator

from .runtime_archive_datasets import _dataset, _source_input
from .runtime_archive_reader import _Unsupported, _file, _json, _authenticated_snapshot, _public_source_identity, _original_path
from .runtime_archive_portfolio import _numbers
from .runtime_contracts import RuntimeCase, digest
from .runtime_custody import verify_case_archive
from .runtime_feature_reference import feature_reference, feature_recommendation_reference

_HASH = r"^[0-9a-f]{64}$"
_ID = r"^[a-zA-Z0-9_-]{1,160}$"
_TOOLS = {"metrics": "feature.compute_feature_metrics", "bins": "feature.analyze_feature_bins",
          "report": "feature.generate_feature_report"}
_PARENTS = {"metrics": (), "bins": ("metrics",), "report": ("metrics", "bins")}
_CHECKS = ("archive_integrity", "recomputation_snapshot", "external_bindings", "native_steps",
           "feature_contract", "registered_dataset", "source_material", "independent_metrics",
           "report_inputs", "artifact_record", "workbook_cells", "download_binding")
_DEFAULT_METRICS = ["iv", "ks", "auc", "coverage"]
_QUALITY_COLUMNS = [("coverage", "覆盖率"), ("missing_rate", "缺失率"), ("mode_rate", "单一值率"),
    ("zero_rate", "零值率"), ("valid_count", "有效样本数"), ("unique_count", "唯一值数"), ("unique_rate", "唯一值率"),
    ("mean", "均值"), ("std", "标准差"), ("min", "最小值"), ("q25", "P25"), ("median", "中位数"), ("q75", "P75"), ("max", "最大值")]


class FrozenFeatureBinding(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True, frozen=True)
    case: RuntimeCase
    run_id: str = Field(min_length=1, max_length=160)
    task_id: str = Field(pattern=_ID)
    plan_id: str = Field(pattern=_ID)
    dataset_id: str = Field(pattern=_ID)
    dataset_sha256: str = Field(pattern=_HASH)
    step_ids: dict[str, str]
    output_sha256: dict[str, str]
    download_sha256: str = Field(pattern=_HASH)
    cases_sha256: str = Field(pattern=_HASH)
    expected_sha256: str = Field(pattern=_HASH)
    source: dict
    model_connection_sha256: str = Field(pattern=_HASH)

    @model_validator(mode="after")
    def required_bindings(self):
        if (self.case.task.task_type != "feature_analysis" or len(self.case.materials) != 1
                or self.case.materials[0].role != "sample" or not self.case.task.feature_columns
                or len(set(self.case.task.feature_columns)) != len(self.case.task.feature_columns)
                or not any(a.kind == "approve_step" and a.tool == _TOOLS["bins"] for a in self.case.actions)
                or not any(a.kind == "download_feature_report" for a in self.case.actions)):
            raise ValueError("explicit feature task, sample, selection and download required")
        if (set(self.step_ids) != set(_TOOLS) or len(set(self.step_ids.values())) != 3
                or any(not re.fullmatch(_ID, v) for v in self.step_ids.values())
                or set(self.output_sha256) != set(_TOOLS)
                or any(not re.fullmatch(_HASH, v) for v in self.output_sha256.values())
                or not re.fullmatch(_HASH, self.source.get("source_sha256", ""))):
            raise ValueError("complete external native execution bindings required")
        return self


def _equal(actual, expected):
    if digest(actual) != digest(expected):
        raise ValueError("feature identity mismatch")


def _native(conn, binding):
    from marvis.orchestrator.evidence import payload_hash

    task = conn.execute("SELECT * FROM tasks WHERE id=?", (binding.task_id,)).fetchone()
    plan = conn.execute("SELECT * FROM plans WHERE id=?", (binding.plan_id,)).fetchone()
    if (task is None or plan is None or task["task_type"] != "feature_analysis" or task["run_mode"] != "agent"
            or plan["task_id"] != binding.task_id or plan["status"] != "done"):
        raise ValueError("feature task or plan mismatch")
    if plan["template_id"] != "feature_analysis":
        raise _Unsupported("feature template requires additional independent domain checks")
    _equal(sorted(s[0] for s in conn.execute("SELECT id FROM plan_steps WHERE plan_id=?", (binding.plan_id,))),
           sorted(binding.step_ids.values()))
    outputs, inputs, evidence, refs = {}, {}, {}, {}
    for key, tool in _TOOLS.items():
        step_id = binding.step_ids[key]
        step = conn.execute("SELECT * FROM plan_steps WHERE id=?", (step_id,)).fetchone()
        if (step is None or step["plan_id"] != binding.plan_id or step["status"] != "done"
                or f"{step['tool_plugin']}.{step['tool_name']}" != tool):
            raise ValueError("feature step identity mismatch")
        _equal(json.loads(step["depends_on_json"]), [binding.step_ids[parent] for parent in _PARENTS[key]])
        if key == "bins" and (step["confirmed"] != 1 or step["needs_confirmation"] != 1
                              or json.loads(step["policy_json"]).get("human_decision_gate") != "required"):
            raise ValueError("feature gate native confirmation missing")
        match = re.fullmatch(r"metrics:" + re.escape(step_id) + r":v([1-9][0-9]*)", step["output_ref"] or "")
        if match is None:
            raise ValueError("feature output is not version bound")
        saved = conn.execute("SELECT * FROM plan_step_output_versions WHERE step_id=? AND version=?",
                             (step_id, int(match[1]))).fetchone()
        if saved is None:
            raise ValueError("feature output missing")
        output, native = json.loads(saved["output_json"]), json.loads(saved["evidence_json"])
        run = conn.execute("SELECT * FROM plan_step_runs WHERE id=?", (native.get("step_run_id"),)).fetchone()
        if run is None:
            raise ValueError("feature producing run missing")
        run_input = json.loads(run["input_json"])
        _equal(digest(output), binding.output_sha256[key])
        if (run["plan_id"] != binding.plan_id or run["step_id"] != step_id or run["tool_ref"] != tool
                or run["status"] != "succeeded" or run["output_ref"] != step["output_ref"]
                or run["output_hash"] != payload_hash(output) or native.get("output_hash") != payload_hash(output)
                or native.get("input_hash") != payload_hash(run_input)
                or native.get("task_id") != binding.task_id or native.get("plan_id") != binding.plan_id
                or native.get("step_id") != step_id or native.get("tool_name") != tool
                or native.get("producer_invocation_id") != run["id"] or run["invocation_id"] != run["id"]
                or native.get("raw_output_hash") != run["raw_output_hash"]
                or native.get("manifest_hash") != run["manifest_hash"]
                or native.get("tool_version") != run["tool_version"]
                or native.get("schema_version") != "evidence.v1" or native.get("output_ref") != step["output_ref"]):
            raise ValueError("feature native run binding mismatch")
        _equal(native.get("source_dataset_refs"), [f"dataset:{binding.dataset_id}"] if key in {"metrics", "bins"} else [])
        _equal(native.get("result_dataset_bindings"), [])
        outputs[key], inputs[key], evidence[key], refs[key] = output, run_input, native, step["output_ref"]
    for key, parents in _PARENTS.items():
        _equal(evidence[key].get("parent_output_refs"), [refs[p] for p in parents])
        _equal(evidence[key].get("parent_output_bindings"), [
            {"step_id": binding.step_ids[p], "output_ref": refs[p], "output_hash": payload_hash(outputs[p])}
            for p in parents])
    return outputs, inputs


def _contract(binding, inputs):
    selected = binding.case.task.metrics if binding.case.task.metrics is not None else _DEFAULT_METRICS
    if not selected or set(selected) - set(_DEFAULT_METRICS):
        raise _Unsupported("additional selected feature metrics require independent calculators")
    _equal(inputs["metrics"], {"dataset_id": binding.dataset_id, "features": binding.case.task.feature_columns,
        "target_col": binding.case.task.target_col, "metrics": selected, "meaning_directions": {}, "bins": 10})
    if inputs["bins"].get("features"):
        raise _Unsupported("selected optional binning requires independent row and workbook checks")
    _equal(inputs["bins"], {"dataset_id": binding.dataset_id, "features": [], "target_col": binding.case.task.target_col, "bins": 10})
    return selected


def _report_artifact(conn, archive, manifest, binding, output, report_inputs):
    from marvis.repositories.task_artifacts import stable_task_artifact_id

    row = conn.execute("SELECT * FROM task_artifacts WHERE id=?", (output["artifact_id"],)).fetchone()
    kind = "feature_report_xlsx"
    if (row is None or row["task_id"] != binding.task_id or row["kind"] != kind
            or row["origin_tool"] != _TOOLS["report"] or row["content_hash"] != output["artifact_content_hash"]
            or row["id"] != stable_task_artifact_id(task_id=binding.task_id, kind=kind, path=row["path"])
            or row["path"] != output["report_path"]):
        raise ValueError("feature report artifact binding mismatch")
    root = Path(manifest["original_workspace"]) / "tasks" / binding.task_id / "feature_reports"
    if Path(row["path"]).parent != root or Path(row["path"]).suffix != ".xlsx":
        raise ValueError("feature report outside task output directory")
    path = _original_path(archive, manifest["original_workspace"], row["path"])
    _equal(digest(path.read_bytes()), row["content_hash"])
    _equal(json.loads(row["provenance_json"]), {"schema_version": "feature-report-artifact.v1",
        "producer_version": "feature.generate_feature_report.v1", "task_id": binding.task_id,
        "input_hash": digest(report_inputs), "feature_count": len(binding.case.task.feature_columns)})
    return path


def _cell(value):
    if value is None:
        return "n/a"
    if isinstance(value, float):
        return round(value, 6)
    if isinstance(value, str):
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", value)
        significant = next((c for c in text if not c.isspace() and not unicodedata.category(c).startswith("C")), "")
        return ("'" + text[:32766]) if significant in {"=", "+", "-", "@"} else text[:32767]
    return value


def _workbook(path, rows, selected):
    from openpyxl import load_workbook
    from .runtime_archive_validation import _ooxml

    _ooxml(path, "excel")
    columns = [("feature", "特征"), ("recommendation", "Agent建议"), ("recommendation_reason", "推荐原因")]
    columns += [(key, key.upper()) for key in ("iv", "ks", "auc") if key in selected]
    if "coverage" in selected:
        columns += _QUALITY_COLUMNS
    expected = [[label for _, label in columns]] + [[_cell(row.get(key)) for key, _ in columns] for row in rows]
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if book.sheetnames != ["特征指标"]:
            raise ValueError("feature report worksheet selection mismatch")
        sheet = book["特征指标"]
        if sheet.max_row != len(expected) or sheet.max_column != len(columns):
            raise ValueError("feature report dimensions mismatch")
        for actual_row, expected_row in zip(sheet.iter_rows(), expected, strict=True):
            for cell, value in zip(actual_row, expected_row, strict=True):
                if cell.data_type in {"f", "e"}:
                    raise ValueError("feature report contains executable or error cells")
                _numbers(cell.value, value)
    finally:
        book.close()
    return len(expected) * len(columns)


def revalidate_feature_archive(archive: Path, *, expected_manifest_sha256: str, frozen_binding: FrozenFeatureBinding):
    from .runtime_runner import _source_identity

    if not isinstance(frozen_binding, FrozenFeatureBinding):
        raise ValueError("an explicit FrozenFeatureBinding is required")
    binding = FrozenFeatureBinding.model_validate(frozen_binding.model_dump())
    archive = original_archive = Path(archive).absolute()
    temporary = ExitStack()
    verifier = _source_identity()
    result = {"schema": "marvis.runtime-feature-revalidation.v1", "case_id": binding.case.id,
              "algorithm": "independent_python_numeric_feature_reference.v1", "verifier_source": verifier,
              "original_source_binding": _public_source_identity(binding.source),
              "original_source_binding_sha256": digest(binding.source),
              "verifier_source_matches_original": verifier["source_sha256"] == binding.source["source_sha256"],
              "source_authentication": "not_established", "acceptance_claim": "not_established",
              "recomputation_isolation": "private_authenticated_copy; excludes_adversarial_same_uid_or_root_access",
              "checks": {key: {"status": "unverified"} for key in _CHECKS},
              "unsupported": {"native_signature_authentication": "original_key_not_available",
                  "business_outcome_truth": "requires_independent_source_verification",
                  "actor_authorization_time_travel_and_duplicate_side_effect_safety": "not_established_by_feature_arithmetic",
                  "optional_metrics_binning_join_and_recovery": "require_separate_domain_checks",
                  "recommendation_business_suitability": "only_declared_deterministic_policy_consistency_is_verified",
                  "original_runtime_cost_and_timing": "not_recomputed"}}
    checks, current = result["checks"], "archive_integrity"
    try:
        verify_case_archive(archive, expected_manifest_sha256=expected_manifest_sha256)
        checks[current] = {"status": "verified", "scope": "bytes_match_caller_held_manifest_digest"}
        current = "recomputation_snapshot"
        snapshot = Path(temporary.enter_context(tempfile.TemporaryDirectory(prefix="marvis-feature-archive-"))).resolve()
        _authenticated_snapshot(archive, snapshot, expected_manifest_sha256)
        archive = snapshot
        scratch = Path(temporary.enter_context(tempfile.TemporaryDirectory(prefix="marvis-feature-input-"))).resolve()
        checks[current] = {"status": "verified", "scope": "all_parser_inputs_in_private_authenticated_copy"}
        current = "external_bindings"
        manifest, bindings = _json(archive, "manifest.json"), _json(archive, "bindings.json")
        for key in ("run_id", "cases_sha256", "expected_sha256", "source", "model_connection_sha256"):
            _equal(bindings.get(key), getattr(binding, key))
        _equal(_json(archive, "case.json"), binding.case.model_dump())
        observed = bindings["execution_before_scoring"]
        if observed["execution"]["task_id"] != binding.task_id or manifest["case_id"] != binding.case.id:
            raise ValueError("feature external task binding mismatch")
        checks[current] = {"status": "verified", "scope": "external_binding_consistency_not_origin_authentication"}
        conn = temporary.enter_context(closing(sqlite3.connect(
            _file(archive, "workspace/marvis.sqlite").as_uri() + "?mode=ro&immutable=1", uri=True)))
        conn.row_factory = sqlite3.Row
        current = "native_steps"
        outputs, inputs = _native(conn, binding)
        checks[current] = {"status": "verified", "steps": 3, "scope": "native_runs_versioned_outputs_and_parent_bindings"}
        current = "feature_contract"
        selected = _contract(binding, inputs)
        checks[current] = {"status": "verified", "scope": "frozen_columns_target_metric_selection_and_default_optional_skip"}
        current = "registered_dataset"
        source, frame = _dataset(conn, archive, binding.task_id, binding.dataset_id, binding.dataset_sha256)
        if source["role"] != "sample":
            raise ValueError("feature source role mismatch")
        checks[current] = {"status": "verified", "rows": len(frame)}
        current = "source_material"
        _source_input(archive, manifest, binding, source, frame, scratch)
        checks[current] = {"status": "verified", "scope": "frozen_upload_and_registered_source_values"}
        current = "independent_metrics"
        rows = []
        for name in binding.case.task.feature_columns:
            row = feature_reference(frame[name].tolist(), frame[binding.case.task.target_col].tolist(), feature=name, selected=selected)
            rows.append({**row, **feature_recommendation_reference(row)})
        _numbers(outputs["metrics"], {"dataset_id": binding.dataset_id, "metrics": rows,
            "collinear": None, "nan_labels_dropped": 0, "selected_metrics": sorted(selected)})
        _equal(outputs["bins"], {"dataset_id": binding.dataset_id, "selected_features": [], "requested_bins": 10,
                                "binning": [], "nan_labels_dropped": 0})
        checks[current] = {"status": "verified", "scope": "selected_numeric_metrics_and_default_recommendation_policy", "features": len(rows)}
        current = "report_inputs"
        _equal(inputs["report"], {"metrics": outputs["metrics"]["metrics"], "collinear": None, "binning": []})
        _numbers(outputs["report"]["metrics"], rows)
        _equal(outputs["report"]["feature_count"], len(rows))
        _equal(outputs["report"]["binning"], [])
        checks[current] = {"status": "verified", "scope": "recomputed_metrics_to_report_input_and_output_echo"}
        current = "artifact_record"
        path = _report_artifact(conn, archive, manifest, binding, outputs["report"], inputs["report"])
        checks[current] = {"status": "verified", "scope": "registered_owner_kind_provenance_and_actual_bytes"}
        current = "workbook_cells"
        cells = _workbook(path, rows, selected)
        checks[current] = {"status": "verified", "cells_compared": cells, "sheets": 1}
        current = "download_binding"
        _equal(digest(path.read_bytes()), binding.download_sha256)
        if not any(e.get("stage") == "download_feature_report" and e.get("status_code") == 200
                   and e.get("sha256") == binding.download_sha256 and e.get("size_bytes") == path.stat().st_size
                   and e.get("step_id") == binding.step_ids["report"] and e.get("artifact_id") == outputs["report"]["artifact_id"]
                   for e in observed["http_events"]):
            raise ValueError("feature actual download binding missing")
        checks[current] = {"status": "verified", "scope": "externally_bound_bytes_native_report_identity_and_HTTP_record"}
        result["values"] = {"metrics": rows, "source_rows": len(frame), "feature_count": len(rows)}
    except _Unsupported as exc:
        checks[current] = {"status": "unsupported", "reason": str(exc)}
    except Exception as exc:
        checks[current] = {"status": "mismatch", "error_type": type(exc).__name__}
    finally:
        if checks["archive_integrity"]["status"] == "verified":
            try:
                verify_case_archive(original_archive, expected_manifest_sha256=expected_manifest_sha256)
                verify_case_archive(archive, expected_manifest_sha256=expected_manifest_sha256)
                result["archive_unchanged_during_revalidation"] = True
            except Exception:
                checks["archive_integrity"] = {"status": "mismatch", "reason": "archive_changed_during_revalidation"}
                result["archive_unchanged_during_revalidation"] = False
        temporary.close()
    result["supported_checks_verified"] = all(item["status"] == "verified" for item in checks.values())
    return result
