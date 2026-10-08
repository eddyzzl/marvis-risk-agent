"""Independent, read-only checks of externally bound portfolio originals.

The normal no-trend workflow is supported. This establishes consistency and
arithmetic, not native signature authentication, actor authority or real outcomes.
"""
from contextlib import ExitStack, closing
import json
import math
from pathlib import Path
import re
import sqlite3
import tempfile

from pydantic import BaseModel, ConfigDict, Field, model_validator

from .runtime_archive_datasets import _dataset, _source_input
from .runtime_archive_reader import (
    _Unsupported, _file, _json, _authenticated_snapshot, _public_source_identity, _original_path,
)
from .runtime_contracts import RuntimeCase, digest
from .runtime_custody import verify_case_archive
from .runtime_portfolio_reference import transition_reference, segment_reference, reference_flag_text

_HASH = r"^[0-9a-f]{64}$"
_ID = r"^[a-zA-Z0-9_-]{1,160}$"
_TOOLS = {"flow": "analysis.flow_rate", "migration": "analysis.bucket_migration",
          "segment": "analysis.segment_profile", "expected_loss": "analysis.expected_loss_estimate",
          "gate": "analysis.portfolio_gate_summary", "report": "analysis.portfolio_report"}
_METRICS = ("flow", "migration", "segment", "expected_loss")
_PARENTS = {"flow": (), "migration": (), "segment": (), "expected_loss": ("migration",),
            "gate": _METRICS, "report": (*_METRICS, "gate")}
_SHEETS = ["组合概览", "桶迁徙", "逐月流量", "细分画像", "稳定性趋势", "预期损失", "数据质量红旗"]
_CHECKS = ("archive_integrity", "recomputation_snapshot", "external_bindings", "native_steps",
           "business_contract", "registered_dataset", "source_material", "independent_metrics",
           "gate_and_report_inputs", "artifact_record", "workbook_cells", "download_binding")


class FrozenPortfolioBinding(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True, frozen=True)
    case: RuntimeCase
    run_id: str = Field(min_length=1, max_length=160)
    task_id: str = Field(pattern=_ID)
    plan_id: str = Field(pattern=_ID)
    dataset_id: str = Field(pattern=_ID)
    dataset_sha256: str = Field(pattern=_HASH)
    state_action_index: int = Field(ge=0)
    states: list[str] = Field(min_length=1, max_length=100)
    step_ids: dict[str, str]
    output_sha256: dict[str, str]
    download_sha256: str = Field(pattern=_HASH)
    cases_sha256: str = Field(pattern=_HASH)
    expected_sha256: str = Field(pattern=_HASH)
    source: dict
    model_connection_sha256: str = Field(pattern=_HASH)

    @model_validator(mode="after")
    def bindings_required(self):
        declarations = [i for i, action in enumerate(self.case.actions) if action.portfolio_request is not None]
        if (self.case.task.task_type != "portfolio" or len(self.case.materials) != 1
                or self.case.materials[0].role != "sample" or len(declarations) != 1
                or not declarations[0] < self.state_action_index < len(self.case.actions)
                or self.case.actions[self.state_action_index].kind != "message"
                or [s.strip() for s in self.case.actions[self.state_action_index].content.split(",")] != self.states
                or len(set(self.states)) != len(self.states) or any(not s for s in self.states)
                or not any(a.kind == "download_portfolio_report" for a in self.case.actions)):
            raise ValueError("explicit portfolio sample, business rule, ordered states and download required")
        if (set(self.step_ids) != set(_TOOLS) or len(set(self.step_ids.values())) != len(_TOOLS)
                or any(not re.fullmatch(_ID, value) for value in self.step_ids.values())
                or set(self.output_sha256) != set(_TOOLS)
                or any(not re.fullmatch(_HASH, value) for value in self.output_sha256.values())
                or not isinstance(self.source.get("source_sha256"), str)
                or not re.fullmatch(_HASH, self.source["source_sha256"])):
            raise ValueError("complete native identities and external source binding required")
        return self


def _equal(actual, expected):
    if digest(actual) != digest(expected):
        raise ValueError("portfolio identity mismatch")


def _numbers(actual, expected):
    if isinstance(expected, dict):
        if not isinstance(actual, dict) or actual.keys() != expected.keys():
            raise ValueError("portfolio metric fields mismatch")
        for key in expected:
            _numbers(actual[key], expected[key])
    elif isinstance(expected, list):
        if not isinstance(actual, list) or len(actual) != len(expected):
            raise ValueError("portfolio metric rows mismatch")
        for left, right in zip(actual, expected):
            _numbers(left, right)
    elif type(expected) is float:
        if (type(actual) not in {int, float} or not math.isfinite(actual)
                or not math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-10)):
            raise ValueError("portfolio arithmetic mismatch")
    elif type(actual) is not type(expected) or actual != expected:
        raise ValueError("portfolio value mismatch")


def _native(conn, binding):
    from marvis.orchestrator.evidence import payload_hash

    task = conn.execute("SELECT * FROM tasks WHERE id=?", (binding.task_id,)).fetchone()
    plan = conn.execute("SELECT * FROM plans WHERE id=?", (binding.plan_id,)).fetchone()
    if (task is None or plan is None or task["task_type"] != "portfolio" or task["run_mode"] != "agent"
            or plan["task_id"] != binding.task_id or plan["status"] != "done"):
        raise ValueError("portfolio task or plan mismatch")
    if plan["template_id"] != "portfolio_analysis_no_trend":
        raise _Unsupported("portfolio template requires additional independent domain checks")
    _equal(sorted(s[0] for s in conn.execute("SELECT id FROM plan_steps WHERE plan_id=?", (binding.plan_id,))),
           sorted(binding.step_ids.values()))
    outputs, inputs, evidence, refs = {}, {}, {}, {}
    for key, tool in _TOOLS.items():
        step_id = binding.step_ids[key]
        step = conn.execute("SELECT * FROM plan_steps WHERE id=?", (step_id,)).fetchone()
        if (step is None or step["plan_id"] != binding.plan_id or step["status"] != "done"
                or f"{step['tool_plugin']}.{step['tool_name']}" != tool):
            raise ValueError("portfolio step identity mismatch")
        _equal(json.loads(step["depends_on_json"]), [binding.step_ids[parent] for parent in _PARENTS[key]])
        if key == "gate" and (step["confirmed"] != 1 or step["needs_confirmation"] != 1
                              or json.loads(step["policy_json"]).get("human_decision_gate") != "required"):
            raise ValueError("portfolio gate native confirmation missing")
        match = re.fullmatch(r"metrics:" + re.escape(step_id) + r":v([1-9][0-9]*)", step["output_ref"] or "")
        if match is None:
            raise ValueError("portfolio output is not version bound")
        saved = conn.execute("SELECT * FROM plan_step_output_versions WHERE step_id=? AND version=?",
                             (step_id, int(match[1]))).fetchone()
        if saved is None:
            raise ValueError("portfolio output missing")
        output, native = json.loads(saved["output_json"]), json.loads(saved["evidence_json"])
        run = conn.execute("SELECT * FROM plan_step_runs WHERE id=?", (native.get("step_run_id"),)).fetchone()
        if run is None:
            raise ValueError("portfolio producing run missing")
        run_input = json.loads(run["input_json"])
        _equal(digest(output), binding.output_sha256[key])
        if (run["plan_id"] != binding.plan_id or run["step_id"] != step_id or run["tool_ref"] != tool
                or run["status"] != "succeeded" or run["output_ref"] != step["output_ref"]
                or run["output_hash"] != payload_hash(output) or native.get("output_hash") != payload_hash(output)
                or native.get("input_hash") != payload_hash(run_input)
                or native.get("task_id") != binding.task_id or native.get("plan_id") != binding.plan_id
                or native.get("step_id") != step_id or native.get("tool_name") != tool
                or native.get("producer_invocation_id") != run["id"] or run["invocation_id"] != run["id"]
                or native.get("raw_output_hash") != run["raw_output_hash"]
                or native.get("manifest_hash") != run["manifest_hash"]
                or native.get("tool_version") != run["tool_version"]
                or native.get("schema_version") != "evidence.v1" or native.get("output_ref") != step["output_ref"]):
            raise ValueError("portfolio native run binding mismatch")
        _equal(native.get("source_dataset_refs"), [f"dataset:{binding.dataset_id}"] if key in _METRICS else [])
        _equal(native.get("result_dataset_bindings"), [])
        outputs[key], inputs[key], evidence[key], refs[key] = output, run_input, native, step["output_ref"]
    for key, parents in _PARENTS.items():
        _equal(evidence[key].get("parent_output_refs"), [refs[p] for p in parents])
        _equal(evidence[key].get("parent_output_bindings"), [
            {"step_id": binding.step_ids[p], "output_ref": refs[p], "output_hash": payload_hash(outputs[p])}
            for p in parents])
    return outputs, inputs


def _contract(binding, inputs):
    request = next(a.portfolio_request for a in binding.case.actions if a.portfolio_request is not None)
    if request.experiment_id is not None or request.score_col is not None:
        raise _Unsupported("trend requires its independent source and domain calculator")
    source = {"dataset_id": binding.dataset_id, "expected_content_hash": binding.dataset_sha256}
    common = {**source, **{key: getattr(request, key) for key in
              ("id_col", "snapshot_col", "bucket_col", "balance_col")}, "states": binding.states}
    _equal(inputs["flow"], common)
    _equal(inputs["migration"], common)
    _equal(inputs["expected_loss"], {**common, **{key: getattr(request, key) for key in
           ("loss_state", "lgd", "horizon_months")}})
    _equal(inputs["segment"], {**source, "segment_col": request.segment_col, "ead_col": request.balance_col})
    return request


def _aggregate(reference):
    flags = [{"source": _TOOLS[key].split(".")[1], **flag}
             for key in _METRICS for flag in reference[key]["red_flags"]]
    loss, segment, migration = reference["expected_loss"], reference["segment"], reference["migration"]
    return {"highlights": {"total_el": loss["total_el"],
            "total_el_basis": "reference_snapshot", "reference_snapshot": loss["assumptions"]["reference_snapshot"],
            "hhi": segment["concentration"]["hhi"], "top1_pct": segment["concentration"]["top1_pct"],
            "migration_months": migration["window_months"]},
            "red_flags": flags, "red_flag_count": len(flags), "checklist": [f["message"] for f in flags]}


def _report_artifact(conn, archive, manifest, binding, output, report_inputs):
    from marvis.repositories.task_artifacts import stable_task_artifact_id

    row = conn.execute("SELECT * FROM task_artifacts WHERE id=?", (output["artifact_id"],)).fetchone()
    kind = "portfolio_report_xlsx"
    if (row is None or row["task_id"] != binding.task_id or row["kind"] != kind
            or row["origin_tool"] != _TOOLS["report"] or row["content_hash"] != output["artifact_content_hash"]
            or row["id"] != stable_task_artifact_id(task_id=binding.task_id, kind=kind, path=row["path"])
            or row["path"] != output["report_path"]):
        raise ValueError("portfolio artifact binding mismatch")
    root = Path(manifest["original_workspace"]) / "tasks" / binding.task_id / "portfolio"
    if Path(row["path"]).parent != root or Path(row["path"]).suffix != ".xlsx":
        raise ValueError("portfolio report outside task directory")
    path = _original_path(archive, manifest["original_workspace"], row["path"])
    _equal(digest(path.read_bytes()), row["content_hash"])
    _equal(output["sheets"], _SHEETS)
    _equal(json.loads(row["provenance_json"]), {
        "schema_version": "portfolio-report-artifact.v1", "producer_version": "analysis.portfolio_report.v1",
        "task_id": binding.task_id, "input_hash": digest(report_inputs), "sheets": _SHEETS,
    })
    return path


def _workbook(path, reference, gate):
    from openpyxl import load_workbook
    from .runtime_archive_validation import _ooxml

    _ooxml(path, "excel")

    def scalar(value):
        return str(value) if isinstance(value, (list, dict, tuple)) else value

    def table(rows):
        if not rows:
            return [["无数据"]]
        headers = list(dict.fromkeys(key for row in rows for key in row))
        return [headers] + [[scalar(row.get(key)) for key in headers] for row in rows]

    loss, segment = reference["expected_loss"], reference["segment"]
    overview = [["项目元数据", None], ["预期损失", None], ["total_el", loss["total_el"]]]
    overview += [[f"假设.{key}", scalar(value)] for key, value in loss["assumptions"].items()]
    overview += [["细分集中度", None], ["concentration_basis", "count"]]
    overview += [[f"concentration.{key}", value] for key, value in segment["concentration"].items()]
    overview += [["ead_concentration_basis", "ead"]]
    overview += [[f"ead_concentration.{key}", value] for key, value in segment["ead_concentration"].items()]
    overview += [["数据质量红旗数", gate["red_flag_count"]]]
    sheets = dict(zip(_SHEETS, [overview, table(reference["migration"]["heat_table"]),
                  table(reference["flow"]["net_flows"]), table(segment["segments"]),
                  [["无数据（未提供 experiment_id，趋势步已剪除）"]],
                  table(loss["el_by_month"]) + [[]] + [["链式吸收概率"]] + table(loss["chain"]),
                  table(gate["red_flags"]) if gate["red_flags"] else [["无红旗"]]]))
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    checked = 0
    try:
        if book.sheetnames != _SHEETS:
            raise ValueError("portfolio report worksheets differ")
        for name, expected in sheets.items():
            sheet = book[name]
            columns = max(map(len, expected))
            if sheet.max_row > 200_000 or sheet.max_column > 128 or sheet.max_row * sheet.max_column > 2_000_000:
                raise _Unsupported("portfolio workbook inspection limit")
            for i, row in enumerate(sheet.iter_rows(max_row=max(len(expected), sheet.max_row),
                                                  max_col=max(columns, sheet.max_column))):
                for j, cell in enumerate(row):
                    value = expected[i][j] if i < len(expected) and j < len(expected[i]) else None
                    if cell.data_type in {"f", "e"}:
                        raise ValueError("portfolio workbook contains a formula or error value")
                    actual = None if cell.value == "" else cell.value
                    _numbers(actual, value)
                    checked += 1
    finally:
        book.close()
    return checked


def revalidate_portfolio_archive(archive: Path, *, expected_manifest_sha256: str, frozen_binding: FrozenPortfolioBinding):
    from .runtime_runner import _source_identity

    if not isinstance(frozen_binding, FrozenPortfolioBinding):
        raise ValueError("an explicit FrozenPortfolioBinding is required")
    # A frozen Pydantic model can still contain mutable lists/dicts, and
    # model_copy(update=...) does not validate updates. Re-establish all
    # declaration invariants before taking the private working copy.
    binding = FrozenPortfolioBinding.model_validate(frozen_binding.model_dump())
    archive = original_archive = Path(archive).absolute()
    temporary = ExitStack()
    verifier = _source_identity()
    result = {"schema": "marvis.runtime-portfolio-revalidation.v1", "case_id": binding.case.id,
              "algorithm": "independent_monthly_portfolio_reference.v1", "verifier_source": verifier,
              "original_source_binding": _public_source_identity(binding.source),
              "original_source_binding_sha256": digest(binding.source),
              "verifier_source_matches_original": verifier["source_sha256"] == binding.source["source_sha256"],
              "source_authentication": "not_established", "acceptance_claim": "not_established",
              "recomputation_isolation": "private_authenticated_copy; excludes_adversarial_same_uid_or_root_access",
              "checks": {key: {"status": "unverified"} for key in _CHECKS},
              "unsupported": {"native_signature_authentication": "original_key_not_available",
                              "business_outcome_truth": "requires_independent_source_verification",
                              "actor_authorization_and_duplicate_side_effect_safety": "not_established_by_portfolio_arithmetic",
                              "trend_vintage_and_recovery": "require_separate_domain_checks",
                              "original_runtime_cost_and_timing": "not_recomputed"}}
    checks, current = result["checks"], "archive_integrity"
    try:
        verify_case_archive(archive, expected_manifest_sha256=expected_manifest_sha256)
        checks[current] = {"status": "verified", "scope": "bytes_match_caller_held_manifest_digest"}
        current = "recomputation_snapshot"
        snapshot = Path(temporary.enter_context(tempfile.TemporaryDirectory(prefix="marvis-portfolio-archive-"))).resolve()
        _authenticated_snapshot(archive, snapshot, expected_manifest_sha256)
        archive = snapshot
        scratch = Path(temporary.enter_context(tempfile.TemporaryDirectory(prefix="marvis-portfolio-input-"))).resolve()
        checks[current] = {"status": "verified", "scope": "all_parser_inputs_in_private_authenticated_copy"}
        current = "external_bindings"
        manifest, bindings = _json(archive, "manifest.json"), _json(archive, "bindings.json")
        for key in ("run_id", "cases_sha256", "expected_sha256", "source", "model_connection_sha256"):
            _equal(bindings.get(key), getattr(binding, key))
        _equal(_json(archive, "case.json"), binding.case.model_dump())
        observed = bindings["execution_before_scoring"]
        if observed["execution"]["task_id"] != binding.task_id or manifest["case_id"] != binding.case.id:
            raise ValueError("portfolio run task binding mismatch")
        checks[current] = {"status": "verified", "scope": "external_binding_consistency_not_origin_authentication"}
        conn = temporary.enter_context(closing(sqlite3.connect(
            _file(archive, "workspace/marvis.sqlite").as_uri() + "?mode=ro&immutable=1", uri=True)))
        conn.row_factory = sqlite3.Row
        current = "native_steps"
        outputs, inputs = _native(conn, binding)
        checks[current] = {"status": "verified", "steps": 6, "scope": "native_runs_versioned_outputs_and_all_parent_bindings"}
        current = "business_contract"
        request = _contract(binding, inputs)
        checks[current] = {"status": "verified", "scope": "frozen_business_declaration_states_and_exact_native_inputs"}
        current = "registered_dataset"
        source, frame = _dataset(conn, archive, binding.task_id, binding.dataset_id, binding.dataset_sha256)
        if source["role"] != "sample":
            raise ValueError("portfolio source role differs")
        checks[current] = {"status": "verified", "rows": len(frame)}
        current = "source_material"
        _source_input(archive, manifest, binding, source, frame, scratch)
        checks[current] = {"status": "verified", "scope": "frozen_upload_and_registered_source_values"}
        current = "independent_metrics"
        rows = frame.to_dict("records")
        fields = {key: getattr(request, key) for key in
                  ("id_col", "snapshot_col", "bucket_col", "balance_col", "loss_state", "lgd", "horizon_months")}
        reference = transition_reference(rows, states=binding.states, **fields)
        reference["segment"] = segment_reference(rows, segment_col=request.segment_col, balance_col=request.balance_col)
        for key in _METRICS:
            for flag in reference[key]["red_flags"]:
                flag["message"] = reference_flag_text(flag)
            _numbers(outputs[key], reference[key])
        checks[current] = {"status": "verified", "scope": "four_independent_metrics_and_semantic_flags", "rows_recomputed": len(rows)}
        current = "gate_and_report_inputs"
        native_metrics = {key: outputs[key] for key in _METRICS}
        _equal(inputs["gate"], native_metrics)
        _equal(inputs["report"], native_metrics)
        gate = _aggregate(reference)
        _numbers(outputs["gate"], gate)
        checks[current] = {"status": "verified", "scope": "all_upstream_values_in_gate_and_report_and_independent_gate_assembly"}
        current = "artifact_record"
        path = _report_artifact(conn, archive, manifest, binding, outputs["report"], inputs["report"])
        checks[current] = {"status": "verified", "scope": "registered_owner_kind_provenance_and_actual_bytes"}
        current = "workbook_cells"
        checked = _workbook(path, reference, gate)
        checks[current] = {"status": "verified", "sheets": len(_SHEETS), "cells_compared": checked}
        current = "download_binding"
        _equal(digest(path.read_bytes()), binding.download_sha256)
        if not any(e.get("stage") == "download_portfolio_report" and e.get("status_code") == 200
                   and e.get("sha256") == binding.download_sha256 and e.get("size_bytes") == path.stat().st_size
                   and e.get("step_id") == binding.step_ids["report"] and e.get("artifact_id") == outputs["report"]["artifact_id"]
                   for e in observed["http_events"]):
            raise ValueError("portfolio actual download binding missing")
        checks[current] = {"status": "verified", "scope": "externally_bound_bytes_native_report_identity_and_HTTP_record"}
        result["values"] = {"metric_sha256": {key: digest(reference[key]) for key in _METRICS},
                            "segment_count": len(reference["segment"]["segments"]),
                            "total_el": reference["expected_loss"]["total_el"],
                            "red_flag_count": gate["red_flag_count"], "source_rows": len(rows)}
    except _Unsupported as exc:
        checks[current] = {"status": "unsupported", "reason": str(exc)}
    except Exception as exc:
        checks[current] = {"status": "mismatch", "error_type": type(exc).__name__}
    finally:
        if checks["archive_integrity"]["status"] == "verified":
            try:
                verify_case_archive(original_archive, expected_manifest_sha256=expected_manifest_sha256)
                verify_case_archive(archive, expected_manifest_sha256=expected_manifest_sha256)
                result["archive_unchanged_during_revalidation"] = True
            except Exception:
                checks["archive_integrity"] = {"status": "mismatch", "reason": "archive_changed_during_revalidation"}
                result["archive_unchanged_during_revalidation"] = False
        temporary.close()
    result["supported_checks_verified"] = all(item["status"] == "verified" for item in checks.values())
    return result
