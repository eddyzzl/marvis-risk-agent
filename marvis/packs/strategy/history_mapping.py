"""Exact local source locations and user-confirmed historical declarations.

These values are reported historical evidence, never platform MetricObservation
objects. Missing locations remain unknown; formulas and prose are not metrics.
"""
from __future__ import annotations

import csv
from datetime import date, datetime
from decimal import Decimal
import hashlib
import io
import json
import math
from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter


SOURCE_KIND = "strategy_history_external_source"
MAPPING_KIND = "strategy_history_mapping"
MAPPING_VERSION = "strategy.history-mapping.v1"
ORIGIN = "strategy.history_mapping.confirm"
METRICS = ("volume", "approval_rate", "bad_rate", "loan_amount", "delinquent_amount", "ks", "auc", "psi", "lift", "net_yield")
METRIC_UNITS = {name: {"unknown", *units} for name, units in {
    "volume": ("count",), "approval_rate": ("ratio", "percent"), "bad_rate": ("ratio", "percent"),
    "loan_amount": ("currency",), "delinquent_amount": ("currency",), "ks": ("ratio", "percent"),
    "auc": ("ratio", "percent"), "psi": ("score",), "lift": ("score",), "net_yield": ("ratio", "percent"),
}.items()}
METRIC_LABELS = {"业务量": "volume", "通过率": "approval_rate", "坏账率": "bad_rate", "贷款金额": "loan_amount", "逾期金额": "delinquent_amount", "净收益率": "net_yield"}
UNITS = ("unknown", "count", "ratio", "percent", "currency", "score")
PERIOD_FORMATS = ("unknown", "date", "month", "date_range")
MAX_BYTES = 16 * 1024 * 1024
MAX_ROWS = 50000
MAX_COLUMNS = 128
MAX_CELLS = 500000


class HistoryMappingError(ValueError):
    def __init__(self, code, message="历史材料映射尚未保存，请核对来源位置、类型和确认版本。", *, status=422):
        self.code, self.status = code, status
        super().__init__(message)


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def digest(value):
    return hashlib.sha256(canonical(value).encode()).hexdigest()


def _json_pairs(pairs):
    value = {}
    for key, item in pairs:
        if key in value:
            raise HistoryMappingError("history_json_duplicate_key", "JSON 含重复键，无法确定唯一位置；请整理为唯一字段后重新登记。")
        value[key] = item
    return value


def load_source_on_connection(conn, settings, task_id, artifact_id, expected_hash=None):
    # Reuse the existing registered, task-owned, no-symlink external report
    # authority. A filename or caller-supplied bytes alone confer no authority.
    from marvis.packs.strategy.project_context_tools import (
        _artifact_row, _read_verified_registered_artifact, _validate_indexed_external_artifact,
    )
    row = conn.execute("SELECT * FROM task_artifacts WHERE id=? AND task_id=? AND kind=?", (artifact_id, task_id, SOURCE_KIND)).fetchone()
    if row is None:
        raise HistoryMappingError("history_source_not_registered", "请选择当前任务已登记的历史材料；不能用本地路径代替来源登记。")
    record = _artifact_row(row, task_id=task_id, tasks_root=settings.tasks_dir)
    if type(record["provenance"].get("content_size")) is int and record["provenance"]["content_size"] > MAX_BYTES:
        raise HistoryMappingError("history_source_too_large", "历史材料超过当前 16 MiB 定位范围；请另存所需静态汇总表并重新登记，内容不会被截断。")
    if expected_hash is not None and record["content_hash"] != expected_hash:
        raise HistoryMappingError("history_source_changed", "历史材料指纹已变化，请重新读取材料并建立映射。", status=409)
    raw = _read_verified_registered_artifact(record, tasks_root=settings.tasks_dir, max_bytes=MAX_BYTES)
    _validate_indexed_external_artifact(record, tasks_root=settings.tasks_dir, task_id=task_id, content_size=len(raw))
    suffix = Path(record["path"]).suffix.lower()
    tables, cells, containers = [], {}, set()
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            try:
                text, encoding = raw.decode("gb18030"), "gb18030"
            except UnicodeDecodeError:
                raise HistoryMappingError("history_csv_encoding", "CSV 编码不可确定；请另存为 UTF-8 CSV 后重新登记。") from None
        reader = csv.reader(io.StringIO(text), strict=True)
        rows, total_cells = [], 0
        try:
            for row_number, row in enumerate(reader, 1):
                total_cells += len(row)
                if row_number > MAX_ROWS or len(row) > MAX_COLUMNS or total_cells > MAX_CELLS:
                    raise HistoryMappingError("history_source_too_large", "材料超过可定位范围；请提供所需历史汇总表，不能静默截断来源。")
                rows.append(row)
                for column, value in enumerate(row, 1):
                    cells[canonical({"format": "csv", "row": row_number, "column": column})] = {"state": "present" if value != "" else "missing", "value": value if value != "" else None}
        except csv.Error:
            raise HistoryMappingError("history_csv_invalid", "CSV 记录结构不完整，请修复后重新登记。") from None
        tables = [{"name": "CSV", "format": "csv", "rows": len(rows), "columns": max(map(len, rows), default=0), "encoding": encoding}]
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                members = archive.infolist()
                if len(members) > 1024 or sum(member.file_size for member in members) > 128 * 1024 * 1024:
                    raise HistoryMappingError("history_source_too_large", "Excel 解压后超过可定位范围，请提供静态历史汇总表。")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            total = 0
            try:
                for sheet in workbook.worksheets:
                    if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS or total + sheet.max_row * sheet.max_column > MAX_CELLS:
                        raise HistoryMappingError("history_source_too_large", "Excel 超过可定位范围，请提供所需静态汇总表；来源不会被截断。")
                    total += sheet.max_row * sheet.max_column
                    tables.append({"name": sheet.title, "format": "xlsx", "rows": sheet.max_row, "columns": sheet.max_column})
                    for row_number, row in enumerate(sheet.iter_rows(), 1):
                        for column_number, cell in enumerate(row, 1):
                            value = cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else cell.value
                            state = "formula" if cell.data_type == "f" else ("missing" if cell.value is None else "present")
                            typed_cell = {"state": state, "value": value}
                            if isinstance(cell.value, datetime):
                                typed_cell["value_type"] = "datetime"
                            elif isinstance(cell.value, date):
                                typed_cell["value_type"] = "date"
                            coordinate = get_column_letter(column_number) + str(row_number)
                            cells[canonical({"format": "xlsx", "sheet": sheet.title, "cell": coordinate})] = typed_cell
            finally:
                workbook.close()
        except HistoryMappingError:
            raise
        except Exception:
            raise HistoryMappingError("history_xlsx_invalid", "Excel 无法安全定位，请另存为标准 XLSX 静态值表后重新登记。") from None
    elif suffix == ".json":
        try:
            document = json.loads(raw.decode("utf-8-sig"), object_pairs_hook=_json_pairs, parse_float=Decimal, parse_constant=lambda _: (_ for _ in ()).throw(ValueError()))
        except HistoryMappingError:
            raise
        except (ValueError, UnicodeError):
            raise HistoryMappingError("history_json_invalid", "JSON 不是唯一且有限的标准结构，请修复后重新登记。") from None
        def walk(value, pointer="", depth=0):
            if depth > 12 or len(cells) >= MAX_CELLS:
                raise HistoryMappingError("history_source_too_large", "JSON 结构超过可定位范围，请提供所需汇总对象。")
            if isinstance(value, dict):
                containers.add(canonical({"format": "json", "pointer": pointer}))
                for key, item in value.items():
                    walk(item, pointer + "/" + key.replace("~", "~0").replace("/", "~1"), depth + 1)
            elif isinstance(value, list):
                containers.add(canonical({"format": "json", "pointer": pointer}))
                for index, item in enumerate(value):
                    walk(item, pointer + "/" + str(index), depth + 1)
            else:
                cell = {"state": "missing" if value is None else "present", "value": str(value) if isinstance(value, Decimal) else value}
                if isinstance(value, Decimal):
                    cell["value_type"] = "decimal"
                cells[canonical({"format": "json", "pointer": pointer})] = cell
        walk(document)
        tables = [{"name": "JSON", "format": "json", "rows": len(cells), "columns": 1}]
    else:
        raise HistoryMappingError("history_source_unlocatable", "这类文档暂不能确定定位字段或指标；请将所需表格另存为 CSV/XLSX，或提供结构化 JSON 后重新登记。")
    return {"artifact_id": record["id"], "content_hash": record["content_hash"], "format": suffix[1:], "tables": tables, "cells": cells, "containers": containers}


def normalize_locator(locator, source):
    if not isinstance(locator, dict) or locator.get("format") != source["format"]:
        raise HistoryMappingError("history_locator_invalid")
    if source["format"] == "csv":
        if set(locator) != {"format", "row", "column"} or type(locator["row"]) is not int or type(locator["column"]) is not int or not 1 <= locator["row"] <= MAX_ROWS or not 1 <= locator["column"] <= MAX_COLUMNS:
            raise HistoryMappingError("history_locator_invalid")
    elif source["format"] == "xlsx":
        if set(locator) != {"format", "sheet", "cell"} or locator["sheet"] not in {table["name"] for table in source["tables"]} or not isinstance(locator["cell"], str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,4}", locator["cell"]):
            raise HistoryMappingError("history_locator_invalid")
        row, column = coordinate_to_tuple(locator["cell"])
        if row > MAX_ROWS or column > MAX_COLUMNS:
            raise HistoryMappingError("history_locator_invalid")
    elif set(locator) != {"format", "pointer"} or not isinstance(locator["pointer"], str) or len(locator["pointer"]) > 2000 or (locator["pointer"] != "" and not locator["pointer"].startswith("/")) or re.search(r"~(?![01])", locator["pointer"]):
        raise HistoryMappingError("history_locator_invalid")
    return dict(locator)


def source_cell(source, locator):
    locator = normalize_locator(locator, source)
    if canonical(locator) in source.get("containers", set()):
        raise HistoryMappingError("history_json_non_scalar", "请选择 JSON 中明确的单个字段值；对象或数组不能当作指标。")
    cell = source["cells"].get(canonical(locator), {"state": "missing", "value": None})
    return {"locator": locator, **cell, "cell_hash": digest(cell)}


def _identifier(value):
    if not isinstance(value, str) or not 1 <= len(value) <= 128 or not re.fullmatch(r"[\w .-]+", value, re.UNICODE):
        raise HistoryMappingError("history_identifier_invalid", "字段或版本名称不能作为任意自由文本；请提供明确的短标识。")
    return value


def typed_binding(request, source):
    if not isinstance(request, dict) or set(request) != {"binding_key", "source_id", "source_hash", "locator", "cell_hash", "target", "period_key", "version_key"}:
        raise HistoryMappingError("history_binding_invalid")
    key = _identifier(request["binding_key"])
    if request["source_id"] != source["artifact_id"] or request["source_hash"] != source["content_hash"]:
        raise HistoryMappingError("history_source_changed", status=409)
    cell = source_cell(source, request["locator"])
    if cell["cell_hash"] != request["cell_hash"]:
        raise HistoryMappingError("history_cell_changed", "来源单元格已变化，请重新读取并确认。", status=409)
    target = request["target"]
    if not isinstance(target, dict) or set(target) != {"kind", "name", "unit", "period_format"} or target["kind"] not in {"field", "metric", "period", "strategy_version"}:
        raise HistoryMappingError("history_target_invalid")
    kind = target["kind"]
    _identifier(target["name"])
    if target["unit"] not in UNITS or target["period_format"] not in PERIOD_FORMATS:
        raise HistoryMappingError("history_target_invalid")
    if kind == "metric" and target["name"] not in METRICS:
        raise HistoryMappingError("history_metric_unknown", "当前指标类型未安装，请明确已有指标名；不能从报告文字推断新指标。")
    if kind == "metric" and target["unit"] not in METRIC_UNITS[target["name"]]:
        raise HistoryMappingError("history_metric_unit_invalid", "指标与单位不相容，请明确件数、比例、百分比、金额或分值口径；不会自动换算。")
    if (kind != "metric" and target["unit"] != "unknown") or (kind != "period" and target["period_format"] != "unknown"):
        raise HistoryMappingError("history_target_invalid")
    if cell["state"] == "formula":
        raise HistoryMappingError("history_formula_unresolved", "公式尚未由平台验证，不能当作历史指标；请另存静态值表后重新登记。")
    value, availability = None, "unknown"
    if cell["state"] == "present":
        raw = cell["value"]
        if kind in {"field", "strategy_version"}:
            value = _identifier(str(raw)) if type(raw) in (str, int) else None
            if value is None:
                raise HistoryMappingError("history_identifier_invalid")
        elif kind == "metric" and target["unit"] != "unknown":
            if type(raw) in (int, float):
                value = raw
            elif isinstance(raw, str) and re.fullmatch(r"[+-]?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?%?", raw):
                if raw.endswith("%") and target["unit"] != "percent":
                    raise HistoryMappingError("history_unit_mismatch", "来源包含百分号，请明确选择 percent；不会自动换算口径。")
                numeric_text = raw.rstrip("%")
                if len(numeric_text) > 64:
                    raise HistoryMappingError("history_numeric_precision_unsupported", "数值精度超过当前可无损表示范围，请明确适用精度后提供静态汇总值；不会自动舍入。")
                decimal = Decimal(numeric_text)
                if abs(decimal) > 10**16:
                    raise HistoryMappingError("history_metric_range_invalid")
                value = int(decimal) if decimal == decimal.to_integral_value() else float(decimal)
                if Decimal(str(value)) != decimal:
                    raise HistoryMappingError("history_numeric_precision_unsupported", "数值精度超过当前可无损表示范围，请明确适用精度后提供静态汇总值；不会自动舍入。")
            else:
                raise HistoryMappingError("history_metric_not_numeric", "来源不是可确定的数值；报告叙述不能提升为实测指标。")
            if type(value) is int and abs(value) > 2**53 - 1:
                raise HistoryMappingError("history_numeric_precision_unsupported", "整数超过界面和表格可无损表示范围，请明确适用精度后提供静态汇总值；不会自动舍入。")
            probability = target["name"] in {"approval_rate", "bad_rate", "ks", "auc"}
            if abs(value) > 10**16 or not math.isfinite(value) or (probability and target["unit"] == "ratio" and not 0 <= value <= 1) or (probability and target["unit"] == "percent" and not 0 <= value <= 100) or (target["unit"] == "count" and (value < 0 or value != int(value))) or (target["name"] in {"psi", "lift"} and value < 0):
                raise HistoryMappingError("history_metric_range_invalid")
        elif kind == "period" and target["period_format"] != "unknown":
            form = target["period_format"]
            try:
                if form == "month" and isinstance(raw, str) and re.fullmatch(r"\d{4}-\d{2}", raw):
                    date.fromisoformat(raw + "-01")
                    value = {"granularity": "month", "month": raw}
                elif form == "date" and isinstance(raw, str):
                    parsed_date = datetime.fromisoformat(raw).date() if cell.get("value_type") == "datetime" else date.fromisoformat(raw)
                    value = {"granularity": "date", "date": parsed_date.isoformat()}
                elif form == "date_range" and isinstance(raw, str) and len(raw.split("..")) == 2:
                    start, end = map(date.fromisoformat, raw.split(".."))
                    if start > end:
                        raise ValueError
                    value = {"granularity": "date_range", "start": start.isoformat(), "end": end.isoformat()}
                else:
                    raise ValueError
            except ValueError:
                raise HistoryMappingError("history_period_invalid", "来源期间不符合明确选择的格式；请提供 YYYY-MM、YYYY-MM-DD 或 YYYY-MM-DD..YYYY-MM-DD。") from None
        availability = "present" if value is not None else "unknown"
    result = {**request, "binding_key": key, "locator": cell["locator"], "value": value, "availability": availability,
              "validation_status": "historical_reported_unvalidated", "source_cell_state": cell["state"]}
    for relation in ("period_key", "version_key"):
        if result[relation] is not None:
            _identifier(result[relation])
    return result


def validate_relations(bindings):
    keys = {item["binding_key"]: item for item in bindings}
    if len(keys) != len(bindings):
        raise HistoryMappingError("history_binding_key_duplicate")
    for binding in bindings:
        for key, kind in (("period_key", "period"), ("version_key", "strategy_version")):
            related = binding[key]
            if related is not None and (related not in keys or keys[related]["target"]["kind"] != kind):
                raise HistoryMappingError("history_relation_invalid", "期间或版本关联必须指向本次明确确认的对应类型，不能自动猜配。")


def parse_local_instruction(instruction, source):
    """Complete, disclosed local grammar; every clause must be understood."""
    if not isinstance(instruction, str) or not 1 <= len(instruction) <= 12000:
        raise HistoryMappingError("history_instruction_invalid")
    clauses = re.split(r"[;；\n]+", instruction)
    bindings = []
    for clause in clauses:
        if not clause.strip():
            continue
        # e.g. 将 CSV 第2行第3列映射为指标 bad_rate，单位 percent
        match = re.fullmatch(r"\s*将\s*(CSV\s*第(\d+)行第(\d+)列|Excel\s*「([^」]+)」\s*([A-Z]{1,3}[1-9][0-9]*)|JSON\s*「([^」]*)」)\s*映射为(字段|指标|期间|策略版本)\s*([\w.-]+)(?:[，,]\s*(单位|格式)\s*([\w_]+))?\s*[。.]?\s*", clause, re.IGNORECASE)
        if match is None:
            raise HistoryMappingError("history_instruction_needs_clarification", "请使用完整位置映射，例如：将 CSV 第2行第3列映射为指标 bad_rate，单位 percent；期间/版本关联请在表单明确选择。其他约束不会被忽略或猜测。")
        _, row, column, sheet, coordinate, pointer, label, name, qualifier_kind, qualifier = match.groups()
        kind = {"字段": "field", "指标": "metric", "期间": "period", "策略版本": "strategy_version"}[label]
        if qualifier_kind and qualifier_kind != {"metric": "单位", "period": "格式"}.get(kind):
            raise HistoryMappingError("history_instruction_needs_clarification")
        if kind == "metric":
            name = METRIC_LABELS.get(name, name.casefold())
        if qualifier:
            qualifier = {"未知": "unknown", "件数": "count", "比例": "ratio", "百分比": "percent", "金额": "currency", "分值": "score", "日期": "date", "月份": "month", "日期区间": "date_range"}.get(qualifier, qualifier.casefold())
        locator = {"format": "csv", "row": int(row), "column": int(column)} if row else ({"format": "xlsx", "sheet": sheet, "cell": coordinate.upper()} if sheet else {"format": "json", "pointer": pointer})
        cell = source_cell(source, locator)
        target = {"kind": kind, "name": name, "unit": qualifier if kind == "metric" and qualifier else "unknown", "period_format": qualifier if kind == "period" and qualifier else "unknown"}
        if qualifier and kind not in {"metric", "period"}:
            raise HistoryMappingError("history_instruction_needs_clarification")
        bindings.append({"binding_key": "b" + str(len(bindings) + 1), "source_id": source["artifact_id"], "source_hash": source["content_hash"], "locator": locator, "cell_hash": cell["cell_hash"], "target": target, "period_key": None, "version_key": None})
    if not bindings:
        raise HistoryMappingError("history_instruction_needs_clarification")
    return bindings
