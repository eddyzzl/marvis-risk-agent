from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
import asyncio
import json
import zipfile

from docx import Document
from fastapi import BackgroundTasks
import pandas as pd
import pytest

from marvis.db import TaskRepository, init_db
from marvis.domain import (
    TASK_TYPE_VALIDATION,
    TASK_TYPE_VALIDATION_BATCH,
    TaskCreate,
    TaskStatus,
)
from marvis.output.validation_batch_excel import write_validation_batch_excel
from marvis.repositories.validation_batches import ValidationBatchRepository
from marvis.repositories.validation_contracts import ValidationContractRepository
from marvis.settings import build_settings
from marvis.validation.input_confirmation import validate_confirmation_against_materials
from marvis.validation_batch_runner import (
    _bounded_public_error,
    _finalize_batch,
    _individual_reports_complete,
    _prepare_failed_child_for_retry,
    compose_batch_report_conclusions,
    mark_validation_batch_parent_running,
    run_validation_batch_job,
    run_validation_batch,
    sync_agent_batch_after_child_report,
    sync_agent_batch_child_progress,
    validation_batch_item_outcome,
    validation_batch_stress_risk,
)
from marvis.validation_materials import resolve_selected_validation_materials
from marvis.validation.results import validation_results_to_dict
from marvis.safe_paths import prepare_task_output_dir
from tests.output.test_excel import _make_pmml_results
from tests.validation_builders import make_validation_confirmation
from tests.validation_material_builders import write_validation_material_bundle


def _batch_with_valid_child(tmp_path: Path, *, child_count: int = 1):
    settings = build_settings(tmp_path / "workspace")
    init_db(settings.db_path)
    report_template_path = settings.workspace / "report_templates" / "default.docx"
    report_template_path.parent.mkdir(parents=True, exist_ok=True)
    template = Document()
    template.add_paragraph("模型：{{TEXT:model_name}}")
    template.add_paragraph("{{TEXT:pressure_test_summary}}")
    template.add_paragraph("{{TEXT:pressure_impact_recommendation}}")
    template.add_paragraph("{{TEXT:final_validation_conclusion}}")
    template.save(report_template_path)
    sample = pd.DataFrame(
        {
            "x1": [0.0, 1.0, 2.0, 3.0] * 3,
            "x2": [0.0, 1.0, 0.0, 1.0] * 3,
            "y": [0, 1, 0, 1] * 3,
            "split": ["train"] * 4 + ["test"] * 4 + ["oot"] * 4,
            "apply_month": ["202601"] * 4
            + ["202602"] * 4
            + ["202603"] * 4,
        }
    )
    child_payloads = []
    for index in range(child_count):
        bundle = write_validation_material_bundle(
            settings.workspace / "materials" / f"model-{index + 1}",
            notebook_source=(
                "RMC_TARGET_COL='y'\nRMC_SPLIT_COL='split'\n"
                "RMC_TIME_COL='apply_month'\n"
                "RMC_PMML_OUTPUT_FIELD='probability_1'\nRMC_MODEL_PARAMS={}\n"
            ),
            sample=sample,
        )
        child_payloads.append(
            TaskCreate(
                task_type=TASK_TYPE_VALIDATION,
                model_name=f"模型{chr(ord('A') + index)}",
                model_version="v1",
                validator="qa",
                source_dir=str(bundle.root),
                run_mode="agent",
                notebook_path=str(bundle.notebook_path),
                sample_path=str(bundle.sample_path),
                pmml_path=str(bundle.pmml_path),
                dictionary_path=str(bundle.dictionary_path),
            )
        )
    repo = ValidationBatchRepository(settings.db_path)
    batch = repo.create_batch(
        TaskCreate(
            task_type=TASK_TYPE_VALIDATION_BATCH,
            model_name="批次",
            model_version="",
            validator="qa",
            source_dir=str(settings.workspace),
            run_mode="agent",
        ),
        child_payloads,
    )
    return settings, repo, batch


def test_batch_runner_stops_only_at_input_contract_and_never_auto_confirms(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    summary_calls: list[dict] = []

    run_validation_batch(
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=lambda **kwargs: summary_calls.append(kwargs),
    )

    updated_batch = repo.get_batch(batch.parent_task_id)
    [item] = repo.list_items(batch.parent_task_id)
    contract = ValidationContractRepository(settings.db_path).get(item.child_task_id)
    assert updated_batch.status == "awaiting_confirmation"
    assert item.status == "awaiting_confirmation"
    assert item.stage == "input_confirmation"
    assert contract is not None
    assert contract.status == "pending_confirmation"
    assert contract.contract.confirmed == {}
    assert summary_calls == []
    assert TaskRepository(settings.db_path).get_task(
        batch.parent_task_id
    ).status is TaskStatus.SCANNED


def test_batch_job_wrapper_closes_job_and_parks_parent_at_scanned_gate(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    task_repo = TaskRepository(settings.db_path)
    job_id = task_repo.start_job(batch.parent_task_id, "validation_batch")

    run_validation_batch_job(
        job_id=job_id,
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=lambda **_kwargs: None,
    )

    job = task_repo.get_job(job_id)
    assert job is not None
    assert job["status"] == "succeeded"
    assert job["started_at"] is not None
    assert job["finished_at"] is not None
    assert task_repo.get_active_job_kind(batch.parent_task_id) is None
    assert task_repo.get_task(batch.parent_task_id).status is TaskStatus.SCANNED
    assert repo.get_batch(batch.parent_task_id).status == "awaiting_confirmation"


def test_batch_job_wrapper_closes_job_and_parent_on_unexpected_exception(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    task_repo = TaskRepository(settings.db_path)
    job_id = task_repo.start_job(batch.parent_task_id, "validation_batch")

    def broken_runner(**_kwargs):
        raise RuntimeError(f"boom at {settings.workspace}/secret.csv")

    run_validation_batch_job(
        job_id=job_id,
        settings=settings,
        parent_task_id=batch.parent_task_id,
        runner=broken_runner,
    )

    job = task_repo.get_job(job_id)
    assert job is not None
    assert job["status"] == "failed"
    assert job["error_name"] == "RuntimeError"
    assert str(settings.workspace) not in (job["error_value"] or "")
    assert job["finished_at"] is not None
    assert task_repo.get_active_job_kind(batch.parent_task_id) is None
    assert task_repo.get_task(batch.parent_task_id).status is TaskStatus.FAILED
    assert repo.get_batch(batch.parent_task_id).status == "failed"


def test_batch_job_wrapper_releases_claim_when_queued_job_never_starts(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    task_repo = TaskRepository(settings.db_path)
    job_id = task_repo.start_job(batch.parent_task_id, "validation_batch")
    repo.claim_start(batch.parent_task_id)
    mark_validation_batch_parent_running(task_repo, batch.parent_task_id)
    task_repo.finish_job(
        job_id,
        status="failed",
        error_name="QueueLost",
        error_value="worker unavailable",
    )

    run_validation_batch_job(
        job_id=job_id,
        settings=settings,
        parent_task_id=batch.parent_task_id,
    )

    recovered = repo.get_batch(batch.parent_task_id)
    parent = task_repo.get_task(batch.parent_task_id)
    assert recovered.status == "partial_failure"
    assert "可修复后重试" in recovered.error_message
    assert parent.status is TaskStatus.FAILED
    assert task_repo.get_active_job_kind(batch.parent_task_id) is None


def test_all_failed_batch_retries_scan_after_transient_failure(
    tmp_path: Path,
    monkeypatch,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)

    def fail_scan(*_args, **_kwargs):
        raise RuntimeError("temporary scanner failure")

    monkeypatch.setattr(
        "marvis.validation_batch_runner.perform_scan_task",
        fail_scan,
    )
    run_validation_batch(
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=lambda **_kwargs: None,
    )
    assert repo.get_batch(batch.parent_task_id).status == "failed"
    [failed] = repo.list_items(batch.parent_task_id)
    assert failed.status == "failed"
    assert failed.stage == "scan"

    monkeypatch.undo()
    run_validation_batch(
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=lambda **_kwargs: None,
    )

    [retried] = repo.list_items(batch.parent_task_id)
    assert repo.get_batch(batch.parent_task_id).status == "awaiting_confirmation", (
        retried.error_message
    )
    assert retried.status == "awaiting_confirmation"
    assert retried.stage == "input_confirmation"


def test_batch_public_error_redacts_all_allowed_material_roots(
    tmp_path: Path,
    monkeypatch,
):
    settings = build_settings(tmp_path / "workspace")
    external_root = tmp_path / "customer-materials"
    external_root.mkdir()
    monkeypatch.setenv("RMC_MATERIAL_ROOTS", str(external_root))

    detail = _bounded_public_error(
        RuntimeError(f"cannot read {external_root}/model-a/sample.csv"),
        settings,
    )

    assert str(external_root) not in detail
    assert "<material-root>" in detail


def test_batch_summary_writer_rejects_parent_outputs_directory_symlink(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    parent_dir = settings.tasks_dir / batch.parent_task_id
    parent_dir.mkdir(parents=True)
    outside = tmp_path / "outside-output"
    outside.mkdir()
    (parent_dir / "outputs").symlink_to(outside, target_is_directory=True)
    writer_calls: list[dict] = []

    final = _finalize_batch(
        settings=settings,
        parent_task_id=batch.parent_task_id,
        batch_repo=repo,
        task_repo=TaskRepository(settings.db_path),
        items=repo.list_items(batch.parent_task_id),
        manual_review_base_url="",
        summary_writer=lambda **kwargs: writer_calls.append(kwargs),
    )

    assert final.status == "failed"
    assert writer_calls == []
    assert not (outside / "validation_batch_summary.xlsx").exists()


def test_individual_report_completeness_rejects_outputs_directory_symlink(
    tmp_path: Path,
):
    tasks_dir = tmp_path / "tasks"
    task_dir = tasks_dir / "task-1"
    outside = tmp_path / "outside-reports"
    outside.mkdir(parents=True)
    for filename in ("validation_report.docx", "validation.xlsx"):
        with zipfile.ZipFile(outside / filename, "w") as archive:
            archive.writestr("placeholder.txt", "valid zip container")
    task_dir.mkdir(parents=True)
    (task_dir / "outputs").symlink_to(outside, target_is_directory=True)

    assert _individual_reports_complete(tasks_dir, "task-1") is False


@pytest.mark.parametrize(
    ("failed_stage", "expected_status"),
    [
        ("scan", TaskStatus.CREATED),
        ("pmml_scoring", TaskStatus.SCANNED),
        ("metrics", TaskStatus.EXECUTED),
        ("report_conclusion", TaskStatus.WRITING_ARTIFACTS),
        ("report", TaskStatus.WRITING_ARTIFACTS),
    ],
)
def test_failed_batch_child_retry_rewinds_to_recorded_stage(
    tmp_path: Path,
    failed_stage: str,
    expected_status: TaskStatus,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    [item] = repo.list_items(batch.parent_task_id)
    task_repo = TaskRepository(settings.db_path)
    child = task_repo.get_task(item.child_task_id)
    task_repo.reset_status_for_agent_rerun(
        child.id,
        TaskStatus.FAILED,
        "simulated failed attempt",
    )
    failed_item = repo.update_item(
        item.id,
        status="failed",
        stage=failed_stage,
        outcome="failed",
        finished=True,
    )

    retried = _prepare_failed_child_for_retry(
        task_repo,
        task_repo.get_task(child.id),
        failed_item,
    )

    assert retried.status is expected_status
    assert "批次修复后重试" in retried.status_message


def test_batch_runner_records_one_failure_and_continues_to_the_next_item(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    [valid_item] = repo.list_items(batch.parent_task_id)
    invalid_root = settings.workspace / "materials" / "model-invalid"
    invalid_root.mkdir(parents=True)
    invalid_paths = {
        "notebook_path": invalid_root / "invalid.ipynb",
        "sample_path": invalid_root / "invalid.csv",
        "pmml_path": invalid_root / "invalid.pmml",
        "dictionary_path": invalid_root / "invalid.xlsx",
    }
    for path in invalid_paths.values():
        path.touch()

    # Recreate the batch so the broken model is ordinal 1 and the valid model is 2.
    valid_task = TaskRepository(settings.db_path).get_task(valid_item.child_task_id)
    isolated_settings = build_settings(tmp_path / "isolated-workspace")
    init_db(isolated_settings.db_path)
    isolated_repo = ValidationBatchRepository(isolated_settings.db_path)
    isolated_batch = isolated_repo.create_batch(
        TaskCreate(
            task_type=TASK_TYPE_VALIDATION_BATCH,
            model_name="隔离批次",
            model_version="",
            validator="qa",
            source_dir=str(isolated_settings.workspace),
            run_mode="agent",
        ),
        [
            TaskCreate(
                task_type=TASK_TYPE_VALIDATION,
                model_name="损坏模型",
                model_version="v1",
                validator="qa",
                source_dir=str(invalid_root),
                run_mode="agent",
                **{key: str(value) for key, value in invalid_paths.items()},
            ),
            TaskCreate(
                task_type=TASK_TYPE_VALIDATION,
                model_name=valid_task.model_name,
                model_version=valid_task.model_version,
                validator=valid_task.validator,
                source_dir=valid_task.source_dir,
                run_mode="agent",
                notebook_path=valid_task.notebook_path,
                sample_path=valid_task.sample_path,
                pmml_path=valid_task.pmml_path,
                dictionary_path=valid_task.dictionary_path,
            ),
        ],
    )

    run_validation_batch(
        settings=isolated_settings,
        parent_task_id=isolated_batch.parent_task_id,
        summary_writer=lambda **_kwargs: None,
    )

    items = isolated_repo.list_items(isolated_batch.parent_task_id)
    assert [item.status for item in items] == ["failed", "awaiting_confirmation"]
    assert items[0].stage == "scan"
    assert items[0].error_code
    messages = TaskRepository(isolated_settings.db_path).list_agent_messages(
        isolated_batch.parent_task_id
    )
    failures = [message for message in messages if message["stage"] == "failure"]
    assert len(failures) == 1
    assert "损坏模型" in failures[0]["content"]
    assert "继续处理后续模型" in failures[0]["content"]


@pytest.mark.pmml_runtime
def test_confirmed_batch_item_waits_for_report_approval_without_failing(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    run_validation_batch(
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=lambda **_kwargs: None,
    )
    [waiting_item] = repo.list_items(batch.parent_task_id)
    task_repo = TaskRepository(settings.db_path)
    child = task_repo.get_task(waiting_item.child_task_id)
    contract_repo = ValidationContractRepository(settings.db_path)
    candidate = contract_repo.get(child.id)
    assert candidate is not None
    materials = resolve_selected_validation_materials(child)
    validated = validate_confirmation_against_materials(
        contract=candidate.contract,
        sample_path=materials.sample,
        dictionary_path=materials.dictionary,
        requested=replace(make_validation_confirmation(), metadata_sheet=None),
    )
    contract_repo.confirm(
        child.id,
        validated.values,
        expected_revision=candidate.revision,
        resolved_sample_schema=validated.sample_schema,
        resolved_feature_metadata=validated.feature_metadata,
    )
    summary_calls: list[dict] = []

    def write_summary(**kwargs):
        summary_calls.append(kwargs)
        return write_validation_batch_excel(**kwargs)

    parent_job_id = task_repo.start_job(
        batch.parent_task_id,
        "validation_batch",
    )
    run_validation_batch_job(
        job_id=parent_job_id,
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=write_summary,
    )

    final_batch = repo.get_batch(batch.parent_task_id)
    [item] = repo.list_items(batch.parent_task_id)
    final_child = task_repo.get_task(child.id)
    report_values, _revision = task_repo.get_report_values(child.id)
    assert final_batch.status == "awaiting_confirmation"
    assert final_batch.summary_path == ""
    assert item.status == "awaiting_confirmation"
    assert item.stage == "report_conclusion"
    assert item.error_message == ""
    assert final_child.status is TaskStatus.WRITING_ARTIFACTS
    assert item.report_complete is False
    assert summary_calls == []
    assert not report_values.get("TEXT:final_validation_conclusion")
    assert not (settings.tasks_dir / child.id / "outputs" / "validation_report.docx").exists()
    messages = task_repo.list_agent_messages(batch.parent_task_id)
    assert messages[-1]["stage"] == "report_conclusion"
    assert messages[-1]["metadata"]["waiting_report_count"] == 1
    assert "等待报告结论确认" in messages[-1]["content"]
    parent = task_repo.get_task(batch.parent_task_id)
    assert parent.status is TaskStatus.SCANNED
    parent_job = task_repo.get_job(parent_job_id)
    assert parent_job is not None
    assert parent_job["status"] == "succeeded"
    assert task_repo.get_active_job_kind(batch.parent_task_id) is None


@pytest.mark.pmml_runtime
def test_two_models_complete_reports_only_after_explicit_confirmation(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path, child_count=2)
    run_validation_batch(
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=lambda **_kwargs: None,
    )
    waiting_items = repo.list_items(batch.parent_task_id)
    assert [item.status for item in waiting_items] == [
        "awaiting_confirmation",
        "awaiting_confirmation",
    ]
    task_repo = TaskRepository(settings.db_path)
    contract_repo = ValidationContractRepository(settings.db_path)
    for item in waiting_items:
        child = task_repo.get_task(item.child_task_id)
        candidate = contract_repo.get(child.id)
        assert candidate is not None
        materials = resolve_selected_validation_materials(child)
        validated = validate_confirmation_against_materials(
            contract=candidate.contract,
            sample_path=materials.sample,
            dictionary_path=materials.dictionary,
            requested=replace(
                make_validation_confirmation(),
                metadata_sheet=None,
            ),
        )
        contract_repo.confirm(
            child.id,
            validated.values,
            expected_revision=candidate.revision,
            resolved_sample_schema=validated.sample_schema,
            resolved_feature_metadata=validated.feature_metadata,
        )

    summary_calls: list[dict] = []

    def write_summary(**kwargs):
        summary_calls.append(kwargs)
        return write_validation_batch_excel(**kwargs)

    parent_job_id = task_repo.start_job(
        batch.parent_task_id,
        "validation_batch",
    )
    run_validation_batch_job(
        job_id=parent_job_id,
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=write_summary,
    )

    pending_items = repo.list_items(batch.parent_task_id)
    assert repo.get_batch(batch.parent_task_id).status == "awaiting_confirmation"
    assert all(item.stage == "report_conclusion" for item in pending_items)
    assert summary_calls == []
    from marvis.agent.service import fallback_word_conclusions
    from marvis.agent.validation_app_service import confirm_agent_report_conclusions
    from marvis.agent.validation_evidence import agent_evidence_from_settings

    for item in pending_items:
        child = task_repo.get_task(item.child_task_id)
        outputs = settings.tasks_dir / child.id / "outputs"
        assert not (outputs / "validation_report.docx").exists()
        background_tasks = BackgroundTasks()
        _values, revision = task_repo.get_report_values(child.id)
        confirm_agent_report_conclusions(
            repo_=task_repo,
            task=child,
            task_id=child.id,
            settings=settings,
            text_values=fallback_word_conclusions(
                task=child,
                evidence=agent_evidence_from_settings(settings, child.id),
            ),
            expected_revision=revision,
            background_tasks=background_tasks,
        )
        asyncio.run(background_tasks())

    final_items = repo.list_items(batch.parent_task_id)
    assert repo.get_batch(batch.parent_task_id).status == "completed"
    assert [item.ordinal for item in final_items] == [1, 2]
    assert all(
        item.status in {"succeeded", "review_required"}
        for item in final_items
    )
    assert Path(repo.get_batch(batch.parent_task_id).summary_path).is_file()
    for item in final_items:
        outputs = settings.tasks_dir / item.child_task_id / "outputs"
        assert (outputs / "validation_report.docx").is_file()
        assert (outputs / "validation.xlsx").is_file()
    completed = [
        message
        for message in task_repo.list_agent_messages(batch.parent_task_id)
        if message["stage"] == "batch_completed"
    ]
    assert len(completed) == 1
    assert completed[0]["metadata"]["manual_review_count"] == sum(
        item.status == "review_required" for item in final_items
    )
    assert "summary_path" not in completed[0]["metadata"]
    assert str(settings.tasks_dir) not in str(completed[0]["metadata"])


def test_batch_outcome_keeps_oot_ks_display_only():
    common = {
        "pmml_status": "pass",
        "oot_psi": 0.03,
        "stress_risk": "low",
        "report_complete": True,
        "task_status": "succeeded",
    }

    low_ks = validation_batch_item_outcome(oot_ks=0.01, **common)
    high_ks = validation_batch_item_outcome(oot_ks=0.99, **common)

    assert low_ks == ("pass", ())
    assert high_ks == low_ks
    assert validation_batch_item_outcome(
        oot_ks=0.99,
        **{**common, "oot_psi": 0.12},
    )[0] == "manual_review"
    assert validation_batch_item_outcome(
        oot_ks=0.99,
        **{**common, "oot_psi": -0.01},
    )[0] == "manual_review"
    assert validation_batch_item_outcome(
        oot_ks=0.99,
        **{**common, "stress_risk": "unknown"},
    )[0] == "manual_review"


def test_batch_stress_risk_requires_every_category_to_be_complete_and_classified():
    baseline = SimpleNamespace(ks=0.4)
    low = SimpleNamespace(
        status="completed",
        ks_after=0.39,
        psi_vs_baseline=0.01,
    )
    incomplete = SimpleNamespace(
        status="failed",
        ks_after=None,
        psi_vs_baseline=None,
    )
    missing_ks = SimpleNamespace(
        status="completed",
        ks_after=None,
        psi_vs_baseline=0.01,
    )
    unclassified = SimpleNamespace(
        status="completed",
        ks_after=0.39,
        psi_vs_baseline=None,
    )

    def results(*categories):
        return SimpleNamespace(
            stress_test=SimpleNamespace(
                status="completed",
                baseline=baseline,
                per_category=list(categories),
            )
        )

    assert validation_batch_stress_risk(results(low)) == "low"
    assert validation_batch_stress_risk(results(low, incomplete)) is None
    assert validation_batch_stress_risk(results(low, missing_ks)) is None
    assert validation_batch_stress_risk(results(low, unclassified)) is None
    assert validation_batch_stress_risk(results()) is None


def test_batch_stress_risk_keeps_known_alarm_when_ks_decay_is_undefined():
    def results(psi: float):
        return SimpleNamespace(
            stress_test=SimpleNamespace(
                status="completed",
                baseline=SimpleNamespace(ks=0.0),
                per_category=[
                    SimpleNamespace(
                        status="completed",
                        ks_after=0.0,
                        psi_vs_baseline=psi,
                    )
                ],
            )
        )

    assert validation_batch_stress_risk(results(0.30)) == "high"
    assert validation_batch_stress_risk(results(0.15)) == "medium"
    assert validation_batch_stress_risk(results(0.01)) is None


def test_batch_report_uses_deterministic_final_verdict_even_with_agent_text():
    fallback = {
        "TEXT:pressure_test_summary": "确定性压力总结",
        "TEXT:pressure_impact_recommendation": "确定性压力建议",
        "TEXT:final_validation_conclusion": "OOT KS 仅展示，不参与通过判定。",
    }
    generated = {
        "TEXT:pressure_test_summary": "Agent 压力总结",
        "TEXT:pressure_impact_recommendation": "Agent 压力建议",
        "TEXT:final_validation_conclusion": "因为 OOT KS 高，所以模型通过。",
    }

    values, source = compose_batch_report_conclusions(
        generated=generated,
        fallback=fallback,
    )

    assert values["TEXT:pressure_test_summary"] == "Agent 压力总结"
    assert values["TEXT:pressure_impact_recommendation"] == "Agent 压力建议"
    assert values["TEXT:final_validation_conclusion"] == fallback[
        "TEXT:final_validation_conclusion"
    ]
    assert source == "agent_generated_with_deterministic_verdict"


def test_partial_failure_retry_reconsiders_failed_item_after_material_repair(
    tmp_path: Path,
):
    settings, repo, batch = _batch_with_valid_child(tmp_path)
    run_validation_batch(
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=lambda **_kwargs: None,
    )
    [item] = repo.list_items(batch.parent_task_id)
    repo.update_item(
        item.id,
        status="failed",
        stage="scan",
        outcome="failed",
        error_code="MaterialError",
        error_message="材料待修复",
        finished=True,
    )
    repo.update_batch(batch.parent_task_id, status="partial_failure", finished=True)

    run_validation_batch(
        settings=settings,
        parent_task_id=batch.parent_task_id,
        summary_writer=lambda **_kwargs: None,
    )

    retried = repo.get_item(item.id)
    assert retried.status == "awaiting_confirmation"
    assert retried.stage == "input_confirmation"
    assert retried.finished_at is None
    retried_batch = repo.get_batch(batch.parent_task_id)
    assert retried_batch.status == "awaiting_confirmation"
    assert retried_batch.finished_at is None


def _advance_child_to_succeeded(task_repo: TaskRepository, task_id: str) -> None:
    task_repo.update_status(
        task_id, TaskStatus.SCANNED, "scanned", expected=TaskStatus.CREATED,
    )
    task_repo.update_status(
        task_id, TaskStatus.RUNNING, "running", expected=TaskStatus.SCANNED,
    )
    task_repo.update_status(
        task_id, TaskStatus.EXECUTED, "executed", expected=TaskStatus.RUNNING,
    )
    task_repo.update_status(
        task_id,
        TaskStatus.COMPUTING_METRICS,
        "metrics",
        expected=TaskStatus.EXECUTED,
    )
    task_repo.update_status(
        task_id,
        TaskStatus.WRITING_ARTIFACTS,
        "artifacts",
        expected=TaskStatus.COMPUTING_METRICS,
    )
    task_repo.update_status(
        task_id,
        TaskStatus.SUCCEEDED,
        "succeeded",
        expected=TaskStatus.WRITING_ARTIFACTS,
    )


def _write_agent_child_outputs(settings, child_task_id: str) -> None:
    output_dir = prepare_task_output_dir(settings.tasks_dir, child_task_id)
    payload = validation_results_to_dict(_make_pmml_results())
    (output_dir / "validation_results.json").write_text(
        json.dumps(payload),
        encoding="utf-8",
    )
    for filename in ("validation_report.docx", "validation.xlsx"):
        with zipfile.ZipFile(output_dir / filename, "w") as archive:
            archive.writestr("placeholder.txt", "report")


def test_agent_child_report_records_item_then_writes_n2_summary(tmp_path: Path):
    settings, repo, batch = _batch_with_valid_child(tmp_path, child_count=2)
    task_repo = TaskRepository(settings.db_path)
    items = repo.list_items(batch.parent_task_id)
    summary_calls: list[dict] = []

    def write_summary(**kwargs):
        summary_calls.append(kwargs)
        return write_validation_batch_excel(**kwargs)

    first = items[0]
    _write_agent_child_outputs(settings, first.child_task_id)
    _advance_child_to_succeeded(task_repo, first.child_task_id)
    after_first = sync_agent_batch_after_child_report(
        settings=settings,
        child_task_id=first.child_task_id,
        summary_writer=write_summary,
    )

    recorded_first = repo.get_item(first.id)
    assert recorded_first.status in {"succeeded", "review_required"}
    assert recorded_first.stage == "completed"
    assert after_first is not None
    assert after_first.status != "completed"
    assert after_first.summary_path in {None, ""}
    assert summary_calls == []

    second = items[1]
    _write_agent_child_outputs(settings, second.child_task_id)
    _advance_child_to_succeeded(task_repo, second.child_task_id)
    after_second = sync_agent_batch_after_child_report(
        settings=settings,
        child_task_id=second.child_task_id,
        summary_writer=write_summary,
    )

    assert after_second is not None
    assert after_second.status in {"completed", "partial_failure"}
    assert after_second.summary_path.endswith("validation_batch_summary.xlsx")
    assert Path(after_second.summary_path).is_file()
    assert len(summary_calls) == 1
    parent = task_repo.get_task(batch.parent_task_id)
    assert parent.status in {TaskStatus.SUCCEEDED, TaskStatus.REVIEW_REQUIRED}


def test_agent_single_child_batch_does_not_write_summary_excel(tmp_path: Path):
    settings, repo, batch = _batch_with_valid_child(tmp_path, child_count=1)
    task_repo = TaskRepository(settings.db_path)
    [item] = repo.list_items(batch.parent_task_id)
    summary_calls: list[dict] = []
    _write_agent_child_outputs(settings, item.child_task_id)
    _advance_child_to_succeeded(task_repo, item.child_task_id)

    result = sync_agent_batch_after_child_report(
        settings=settings,
        child_task_id=item.child_task_id,
        summary_writer=lambda **kwargs: summary_calls.append(kwargs),
    )

    recorded = repo.get_item(item.id)
    assert recorded.status == "queued"
    assert summary_calls == []
    assert result is not None
    assert result.summary_path in {None, ""}
    assert not (
        settings.tasks_dir
        / batch.parent_task_id
        / "outputs"
        / "validation_batch_summary.xlsx"
    ).exists()


@pytest.mark.parametrize("all_failed", [False, True])
def test_agent_failed_children_are_recorded_and_batch_summary_finishes(tmp_path: Path, all_failed):
    settings, batch_repo, batch = _batch_with_valid_child(tmp_path, child_count=2)
    task_repo = TaskRepository(settings.db_path)
    items = batch_repo.list_items(batch.parent_task_id)
    for index, item in enumerate(items):
        sync_agent_batch_child_progress(
            db_path=settings.db_path, child_task_id=item.child_task_id,
            status="running", stage="scan",
        )
        if index == 0 or all_failed:
            task_repo.update_status(
                item.child_task_id, TaskStatus.FAILED, "scan failed",
                expected=TaskStatus.CREATED,
            )
        else:
            _write_agent_child_outputs(settings, item.child_task_id)
            _advance_child_to_succeeded(task_repo, item.child_task_id)

    result = sync_agent_batch_after_child_report(
        settings=settings, child_task_id=items[-1].child_task_id,
    )

    assert result.status == ("failed" if all_failed else "partial_failure")
    assert Path(result.summary_path).is_file()
    first = batch_repo.get_item(items[0].id)
    assert first.status == "failed"
    assert first.report_complete is False
    assert "scan failed" in first.error_message


def test_failed_agent_jobs_finalize_all_failed_batch(tmp_path: Path, monkeypatch):
    from marvis.agent.validation_app_service import run_agent_validation_job

    settings, batch_repo, batch = _batch_with_valid_child(tmp_path, child_count=2)
    task_repo = TaskRepository(settings.db_path)
    monkeypatch.setattr("marvis.agent.validation_app_service.open_agent_stage", lambda *_args, **_kwargs: None)
    def fail_scan(repo, _settings, task_id, _profile, **_kwargs):
        repo.update_status(task_id, TaskStatus.FAILED, "scan failed", expected=TaskStatus.CREATED)
        return False
    monkeypatch.setattr("marvis.agent.validation_app_service.run_agent_scan_stage", fail_scan)
    for item in batch_repo.list_items(batch.parent_task_id):
        job_id = task_repo.start_job(item.child_task_id, "agent")
        run_agent_validation_job(
            job_id, settings, item.child_task_id, {"model_id": "test"},
            stage="scan", acceptance_mode="auto_accept",
        )
        assert task_repo.get_active_job_kind(item.child_task_id) is None
    final = batch_repo.get_batch(batch.parent_task_id)
    assert final.status == "failed"
    assert Path(final.summary_path).is_file()


def test_agent_child_progress_updates_batch_item_without_finalizing(tmp_path: Path):
    settings, repo, batch = _batch_with_valid_child(tmp_path, child_count=2)
    items = repo.list_items(batch.parent_task_id)
    first = items[0]
    sync_agent_batch_child_progress(
        db_path=settings.db_path,
        child_task_id=first.child_task_id,
        status="running",
        stage="scan",
    )
    recorded = repo.get_item(first.id)
    assert recorded.status == "running"
    assert recorded.stage == "scan"
    other = repo.get_item(items[1].id)
    assert other.status == "queued"
    parent = repo.get_batch(batch.parent_task_id)
    assert parent.status == "created"
