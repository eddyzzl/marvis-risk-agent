from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from marvis.packs.strategy.asset_stability import (
    FREEZE_SCHEMA,
    measure_asset_stability,
    seal,
    validate_evidence,
)
from marvis.packs.strategy.asset_stability_tools import (
    ARTIFACT_KIND,
    FREEZE_KIND,
    candidate_asset_stability_catalog,
    load_asset_stability_artifact,
    run_freeze_candidate_asset_stability,
    run_measure_candidate_asset_stability,
)
from marvis.packs.strategy.errors import StrategyError
from tests.test_strategy_pool_tools import _setup


def _fixture(tmp_path):
    fx = _setup(tmp_path, native_sample=True)
    catalog = candidate_asset_stability_catalog(fx["runtime"], task_id=fx["task"].id)
    assert catalog["sources"] and catalog["samples"]
    return fx, {
        "source_ref": catalog["sources"][0]["ref"],
        "sample_design_ref": catalog["samples"][0]["ref"],
        "partitions": ["validation", "oot"],
    }


def test_fixed_buckets_bad_zero_missing_labels_and_monthly_psi():
    frozen = seal(
        {"schema_version": FREEZE_SCHEMA, "partitions": ["validation", "oot"]},
        prefix="asset-freeze",
    )
    output = measure_asset_stability(
        freeze=frozen,
        frozen_ref={},
        buckets=np.array(["hit", "miss", "hit", "miss", "hit", "hit", "miss", "miss"]),
        bucket_ids=["hit", "miss"],
        labels=np.array([0, 1, 1, np.nan, 1, 0, 1, 0]),
        months=pd.Series(["202601"] * 4 + ["202602"] * 2 + ["202603"] * 2),
        masks={
            "development": np.array([1, 1, 1, 1, 0, 0, 0, 0], dtype=bool),
            "validation": np.array([0, 0, 0, 0, 1, 1, 0, 0], dtype=bool),
            "oot": np.array([0, 0, 0, 0, 0, 0, 1, 1], dtype=bool),
        },
    )
    evidence = validate_evidence(output)
    assert evidence["baseline"]["labeled_count"] == 3
    assert evidence["baseline"]["buckets"][0]["share"] == 0.5
    assert evidence["partitions"][0]["buckets"][0]["bad_rate"] == 0.5
    assert evidence["partitions"][0]["psi_vs_development"] > 1
    assert evidence["lifecycle"]["candidate_promoted"] is False


def test_native_asset_freeze_then_independent_measurement_exact_retry(
    tmp_path, monkeypatch
):
    fx, request = _fixture(tmp_path)
    frozen = run_freeze_candidate_asset_stability(request, fx["ctx"], fx["runtime"])
    assert frozen["evidence"]["membership"]["development"]["count"] == 6
    assert frozen["evidence"]["partitions"] == ["validation", "oot"]
    assert not [
        r
        for r in fx["runtime"].task_artifacts.list_for_task(fx["task"].id)
        if r["kind"] == ARTIFACT_KIND
    ]
    output = run_measure_candidate_asset_stability(
        {"frozen_ref": frozen["ref"]}, fx["ctx"], fx["runtime"]
    )
    assert [p["sample_count"] for p in output["evidence"]["partitions"]] == [2, 2]
    assert output["evidence"]["baseline"]["sample_count"] == 6
    assert output["evidence"]["freeze"] == frozen["evidence"]
    assert all(p["small_sample"] for p in output["evidence"]["partitions"])
    import marvis.packs.strategy.asset_stability_tools as tools

    def forbid_snapshot(*args, **kwargs):
        raise AssertionError("exact measured evidence replay must not read source rows")

    monkeypatch.setattr(tools, "_read_authenticated_parquet_snapshot", forbid_snapshot)
    assert (
        run_measure_candidate_asset_stability(
            {"frozen_ref": frozen["ref"]}, fx["ctx"], fx["runtime"]
        )
        == output
    )
    assert (
        load_asset_stability_artifact(
            fx["runtime"], task_id=fx["task"].id, ref=output["ref"]
        ).evidence
        == output["evidence"]
    )


def test_exact_source_sample_refs_are_required_before_any_snapshot_read(
    tmp_path, monkeypatch
):
    fx, request = _fixture(tmp_path)
    import marvis.packs.strategy.asset_stability_tools as tools

    calls = []
    monkeypatch.setattr(
        tools, "_read_authenticated_parquet_snapshot", lambda *a, **kw: calls.append(kw)
    )
    for field in ["asset_hash", "artifact_content_hash"]:
        bad = deepcopy(request)
        bad["source_ref"][field] = "a" * 64
        with pytest.raises(StrategyError):
            run_freeze_candidate_asset_stability(bad, fx["ctx"], fx["runtime"])
    bad = deepcopy(request)
    bad["sample_design_ref"]["expected_sample_design_content_hash"] = "b" * 64
    with pytest.raises(StrategyError):
        run_freeze_candidate_asset_stability(bad, fx["ctx"], fx["runtime"])
    assert calls == []
    assert not [
        r
        for r in fx["runtime"].task_artifacts.list_for_task(fx["task"].id)
        if r["kind"] == FREEZE_KIND
    ]


def test_freeze_and_measure_do_not_add_or_adopt_candidate(tmp_path):
    fx, request = _fixture(tmp_path)
    frozen = run_freeze_candidate_asset_stability(request, fx["ctx"], fx["runtime"])
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(
            executor.map(
                lambda _: run_measure_candidate_asset_stability(
                    {"frozen_ref": frozen["ref"]}, fx["ctx"], fx["runtime"]
                ),
                range(2),
            )
        )
    assert results[0] == results[1]
    assert (
        len(
            [
                r
                for r in fx["runtime"].task_artifacts.list_for_task(fx["task"].id)
                if r["kind"] == ARTIFACT_KIND
            ]
        )
        == 1
    )
    with fx["runtime"].task_artifacts.transaction() as conn:
        assert (
            conn.execute(
                "SELECT COUNT(*) FROM strategies WHERE task_id=?", (fx["task"].id,)
            ).fetchone()[0]
            == 0
        )


def _measure_selected(fx, source_ref, sample_ref):
    frozen = run_freeze_candidate_asset_stability(
        {
            "source_ref": source_ref,
            "sample_design_ref": sample_ref,
            "partitions": ["validation", "oot"],
        },
        fx["ctx"],
        fx["runtime"],
    )
    return run_measure_candidate_asset_stability(
        {"frozen_ref": frozen["ref"]}, fx["ctx"], fx["runtime"]
    )


@pytest.mark.parametrize("grouped", [False, True])
def test_interactive_frontier_and_group_are_independent_assets(tmp_path, grouped):
    from tests.test_strategy_interactive_tree_downstream import _frontier_pool

    fx = _frontier_pool(tmp_path, grouped=grouped)
    selected = fx["selection"]
    source_ref = {
        "artifact_id": selected["artifacts"][0]["artifact_id"],
        "artifact_content_hash": selected["artifacts"][0]["content_hash"],
        "asset_id": selected["semantic_tree_id"],
        "asset_hash": selected["tree_hash"],
    }
    output = _measure_selected(fx, source_ref, fx["sample_v2_ref"])
    assert len(output["evidence"]["baseline"]["buckets"]) == 2
    assert [p["partition"] for p in output["evidence"]["partitions"]] == [
        "validation",
        "oot",
    ]
    revision_ref = {
        "artifact_id": fx["revision_record"]["id"],
        "artifact_content_hash": fx["revision_record"]["content_hash"],
        "asset_id": fx["revision"]["semantic_tree_id"],
        "asset_hash": fx["revision"]["tree"]["tree_hash"],
    }
    complete = _measure_selected(fx, revision_ref, fx["sample_v2_ref"])
    assert (
        complete["evidence"]["freeze"]["source_kind"] == "complete_interactive_frontier"
    )
    assert (
        len(complete["evidence"]["baseline"]["buckets"])
        == len(fx["revision"]["fragments"]) + 1
    )


def test_complete_native_tree_keeps_every_leaf_as_fixed_bucket(tmp_path):
    from tests.test_strategy_automatic_tree_tool import (
        _runtime,
        _materialize_native_sample_design_ref,
        _inputs,
        _tool_context,
    )
    from marvis.packs.strategy import tools

    settings, runner, _, task, _, dataset, workspace, mapping = _runtime(
        tmp_path, with_split=True
    )
    ctx = _tool_context(settings, task)
    native = _materialize_native_sample_design_ref(
        settings, task, dataset, workspace, mapping, target_bad_value=1
    )
    tree = tools.tool_build_automatic_tree_candidate(
        _inputs(dataset, workspace, mapping, native), ctx
    )
    runtime = tools._runtime(ctx)
    artifact = next(
        a
        for a in tree["artifacts"]
        if a["kind"] == "strategy_automatic_tree_asset_json"
    )
    catalog = candidate_asset_stability_catalog(runtime, task_id=task.id)
    ref = {
        "artifact_id": artifact["artifact_id"],
        "artifact_content_hash": artifact["content_hash"],
        "asset_id": tree["summary"]["asset_id"],
        "asset_hash": tree["summary"]["asset_hash"],
    }
    from marvis.plugins.manifest import ToolRef

    native_freeze = runner.invoke(
        ToolRef("strategy", "freeze_candidate_asset_stability"),
        {
            "source_ref": ref,
            "sample_design_ref": catalog["samples"][0]["ref"],
            "partitions": ["validation", "oot"],
        },
        task_id=task.id,
    )
    assert native_freeze.ok, native_freeze.error
    native_measure = runner.invoke(
        ToolRef("strategy", "measure_candidate_asset_stability"),
        {"frozen_ref": native_freeze.output["ref"]},
        task_id=task.id,
    )
    assert native_measure.ok, native_measure.error
    output = _measure_selected(
        {"ctx": ctx, "runtime": runtime}, ref, catalog["samples"][0]["ref"]
    )
    assert output["evidence"]["freeze"]["source_kind"] == "complete_tree"
    assert len(output["evidence"]["baseline"]["buckets"]) == len(tree["leaf_index"]) + 1
    assert output["evidence"]["baseline"]["buckets"][-1]["count"] == 0
    assert native_measure.output == output


@pytest.mark.slow
def test_complete_native_scorecard_uses_authenticated_frozen_raw_pd_bins(tmp_path):
    from tests.test_strategy_scorecard_candidate_tools import (
        _real_native_scorecard,
        _build_inputs,
    )
    from marvis.packs.strategy.scorecard_candidate_tools import (
        run_build_scorecard_band_asset,
    )

    real = _real_native_scorecard(tmp_path)
    band = run_build_scorecard_band_asset(
        {
            **_build_inputs(real),
            "banding": {"method": "equal_frequency", "bin_count": 5},
        },
        real["fx"]["ctx"],
        real["runtime"],
    )
    fx = {**real["fx"], "runtime": real["runtime"]}
    ref = {
        "artifact_id": band["artifacts"][0]["artifact_id"],
        "artifact_content_hash": band["artifacts"][0]["content_hash"],
        "asset_id": band["asset_id"],
        "asset_hash": band["asset_hash"],
    }
    output = _measure_selected(fx, ref, real["fx"]["sample_ref"])
    evidence = output["evidence"]
    assert evidence["freeze"]["source_kind"] == "scorecard_bands"
    assert evidence["freeze"]["requirements"][0]["type"] == "model_score_vector.v1"
    assert evidence["baseline"]["performance"]["status"] == "available"
    assert (
        evidence["baseline"]["performance"]["auc"]
        == band["scorecard_band_asset"]["performance"]["auc"]
    )
    assert [b["count"] for b in evidence["baseline"]["buckets"][:-1]] == [
        b["count"] for b in band["scorecard_band_asset"]["bands"]
    ]


def test_complete_native_cross_matrix_keeps_all_cells_and_bad_zero_polarity(tmp_path):
    from tests.test_strategy_cross_matrix_cell_selection_tool import (
        _fixture as matrix_fixture,
    )

    fx = matrix_fixture(tmp_path, native=True)
    catalog = candidate_asset_stability_catalog(fx.runtime, task_id=fx.task.id)
    ref = {
        "artifact_id": fx.matrix_artifact["artifact_id"],
        "artifact_content_hash": fx.matrix_artifact["content_hash"],
        "asset_id": fx.matrix["asset_id"],
        "asset_hash": fx.matrix["asset_hash"],
    }
    output = _measure_selected(
        {"ctx": fx.ctx, "runtime": fx.runtime}, ref, catalog["samples"][0]["ref"]
    )
    evidence = output["evidence"]
    assert evidence["freeze"]["source_kind"] == "cross_matrix"
    assert len(evidence["baseline"]["buckets"]) == fx.matrix["matrix"]["cell_count"] + 1
    assert evidence["freeze"]["target"]["bad_value"] == 0
    assert evidence["baseline"]["bad_count"] == sum(
        c["effect"]["bad"] for c in fx.matrix["matrix"]["cells"]
    )


def test_voting_asset_replays_parents_without_current_pool_selection(tmp_path):
    from marvis.packs.strategy.pool_tools import run_add_candidate_to_pool
    from marvis.packs.strategy.voting_candidate_tools import run_build_voting_candidate
    from tests.test_strategy_pool_tools import _add_inputs
    from marvis.packs.strategy.pool import ABSENT_POOL_SNAPSHOT_HASH

    fx = _setup(tmp_path, native_sample=True)
    p = None
    for candidate in [fx["first"], fx["refine"](2)]:
        p = run_add_candidate_to_pool(
            _add_inputs(
                candidate,
                expected_revision=0 if p is None else p["revision"],
                expected_hash=ABSENT_POOL_SNAPSHOT_HASH
                if p is None
                else p["snapshot_hash"],
            ),
            fx["ctx"],
            fx["runtime"],
        )
    vote = run_build_voting_candidate(
        {
            "strategy_type": "approval",
            "expected_pool_revision": p["revision"],
            "expected_pool_snapshot_hash": p["snapshot_hash"],
            "selected_entry_ids": [e["entry_id"] for e in p["entries"]],
            "n": 1,
        },
        fx["ctx"],
        fx["runtime"],
    )
    catalog = candidate_asset_stability_catalog(fx["runtime"], task_id=fx["task"].id)
    ref = {
        "artifact_id": vote["artifacts"][0]["artifact_id"],
        "artifact_content_hash": vote["artifacts"][0]["content_hash"],
        "asset_id": vote["asset_id"],
        "asset_hash": vote["asset_hash"],
    }
    output = _measure_selected(fx, ref, catalog["samples"][0]["ref"])
    assert "voting" in output["evidence"]["freeze"]["source_kind"]
    assert output["evidence"]["freeze"]["source_ref"] == ref


def test_report_publishes_same_exact_asset_metrics_in_four_formats(tmp_path):
    from tests.test_strategy_report_bundle_tools import (
        _setup_pool_validation_report,
        _run,
    )
    from marvis.packs.strategy.asset_stability_tools import _load_source

    fx = _setup_pool_validation_report(tmp_path)
    rows = fx["runtime"].task_artifacts.list_for_task(fx["task"].id)
    source = next(r for r in rows if r["kind"] == "strategy_candidate_asset_json")
    ref = {
        "artifact_id": source["id"],
        "artifact_content_hash": source["content_hash"],
        "asset_id": source["provenance"]["asset_id"],
        "asset_hash": source["provenance"]["asset_hash"],
    }
    assert _load_source(fx["runtime"], task_id=fx["task"].id, ref=ref)
    measured = _measure_selected(fx, ref, fx["sample_ref"])
    fx["request"]["asset_stability_refs"] = [measured["ref"]]
    output = _run(fx)
    bundle = output["bundle"]
    section = next(s for s in bundle["sections"] if s["key"] == "univariate_and_models")
    table = next(t for t in section["tables"] if t["table_id"] == "asset_stability_0")
    assert len(table["rows"]) == len(measured["evidence"]["baseline"]["buckets"]) * 3
    assert any(
        r["ref_id"] == measured["ref"]["artifact_id"]
        for r in bundle["strategy_artifact_refs"]
    )
    artifacts = output["artifacts"]
    assert {a["format"] for a in artifacts} == {"json", "markdown", "xlsx", "docx"}
    from openpyxl import load_workbook
    from docx import Document

    records = fx["runtime"].task_artifacts.list_for_task(fx["task"].id)
    paths = {
        a["format"]: Path(
            next(r["path"] for r in records if r["id"] == a["artifact_id"])
        )
        for a in artifacts
    }
    assert ref["asset_id"] in paths["markdown"].read_text()
    assert measured["ref"]["artifact_id"] in paths["json"].read_text()
    workbook = load_workbook(paths["xlsx"], read_only=True, data_only=True)
    assert "appendix_asset_stab_0" in workbook.sheetnames
    document = Document(paths["docx"])
    assert any(ref["asset_id"] in p.text for p in document.paragraphs) or any(
        ref["asset_id"] in c.text
        for t in document.tables
        for row in t.rows
        for c in row.cells
    )


def test_manual_api_and_agent_commands_share_exact_native_contract(tmp_path):
    from fastapi import FastAPI
    from starlette.testclient import TestClient
    from types import SimpleNamespace
    from marvis.agent.asset_stability_request import handle_asset_stability_command
    from marvis.repositories.tasks import TaskRepository
    from marvis.routers.strategy_asset_stability import router

    fx, request = _fixture(tmp_path)
    app = FastAPI()
    app.state.settings = fx["settings"]
    app.include_router(router)
    with TestClient(app) as client:
        base = f"/api/tasks/{fx['task'].id}/strategy-asset-stability"
        catalog = client.get(base)
        assert catalog.status_code == 200
        frozen = client.post(base + "/freeze", json=request)
        assert frozen.status_code == 200
        result = client.post(
            base + "/measure", json={"frozen_ref": frozen.json()["ref"]}
        )
        assert result.status_code == 200
        assert result.json()["evidence"]["partitions"][1]["partition"] == "oot"
        wrong = deepcopy(request)
        wrong["source_ref"]["asset_hash"] = "a" * 64
        assert client.post(base + "/freeze", json=wrong).status_code == 422
    runtime = SimpleNamespace(
        settings=fx["settings"],
        plan_repo=SimpleNamespace(list_plans_for_task=lambda _: []),
    )
    repo = TaskRepository(fx["settings"].db_path)
    command = f"冻结候选资产稳定性 {request['source_ref']['artifact_id']} {request['sample_design_ref']['bundle_artifact_id']} validation,oot"
    frozen_response = handle_asset_stability_command(runtime, repo, fx["task"], command)
    assert frozen_response["status"] == "done"
    command = (
        f"测量候选资产稳定性 {frozen_response['asset_stability']['ref']['artifact_id']}"
    )
    measured = handle_asset_stability_command(runtime, repo, fx["task"], command)
    assert measured["status"] == "done"
    assert measured["asset_stability"] == result.json()
    assert (
        handle_asset_stability_command(
            runtime, repo, fx["task"], "不要测量候选资产稳定性"
        )
        is None
    )


def test_report_proposal_and_publication_keep_explicit_asset_ref(tmp_path):
    from tests.test_strategy_report_bundle_tools import _setup_pool_validation_report
    from marvis.packs.strategy.asset_stability_report_service import (
        prepare_asset_stability_report,
        publish_asset_stability_report,
    )

    fx = _setup_pool_validation_report(tmp_path)
    source = next(
        r
        for r in fx["runtime"].task_artifacts.list_for_task(fx["task"].id)
        if r["kind"] == "strategy_candidate_asset_json"
    )
    ref = {
        "artifact_id": source["id"],
        "artifact_content_hash": source["content_hash"],
        "asset_id": source["provenance"]["asset_id"],
        "asset_hash": source["provenance"]["asset_hash"],
    }
    measured = _measure_selected(fx, ref, fx["sample_ref"])
    prepared = prepare_asset_stability_report(
        {
            "title": "明确资产稳定性报告",
            "status": "partial",
            "asset_stability_refs": [measured["ref"]],
        },
        task=fx["task"],
        runtime=fx["runtime"],
    )
    assert prepared["proposal"]["asset_stability_refs"] == [measured["ref"]]
    output = publish_asset_stability_report(prepared, fx["ctx"], fx["runtime"])
    assert len(output["artifacts"]) == 4
    changed = deepcopy(prepared)
    changed["proposal"]["asset_stability_refs"] = []
    with pytest.raises(StrategyError, match="proposal changed"):
        publish_asset_stability_report(changed, fx["ctx"], fx["runtime"])


def test_stale_or_tampered_freeze_fails_before_snapshot_and_rejected_publish_rolls_back(
    tmp_path, monkeypatch
):
    fx, request = _fixture(tmp_path)
    frozen = run_freeze_candidate_asset_stability(request, fx["ctx"], fx["runtime"])
    import marvis.packs.strategy.asset_stability_tools as tools

    original = tools._read_authenticated_parquet_snapshot
    calls = []

    def read_then_drift(*args, **kwargs):
        frame = original(*args, **kwargs)
        calls.append(1)
        record = fx["runtime"].task_artifacts.get_for_task(
            fx["task"].id, frozen["ref"]["artifact_id"]
        )
        Path(record["path"]).write_bytes(b"tampered frozen artifact")
        return frame

    # Deliberately drift the freeze during private snapshot computation. No result
    # artifact/file may survive the writer-lock CAS failure.
    monkeypatch.setattr(tools, "_read_authenticated_parquet_snapshot", read_then_drift)
    with pytest.raises(StrategyError):
        run_measure_candidate_asset_stability(
            {"frozen_ref": frozen["ref"]}, fx["ctx"], fx["runtime"]
        )
    assert calls == [1]
    assert not [
        r
        for r in fx["runtime"].task_artifacts.list_for_task(fx["task"].id)
        if r["kind"] == ARTIFACT_KIND
    ]
    assert not list(
        (
            Path(fx["settings"].tasks_dir)
            / fx["task"].id
            / "outputs"
            / "asset_stability"
        ).glob("asset-stability-*.json")
    )


@pytest.mark.parametrize("defect", ["unknown_maturity", "unknown_time", "cell_budget"])
def test_unknown_governance_or_over_budget_refuses_before_snapshot(
    tmp_path, monkeypatch, defect
):
    from dataclasses import replace
    import marvis.packs.strategy.asset_stability_tools as tools

    fx, request = _fixture(tmp_path)
    real = tools.load_requested_strategy_sample_design_v2_artifacts(
        fx["runtime"], fx["task"].id, request
    )
    bundle = deepcopy(real.bundle)
    source = real.source_binding
    if defect == "unknown_maturity":
        next(p for p in bundle["populations"] if p["role"] == "risk")[
            "maturity_evidence"
        ]["status"] = "unknown"
    elif defect == "unknown_time":
        bundle["sample_design"]["sample_semantics"]["field_bindings"]["month_field"] = (
            None
        )
        bundle["sample_design"]["sample_semantics"]["field_bindings"]["time_field"] = (
            None
        )
    else:
        monkeypatch.setattr(tools, "MAX_EVALUATION_CELLS", 1)
    admitted = replace(real, bundle=bundle, source_binding=source)
    monkeypatch.setattr(
        tools,
        "load_requested_strategy_sample_design_v2_artifacts",
        lambda *a, **k: admitted,
    )
    calls = []
    monkeypatch.setattr(
        tools, "_read_authenticated_parquet_snapshot", lambda *a, **k: calls.append(1)
    )
    with pytest.raises(StrategyError):
        run_freeze_candidate_asset_stability(request, fx["ctx"], fx["runtime"])
    assert calls == []
    assert not [
        r
        for r in fx["runtime"].task_artifacts.list_for_task(fx["task"].id)
        if r["kind"] in {FREEZE_KIND, ARTIFACT_KIND}
    ]
