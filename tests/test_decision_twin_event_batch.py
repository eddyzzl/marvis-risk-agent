"""Native event evidence travels through the real reviewed historical batch worker."""

from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
import json
import sqlite3
from types import SimpleNamespace

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pandas as pd
import pytest

from marvis.app import create_app
from marvis.db_schema import connect
from marvis.decision_twin.batch import replay_batch
from marvis.decision_twin.batch_contracts import HistoricalReplayRequest
from marvis.decision_twin.batch_material import BatchMaterial, REPLAY_KIND
from marvis.domain import TaskCreate, TASK_TYPE_MODELING
from marvis.plugins.manifest import ToolRef
from marvis.reference_decision.contracts import DecisionError
from marvis.repositories.tasks import TaskRepository
from test_reference_event_binding import build, event_runtime as runtime_fixture


@pytest.fixture(name="event_runtime")
def native_runtime(tmp_path):
    return runtime_fixture.__wrapped__(tmp_path)


def batch(rt, *, complete=True):
    _, reference, package, evidence = build(rt, complete=complete)
    task = TaskRepository(rt.settings.db_path).create_task(
        TaskCreate(
            model_name="原生事件历史申请回放",
            model_version="test",
            validator="qa",
            source_dir=str(rt.root),
            run_mode="manual",
            task_type=TASK_TYPE_MODELING,
        )
    )
    material = BatchMaterial(rt.settings, task.id, actor_id=rt.principal["id"])
    frame = pd.DataFrame(
        {
            "id": ["application-1", "application-2"],
            "decision_at": reference.contract.decision_at,
            "subject_namespace": reference.contract.focus.namespace,
            "subject_token": reference.contract.focus.token,
            "event_ref": json.dumps(reference.model_dump(exclude={"grant_id"})),
            "actual_action": "approval",
            "action_at": reference.contract.decision_at,
        }
    )
    ctx = SimpleNamespace(
        rt=rt,
        material=material,
        reference=reference,
        package=package,
        evidence=evidence,
        frame=frame,
        serial=0,
    )
    contract = {
        "dataset_id": "pending",
        "expected_content_hash": "0" * 64,
        "record_id_col": "id",
        "decision_at_col": "decision_at",
        "as_of": datetime.now(UTC).isoformat(),
        "source_ref": "synthetic native event source",
        "population": "all applications",
        "features": [],
        "scenarios": [
            {"name": kind, "kind": kind, "package_hash": package["package_hash"]}
            for kind in ("baseline", "challenger")
        ],
        "event_mappings": [
            {
                "scenario_kind": kind,
                "reference_col": "event_ref",
                "subject_namespace_col": "subject_namespace",
                "subject_token_col": "subject_token",
                "grant_id": reference.grant_id,
            }
            for kind in ("baseline", "challenger")
        ],
        "observed_actions": {
            "action_col": "actual_action",
            "recorded_at_col": "action_at",
            "source_ref": "external ledger",
        },
    }
    ctx.contract = HistoricalReplayRequest.model_validate(contract)
    ctx.contract = register(ctx, frame)
    return ctx


def register(ctx, frame):
    path = ctx.rt.root / f"event-history-{ctx.serial}.parquet"
    ctx.serial += 1
    frame.to_parquet(path, index=False)
    record = ctx.material.registry.register_existing(
        path, task_id=ctx.material.task_id, role="historical_replay"
    )
    return ctx.contract.model_copy(
        update={"dataset_id": record.id, "expected_content_hash": record.content_hash}
    )


def proposal(ctx, *, client=None, contract=None):
    response = (client or ctx.rt.maker).post(
        f"/api/tasks/{ctx.material.task_id}/decision-twin/proposal",
        json=(contract or ctx.contract).model_dump(),
    )
    assert response.status_code == 200, response.text
    return response.json()


def run(ctx, proposed, *, goal="历史决策回放", slot="replay_contract"):
    client = ctx.rt.maker
    response = client.post(
        f"/api/tasks/{ctx.material.task_id}/plans",
        json={
            "goal": goal,
            "slots": {
                slot: proposed["contract"],
                "proposal_hash": proposed["proposal_hash"],
            },
        },
    )
    assert response.status_code == 201, response.text
    plan = response.json()["plan"]
    response = client.post(
        f"/api/plans/{plan['id']}/confirm", json=plan["confirmation_snapshot"]
    )
    assert response.status_code == 200, response.text
    assert client.post(f"/api/plans/{plan['id']}/run").status_code == 202
    current = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert current["status"] == "awaiting_confirm", current
    step = current["steps"][0]
    response = client.post(
        f"/api/plans/{plan['id']}/steps/{step['id']}/decisions",
        json={
            "decision": "approve",
            "reason": "确认每笔原生事件引用及完整人口，未知不补零",
            **step["confirmation_snapshot"],
        },
    )
    assert response.status_code == 202, response.text
    return client.get(f"/api/plans/{plan['id']}").json()["plan"]


def receipts(ctx):
    return [
        r
        for r in ctx.material.artifacts.list_for_task(ctx.material.task_id)
        if r["kind"] == REPLAY_KIND
    ]


def replay(ctx):
    prepared = proposal(ctx)
    result = replay_batch(ctx.material, ctx.contract, prepared["proposal_hash"])
    return result


def test_real_native_event_http_human_toolrunner_and_all_export_carriers(event_runtime):
    ctx = batch(event_runtime)
    prepared = proposal(ctx)
    assert prepared["event_authorization"]["actor_id"] == ctx.rt.principal["id"]
    assert ctx.reference.contract.focus.token not in json.dumps(prepared)
    result = run(ctx, prepared)
    assert result["status"] == "done", result
    saved = receipts(ctx)
    assert len(saved) == 1
    url = f"/api/tasks/{ctx.material.task_id}/decision-twin/{saved[0]['content_hash']}"
    detail = ctx.rt.maker.get(url)
    assert detail.status_code == 200, detail.text
    detail = detail.json()
    assert detail == ctx.rt.maker.get(url + "/export/json").json()
    for scenario in detail["payload"]["scenarios"]:
        assert scenario["metrics"]["count"] == 2
        assert scenario["metrics"]["event_evidence"]["measured_rows"] == 2
        for row in scenario["decisions"]:
            assert row["action"]["reason_code"] == "EVENT_REVIEW"
            assert row["score"] is None
            assert row["event_evidence"]["content_hash"] == ctx.evidence["content_hash"]
            assert (
                row["event_evidence"]["snapshot_hash"]
                == ctx.evidence["receipt"]["result"]["snapshot_hash"]
            )
    assert ctx.reference.contract.focus.token not in json.dumps(detail)
    xlsx = ctx.rt.maker.get(url + "/export/xlsx")
    assert xlsx.status_code == 200, xlsx.text
    workbook = load_workbook(BytesIO(xlsx.content))
    rows = list(workbook["原生事件证据"].values)
    assert len(rows) == 5
    assert rows[1][2] == "measured" and rows[1][6] == ctx.evidence["content_hash"]
    docx = ctx.rt.maker.get(url + "/export/docx")
    assert docx.status_code == 200, docx.text
    (ctx.rt.root / "verified-event-batch.json").write_text(
        json.dumps(detail, ensure_ascii=False, indent=2)
    )
    (ctx.rt.root / "verified-event-batch.xlsx").write_bytes(xlsx.content)
    (ctx.rt.root / "verified-event-batch.docx").write_bytes(docx.content)
    text = "\n".join(p.text for p in Document(BytesIO(docx.content)).paragraphs)
    assert "逐笔原生事件收据引用" in text and ctx.evidence["content_hash"] in text
    fresh = create_app(ctx.rt.settings)
    authenticated = BatchMaterial(
        fresh.state.settings, ctx.material.task_id, actor_id=ctx.rt.principal["id"]
    )
    assert authenticated.load(detail["artifact_id"]) == detail
    # Same reviewed input and source snapshots retain exact deterministic identity.
    assert (
        replay_batch(authenticated, ctx.contract, prepared["proposal_hash"])[
            "artifact_id"
        ]
        == detail["artifact_id"]
    )


def assert_private(ctx, record, clients):
    relative = Path(record["path"]).relative_to(ctx.rt.settings.workspace).as_posix()
    ordinary = f"/api/tasks/{ctx.material.task_id}/task-artifacts/{record['id']}/download?expected_content_hash={record['content_hash']}"
    for client in clients:
        for route in (
            ordinary,
            "/api/artifacts/" + relative,
            "/api/artifacts/" + relative + "/preview",
        ):
            response = client.get(route)
            assert response.status_code == (403 if route == ordinary else 404), (
                route,
                response.status_code,
                response.text[:100],
            )
            assert ctx.evidence["content_hash"] not in response.text


def test_current_grant_guards_detail_exports_and_both_generic_downloads(event_runtime):
    ctx = batch(event_runtime)
    result = replay(ctx)
    record = receipts(ctx)[0]
    url = f"/api/tasks/{ctx.material.task_id}/decision-twin/{result['artifact_id']}"
    anonymous = TestClient(ctx.rt.app)
    assert_private(ctx, record, [anonymous, ctx.rt.other, ctx.rt.maker])
    for client in (anonymous, ctx.rt.other):
        for suffix in ("", "/export/json", "/export/xlsx", "/export/docx"):
            assert client.get(url + suffix).status_code == 403
    revoke = ctx.rt.admin.post(
        f"/api/tasks/{ctx.rt.task.id}/risk-events/grants/{ctx.reference.grant_id}/revoke"
    )
    assert revoke.status_code == 200, revoke.text
    for suffix in ("", "/export/json", "/export/xlsx", "/export/docx"):
        assert ctx.rt.maker.get(url + suffix).status_code == 403
    assert_private(ctx, record, [ctx.rt.maker])
    listing = ctx.rt.maker.get(
        f"/api/tasks/{ctx.material.task_id}/decision-twin"
    ).json()
    assert listing["artifacts"][0]["status"] == "unauthorized"


def test_expired_current_grant_cannot_read_any_carrier(event_runtime, monkeypatch):
    import marvis.risk_context.event_repository as module

    ctx = batch(event_runtime)
    result = replay(ctx)
    monkeypatch.setattr(
        module, "_now", lambda: (datetime.now(UTC) + timedelta(hours=2)).isoformat()
    )
    url = f"/api/tasks/{ctx.material.task_id}/decision-twin/{result['artifact_id']}"
    for suffix in ("", "/export/json", "/export/xlsx", "/export/docx"):
        assert ctx.rt.maker.get(url + suffix).status_code == 403
    assert_private(ctx, receipts(ctx)[0], [ctx.rt.maker])


@pytest.mark.parametrize(
    "change,code",
    [
        ({"event_ref": None}, "native_event_reference_required"),
        ({"event_ref": "{}"}, "native_event_reference_invalid"),
        ({"subject_token": "b" * 64}, "subject_mismatch"),
        ({"subject_namespace": "other"}, "subject_mismatch"),
        ({"decision_at": "2020-01-01T00:00:00Z"}, "time_mismatch"),
    ],
)
def test_each_row_requires_native_reference_matching_independent_subject_and_time(
    event_runtime, change, code
):
    ctx = batch(event_runtime)
    frame = ctx.frame.copy()
    for key, value in change.items():
        frame.loc[1, key] = value
    contract = register(ctx, frame)
    response = ctx.rt.maker.post(
        f"/api/tasks/{ctx.material.task_id}/decision-twin/proposal",
        json=contract.model_dump(),
    )
    assert response.status_code == 422, response.text
    assert code in response.text
    assert receipts(ctx) == []


def test_no_manual_event_features_mapping_omission_or_foreign_actor(event_runtime):
    ctx = batch(event_runtime)
    base = f"/api/tasks/{ctx.material.task_id}/decision-twin/proposal"
    for client in (TestClient(ctx.rt.app), ctx.rt.other):
        response = client.post(base, json=ctx.contract.model_dump())
        assert response.status_code == 403, response.text
    bad = ctx.contract.model_dump()
    bad["event_mappings"] = []
    response = ctx.rt.maker.post(base, json=bad)
    assert response.status_code == 422 and "mapping_required" in response.text
    bad = ctx.contract.model_dump()
    bad["actor_id"] = ctx.rt.principal["id"]
    assert ctx.rt.maker.post(base, json=bad).status_code == 422
    frame = ctx.frame.copy()
    frame["transactions"] = 123
    contract = register(ctx, frame).model_dump()
    contract["features"] = [
        {
            "name": "transactions",
            "value_col": "transactions",
            "event_at_col": "decision_at",
            "available_at_col": "decision_at",
        }
    ]
    response = ctx.rt.maker.post(base, json=contract)
    assert response.status_code == 422 and "raw_schema_mismatch" in response.text


def test_unknown_event_coverage_falls_back_without_reducing_denominator(event_runtime):
    ctx = batch(event_runtime, complete=False)
    at = datetime.fromisoformat(ctx.reference.contract.decision_at)
    payload = ctx.contract.model_dump()
    payload["temporal_stability"] = {
        "timezone": "UTC",
        "reference_window": {
            "name": "before",
            "start": (at - timedelta(minutes=2)).isoformat(),
            "end": (at - timedelta(minutes=1)).isoformat(),
        },
        "comparison_windows": [
            {
                "name": "current",
                "start": (at - timedelta(minutes=1)).isoformat(),
                "end": (at + timedelta(seconds=1)).isoformat(),
            }
        ],
        "minimum_reference_rows": 2,
        "minimum_comparison_rows": 2,
        "bin_count": 2,
        "thresholds": {
            "max_score_psi": 0.1,
            "max_action_psi": 0.1,
            "max_absolute_approval_rate_delta": 0.1,
        },
        "policy_source_ref": "synthetic policy",
    }
    payload["constraints"] = [
        {"metric": "approval_rate", "operator": ">=", "threshold": 0, "unit": "rate"}
    ]
    ctx.contract = HistoricalReplayRequest.model_validate(payload)
    current = run(ctx, proposal(ctx))
    assert current["status"] == "done", current
    result = ctx.material.load(receipts(ctx)[0]["content_hash"])
    for scenario in result["payload"]["scenarios"]:
        metrics = scenario["metrics"]
        assert metrics["count"] == metrics["review_count"] == 2
        assert metrics["event_evidence"]["unknown_rows"] == 2
        assert metrics["stability"]["reason"] == "event_evidence_incomplete"
        assert (
            metrics["stability"]["score_missing_cause"]
            == "authenticated_event_features_unknown"
        )
        assert scenario["constraints"]["status"] == "insufficient_evidence"
        assert all(
            row["error_code"] == "event_features_unknown" and row["score"] is None
            for row in scenario["decisions"]
        )


@pytest.mark.parametrize(
    "failure", ["content_hash", "contract", "revocation", "native_file"]
)
def test_integrity_authority_or_context_failure_aborts_whole_real_worker_batch(
    event_runtime, failure
):
    ctx = batch(event_runtime)
    if failure in {"content_hash", "contract"}:
        native = ctx.reference.model_dump(exclude={"grant_id"})
        if failure == "content_hash":
            native["expected_content_hash"] = "0" * 64
        else:
            # Same subject/decision and recipe, different dynamic context.
            native["contract"]["knowledge_cutoff"] = native["contract"]["decision_at"]
        frame = ctx.frame.copy()
        frame.loc[1, "event_ref"] = json.dumps(native)
        ctx.contract = register(ctx, frame)
    prepared = proposal(ctx)
    if failure == "revocation":
        response = ctx.rt.admin.post(
            f"/api/tasks/{ctx.rt.task.id}/risk-events/grants/{ctx.reference.grant_id}/revoke"
        )
        assert response.status_code == 200
    if failure == "native_file":
        record = ctx.material.artifacts.get_for_task(
            ctx.rt.task.id, ctx.evidence["artifact_id"]
        )
        Path(record["path"]).write_text("{}")
    result = run(ctx, prepared)
    assert result["status"] != "done", result
    assert receipts(ctx) == []


def test_signed_actor_intent_is_immutable_and_toolrunner_cannot_skip_human(
    event_runtime,
):
    ctx = batch(event_runtime)
    prepared = proposal(ctx)
    result = ctx.rt.app.state.tool_runner.invoke(
        ToolRef("decision_twin", "replay_history"),
        {
            "contract": ctx.contract.model_dump(),
            "proposal_hash": prepared["proposal_hash"],
        },
        task_id=ctx.material.task_id,
    )
    assert not result.ok and "governance" in result.error
    with connect(ctx.rt.settings.db_path) as conn:
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            conn.execute(
                "UPDATE historical_event_intents SET body='{}' WHERE proposal_hash=?",
                (prepared["proposal_hash"],),
            )
        with pytest.raises(sqlite3.IntegrityError, match="retained"):
            conn.execute(
                "DELETE FROM historical_event_intents WHERE proposal_hash=?",
                (prepared["proposal_hash"],),
            )
    fresh = BatchMaterial(ctx.rt.settings, ctx.material.task_id)
    with pytest.raises(DecisionError, match="reviewed_proposal_required"):
        replay_batch(fresh, ctx.contract, "0" * 64)
    assert receipts(ctx) == []


def test_event_reconciliation_uses_signed_server_authority_and_same_worker(
    event_runtime,
):
    rt = event_runtime
    rt.anchor -= timedelta(days=120)
    ctx = batch(rt)
    result = replay(ctx)
    path = rt.root / "event-outcomes.parquet"
    pd.DataFrame(
        {
            "id": ctx.frame["id"],
            "observed_at": datetime.now(UTC).isoformat(),
            "loss": [5.0, 7.0],
            "profit": [20.0, 30.0],
        }
    ).to_parquet(path, index=False)
    record = ctx.material.registry.register_existing(
        path, task_id=ctx.material.task_id, role="historical_outcomes"
    )
    contract = {
        "replay_artifact_id": result["artifact_id"],
        "dataset_id": record.id,
        "expected_content_hash": record.content_hash,
        "record_id_col": "id",
        "observed_at_col": "observed_at",
        "actual_loss_col": "loss",
        "actual_profit_col": "profit",
        "currency": "CNY",
        "maturity_days": 90,
        "maturity_source_ref": "90 day observed horizon",
        "source_ref": "external cashflow history",
        "reconciled_at": datetime.now(UTC).isoformat(),
    }
    route = f"/api/tasks/{ctx.material.task_id}/decision-twin/reconciliation-proposal"
    assert rt.other.post(route, json=contract).status_code == 403
    prepared = rt.maker.post(route, json=contract)
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()["event_authorization"]["actor_id"] == rt.principal["id"]
    plan = run(
        ctx, prepared.json(), goal="历史决策现金流对账", slot="reconciliation_contract"
    )
    assert plan["status"] == "done", plan
    saved = [
        r
        for r in ctx.material.artifacts.list_for_task(ctx.material.task_id)
        if r["kind"] == "decision_twin_batch_reconciliation"
    ]
    assert len(saved) == 1
    url = f"/api/tasks/{ctx.material.task_id}/decision-twin/{saved[0]['content_hash']}"
    result = rt.maker.get(url).json()
    assert result["payload"]["actual_loss"] == 12
    assert result["payload"]["actual_profit"] == 50
    assert result["payload"]["marvis_execution_verified"] is False
    assert result["payload"]["causal_gain_verified"] is False
    assert_private(ctx, saved[0], [rt.other, rt.maker, TestClient(rt.app)])
    for suffix in ("", "/export/json", "/export/xlsx", "/export/docx"):
        assert rt.maker.get(url + suffix).status_code == 200
        assert rt.other.get(url + suffix).status_code == 403
    assert (
        rt.admin.post(
            f"/api/tasks/{rt.task.id}/risk-events/grants/{ctx.reference.grant_id}/revoke"
        ).status_code
        == 200
    )
    assert rt.maker.get(url + "/export/json").status_code == 403


def test_each_later_row_recipe_task_and_cutoff_is_validated(event_runtime):
    ctx = batch(event_runtime)
    for key in ("recipe", "task", "cutoff", "embedded_grant"):
        native = ctx.reference.model_dump(exclude={"grant_id"})
        if key == "recipe":
            native["contract"]["window_seconds"] += 1
        elif key == "task":
            native["task_id"] = ctx.material.task_id
        elif key == "embedded_grant":
            native["grant_id"] = ctx.reference.grant_id
        else:
            native["contract"]["knowledge_cutoff"] = (
                datetime.now(UTC) + timedelta(hours=1)
            ).isoformat()
        frame = ctx.frame.copy()
        frame.loc[1, "event_ref"] = json.dumps(native)
        response = ctx.rt.maker.post(
            f"/api/tasks/{ctx.material.task_id}/decision-twin/proposal",
            json=register(ctx, frame).model_dump(),
        )
        assert response.status_code in {409, 422}, response.text
    assert receipts(ctx) == []


def test_native_reference_contract_is_discoverable(event_runtime):
    response = event_runtime.maker.get("/api/decision-twin/capabilities")
    assert response.status_code == 200
    schema = response.json()["native_event_reference_schema"]
    assert set(schema["required"]) == {
        "task_id",
        "request_id",
        "expected_content_hash",
        "contract",
    }
    assert schema["additionalProperties"] is False
    assert (
        response.json()["native_event_policy"]["derived_event_values"] == "not_accepted"
    )
