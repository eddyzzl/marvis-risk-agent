from dataclasses import replace
from io import BytesIO
import json
import sqlite3

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from marvis.agent.plan_message_composer import PlanMessageComposer
from marvis.api_schemas import StrategyTaskInputRequest
from marvis.db import TaskRepository, connect
from marvis.domain import StrategyTaskInput
from marvis.orchestrator.business_acceptance import (
    adopted_evidence,
    stored_business_review,
)
from marvis.orchestrator.contracts import PlanStatus
from marvis.orchestrator.reviewer import FinalReview, Reviewer
from marvis.orchestrator.templates import load_builtin_templates
from marvis.routers.plans import router
from tests.test_business_acceptance import objective
from tests.test_economic_assumptions import contract
from tests.test_db import _task_create
from tests.test_orch_api import _client
from tests.test_orch_executor import (
    _executor,
    _ok,
    _plan,
    _repo,
    _step,
    FakeRunner,
    FakeLLM,
)


def execute_selection(tmp_path, *, criteria=None, reply=None):
    plan = _plan(
        _step("selected", plugin="modeling", tool="select_experiment"),
        success_criteria=criteria,
    )
    repo = _repo(tmp_path, plan)
    output = {
        "selected_experiment_id": "model-b",
        "artifact_id": "v2",
        "metrics": {"oot_ks": 0.2},
        "experiments": [{"experiment_id": "model-a", "metrics": {"oot_ks": 0.95}}],
        "measurement_context": {"effect_stage": "oot_validated", "labels_mature": True},
    }
    reviewer = Reviewer(
        lambda: FakeLLM(json.dumps(reply or {"goal_met": False, "goal_doubt": True})),
        plan_repository=repo,
    )
    runner = FakeRunner([_ok(output)])
    result = _executor(repo, runner, reviewer=reviewer).run(plan.id)
    return repo, repo.load_plan(plan.id), result, runner


def test_real_execution_done_with_missing_business_metadata_and_no_candidate_max(
    tmp_path,
):
    repo, plan, result, runner = execute_selection(
        tmp_path, criteria=[objective().to_dict()]
    )
    assert result.status == PlanStatus.DONE
    assert len(runner.calls) == 1
    evidence = adopted_evidence(plan, repo, "model")
    assert evidence.target_id == "model-b"
    assert evidence.metrics["oot_ks"] == 0.2
    assert (
        evidence.effect_stage is None
    )  # arbitrary output declarations cannot certify maturity
    assert result.final_review.execution_completed is True
    assert result.final_review.business_acceptance["status"] == "insufficient_evidence"
    assert result.final_review.goal_doubt is False


def test_tampered_immutable_result_cannot_be_used_for_business_acceptance(tmp_path):
    repo, plan, _, _ = execute_selection(tmp_path, criteria=[objective().to_dict()])
    with (
        pytest.raises(sqlite3.IntegrityError, match="immutable"),
        connect(repo.db_path) as conn,
    ):
        conn.execute(
            "UPDATE plan_step_output_versions SET output_json = ? WHERE step_id = ?",
            ('{"selected_experiment_id":"forged"}', plan.steps[0].id),
        )
    # Corrupt the mutable pointer instead: the authenticated loader must reject
    # the stale/fabricated version rather than recover a candidate by name.
    with connect(repo.db_path) as conn:
        conn.execute(
            "UPDATE plan_steps SET output_ref = ? WHERE id = ?",
            ("metrics:selected:v999", plan.steps[0].id),
        )
    plan = repo.load_plan(plan.id)
    assert adopted_evidence(plan, repo, "model") is None
    assert stored_business_review(repo, plan.id) == {}


def test_api_agent_and_document_share_same_persisted_verdict(tmp_path):
    repo, plan, _, _ = execute_selection(tmp_path, criteria=[objective().to_dict()])
    app = FastAPI()
    app.include_router(router)
    app.state.plan_repo = repo
    client = TestClient(app)
    stored = stored_business_review(repo, plan.id)
    payload = client.get(f"/api/plans/{plan.id}").json()["plan"]
    message = PlanMessageComposer(load_output=repo.load_step_output).done_message(
        plan, run_seq=1
    )
    assert (
        payload["business_acceptance"]
        == message.metadata["business_acceptance"]
        == stored["business_acceptance"]
    )
    assert "业务验收证据不足" in message.content
    xlsx = client.get(f"/api/plans/{plan.id}/business-acceptance/xlsx")
    assert xlsx.status_code == 200
    assert (
        client.get(
            f"/api/plans/{plan.id}/business-acceptance/xlsx",
            params={"expected_summary_ref": "stale"},
        ).status_code
        == 409
    )
    assert (
        client.get(
            f"/api/plans/{plan.id}/business-acceptance/xlsx",
            params={"expected_summary_ref": stored["summary_ref"]},
        ).status_code
        == 200
    )
    workbook = load_workbook(BytesIO(xlsx.content))
    assert workbook.active["B2"].value == "insufficient_evidence"
    assert (
        workbook.active["B4"].value == stored["business_acceptance"]["objective_hash"]
    )
    docx = client.get(f"/api/plans/{plan.id}/business-acceptance/docx")
    assert docx.status_code == 200
    document = Document(BytesIO(docx.content))
    assert document.tables[0].rows[1].cells[1].text == "insufficient_evidence"
    assert (
        document.tables[0].rows[3].cells[1].text
        == stored["business_acceptance"]["objective_hash"]
    )


def test_no_criteria_plan_completes_without_claiming_business_pass(tmp_path):
    _, _, result, _ = execute_selection(tmp_path)
    assert result.status == PlanStatus.DONE
    assert result.final_review.business_acceptance["status"] == "not_configured"


def test_http_explicit_objective_is_frozen_for_any_workflow(tmp_path):
    load_builtin_templates()
    client = _client(tmp_path)
    response = client.post(
        "/api/tasks/task-1/plans",
        json={"goal": "echo", "business_objective": objective().to_dict()},
    )
    assert response.status_code == 201, response.text
    assert response.json()["plan"]["success_criteria"] == [objective().to_dict()]


def test_http_rejects_forged_objective_fields_without_creating_plan(tmp_path):
    load_builtin_templates()
    client = _client(tmp_path)
    goal = {**objective().to_dict(), "success": True}
    response = client.post(
        "/api/tasks/task-1/plans", json={"goal": "echo", "business_objective": goal}
    )
    assert response.status_code == 422
    assert client.get("/api/tasks/task-1/plans").json()["plans"] == []


def test_task_objective_and_assumption_metadata_round_trip_without_migration(tmp_path):
    repo = _repo(tmp_path, _plan(_step("placeholder")))
    value = StrategyTaskInput(
        business_objective=objective(target_kind="strategy"), profit=contract()
    )
    task = TaskRepository(repo.db_path).create_task(
        _task_create(task_type="strategy", strategy_input=value)
    )
    assert TaskRepository(repo.db_path).get_task(task.id).strategy_input == value
    request = StrategyTaskInputRequest.model_validate(
        {"business_objective": value.business_objective.to_dict()}
    )
    assert request.business_objective == value.business_objective.to_dict()
    with pytest.raises(ValueError):
        StrategyTaskInputRequest.model_validate(
            {
                "business_objective": {
                    **value.business_objective.to_dict(),
                    "require_mature_labels": "false",
                }
            }
        )


def test_planner_cannot_replace_task_objective_with_generated_exemption():
    from marvis.orchestrator.planner import Planner

    goal = objective()
    planner = Planner(None, None, None, business_objective_loader=lambda task_id: goal)
    forged = replace(
        goal,
        applicable=False,
        allow_not_applicable=True,
        not_applicable_reason="LLM says skip",
    )
    assert planner._business_criteria("task", [forged.to_dict()]) == [goal.to_dict()]
    assert (
        Planner(None, None, None)._business_criteria("task", [forged.to_dict()]) == []
    )


def test_execution_can_finish_with_deterministic_business_failure(tmp_path):
    from marvis.business_acceptance import evaluate_business_acceptance
    from tests.test_business_acceptance import evidence

    measured = evaluate_business_acceptance(
        objective(), evidence(metrics={"oot_ks": 0.2})
    )
    assert measured["status"] == "failed"

    class CompletedReview(Reviewer):
        def final_review(self, plan, outputs, goal):
            return FinalReview(
                False,
                "business threshold missed",
                [],
                execution_completed=True,
                business_acceptance=measured,
            )

    plan = _plan(_step("work"))
    repo = _repo(tmp_path, plan)
    runner = FakeRunner([_ok({"ok": True})])
    result = _executor(repo, runner, reviewer=CompletedReview(lambda: FakeLLM())).run(
        plan.id
    )
    assert result.status == PlanStatus.DONE
    assert (
        repo.load_plan_summary(result.summary_ref)["business_acceptance"]["status"]
        == "failed"
    )


def test_final_narrative_receives_bounded_authoritative_adopted_verdict(tmp_path):
    repo, plan, _, _ = execute_selection(tmp_path, criteria=[objective().to_dict()])
    llm = FakeLLM(
        json.dumps({"summary": "LLM incorrectly claims success", "goal_met": True})
    )
    review = Reviewer(lambda: llm, plan_repository=repo).final_review(
        plan, {"candidate": {"oot_ks": 0.99}}, plan.goal
    )
    prompt = json.loads(llm.calls[-1]["user_prompt"])
    bound = prompt["business_acceptance"]
    assert bound["objective_hash"] == review.business_acceptance["objective_hash"]
    assert bound["target"] == review.business_acceptance["target"]
    assert bound["target"]["id"] == "model-b"
    assert (
        bound["status"]
        == review.business_acceptance["status"]
        == "insufficient_evidence"
    )
    assert "objective" not in bound
    assert review.goal_met is False
    assert review.execution_completed is True
    assert "不得用候选指标" in prompt["acceptance_authority"]


def test_real_profit_producer_consumes_persisted_assumptions_and_refuses_mismatch(
    tmp_path,
):
    from marvis.packs.strategy.tools import tool_profit_calc
    from marvis.settings import build_settings
    from tests.test_strategy_pack import _runtime, _register_strategy_sample
    from types import SimpleNamespace

    _runner, _, registry, task = _runtime(tmp_path)
    settings = build_settings(tmp_path / "workspace")
    # The test fixture uses a validation task by default; create the actual
    # strategy task contract via the public repository boundary.
    task = TaskRepository(settings.db_path).create_task(
        _task_create(
            task_type="strategy", strategy_input=StrategyTaskInput(profit=contract())
        )
    )
    dataset = _register_strategy_sample(registry, tmp_path, task.id)
    ctx = SimpleNamespace(
        workspace=settings.workspace,
        datasets_root=settings.datasets_dir,
        task_id=task.id,
    )
    params = {
        name: getattr(contract(), name)
        for name in (
            "annual_rate",
            "funding_rate",
            "lgd",
            "operating_cost_per_loan",
            "term_months",
        )
    }
    inputs = {
        "dataset_id": dataset.id,
        "ead_col": "ead",
        "pd_col": "pd",
        "params": params,
    }
    result = tool_profit_calc(inputs, ctx)
    assert result["assumptions"]["assessment"]["status"] == "estimated"
    assert result["assumptions"]["assessment"]["assumptions"]["currency"] == "CNY"
    mismatched = tool_profit_calc(
        {**inputs, "params": {**params, "annual_rate": 0.2}}, ctx
    )
    assert mismatched["assumptions"]["assessment"]["economics"] is None
    assert mismatched["assumptions"]["assessment"]["missing"] == ["profit_contract"]


def test_legacy_stricter_threshold_cannot_be_hidden_by_another_contract(tmp_path):
    repo, plan, result, _ = execute_selection(
        tmp_path,
        criteria=[
            objective().to_dict(),
            {"metric": "oot_ks", "min": 0.9, "aggregate": "max"},
        ],
    )
    assert result.final_review.business_acceptance["status"] == "insufficient_evidence"
    assert "旧阈值" in " ".join(result.final_review.business_acceptance["reasons"])
    assert len(plan.success_criteria) == 2
