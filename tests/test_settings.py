from pathlib import Path

from docx import Document

from marvis.settings import build_settings


def _write_template(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path


def test_fresh_workspace_uses_packaged_default_report_template(tmp_path: Path) -> None:
    settings = build_settings(tmp_path / "workspace")

    template_path = settings.report_template_path

    assert template_path == Path(__file__).parents[1] / "marvis/report_templates/default.docx"
    assert template_path.is_file()
    assert Document(template_path).paragraphs


def test_workspace_report_templates_override_packaged_default(tmp_path: Path) -> None:
    settings = build_settings(tmp_path / "workspace")
    workspace_default = _write_template(
        settings.workspace / "report_templates/default.docx",
        "workspace default",
    )

    assert settings.report_template_path == workspace_default


def test_workspace_legacy_report_template_precedes_packaged_default(tmp_path: Path) -> None:
    settings = build_settings(tmp_path / "workspace")
    workspace_legacy = _write_template(
        settings.workspace
        / "report_templates/04_贷前评分卡MOB3验证模板_带占位符.docx",
        "workspace legacy",
    )

    assert settings.report_template_path == workspace_legacy


def test_missing_explicit_custom_report_template_does_not_silently_fallback(
    tmp_path: Path,
) -> None:
    settings = build_settings(
        tmp_path / "workspace",
        report_template_name="institution-template.docx",
    )

    expected = settings.workspace / "report_templates/institution-template.docx"
    assert settings.report_template_path == expected
    assert not settings.report_template_path.exists()
