from dataclasses import asdict
import shutil
from types import SimpleNamespace

import joblib
import numpy as np
import pytest

from marvis.packs.modeling._runtime import _Runtime
from marvis.packs.modeling.development_refit import development_refit
from marvis.packs.modeling.development_refit_evidence import (
    _mask,
    validate_development_refit,
)
from marvis.packs.modeling.producer_receipts import load_receipt
from marvis.plugins.manifest import ToolRef
from marvis.settings import build_settings
from tests.test_fold_tuning import _source


def test_development_baseline_repeated_and_constant_feature_edges_replay_native_monitor():
    import pandas as pd
    from marvis.packs.modeling.development_refit_baseline import development_baseline
    from marvis.packs.modeling.monitor_tools import _monitor_run_feature_csi_checks
    frame = pd.DataFrame({"repeated": [0.] * 80 + [1.] * 20, "constant": [7.] * 100})
    baseline = development_baseline(frame, np.linspace(.1, .9, 100), ["repeated", "constant"],
        source=SimpleNamespace(id="ds-native", content_hash="a" * 64), membership={"sha256": "b" * 64})
    assert baseline["feature_distributions"]["repeated"]["quantile_edges"] == pytest.approx([0., .2, 1.])
    assert baseline["feature_distributions"]["repeated"]["bin_proportions"] == [.8, .2]
    assert baseline["feature_distributions"]["constant"]["quantile_edges"] == [-np.inf, np.inf]
    spec = {"label": "CSI", "metric": "feature_csi_max", "direction": "max", "warn": .1, "fail": .25}
    checks, features = _monitor_run_feature_csi_checks(frame, ("repeated", "constant"), baseline, spec)
    assert checks[0]["level"] == "green" and all(row["csi"] == 0 for row in features)
    shifted = frame.copy()
    shifted["repeated"] = 1.
    checks, _ = _monitor_run_feature_csi_checks(shifted, ("repeated", "constant"), baseline, spec)
    assert checks[0]["level"] == "red" and checks[0]["value"] > .25


@pytest.fixture(scope="module")
def trained(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("native_development_refit")

    def mutate(frame):
        frame.loc[frame["split"] == "oot", "x"] += 1_000_000
        # A source-owned score is not the final model's scoring output.
        frame["score"] = 0.125

    runner, registry, plan, frame = _source(tmp, mutate=mutate, return_runtime=True)
    tuning = runner.invoke(
        ToolRef("modeling", "tune_hyperparameters"),
        {
            "dataset_id": plan["training_dataset_id"],
            "features": ["x", "z"],
            "target_col": "y",
            "split_col": "split",
            "split_values": plan["split_values"],
            "recipe": "lr",
            "seed": 19,
            "n_trials": 1,
            "cv_folds": 3,
            "selection_evidence_refs": plan["references"],
            "fold_selection": {
                "source_dataset_id": plan["source_dataset_id"],
                "candidates": ["x", "z"],
            },
        },
        task_id="task-feature",
    )
    assert tuning.ok, tuning.error
    result = runner.invoke(
        ToolRef("modeling", "train_models"),
        {
            "dataset_id": plan["training_dataset_id"],
            "features": ["x", "z"],
            "target_col": "y",
            "split_col": "split",
            "split_values": plan["split_values"],
            "recipes": ["lr"],
            "seed": 19,
            "params": tuning.output["best_params"],
            "features_by_recipe": tuning.output["features_by_recipe"],
            "fold_tuning_evidence_refs": tuning.output["fold_tuning_evidence_refs"],
            "selection_evidence_refs": plan["references"],
        },
        task_id="task-feature",
    )
    assert result.ok, result.error
    selected = runner.invoke(
        ToolRef("modeling", "select_experiment"),
        {
            "experiment_ids": [result.output["best_experiment_id"]],
            "selected_experiment_id": result.output["best_experiment_id"],
            "refit_on_train_plus_test": True,
        },
        task_id="task-feature",
    )
    assert selected.ok and selected.output["refit"]["applied"], selected.error
    return runner, registry, frame, result.output["best_experiment_id"]


def test_native_select_produces_complete_development_refit_and_authentic_prior_model(
    trained,
):
    runner, registry, frame, identity = trained
    settings = build_settings(registry._repo.db_path.parent)
    runtime = _Runtime(
        SimpleNamespace(
            workspace=settings.workspace, datasets_root=settings.datasets_dir
        )
    )
    parent_before = runtime.experiments.get(identity)
    source_hash = registry.get(parent_before.config.dataset_id).content_hash
    parent_bytes = (
        settings.tasks_dir
        / "task-feature/modeling_artifacts"
        / runtime.modeling_repo.get_model_artifact(parent_before.artifact_id).model_path
    ).read_bytes()
    selected = runner.invoke(
        ToolRef("modeling", "select_experiment"),
        {
            "experiment_ids": [identity],
            "selected_experiment_id": identity,
            "refit_on_train_plus_test": True,
        },
        task_id="task-feature",
    )
    assert selected.ok, selected.error
    assert selected.output["refit"]["applied"] is True
    assert selected.output["selected_experiment_id"] != identity
    assert selected.output["metrics"]["test_ks"] is None
    assert selected.output["metrics"]["overfit_flag"] is None
    assert selected.output["selection_evidence"]["experiment_id"] == identity
    assert (
        selected.output["selection_evidence"]["metrics"]["test_ks"]
        == parent_before.metrics.test_ks
    )
    child = runtime.experiments.get(selected.output["selected_experiment_id"])
    artifact = runtime.modeling_repo.get_model_artifact(child.artifact_id)
    directory = settings.tasks_dir / "task-feature/modeling_artifacts"
    receipt, _ = load_receipt(settings.db_path, child, artifact, directory)
    assert receipt["provenance"]["experiment_id"] == child.id
    proof = validate_development_refit(settings.db_path, child, artifact, directory)
    assert np.array_equal(
        np.flatnonzero(_mask(proof["memberships"]["development_fit"], len(frame))),
        np.arange(300),
    )
    assert np.array_equal(
        np.flatnonzero(_mask(proof["memberships"]["excluded"], len(frame))),
        np.arange(300, 360),
    )
    model = joblib.load(directory / artifact.model_path)
    assert model.named_steps["scale"].mean_ == pytest.approx(
        frame.iloc[:300][list(artifact.feature_list)].mean().to_numpy()
    )
    assert (
        child.metrics.test_ks
        is child.metrics.oot_ks
        is child.metrics.overfit_train_test_gap
        is child.metrics.overfit_flag
        is None
    )
    baseline = artifact.baseline_distributions
    assert baseline["baseline_scope"] == "development_fit"
    assert baseline["score_distribution"]["train"]["sample_count"] == 300
    assert set(baseline["score_distribution"]) == {"train"}
    assert max(baseline["feature_distributions"]["x"]["quantile_edges"]) < 1_000_000
    assert proof["baseline_scope"] == "development_fit"
    assert baseline["oot_evaluation"] == "not_performed"
    assert asdict(runtime.experiments.get(identity).metrics) == asdict(
        parent_before.metrics
    )
    assert registry.get(parent_before.config.dataset_id).content_hash == source_hash
    parent_artifact = runtime.modeling_repo.get_model_artifact(
        parent_before.artifact_id
    )
    assert (directory / parent_artifact.model_path).read_bytes() == parent_bytes
    delivered = runner.invoke(
        ToolRef("modeling", "post_training_action"),
        {
            "experiment_id": child.id,
            "sample_dataset_id": child.config.dataset_id,
            "actions": ["create_challenger_backtest"],
        },
        task_id="task-feature",
    )
    assert delivered.ok, delivered.error
    champion = delivered.output["challenger_comparison"]["champion"]
    assert (
        champion["label"] == "previous_selected_experiment"
        and champion["experiment_id"] == identity
    )
    assert champion["metrics"]["test_ks"] == parent_before.metrics.test_ks


def test_completed_retry_never_fits_again(trained, monkeypatch):
    _, registry, _, identity = trained
    settings = build_settings(registry._repo.db_path.parent)
    runtime = _Runtime(
        SimpleNamespace(
            workspace=settings.workspace, datasets_root=settings.datasets_dir
        )
    )
    parent = runtime.experiments.get(identity)
    from sklearn.pipeline import Pipeline

    monkeypatch.setattr(
        Pipeline,
        "fit",
        lambda *args, **kwargs: pytest.fail("completed replay refitted"),
    )
    result = development_refit(
        runtime,
        experiment=parent,
        recipe="lr",
        pre_refit_metrics={"test_ks": 1.0, "oot_auc": 1.0},
    )
    assert result["applied"] and result["fit_rows"] == 300
    assert result["metrics_before_refit"] == asdict(parent.metrics)


def test_child_actual_monitoring_baseline_and_drift_and_immutable_replay(trained, tmp_path):
    from marvis.packs.modeling.monitor_binding import capture_monitoring_binding
    from marvis.repositories.modeling import ModelingRepository
    runner, registry, frame, identity = trained
    settings = build_settings(registry._repo.db_path.parent)
    selected = runner.invoke(ToolRef("modeling", "select_experiment"), {
        "experiment_ids": [identity], "selected_experiment_id": identity,
        "refit_on_train_plus_test": True,
    }, task_id="task-feature")
    assert selected.ok, selected.error
    child_id = selected.output["selected_experiment_id"]
    artifact = ModelingRepository(settings.db_path).get_model_artifact(selected.output["artifact_id"])
    before = asdict(artifact)
    for shifted in (False, True):
        source = frame.iloc[:300].drop(columns=["split", "score", "y"]).copy()
        if shifted:
            source["x"] += 1_000
            source["z"] += 1_000
        path = tmp_path / f"monitor-{shifted}.parquet"
        source.to_parquet(path, index=False)
        dataset = registry.register_from_upload("task-feature", path, role="monitoring.input")
        binding = capture_monitoring_binding(settings, "task-feature", child_id, dataset.id, dataset.content_hash)
        scores = runner.invoke(ToolRef("modeling", "score_dataset"), {
            "experiment_id": child_id, "dataset_id": dataset.id, "monitoring_binding": binding,
        }, task_id="task-feature")
        assert scores.ok, scores.error
        result = runner.invoke(ToolRef("modeling", "monitor_run"), {
            "experiment_id": child_id, "scored_dataset_id": scores.output["result_dataset_id"],
            "score_col": scores.output["score_col"], "monitoring_binding": binding,
        }, task_id="task-feature")
        assert result.ok, result.error
        assert result.output["row_count"] == 300
        check = next(c for c in result.output["checks"] if c["id"] == "score_psi")
        assert check["value"] > .25 if shifted else check["value"] == pytest.approx(0., abs=1e-12)
        assert result.output["overall_level"] == ("red" if shifted else "green")
        assert result.output["label_maturity_assurance"] == "unknown"
        replay = runner.invoke(ToolRef("modeling", "monitor_run"), {
            "experiment_id": child_id, "scored_dataset_id": scores.output["result_dataset_id"],
            "score_col": scores.output["score_col"], "monitoring_binding": binding,
        }, task_id="task-feature")
        assert replay.ok and replay.output == result.output
    assert asdict(ModelingRepository(settings.db_path).get_model_artifact(artifact.id)) == before


def test_child_monitoring_source_change_refuses_before_prediction(trained, tmp_path, monkeypatch):
    from marvis.packs.modeling.monitor_binding import capture_monitoring_binding
    from marvis.packs.modeling import scoring
    runner, registry, frame, identity = trained
    settings = build_settings(registry._repo.db_path.parent)
    selected = runner.invoke(ToolRef("modeling", "select_experiment"), {
        "experiment_ids": [identity], "selected_experiment_id": identity,
        "refit_on_train_plus_test": True,
    }, task_id="task-feature")
    child_id = selected.output["selected_experiment_id"]
    path = tmp_path / "changed-monitoring-source.parquet"
    frame.iloc[:20].drop(columns=["split"]).to_parquet(path, index=False)
    dataset = registry.register_from_upload("task-feature", path, role="monitoring.input")
    binding = capture_monitoring_binding(settings, "task-feature", child_id, dataset.id, dataset.content_hash)
    original = registry.resolve_path(dataset.id)
    original.write_bytes(original.read_bytes() + b"source-changed")
    monkeypatch.setattr(scoring._ModelArtifactScorer, "score", lambda *args, **kwargs: pytest.fail("source CAS must run before prediction"))
    from marvis.packs.modeling.monitor_tools import tool_score_dataset
    ctx = SimpleNamespace(workspace=settings.workspace, datasets_root=settings.datasets_dir, task_id="task-feature")
    from marvis.packs.modeling.errors import ModelingError
    with pytest.raises((ModelingError, ValueError)):
        tool_score_dataset({"experiment_id": child_id, "dataset_id": dataset.id, "monitoring_binding": binding}, ctx)


def test_failed_publication_restores_original_rows_and_model_pointer(
    trained, tmp_path, monkeypatch
):
    _, registry, _, identity = trained
    destination = tmp_path / "workspace"
    shutil.copytree(registry._repo.db_path.parent, destination)
    settings = build_settings(destination)
    runtime = _Runtime(
        SimpleNamespace(
            workspace=settings.workspace, datasets_root=settings.datasets_dir
        )
    )
    from marvis.packs.modeling.errors import ModelingError

    # Remove only this copied fixture's completed child so the publication
    # branch can be exercised against an independently retained parent.
    from marvis.db_schema import connect

    with connect(settings.db_path) as conn:
        child = conn.execute(
            "SELECT id,artifact_id FROM experiments WHERE id!=?", (identity,)
        ).fetchone()
        conn.execute(
            "DELETE FROM task_artifacts WHERE kind IN ('modeling_development_refit','modeling_native_model_closure') AND provenance_json LIKE ?",
            (f"%{child[0]}%",),
        )
        conn.execute("DELETE FROM model_artifacts WHERE experiment_id=?", (child[0],))
        conn.execute("DELETE FROM experiments WHERE id=?", (child[0],))
    directory = settings.tasks_dir / "task-feature/modeling_artifacts"
    before = {p.name: p.read_bytes() for p in directory.iterdir() if p.is_file()}
    before_count = len(runtime.experiments.list_for_task("task-feature"))
    monkeypatch.setattr(
        runtime.experiments,
        "attach_result_on_connection",
        lambda *args: (_ for _ in ()).throw(ModelingError("test publication failure")),
    )
    with pytest.raises(ModelingError, match="test publication failure"):
        development_refit(
            runtime,
            experiment=runtime.experiments.get(identity),
            recipe="lr",
            pre_refit_metrics={},
        )
    assert {
        p.name: p.read_bytes() for p in directory.iterdir() if p.is_file()
    } == before
    assert len(runtime.experiments.list_for_task("task-feature")) == before_count


def _copied_runtime(trained, tmp_path, *, remove_child=False):
    _, registry, _, identity = trained
    destination = tmp_path / "workspace"
    shutil.copytree(registry._repo.db_path.parent, destination)
    settings = build_settings(destination)
    runtime = _Runtime(
        SimpleNamespace(
            workspace=settings.workspace, datasets_root=settings.datasets_dir
        )
    )
    if remove_child:
        from marvis.db_schema import connect

        with connect(settings.db_path) as conn:
            child = conn.execute(
                "SELECT id,artifact_id FROM experiments WHERE id!=?", (identity,)
            ).fetchone()
            conn.execute(
                "DELETE FROM task_artifacts WHERE kind IN ('modeling_development_refit','modeling_native_model_closure') AND provenance_json LIKE ?",
                (f"%{child[0]}%",),
            )
            conn.execute(
                "DELETE FROM model_artifacts WHERE experiment_id=?", (child[0],)
            )
            conn.execute("DELETE FROM experiments WHERE id=?", (child[0],))
    return settings, runtime, identity


def test_corrupt_parent_source_or_refit_controls_are_rejected_before_fit(
    trained, tmp_path, monkeypatch
):
    settings, runtime, identity = _copied_runtime(trained, tmp_path)
    from marvis.packs.modeling.errors import ModelingError
    from sklearn.pipeline import Pipeline

    monkeypatch.setattr(
        Pipeline,
        "fit",
        lambda *args, **kwargs: pytest.fail("invalid evidence entered fit"),
    )
    parent = runtime.experiments.get(identity)
    parent_artifact = runtime.modeling_repo.get_model_artifact(parent.artifact_id)
    params = dict(parent_artifact.params)
    params["C"] = params["C"] * 2
    runtime.modeling_repo.set_model_artifact_params(parent_artifact.id, params)
    with pytest.raises(ModelingError, match="authentication"):
        development_refit(runtime, experiment=parent, recipe="lr", pre_refit_metrics={})


def test_unknown_metrics_are_preserved_and_legacy_missing_keys_stay_compatible():
    from marvis.repositories.modeling import (
        _model_metrics_from_dict,
        _model_metrics_to_dict,
    )

    unknown = _model_metrics_from_dict(
        {"overfit_train_test_gap": None, "overfit_flag": None}
    )
    assert unknown.overfit_train_test_gap is unknown.overfit_flag is None
    assert _model_metrics_to_dict(unknown)["overfit_flag"] is None
    legacy = _model_metrics_from_dict({})
    assert legacy.overfit_train_test_gap == 0.0 and legacy.overfit_flag is False


@pytest.mark.parametrize("committed", [False, True])
def test_actual_process_crash_reconciles_only_owned_publication(
    trained, tmp_path, monkeypatch, committed
):
    import json
    import subprocess
    import sys

    settings, runtime, identity = _copied_runtime(trained, tmp_path, remove_child=True)
    directory = settings.tasks_dir / "task-feature/modeling_artifacts"
    before = {p.name: p.read_bytes() for p in directory.iterdir() if p.is_file()}
    lock = directory / ".development-refit.lock"
    assert before.get(lock.name, b"") == b""
    script = """
from pathlib import Path
from types import SimpleNamespace
import os,sys
from marvis.packs.modeling._runtime import _Runtime
from marvis.packs.modeling.development_refit import development_refit
from marvis.settings import build_settings
settings=build_settings(sys.argv[1]); runtime=_Runtime(SimpleNamespace(workspace=settings.workspace,datasets_root=settings.datasets_dir))
if sys.argv[3] == "committed":
    old=Path.unlink
    def unlink(path,*args,**kwargs):
        if path.name.endswith(".pending.json"): os._exit(23)
        return old(path,*args,**kwargs)
    Path.unlink=unlink
else:
    runtime.experiments.attach_result_on_connection=lambda *args: os._exit(23)
development_refit(runtime,experiment=runtime.experiments.get(sys.argv[2]),recipe="lr",pre_refit_metrics={})
"""
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            script,
            str(settings.workspace),
            identity,
            "committed" if committed else "uncommitted",
        ],
        timeout=45,
        capture_output=True,
        text=True,
    )
    assert result.returncode == 23, result.stderr
    assert lock.is_file() and not lock.is_symlink()
    assert lock.read_bytes() == b""
    journals = list(directory.glob(".development-refit-*.pending.json"))
    assert len(journals) == 1
    body = json.loads(journals[0].read_bytes())
    from marvis.packs.modeling.development_refit_recovery import reconcile_journal

    reconcile_journal(
        runtime, directory, body["identity"], body["receipt_id"], "task-feature"
    )
    assert not journals[0].exists()
    if committed:
        from sklearn.pipeline import Pipeline

        monkeypatch.setattr(
            Pipeline,
            "fit",
            lambda *args, **kwargs: pytest.fail("committed crash replay refitted"),
        )
        replay = development_refit(
            runtime,
            experiment=runtime.experiments.get(identity),
            recipe="lr",
            pre_refit_metrics={},
        )
        assert (
            replay["applied"]
            and len(runtime.experiments.list_for_task("task-feature")) == 2
        )
    else:
        # FileLock releases may unlink their control file, but a hard process
        # exit leaves the empty lock. All original bytes and every other name
        # must still match; no other publication residue is permitted.
        expected = before | {lock.name: b""}
        assert {
            p.name: p.read_bytes() for p in directory.iterdir() if p.is_file()
        } == expected
        assert len(runtime.experiments.list_for_task("task-feature")) == 1
        replay = development_refit(
            runtime,
            experiment=runtime.experiments.get(identity),
            recipe="lr",
            pre_refit_metrics={},
        )
        assert replay["applied"]


def test_concurrent_native_processes_publish_one_child_and_fit_once(trained, tmp_path):
    import json
    import subprocess
    import sys

    settings, runtime, identity = _copied_runtime(trained, tmp_path, remove_child=True)
    marker = settings.workspace / "fit-count"
    script = """
from types import SimpleNamespace
import json,os,sys
from sklearn.pipeline import Pipeline
from marvis.packs.modeling._runtime import _Runtime
from marvis.packs.modeling.development_refit import development_refit
from marvis.settings import build_settings
settings=build_settings(sys.argv[1]); runtime=_Runtime(SimpleNamespace(workspace=settings.workspace,datasets_root=settings.datasets_dir))
original=Pipeline.fit
def fit(self,*args,**kwargs):
    fd=os.open(sys.argv[3],os.O_CREAT|os.O_WRONLY|os.O_APPEND,0o600)
    try: os.write(fd,b"fit\\n")
    finally: os.close(fd)
    return original(self,*args,**kwargs)
Pipeline.fit=fit
result=development_refit(runtime,experiment=runtime.experiments.get(sys.argv[2]),recipe="lr",pre_refit_metrics={})
print(json.dumps({"experiment_id":result["experiment_id"],"artifact_id":result["artifact_id"]}))
"""
    processes = [
        subprocess.Popen(
            [
                sys.executable,
                "-c",
                script,
                str(settings.workspace),
                identity,
                str(marker),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        for _ in range(2)
    ]
    outputs = []
    try:
        for process in processes:
            output, errors = process.communicate(timeout=45)
            assert process.returncode == 0, errors
            outputs.append(json.loads(output.strip().splitlines()[-1]))
    finally:
        for process in processes:
            if process.poll() is None:
                process.kill()
                process.wait()
    assert outputs[0] == outputs[1]
    assert marker.read_text().splitlines() == ["fit"]
    assert len(runtime.experiments.list_for_task("task-feature")) == 2


def test_source_drift_rejects_before_fit_and_cached_replay(
    trained, tmp_path, monkeypatch
):
    settings, runtime, identity = _copied_runtime(trained, tmp_path)
    from marvis.data.errors import DatasetContentDriftError
    from sklearn.pipeline import Pipeline

    parent = runtime.experiments.get(identity)
    source = runtime.registry.resolve_verified_path(parent.config.dataset_id)
    source.write_bytes(source.read_bytes() + b"changed")
    monkeypatch.setattr(
        Pipeline, "fit", lambda *args, **kwargs: pytest.fail("source drift entered fit")
    )
    with pytest.raises(DatasetContentDriftError):
        development_refit(runtime, experiment=parent, recipe="lr", pre_refit_metrics={})


def test_source_change_during_fit_rolls_back_all_producer_outputs(
    trained, tmp_path, monkeypatch
):
    settings, runtime, identity = _copied_runtime(trained, tmp_path, remove_child=True)
    from marvis.feature.errors import FeatureError
    from sklearn.pipeline import Pipeline

    parent = runtime.experiments.get(identity)
    source = runtime.registry.resolve_verified_path(parent.config.dataset_id)
    original_fit = Pipeline.fit
    directory = settings.tasks_dir / "task-feature/modeling_artifacts"
    before = {p.name: p.read_bytes() for p in directory.iterdir() if p.is_file()}

    def fit(self, *args, **kwargs):
        result = original_fit(self, *args, **kwargs)
        source.write_bytes(source.read_bytes() + b"changed")
        return result

    monkeypatch.setattr(Pipeline, "fit", fit)
    with pytest.raises(FeatureError, match="source bytes changed"):
        development_refit(runtime, experiment=parent, recipe="lr", pre_refit_metrics={})
    assert {
        p.name: p.read_bytes() for p in directory.iterdir() if p.is_file()
    } == before
    assert len(runtime.experiments.list_for_task("task-feature")) == 1


def test_refit_report_uses_real_fit_scope_and_authentic_parent_test(
    trained, monkeypatch
):
    from openpyxl import load_workbook
    from marvis.packs.modeling import report_tools

    monkeypatch.setattr(
        report_tools,
        "_report_llm_factory",
        lambda *args: pytest.fail("changed fit scope entered ordinary report DTO"),
    )
    runner, registry, _, identity = trained
    runtime = _Runtime(
        SimpleNamespace(
            workspace=registry._repo.db_path.parent,
            datasets_root=registry.datasets_root,
        )
    )
    child = next(
        item
        for item in runtime.experiments.list_for_task("task-feature")
        if item.id != identity
    )
    parent = runtime.experiments.get(identity)
    output = runner.invoke(
        ToolRef("modeling", "generate_model_report"),
        {
            "experiment_id": child.id,
            "dataset_id": child.config.dataset_id,
        },
        task_id="task-feature",
    )
    assert output.ok, output.error
    rows = {row["split"]: row for row in output.output["dataset_split"]}
    assert rows["development_train_plus_test"]["sample_count"] == 300
    assert rows["pre_refit_test"]["sample_count"] == 60
    assert rows["pre_refit_test"]["experiment_id"] == identity
    assert rows["pre_refit_test"]["ks"] == parent.metrics.test_ks
    assert rows["final_independent_test"]["ks"] is None
    test_bands = [row for row in output.output["score_bands"] if row["split"] == "test"]
    assert test_bands and all(
        row["evaluation_purpose"] == "consumed_in_fit_diagnostic" for row in test_bands
    )
    workbook = load_workbook(output.output["report_path"], read_only=True)
    cells = [
        value for row in workbook["汇总"].iter_rows(values_only=True) for value in row
    ]
    assert "development_train_plus_test" in cells and "pre_refit_test" in cells
    assert any(
        isinstance(value, str) and "不能作为最终模型的独立测试" in value
        for value in cells
    )
    workbook.close()
    import pandas as pd
    from marvis.packs.modeling._common import MODEL_REPORT_SCORE_COL

    scored = pd.read_parquet(
        runtime.settings.tasks_dir / "task-feature/outputs/model_report_scored.parquet"
    )
    assert MODEL_REPORT_SCORE_COL in scored
    assert not np.allclose(scored[MODEL_REPORT_SCORE_COL], 0.125)
