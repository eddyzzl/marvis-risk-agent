from copy import deepcopy
from io import BytesIO
import json

from docx import Document
from fastapi.testclient import TestClient
import numpy as np
from openpyxl import load_workbook
import pytest

from marvis.decision_twin._canonical import content_hash
from marvis.decision_twin.batch import replay_batch
from marvis.decision_twin.batch_contracts import (
    HistoricalReplayRequest,
    HistoricalTemporalStability,
)
from marvis.decision_twin.temporal import (
    bind_temporal_population,
    measure_temporal_stability,
)
from marvis.feature.binning import equal_frequency_edges
from marvis.feature.metrics import compute_psi
from marvis.reference_decision.contracts import DecisionError
from marvis.validation.binning import bin_distribution

from test_decision_twin_batch import batch, _changed_dataset  # noqa: F401
from test_reference_decision import packaged  # noqa: F401


def _policy(**updates):
    policy = {
        "timezone": "UTC",
        "reference_window": {
            "name": "reference",
            "start": "2026-01-01T00:00:00Z",
            "end": "2026-02-01T00:00:00Z",
        },
        "comparison_windows": [
            {
                "name": "comparison",
                "start": "2026-02-01T00:00:00Z",
                "end": "2026-03-01T00:00:00Z",
            }
        ],
        "minimum_reference_rows": 4,
        "minimum_comparison_rows": 4,
        "bin_count": 2,
        "thresholds": {
            "max_score_psi": 0.25,
            "max_action_psi": 0.25,
            "max_absolute_approval_rate_delta": 0.15,
        },
        "policy_source_ref": "human-approved retrospective monitoring policy",
    }
    policy.update(updates)
    return HistoricalTemporalStability.model_validate(policy)


def _native_rows(reference=(0.1, 0.2, 0.3, 0.4), comparison=(0.7, 0.8, 0.9, 1.0)):
    records, decisions = [], []
    for i, score in enumerate([*reference, *comparison]):
        record = {
            "record_id": content_hash(f"record-{i}"),
            "facts_hash": content_hash(f"facts-{i}"),
            "decision_at": "2026-01-02T00:00:00Z"
            if i < len(reference)
            else "2026-02-02T00:00:00Z",
        }
        records.append(record)
        decisions.append(
            {
                **record,
                "score": score,
                "score_product": "raw_pd",
                "package_hash": "a" * 64,
                "action": {"type": "approval" if score < 0.5 else "reject"},
            }
        )
    return records, decisions


def _measure(records, decisions, policy=None, population=None):
    policy = policy or _policy()
    return measure_temporal_stability(
        records,
        decisions,
        policy,
        population=population or bind_temporal_population(records, policy),
        package_hash="a" * 64,
        score_product="raw_pd",
    )


def test_reference_fitted_shared_kernel_detects_shift_without_new_claims():
    records, decisions = _native_rows()
    result = _measure(records, decisions)
    edges = equal_frequency_edges(np.array([0.1, 0.2, 0.3, 0.4]), 2)
    assert result["reference"]["score_bins"]["edges"] == ["-inf", 0.25, "inf"]
    expected = compute_psi(
        bin_distribution([0.1, 0.2, 0.3, 0.4], edges),
        bin_distribution([0.7, 0.8, 0.9, 1.0], edges),
    )
    checks = result["comparisons"][0]["checks"]
    assert checks[0]["value"] == pytest.approx(expected)
    assert checks[1]["value"] == pytest.approx(
        compute_psi(np.array([1.0, 0.0, 0.0]), np.array([0.0, 1.0, 0.0]))
    )
    assert checks[2]["value"] == 1.0
    assert result["status"] == "measured"
    assert result["verdict"] == "failed"
    assert result["reference"]["score_bins"]["fit_scope"] == "reference_window_only"
    assert result["historical_package_availability"] == "not_established"
    assert result["out_of_training_validation"] == "not_established"
    assert result["causal_gain_verified"] is False
    assert result["automatic_action_permitted"] is False
    # Moving only comparison scores cannot refit the independently bound bins.
    changed = deepcopy(decisions)
    for row in changed[4:]:
        row["score"] = 0.01
    assert (
        _measure(records, changed)["reference"]["score_bins"]
        == result["reference"]["score_bins"]
    )


def test_supported_equal_distributions_pass_and_inclusive_threshold():
    records, decisions = _native_rows(comparison=(0.1, 0.2, 0.3, 0.4))
    result = _measure(
        records,
        decisions,
        _policy(
            thresholds={
                "max_score_psi": 0.0,
                "max_action_psi": 0.0,
                "max_absolute_approval_rate_delta": 0.0,
            }
        ),
    )
    assert result["verdict"] == "passed"
    assert set(result["value"].values()) == {0.0}


@pytest.mark.parametrize("score_product", ["calibrated_pd", "scorecard_points"])
def test_native_score_products_keep_separate_reference_scale(score_product):
    records, decisions = _native_rows()
    for row in decisions:
        row["score_product"] = score_product
        if score_product == "scorecard_points":
            row["score"] = 900 - row["score"] * 500
    policy = _policy()
    result = measure_temporal_stability(
        records,
        decisions,
        policy,
        population=bind_temporal_population(records, policy),
        package_hash="a" * 64,
        score_product=score_product,
    )
    assert result["score_product"] == score_product
    assert result["status"] == "measured"
    assert result["reference"]["score_bins"]["edges"][1] == (
        775 if score_product == "scorecard_points" else 0.25
    )


def test_half_open_membership_and_excluded_population_are_bound():
    records, _ = _native_rows()
    timestamps = [
        "2025-12-31T23:59:59Z",
        "2026-01-01T00:00:00Z",
        "2026-01-31T23:59:59Z",
        "2026-02-01T00:00:00Z",
        "2026-02-28T23:59:59Z",
        "2026-03-01T00:00:00Z",
        "2026-03-02T00:00:00Z",
        "2026-02-15T00:00:00Z",
    ]
    for row, timestamp in zip(records, timestamps, strict=True):
        row["decision_at"] = timestamp
    population = bind_temporal_population(records, _policy())
    assert population.reference.positions == (1, 2)
    assert population.comparisons[0].positions == (3, 4, 7)
    assert population.excluded_count == 3
    assert population.reference.members_hash == content_hash([records[1], records[2]])
    assert population.excluded_members_hash == content_hash(
        [records[0], records[5], records[6]]
    )
    assert population.to_dict()["interval_convention"] == "[start,end)"


@pytest.mark.parametrize(
    "reference,bin_count", [((0.2,) * 4, 2), ((0.2, 0.2, 0.2, 0.3), 4)]
)
def test_degraded_reference_bins_unknown_but_supported_actions_retained(
    reference, bin_count
):
    records, decisions = _native_rows(reference=reference, comparison=reference)
    result = _measure(records, decisions, _policy(bin_count=bin_count))
    assert result["reference"]["score_bins"]["reason"] == "degraded_reference_binning"
    assert result["comparisons"][0]["checks"][0]["value"] is None
    assert result["comparisons"][0]["checks"][1]["value"] == 0.0
    assert result["value"]["max_score_psi"] is None
    assert result["status"] == "partial"
    assert result["verdict"] == "insufficient_evidence"


@pytest.mark.parametrize(
    "reference,comparison,reason",
    [
        ((), (0.1,) * 4, "insufficient_reference_support"),
        ((0.1,) * 3, (0.1,) * 4, "insufficient_reference_support"),
        ((0.1, 0.2, 0.3, 0.4), (), "insufficient_comparison_support"),
        ((0.1, 0.2, 0.3, 0.4), (0.1,) * 3, "insufficient_comparison_support"),
    ],
)
def test_missing_support_never_becomes_zero(reference, comparison, reason):
    records, decisions = _native_rows(reference, comparison)
    result = _measure(records, decisions)
    assert result["value"] is None
    assert result["status"] == "unknown"
    assert result["verdict"] == "insufficient_evidence"
    assert all(
        check["reason"] == reason and check["value"] is None and check["passed"] is None
        for check in result["comparisons"][0]["checks"]
    )


def test_missing_period_cannot_be_hidden_by_maximum_of_other_periods():
    records, decisions = _native_rows(comparison=(0.1, 0.2, 0.3, 0.4))
    policy = _policy()
    data = policy.model_dump()
    data["comparison_windows"].append(
        {
            "name": "empty",
            "start": "2026-03-01T00:00:00Z",
            "end": "2026-04-01T00:00:00Z",
        }
    )
    result = _measure(
        records, decisions, HistoricalTemporalStability.model_validate(data)
    )
    assert result["status"] == "partial"
    assert result["verdict"] == "insufficient_evidence"
    assert set(result["value"].values()) == {None}


@pytest.mark.parametrize(
    "field,value",
    [
        ("record_id", "b" * 64),
        ("facts_hash", "b" * 64),
        ("decision_at", "2026-01-03T00:00:00Z"),
        ("package_hash", "b" * 64),
        ("score_product", "calibrated_pd"),
        ("score", float("nan")),
        ("score", float("inf")),
        ("score", True),
        ("score", 1.01),
        ("action", {"type": "grant"}),
    ],
)
def test_decision_binding_and_native_score_errors_rejected(field, value):
    records, decisions = _native_rows()
    decisions[0][field] = value
    with pytest.raises(ValueError):
        _measure(records, decisions)


@pytest.mark.parametrize(
    "timestamp", [None, "2026-01-01", "2026-01-01T00:00:00", "unknown"]
)
def test_unknown_or_naive_decision_times_rejected(timestamp):
    records, _ = _native_rows()
    records[0]["decision_at"] = timestamp
    with pytest.raises(ValueError):
        bind_temporal_population(records, _policy())


def test_population_and_decision_order_drift_rejected():
    records, decisions = _native_rows()
    population = bind_temporal_population(records, _policy())
    with pytest.raises(ValueError, match="population binding"):
        _measure(records[::-1], decisions[::-1], population=population)
    with pytest.raises(ValueError, match="record binding"):
        _measure(records, decisions[::-1])
    records[1]["record_id"] = records[0]["record_id"]
    with pytest.raises(ValueError, match="unique"):
        bind_temporal_population(records, _policy())


@pytest.mark.parametrize(
    "mutation",
    [
        lambda p: p.update(timezone="unknown/zone"),
        lambda p: p.update(timezone="Asia/Shanghai"),
        lambda p: p["reference_window"].update(start="2026-01-01T00:00:00"),
        lambda p: p["reference_window"].update(end="2026-01-01T00:00:00Z"),
        lambda p: p["comparison_windows"][0].update(start="2026-01-15T00:00:00Z"),
        lambda p: p["comparison_windows"][0].update(name="reference"),
        lambda p: p["comparison_windows"].append(
            {
                "name": "overlap",
                "start": "2026-02-15T00:00:00Z",
                "end": "2026-04-01T00:00:00Z",
            }
        ),
        lambda p: p.update(minimum_reference_rows=True),
        lambda p: p.update(minimum_comparison_rows=1),
        lambda p: p.update(bin_count=21),
        lambda p: p["thresholds"].update(max_score_psi=float("nan")),
        lambda p: p["thresholds"].update(max_action_psi=-0.1),
        lambda p: p["thresholds"].update(max_absolute_approval_rate_delta=1.1),
        lambda p: p.pop("policy_source_ref"),
        lambda p: p.pop("thresholds"),
        lambda p: p.pop("minimum_comparison_rows"),
        lambda p: p.update(comparison_windows=[]),
    ],
)
def test_ambiguous_or_incomplete_policy_rejected(mutation):
    data = _policy().model_dump()
    mutation(data)
    with pytest.raises(ValueError):
        HistoricalTemporalStability.model_validate(data)


def test_timezone_boundaries_and_dst_are_explicit():
    data = _policy().model_dump()
    data["timezone"] = "Asia/Shanghai"
    for window in [data["reference_window"], *data["comparison_windows"]]:
        for key in ("start", "end"):
            window[key] = window[key].replace("Z", "+08:00")
    HistoricalTemporalStability.model_validate(data)
    data["timezone"] = "America/New_York"
    data["reference_window"] = {
        "name": "spring",
        "start": "2026-03-08T02:30:00-05:00",
        "end": "2026-04-01T00:00:00-04:00",
    }
    data["comparison_windows"] = [
        {
            "name": "fall",
            "start": "2026-11-01T01:30:00-05:00",
            "end": "2026-12-01T00:00:00-05:00",
        }
    ]
    with pytest.raises(ValueError, match="offset"):
        HistoricalTemporalStability.model_validate(data)
    data["reference_window"]["start"] = "2026-03-08T01:30:00-05:00"
    # An explicit valid offset resolves the repeated autumn wall-clock hour.
    HistoricalTemporalStability.model_validate(data)
    data["comparison_windows"][0]["start"] = "2026-11-01T01:30:00"
    with pytest.raises(ValueError, match="timezone-aware"):
        HistoricalTemporalStability.model_validate(data)


def _legacy_wire():
    # Explicit wire contract published before temporal_stability; do not generate
    # this fixture from the current model fields or exclude defaults.
    return {
        "schema_version": "decision_twin.batch_request.v2",
        "dataset_id": "legacy-data",
        "expected_content_hash": "0" * 64,
        "record_id_col": "id",
        "decision_at_col": "at",
        "as_of": "2026-08-01T00:00:00Z",
        "source_ref": "export",
        "population": "all",
        "features": [
            {
                "name": "x",
                "value_col": "x",
                "available_at_col": "available",
                "event_at_col": "event",
            }
        ],
        "scenarios": [
            {"name": "baseline", "kind": "baseline", "package_hash": "a" * 64},
            {"name": "challenger", "kind": "challenger", "package_hash": "a" * 64},
        ],
        "economics": None,
        "protected_group": None,
        "capacity": None,
        "observed_actions": None,
        "constraints": [],
    }


def test_absent_and_null_policy_preserve_published_v2_wire_hash():
    legacy = _legacy_wire()
    for data in (legacy, {**legacy, "temporal_stability": None}):
        contract = HistoricalReplayRequest.model_validate(data)
        assert contract.model_dump() == legacy
        assert json.loads(contract.model_dump_json()) == legacy
        assert contract.contract_hash == content_hash(legacy)
        # Produced with the predecessor contract class at a283c15e.
        assert (
            contract.contract_hash
            == "29ec8409e62d4ec55fe233d7a814d3ba45ce9df7726070d7613fdcbf717cab7b"
        )


def test_full_contract_rejects_window_after_as_of():
    with pytest.raises(ValueError, match="as_of"):
        HistoricalReplayRequest.model_validate(
            {
                **_legacy_wire(),
                "as_of": "2026-02-15T00:00:00Z",
                "temporal_stability": _policy().model_dump(),
            }
        )


def test_legacy_native_receipt_identity_survives_explicit_null_policy(batch):  # noqa: F811
    app, material, contract, _, _ = batch
    before = replay_batch(
        material, contract, material.prepare(contract)[2]["proposal_hash"]
    )
    null_contract = HistoricalReplayRequest.model_validate(
        {**contract.model_dump(), "temporal_stability": None}
    )
    after = replay_batch(
        material, null_contract, material.prepare(null_contract)[2]["proposal_hash"]
    )
    assert before == after
    assert "temporal_stability" not in after["payload"]["contract"]
    assert material.load(before["artifact_id"])["payload"] == before["payload"]
    client = TestClient(app)
    response = client.get(
        f"/api/tasks/{material.task_id}/decision-twin/{before['artifact_id']}/export/xlsx"
    )
    workbook = load_workbook(BytesIO(response.content))
    assert "时间稳定性" not in workbook.sheetnames


def test_real_temporal_workflow_gate_native_execution_and_canonical_exports(batch):  # noqa: F811
    app, material, _, _, _ = batch
    contract = _changed_dataset(
        batch, decision_at=["2026-01-02T00:00:00Z"] * 6 + ["2026-02-02T00:00:00Z"] * 6
    )
    contract = HistoricalReplayRequest.model_validate(
        {**contract.model_dump(), "temporal_stability": _policy().model_dump()}
    )
    client = TestClient(app)
    base = f"/api/tasks/{material.task_id}"
    response = client.post(base + "/decision-twin/proposal", json=contract.model_dump())
    assert response.status_code == 200, response.text
    proposal = response.json()
    assert proposal["temporal_population"]["reference"]["sample_count"] == 6
    assert proposal["temporal_population"]["comparisons"][0]["sample_count"] == 6
    altered = contract.model_dump()
    altered["temporal_stability"]["thresholds"]["max_score_psi"] = 0.9
    with pytest.raises(DecisionError, match="proposal_stale"):
        replay_batch(
            material,
            HistoricalReplayRequest.model_validate(altered),
            proposal["proposal_hash"],
        )
    created = client.post(
        base + "/plans",
        json={
            "goal": "历史决策回放",
            "slots": {
                "replay_contract": contract.model_dump(),
                "proposal_hash": proposal["proposal_hash"],
            },
        },
    )
    assert created.status_code == 201, created.text
    plan = created.json()["plan"]
    assert (
        client.post(
            f"/api/plans/{plan['id']}/confirm", json=plan["confirmation_snapshot"]
        ).status_code
        == 200
    )
    assert client.post(f"/api/plans/{plan['id']}/run").status_code == 202
    current = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert current["status"] == "awaiting_confirm"
    assert material.artifacts.list_for_task(material.task_id) == []
    step = current["steps"][0]
    approved = client.post(
        f"/api/plans/{plan['id']}/steps/{step['id']}/decisions",
        json={
            "decision": "approve",
            "reason": "核对参考和对照窗口、时区、样本支持及阈值",
            **step["confirmation_snapshot"],
        },
    )
    assert approved.status_code == 202, approved.text
    current = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert current["status"] == "done", current
    artifacts = client.get(base + "/decision-twin").json()["artifacts"]
    assert len(artifacts) == 1
    identity = artifacts[0]["artifact_id"]
    url = base + "/decision-twin/" + identity
    receipt = client.get(url).json()
    assert receipt == client.get(url + "/export/json").json()
    results = [
        scenario["metrics"]["stability"] for scenario in receipt["payload"]["scenarios"]
    ]
    assert (
        results[0]["population"]
        == results[1]["population"]
        == proposal["temporal_population"]
    )
    assert results[0]["status"] == "measured"
    assert (
        results[0] == results[1]
    )  # fixture intentionally selects the same frozen package
    # Native receipt binds each score into the same reference-only kernel result.
    native = receipt["payload"]["scenarios"][0]["decisions"]
    edges = equal_frequency_edges(np.array([row["score"] for row in native[:6]]), 2)
    expected = compute_psi(
        bin_distribution([row["score"] for row in native[:6]], edges),
        bin_distribution([row["score"] for row in native[6:]], edges),
    )
    assert results[0]["comparisons"][0]["checks"][0]["value"] == pytest.approx(expected)
    workbook = load_workbook(BytesIO(client.get(url + "/export/xlsx").content))
    rows = list(workbook["时间稳定性"].values)
    assert len(rows) == 7
    for i, check in enumerate(results[0]["comparisons"][0]["checks"], 1):
        assert rows[i][10] == check["metric"]
        assert rows[i][11] == pytest.approx(check["value"])
        assert rows[i][12] == check["threshold"]
        assert rows[i][14] == check["status"]
    assert rows[1][4] == results[0]["population"]["reference"]["members_hash"]
    document = Document(BytesIO(client.get(url + "/export/docx").content))
    assert len(document.tables[0].rows) == 7
    assert document.tables[0].rows[1].cells[3].text == str(
        results[0]["comparisons"][0]["checks"][0]["value"]
    )
    assert identity in "\n".join(p.text for p in document.paragraphs)
