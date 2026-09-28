"""Rule-only replay has no score PSI, but retains measured action evidence."""

from copy import deepcopy
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pandas as pd
import pytest

from marvis.decision_twin.batch_contracts import HistoricalReplayRequest
from marvis.decision_twin.batch_material import BatchMaterial
from marvis.decision_twin.temporal import (
    bind_temporal_population,
    measure_temporal_stability,
)
from marvis.domain import TaskCreate
from marvis.repositories.tasks import TaskRepository
from test_decision_twin_temporal import _native_rows, _policy
from test_reference_rule_packages import rule_runtime  # noqa: F401


def _rule_rows():
    records, decisions = _native_rows()
    for decision in decisions:
        decision.update(score=None, score_product=None)
    return records, decisions


def _measure_rule(records, decisions, **overrides):
    policy = _policy()
    args = dict(
        population=bind_temporal_population(records, policy),
        package_hash="a" * 64,
        score_product=None,
        package_kind="rule_only",
    )
    args.update(overrides)
    return measure_temporal_stability(records, decisions, policy, **args)


def test_rule_only_retains_action_drift_without_fabricating_score_or_overall_pass():
    records, decisions = _rule_rows()
    result = _measure_rule(records, decisions)
    assert result["status"] == "partial" and result["verdict"] == "failed"
    assert result["value"]["max_score_psi"] is None
    assert result["value"]["max_absolute_approval_rate_delta"] == 1.0
    assert result["reference"]["score_bins"]["reason"] == "package_has_no_score"
    assert result["reference"]["score_bins"]["fit_sample_count"] == 0
    checks = result["comparisons"][0]["checks"]
    assert checks[0]["value"] is None and checks[0]["passed"] is None
    assert checks[1]["value"] > 0 and checks[1]["passed"] is False
    for row in decisions:
        row["action"] = {"type": "approval"}
    equal = _measure_rule(records, decisions)
    assert equal["status"] == "partial" and equal["verdict"] == "insufficient_evidence"
    assert equal["value"]["max_action_psi"] == 0.0


@pytest.mark.parametrize(
    "kind,product,score",
    [
        ("model", None, None),
        ("other", None, None),
        ("rule_only", "raw_pd", None),
        ("rule_only", None, 0.0),
    ],
)
def test_missing_model_scores_or_invented_rule_scores_are_rejected(
    kind, product, score
):
    records, decisions = _rule_rows()
    for row in decisions:
        row.update(score=score, score_product=product)
    with pytest.raises(ValueError):
        _measure_rule(records, decisions, package_kind=kind, score_product=product)


@pytest.mark.parametrize(
    "key,value",
    [
        ("package_hash", "b" * 64),
        ("facts_hash", "b" * 64),
        ("record_id", "b" * 64),
        ("decision_at", "2026-01-03T00:00:00Z"),
    ],
)
def test_rule_temporal_still_requires_original_member_and_package_binding(key, value):
    records, decisions = _rule_rows()
    decisions = deepcopy(decisions)
    decisions[0][key] = value
    with pytest.raises(ValueError, match="drifted"):
        _measure_rule(records, decisions)


@pytest.mark.parametrize("key", ["score", "score_product"])
def test_rule_receipt_must_explicitly_declare_missing_score(key):
    records, decisions = _rule_rows()
    decisions[0].pop(key)
    with pytest.raises(ValueError):
        _measure_rule(records, decisions)


def test_rule_package_real_workflow_and_exports_keep_missing_score_visible(
    rule_runtime, tmp_path
):  # noqa: F811
    app, _, _, _, _, request, _ = rule_runtime
    package_hash, _ = app.state.reference_decision.packages.build(
        request, actor_id="maker"
    )
    task = TaskRepository(app.state.settings.db_path).create_task(
        TaskCreate(
            model_name="规则时间窗回放",
            model_version="test",
            validator="qa",
            source_dir=str(tmp_path),
            run_mode="manual",
            task_type="modeling",
        )
    )
    material = BatchMaterial(app.state.settings, task.id)
    frame = pd.DataFrame(
        {
            "record_id": [f"loan-{i}" for i in range(12)],
            "decision_at": ["2026-01-02T00:00:00Z"] * 6 + ["2026-02-02T00:00:00Z"] * 6,
            "event_at": ["2026-01-01T00:00:00Z"] * 12,
            "available_at": ["2026-01-01T12:00:00Z"] * 12,
            "x1": [0.1] * 6 + [0.9] * 6,
            "x2": [0.2] * 12,
        }
    )
    path = tmp_path / "rule-history.parquet"
    frame.to_parquet(path, index=False)
    dataset = material.registry.register_existing(
        path, task_id=task.id, role="historical_replay"
    )
    contract = HistoricalReplayRequest(
        dataset_id=dataset.id,
        expected_content_hash=dataset.content_hash,
        record_id_col="record_id",
        decision_at_col="decision_at",
        as_of="2026-08-01T00:00:00Z",
        source_ref="public synthetic reference",
        population="all synthetic applicants",
        features=[
            {
                "name": name,
                "value_col": name,
                "event_at_col": "event_at",
                "available_at_col": "available_at",
            }
            for name in ("x1", "x2")
        ],
        scenarios=[
            {"name": "baseline", "kind": "baseline", "package_hash": package_hash},
            {"name": "challenger", "kind": "challenger", "package_hash": package_hash},
        ],
        temporal_stability=_policy(),
    )
    client = TestClient(app)
    base = f"/api/tasks/{task.id}"
    proposal = client.post(base + "/decision-twin/proposal", json=contract.model_dump())
    assert proposal.status_code == 200, proposal.text
    created = client.post(
        base + "/plans",
        json={
            "goal": "历史决策回放",
            "slots": {
                "replay_contract": contract.model_dump(),
                "proposal_hash": proposal.json()["proposal_hash"],
            },
        },
    )
    assert created.status_code == 201, created.text
    plan = created.json()["plan"]
    assert (
        client.post(
            f"/api/plans/{plan['id']}/confirm", json=plan["confirmation_snapshot"]
        ).status_code
        == 200
    )
    assert client.post(f"/api/plans/{plan['id']}/run").status_code == 202
    plan = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert plan["status"] == "awaiting_confirm"
    step = plan["steps"][0]
    approved = client.post(
        f"/api/plans/{plan['id']}/steps/{step['id']}/decisions",
        json={
            "decision": "approve",
            "reason": "确认规则包无分数，仅核验动作分布与真实未知项",
            **step["confirmation_snapshot"],
        },
    )
    assert approved.status_code == 202, approved.text
    final = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert final["status"] == "done", final
    artifacts = client.get(base + "/decision-twin").json()["artifacts"]
    url = base + "/decision-twin/" + artifacts[0]["artifact_id"]
    receipt = client.get(url).json()
    assert receipt == client.get(url + "/export/json").json()
    scenario = receipt["payload"]["scenarios"][0]
    assert all(
        row["score"] is None and row["score_product"] is None
        for row in scenario["decisions"]
    )
    measured = scenario["metrics"]["stability"]
    assert (
        measured["status"] == "partial" and measured["value"]["max_score_psi"] is None
    )
    assert measured["reference"]["actions"]["counts"] == [6, 0, 0]
    assert measured["comparisons"][0]["actions"]["counts"] == [0, 0, 6]
    assert measured["value"]["max_absolute_approval_rate_delta"] == 1.0
    workbook = load_workbook(BytesIO(client.get(url + "/export/xlsx").content))
    score_row = list(workbook["时间稳定性"].values)[1]
    assert (
        score_row[10] == "score_psi"
        and score_row[11] is None
        and score_row[14] == "unknown"
    )
    document = Document(BytesIO(client.get(url + "/export/docx").content))
    assert "unknown" in " ".join(
        c.text for r in document.tables[0].rows for c in r.cells
    )
