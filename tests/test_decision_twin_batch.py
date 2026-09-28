from io import BytesIO
import json

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pandas as pd
import pytest

from marvis.decision_twin.batch import reconcile_batch, replay_batch
from marvis.decision_twin.batch_contracts import (
    HistoricalReconciliationRequest,
    HistoricalReplayRequest,
)
from marvis.decision_twin.batch_material import BatchMaterial
from marvis.domain import TaskCreate, TASK_TYPE_MODELING
from marvis.plugins.manifest import ToolRef
from marvis.reference_decision.contracts import DecisionError
from marvis.reference_decision.evaluation import evaluate
from marvis.repositories.tasks import TaskRepository

from test_reference_decision import packaged  # noqa: F401


@pytest.fixture
def batch(packaged, tmp_path):  # noqa: F811 - imported pytest fixture
    app, _, _, package_hash, _, _, frame, _ = packaged
    task = TaskRepository(app.state.settings.db_path).create_task(
        TaskCreate(
            model_name="历史回放",
            model_version="test",
            validator="qa",
            source_dir=str(tmp_path),
            run_mode="manual",
            task_type=TASK_TYPE_MODELING,
        )
    )
    material = BatchMaterial(app.state.settings, task.id)
    sample = frame.iloc[:12][["x1", "x2"]].copy().reset_index(drop=True)
    sample["record_id"] = [f"loan-{i}" for i in range(len(sample))]
    sample["decision_at"] = "2026-01-02T00:00:00Z"
    sample["available_at"] = "2026-01-01T12:00:00Z"
    sample["event_at"] = "2026-01-01T00:00:00Z"
    sample["action_at"] = "2026-01-02T00:00:00Z"
    sample["ead"] = 1000.0
    sample["group"] = ["A", "B"] * 6
    sample["actual_action"] = "approval"
    path = tmp_path / "history.parquet"
    sample.to_parquet(path, index=False)
    dataset = material.registry.register_existing(
        path, task_id=task.id, role="historical_replay"
    )
    contract = HistoricalReplayRequest(
        dataset_id=dataset.id,
        expected_content_hash=dataset.content_hash,
        record_id_col="record_id",
        decision_at_col="decision_at",
        as_of="2026-08-01T00:00:00Z",
        source_ref="脱敏历史导出",
        population="完整申请人群",
        features=[
            {
                "name": name,
                "value_col": name,
                "event_at_col": "event_at",
                "available_at_col": "available_at",
            }
            for name in ("x1", "x2")
        ],
        scenarios=[
            {"name": "基线冻结包", "kind": "baseline", "package_hash": package_hash},
            {"name": "挑战冻结包", "kind": "challenger", "package_hash": package_hash},
        ],
    )
    return app, material, contract, sample, tmp_path


def _replay(material, contract):
    proposal = material.prepare(contract)[2]
    return replay_batch(material, contract, proposal["proposal_hash"])


def _changed_dataset(batch, **changes):
    _, material, contract, sample, path = batch
    sample = sample.copy()
    for key, value in changes.items():
        sample[key] = value
    target = path / "changed.parquet"
    sample.to_parquet(target, index=False)
    dataset = material.registry.register_existing(
        target, task_id=material.task_id, role="historical_replay"
    )
    return contract.model_copy(
        update={"dataset_id": dataset.id, "expected_content_hash": dataset.content_hash}
    )


def _economic_contract(contract):
    return HistoricalReplayRequest.model_validate(
        {
            **contract.model_dump(),
            "economics": {
                "currency": "CNY",
                "ead_col": "ead",
                "available_at_col": "available_at",
                "annual_rate": 0.18,
                "funding_rate": 0.03,
                "lgd": 0.5,
                "term_months": 12.0,
                "operating_cost_per_loan": 8.0,
                "assumption_sources": {
                    key: "人工确认的固定情景假设"
                    for key in (
                        "ead",
                        "pd",
                        "annual_rate",
                        "funding_rate",
                        "lgd",
                        "term_months",
                        "operating_cost_per_loan",
                    )
                },
            },
            "protected_group": {
                "column": "group",
                "available_at_col": "available_at",
                "governance_ref": "测试授权分组口径",
                "minimum_group_size": 2,
            },
            "capacity": {
                "unit": "review_minutes",
                "per_action": {"approval": 1.0, "reject": 1.0, "review": 15.0},
                "source_ref": "人工工时假设",
            },
            "observed_actions": {
                "action_col": "actual_action",
                "recorded_at_col": "action_at",
                "source_ref": "外部历史审批流水",
            },
        }
    )


def test_batch_native_scores_reproducible_receipt_and_honest_unknowns(batch):
    _, material, contract, sample, _ = batch
    receipt = _replay(material, contract)
    assert _replay(material, contract)["artifact_id"] == receipt["artifact_id"]
    verified = material.load(receipt["artifact_id"])
    assert verified["payload"] == receipt["payload"]
    output = verified["payload"]
    assert output["historical_package_availability"] == "not_established"
    assert output["causal_assessment"]["identifiability"] == "unidentified"
    assert output["observed_actions"]["status"] == "unknown"
    for scenario in output["scenarios"]:
        assert scenario["observed_execution"] is False
        assert scenario["comparison"]["decision_change_count"] == 0
        assert scenario["metrics"]["count"] == len(sample)
        for key in (
            "economics",
            "protected_groups",
            "operations_capacity",
            "stability",
        ):
            assert scenario["metrics"][key]["value"] is None
        with material.packages.snapshot(scenario["package_hash"]) as (
            manifest,
            directory,
        ):
            expected = evaluate(
                manifest, directory, sample.iloc[0][["x1", "x2"]].to_dict()
            )
        assert scenario["decisions"][0]["score"] == expected["score"]
        assert scenario["decisions"][0]["action"] == expected["action"]
    assert "loan-0" not in json.dumps(output)
    assert "features" not in output["scenarios"][0]["decisions"][0]


@pytest.mark.parametrize(
    "changes, code",
    [
        ({"available_at": "2026-01-03T00:00:00Z"}, "not_available"),
        ({"event_at": "2026-01-01T13:00:00Z"}, "not_available"),
        ({"available_at": None}, "invalid_timestamp"),
        ({"available_at": "2026-01-01"}, "invalid_timestamp"),
        ({"record_id": "duplicate"}, "duplicate"),
        ({"decision_at": "2026-09-01T00:00:00Z"}, "cutoff"),
    ],
)
def test_future_missing_naive_and_duplicate_history_rejected(batch, changes, code):
    _, material, _, _, _ = batch
    contract = _changed_dataset(batch, **changes)
    with pytest.raises(DecisionError, match=code):
        _replay(material, contract)
    assert material.artifacts.list_for_task(material.task_id) == []


def test_stale_proposal_cross_task_and_schema_mismatch_rejected(batch):
    app, material, contract, _, path = batch
    with pytest.raises(DecisionError, match="proposal_stale"):
        replay_batch(material, contract, "0" * 64)
    other = TaskRepository(app.state.settings.db_path).create_task(
        TaskCreate(
            model_name="other",
            model_version="test",
            validator="qa",
            source_dir=str(path),
        )
    )
    with pytest.raises(Exception, match="task|owned"):
        BatchMaterial(app.state.settings, other.id).prepare(contract)
    bad = contract.model_copy(update={"features": contract.features[:1]})
    with pytest.raises(DecisionError, match="schema_mismatch"):
        material.prepare(bad)


def test_economic_values_use_existing_row_engine_and_do_not_assert_observation(batch):
    _, material, contract, _, _ = batch
    contract = _economic_contract(contract)
    receipt = _replay(material, contract)["payload"]
    first = receipt["scenarios"][0]
    approved = [r for r in first["decisions"] if r["action"]["type"] == "approval"]
    measured = first["metrics"]["economics"]
    assert measured["status"] == "estimated"
    assert measured["value"]["ead"] == 1000 * len(approved)
    assert measured["value"]["expected_loss"] == pytest.approx(
        sum(r["score"] * 500 for r in approved)
    )
    assert measured["value"]["profit"] == pytest.approx(
        sum(150 - r["score"] * 500 - 8 for r in approved)
    )
    assert first["metrics"]["protected_groups"]["status"] == "measured"
    assert first["metrics"]["operations_capacity"]["value"] == 12
    assert receipt["observed_actions"]["origin"] == "external_history_import"
    assert receipt["observed_actions"]["marvis_execution_verified"] is False


def test_tampered_package_and_receipt_rejected(batch):
    _, material, contract, _, _ = batch
    receipt = _replay(material, contract)
    manifest = material.packages.get(contract.scenarios[0].package_hash)
    path = (
        material.packages.root
        / contract.scenarios[0].package_hash
        / manifest["files"][0]["path"]
    )
    original = path.read_bytes()
    try:
        path.write_bytes(b"tampered")
        with pytest.raises(DecisionError, match="integrity"):
            _replay(material, contract)
    finally:
        path.write_bytes(original)
    path = material.store.artifact_path(receipt["artifact_id"])
    original = path.read_bytes()
    try:
        path.write_text("{}")
        with pytest.raises(RuntimeError, match="hash|artifact"):
            material.load(receipt["artifact_id"])
    finally:
        path.write_bytes(original)


def test_unsupported_constraints_are_unknown_and_missing_sources_rejected(batch):
    _, material, contract, _, _ = batch
    modified = HistoricalReplayRequest.model_validate(
        {
            **contract.model_dump(),
            "constraints": [
                {"metric": "profit", "operator": ">=", "threshold": 0.0, "unit": "CNY"}
            ],
        }
    )
    for scenario in _replay(material, modified)["payload"]["scenarios"]:
        assert scenario["constraints"]["status"] == "insufficient_evidence"
        assert scenario["constraints"]["checks"][0]["value"] is None
    modified = _economic_contract(contract).model_dump()
    modified["economics"]["assumption_sources"] = {}
    with pytest.raises(ValueError, match="source"):
        HistoricalReplayRequest.model_validate(modified)


def _outcomes(batch, receipt):
    _, material, _, sample, path = batch
    frame = pd.DataFrame(
        {
            "id": sample["record_id"],
            "observed_at": "2026-05-01T00:00:00Z",
            "loss": 0.0,
            "profit": 80.0,
        }
    )
    target = path / "outcomes.parquet"
    frame.to_parquet(target, index=False)
    data = material.registry.register_existing(
        target, task_id=material.task_id, role="historical_outcomes"
    )
    contract = HistoricalReconciliationRequest(
        replay_artifact_id=receipt["artifact_id"],
        dataset_id=data.id,
        expected_content_hash=data.content_hash,
        record_id_col="id",
        observed_at_col="observed_at",
        actual_loss_col="loss",
        actual_profit_col="profit",
        currency="CNY",
        maturity_days=90,
        maturity_source_ref="90d cashflow contract",
        source_ref="external ledger export",
        reconciled_at="2026-06-01T00:00:00Z",
    )
    return contract


def test_mature_outcomes_match_imported_approvals_and_preserve_causal_boundary(batch):
    _, material, contract, _, _ = batch
    replay = _replay(material, _economic_contract(contract))
    request = _outcomes(batch, replay)
    receipt = reconcile_batch(material, request, request.contract_hash)["payload"]
    assert receipt["actual_loss"] == 0
    assert receipt["actual_profit"] == 12 * 80
    assert receipt["approved_denominator"] == 12
    assert receipt["comparable_denominator"] < 12
    assert receipt["status"] == "partial_comparability"
    assert receipt["marvis_execution_verified"] is False
    assert receipt["causal_gain_verified"] is False
    immature = request.model_copy(update={"maturity_days": 365})
    with pytest.raises(DecisionError, match="not_mature"):
        reconcile_batch(material, immature, immature.contract_hash)
    missing = request.model_copy(
        update={"replay_artifact_id": _replay(material, contract)["artifact_id"]}
    )
    with pytest.raises(DecisionError, match="observed_historical_actions_required"):
        reconcile_batch(material, missing, missing.contract_hash)


def test_real_api_workflow_human_gate_toolrunner_and_canonical_exports(batch):
    app, material, contract, _, _ = batch
    client = TestClient(app)
    proposal = client.post(
        f"/api/tasks/{material.task_id}/decision-twin/proposal",
        json=contract.model_dump(),
    )
    assert proposal.status_code == 200, proposal.text
    created = client.post(
        f"/api/tasks/{material.task_id}/plans",
        json={
            "goal": "历史决策回放",
            "slots": {
                "replay_contract": contract.model_dump(),
                "proposal_hash": proposal.json()["proposal_hash"],
            },
        },
    )
    assert created.status_code == 201, created.text
    plan = created.json()["plan"]
    confirmation = client.post(
        f"/api/plans/{plan['id']}/confirm", json=plan["confirmation_snapshot"]
    )
    assert confirmation.status_code == 200, confirmation.text
    started = client.post(f"/api/plans/{plan['id']}/run")
    assert started.status_code == 202, started.text
    current = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert current["status"] == "awaiting_confirm", current
    assert material.artifacts.list_for_task(material.task_id) == []
    step = current["steps"][0]
    response = client.post(
        f"/api/plans/{plan['id']}/steps/{step['id']}/decisions",
        json={
            "decision": "approve",
            "reason": "已核对历史时点、包版本与未知证据边界",
            **step["confirmation_snapshot"],
        },
    )
    assert response.status_code == 202, response.text
    current = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert current["status"] == "done", current
    listing = client.get(f"/api/tasks/{material.task_id}/decision-twin").json()[
        "artifacts"
    ]
    assert len(listing) == 1
    identity = listing[0]["artifact_id"]
    url = f"/api/tasks/{material.task_id}/decision-twin/{identity}"
    detail = client.get(url).json()
    assert detail == client.get(url + "/export/json").json()
    xlsx = client.get(url + "/export/xlsx")
    workbook = load_workbook(BytesIO(xlsx.content))
    assert workbook["回放证据"]["B1"].value == identity
    assert workbook["逐笔回放"].max_row == 25
    document = Document(BytesIO(client.get(url + "/export/docx").content))
    assert identity in "\n".join(p.text for p in document.paragraphs)
    assert "loan-0" not in json.dumps(detail)


def test_toolrunner_cannot_bypass_human_gate(batch):
    app, material, contract, _, _ = batch
    runner = app.state.tool_runner
    proposal = material.prepare(contract)[2]
    result = runner.invoke(
        ToolRef("decision_twin", "replay_history"),
        {"contract": contract.model_dump(), "proposal_hash": proposal["proposal_hash"]},
        task_id=material.task_id,
    )
    assert not result.ok
    assert "governance" in result.error


def test_distinct_native_packages_compare_identical_population(batch, packaged):  # noqa: F811 - imported pytest fixture
    from dataclasses import replace
    from marvis.db_schema import connect
    from marvis.packs.strategy.dsl import StrategyAction, StrategyRuleSpec, StrategySpec
    from marvis.packs.strategy.strategy import build_strategy_from_spec
    from marvis.repositories.strategy import StrategyRepository

    _, material, contract, _, _ = batch
    _, store, request, _, _, _, _, _ = packaged
    spec = StrategySpec(
        strategy_type="approval",
        rules=(
            StrategyRuleSpec(
                rule_id="strict_pd",
                priority=1,
                condition={
                    "op": "compare",
                    "field": "pd",
                    "operator": "<",
                    "value": 0.1,
                },
                action=StrategyAction(type="approval", reason_code="STRICT_PD"),
            ),
        ),
        default_action=StrategyAction(type="reject", reason_code="HIGH_PD"),
    )
    strategy = replace(
        build_strategy_from_spec(
            spec, score_col="pd", description="distinct native replay"
        ),
        id="replay-strict-" + material.task_id,
    )
    StrategyRepository(material.settings.db_path).create_strategy(
        material.task_id, strategy
    )
    with connect(material.settings.db_path) as conn:
        conn.execute(
            "UPDATE strategies SET status='validated',asset_status='validated' WHERE id=?",
            (strategy.id,),
        )
    identity, _ = store.build(
        request.model_copy(update={"strategy_id": strategy.id}), actor_id="test-maker"
    )
    payload = _economic_contract(contract).model_dump()
    payload["scenarios"][1]["package_hash"] = identity
    output = _replay(material, HistoricalReplayRequest.model_validate(payload))[
        "payload"
    ]
    baseline, challenger = output["scenarios"]
    assert baseline["population_hash"] == challenger["population_hash"]
    assert baseline["metrics"]["count"] == challenger["metrics"]["count"] == 12
    assert challenger["comparison"]["decision_change_count"] > 0
    assert challenger["comparison"]["approval_rate_delta"] < 0
    assert challenger["comparison"]["estimated_economics_delta"]["ead"] < 0
    assert challenger["comparison"]["causal_gain_verified"] is False


def test_missing_group_support_stays_unknown(batch):
    _, material, contract, _, _ = batch
    payload = _economic_contract(contract).model_dump()
    payload["protected_group"]["minimum_group_size"] = 100
    result = _replay(material, HistoricalReplayRequest.model_validate(payload))
    assert (
        result["payload"]["scenarios"][0]["metrics"]["protected_groups"]["value"]
        is None
    )


def test_source_hash_and_row_budget_fail_before_model_execution(batch, monkeypatch):
    from marvis.data.errors import DatasetContentDriftError
    import marvis.decision_twin.batch_material as module

    _, material, contract, _, _ = batch
    with pytest.raises(DatasetContentDriftError):
        _replay(
            material, contract.model_copy(update={"expected_content_hash": "0" * 64})
        )
    monkeypatch.setattr(module, "MAX_BATCH_ROWS", 5)
    with pytest.raises(DecisionError, match="row_budget"):
        material.prepare(contract)
    assert material.artifacts.list_for_task(material.task_id) == []


def test_forged_audit_entry_does_not_become_platform_execution(batch):
    from marvis.decision_twin.batch_material import PRODUCER, REPLAY_KIND

    _, material, contract, _, _ = batch
    real = _replay(material, contract)
    envelope = material.store.get(real["artifact_id"])
    envelope["body"]["payload"]["scenarios"][0]["decisions"][0]["score"] = 0.123
    receipt = material.store.put(
        REPLAY_KIND, envelope, idempotency_key="forged-recomputed-unkeyed-hash"
    )
    material.artifacts.register(
        task_id=material.task_id,
        kind=REPLAY_KIND,
        path=str(material.store.artifact_path(receipt.artifact_hash)),
        content_hash=receipt.artifact_hash,
        origin_tool=PRODUCER,
        provenance={"artifact_id": receipt.artifact_hash},
    )
    with pytest.raises(DecisionError, match="authentication_failed"):
        material.load(receipt.artifact_hash)


def test_source_mutation_before_publication_cannot_publish_a_receipt(batch):
    from marvis.data.errors import DatasetContentDriftError
    from marvis.decision_twin.batch_material import REPLAY_KIND

    _, material, contract, _, _ = batch
    binding, _, _ = material.prepare(contract)
    original = binding.path.read_bytes()
    try:
        binding.path.chmod(0o600)
        binding.path.write_bytes(b"changed after snapshot")
        with pytest.raises(DatasetContentDriftError):
            material.publish(
                REPLAY_KIND, {"contract_hash": contract.contract_hash}, binding=binding
            )
        assert material.artifacts.list_for_task(material.task_id) == []
    finally:
        binding.path.write_bytes(original)
        binding.path.chmod(0o444)


def test_agent_start_enters_the_same_explicit_human_gate(batch):
    from marvis.db_schema import connect

    app, material, contract, _, _ = batch
    with connect(material.settings.db_path) as conn:
        conn.execute(
            "UPDATE tasks SET run_mode='agent' WHERE id=?", (material.task_id,)
        )
    client = TestClient(app)
    proposal = client.post(
        f"/api/tasks/{material.task_id}/decision-twin/proposal",
        json=contract.model_dump(),
    )
    created = client.post(
        f"/api/tasks/{material.task_id}/plans",
        json={
            "goal": "历史决策回放",
            "slots": {
                "replay_contract": contract.model_dump(),
                "proposal_hash": proposal.json()["proposal_hash"],
            },
        },
    )
    assert created.status_code == 201, created.text
    plan = created.json()["plan"]
    started = client.post(
        f"/api/tasks/{material.task_id}/agent/messages",
        json={
            "content": "开始",
            "ui_action": "start_plan",
            "expected_plan_id": plan["id"],
            **plan["confirmation_snapshot"],
        },
    )
    assert started.status_code == 202, started.text
    current = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert current["status"] == "awaiting_confirm", current
    step = current["steps"][0]
    response = client.post(
        f"/api/plans/{plan['id']}/steps/{step['id']}/decisions",
        json={
            "decision": "approve",
            "reason": "核对提案后授权这一次历史回放",
            **step["confirmation_snapshot"],
        },
    )
    assert response.status_code == 202, response.text
    completed = client.get(f"/api/plans/{plan['id']}").json()["plan"]
    assert completed["status"] == "done", completed
