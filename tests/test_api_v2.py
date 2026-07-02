from dataclasses import asdict
import json
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import unquote

import nbformat
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from marvis.api import _agent_has_stop_ack_message, router
from marvis.domain import (
    TASK_STATUS_REASON_SERVER_RESTART,
    TASK_STATUS_REASON_USER_CANCELLED,
    TaskRecord,
    TaskStatus,
)
from marvis.execution_environment import ExecutionEnvironmentOption
from marvis.routers.branding import router as branding_router
from marvis.routers.evidence import router as evidence_router
from marvis.routers.materials import router as materials_router
from marvis.routers.report_fields import router as report_fields_router
from marvis.routers.reports import router as reports_router
from marvis.routers.scans import router as scans_router
from marvis.routers.stage_controls import router as stage_controls_router
from marvis.routers.tasks import router as tasks_router
from marvis.routers.validation_agent import router as validation_agent_router
from marvis.routers.validation_stages import router as validation_stages_router
from marvis.pipeline import LEGACY_LIVE_NOTEBOOK_ENV_VAR, PipelineSettings


class FakeTaskRepository:
    tasks: dict[str, TaskRecord] = {}
    deleted: list[str] = []
    report_values: dict[str, tuple[dict[str, str], int]] = {}
    jobs: dict[str, dict[str, str]] = {}
    audits: list[dict] = []

    def __init__(self, _db_path: Path):
        pass

    def create_task(self, payload):
        task_id = f"task-{len(self.tasks) + 1}"
        task = TaskRecord(
            id=task_id,
            task_type=payload.task_type,
            model_name=payload.model_name,
            model_version=payload.model_version,
            validator=payload.validator,
            source_dir=payload.source_dir,
            algorithm=payload.algorithm,
            run_mode=payload.run_mode,
            target_col=payload.target_col,
            score_col=payload.score_col,
            split_col=payload.split_col,
            time_col=payload.time_col,
            feature_columns=payload.feature_columns,
            notebook_path=payload.notebook_path,
            sample_path=payload.sample_path,
            pmml_path=payload.pmml_path,
            dictionary_path=payload.dictionary_path,
            report_values_revision=0,
            status=TaskStatus.CREATED,
            status_message="created",
            created_at="2026-05-21T00:00:00+00:00",
            updated_at="2026-05-21T00:00:00+00:00",
        )
        self.tasks[task_id] = task
        self.report_values[task_id] = (dict(payload.report_values), 0)
        return task

    def start_job(self, task_id: str, kind: str) -> str:
        from marvis.state_machine import ConflictError

        self.get_task(task_id)
        if self.task_has_active_job(task_id):
            raise ConflictError(f"task {task_id} already has an active job")
        job_id = f"job-{len(self.jobs) + 1}"
        self.jobs[job_id] = {
            "id": job_id,
            "task_id": task_id,
            "kind": kind,
            "status": "queued",
        }
        return job_id

    def mark_job_running(self, job_id: str) -> None:
        self.jobs[job_id]["status"] = "running"

    def finish_job(
        self,
        job_id: str,
        *,
        status: str,
        error_name: str | None = None,
        error_value: str | None = None,
        traceback: str | None = None,
    ) -> None:
        self.jobs[job_id].update(
            {
                "status": status,
                "error_name": error_name or "",
                "error_value": error_value or "",
                "traceback": traceback or "",
            }
        )

    def task_has_active_job(self, task_id: str) -> bool:
        self.get_task(task_id)
        return any(
            job["task_id"] == task_id and job["status"] in {"queued", "running"}
            for job in self.jobs.values()
        )

    def get_active_job_kind(self, task_id: str) -> str | None:
        self.get_task(task_id)
        for job in self.jobs.values():
            if job["task_id"] == task_id and job["status"] in {"queued", "running"}:
                return job["kind"]
        return None

    def get_latest_failed_job_kind(self, task_id: str) -> str | None:
        self.get_task(task_id)
        for job in reversed(list(self.jobs.values())):
            if job["task_id"] == task_id and job["status"] == "failed":
                return job["kind"]
        return None

    def get_latest_job(self, task_id: str, *, kind: str | None = None) -> dict | None:
        self.get_task(task_id)
        for job in reversed(list(self.jobs.values())):
            if job["task_id"] != task_id:
                continue
            if kind and job["kind"] != kind:
                continue
            return dict(job)
        return None

    def get_task(self, task_id: str):
        try:
            return self.tasks[task_id]
        except KeyError as exc:
            raise KeyError(f"Task not found: {task_id}") from exc

    def list_tasks(self, *, limit: int | None = None, offset: int = 0):
        tasks = list(self.tasks.values())
        start = max(0, int(offset))
        if limit is None:
            return tasks[start:]
        return tasks[start:start + max(1, int(limit))]

    def delete_task(self, task_id: str):
        self.get_task(task_id)
        self.deleted.append(task_id)
        del self.tasks[task_id]

    def update_status(
        self,
        task_id: str,
        status,
        message,
        *,
        expected=None,
        reason_code: str = "",
    ):
        from marvis.state_machine import IllegalTransition

        task = self.get_task(task_id)
        if expected is not None:
            expected_set = {expected} if isinstance(expected, TaskStatus) else set(expected)
            if task.status not in expected_set:
                raise IllegalTransition(task.status, status)
        self.tasks[task_id] = TaskRecord(
            **{
                **asdict(task),
                "status": status,
                "status_message": message,
                "status_reason_code": reason_code,
                "updated_at": "2026-05-21T00:01:00+00:00",
            }
        )

    def update_status_message(
        self,
        task_id: str,
        message: str,
        *,
        reason_code: str | None = None,
    ):
        task = self.get_task(task_id)
        values = {**asdict(task), "status_message": message}
        if reason_code is not None:
            values["status_reason_code"] = reason_code
        self.tasks[task_id] = TaskRecord(**values)

    def get_report_values(self, task_id: str):
        self.get_task(task_id)
        return self.report_values[task_id]

    def update_report_values(self, task_id: str, values, expected_revision: int):
        from marvis.state_machine import ConflictError
        from marvis.report_texts import COMPUTED_REPORT_TEXT_KEYS

        self.get_task(task_id)
        computed_keys = sorted(set(values) & COMPUTED_REPORT_TEXT_KEYS)
        if computed_keys:
            raise ValueError("platform-computed report values cannot be updated")
        current, revision = self.report_values[task_id]
        if revision != expected_revision:
            raise ConflictError("stale report values revision")
        merged = {**current, **values}
        new_revision = revision + 1
        self.report_values[task_id] = (merged, new_revision)
        current_task = self.tasks[task_id]
        self.tasks[task_id] = TaskRecord(
            **{**asdict(current_task), "report_values_revision": new_revision}
        )
        return new_revision

    def update_report_values_with_audit(
        self, task_id: str, values, expected_revision: int, *, audit: dict
    ) -> int:
        new_revision = self.update_report_values(task_id, values, expected_revision)
        self.audits.append(audit)
        return new_revision

    def update_agent_report_conclusions(self, task_id: str, values, expected_revision: int) -> int:
        return self.update_report_values(task_id, values, expected_revision)

    def update_agent_report_conclusions_with_audit(
        self, task_id: str, values, expected_revision: int, *, audit: dict
    ) -> int:
        new_revision = self.update_agent_report_conclusions(task_id, values, expected_revision)
        self.audits.append(audit)
        return new_revision


class FakeHookDispatcher:
    def __init__(self):
        self.calls = []

    def dispatch(self, event: str, payload: dict, *, task_id: str) -> None:
        self.calls.append((event, payload, task_id))


def _client(tmp_path: Path, monkeypatch) -> TestClient:
    FakeTaskRepository.tasks = {}
    FakeTaskRepository.deleted = []
    FakeTaskRepository.report_values = {}
    FakeTaskRepository.jobs = {}
    FakeTaskRepository.audits = []
    monkeypatch.setattr("marvis.api.TaskRepository", FakeTaskRepository)
    monkeypatch.setattr("marvis.api_stage_helpers.TaskRepository", FakeTaskRepository)
    monkeypatch.setattr("marvis.routers.evidence.TaskRepository", FakeTaskRepository)
    monkeypatch.setattr("marvis.routers.report_fields.TaskRepository", FakeTaskRepository)
    monkeypatch.setattr("marvis.routers.reports.TaskRepository", FakeTaskRepository)
    monkeypatch.setattr("marvis.routers.scans.TaskRepository", FakeTaskRepository)
    monkeypatch.setattr("marvis.routers.stage_controls.TaskRepository", FakeTaskRepository)
    monkeypatch.setattr("marvis.routers.tasks.TaskRepository", FakeTaskRepository)
    monkeypatch.setattr(
        "marvis.routers.validation_stages.TaskRepository",
        FakeTaskRepository,
    )

    app = FastAPI()
    app.state.settings = SimpleNamespace(
        workspace=tmp_path,
        tasks_dir=tmp_path / "tasks",
        db_path=tmp_path / "marvis.sqlite",
        report_template_path=tmp_path / "report_templates" / "default.docx",
    )
    app.state.settings.tasks_dir.mkdir()
    app.include_router(router)
    app.include_router(evidence_router)
    app.include_router(materials_router)
    app.include_router(report_fields_router)
    app.include_router(reports_router)
    app.include_router(scans_router)
    app.include_router(stage_controls_router)
    app.include_router(tasks_router)
    app.include_router(validation_agent_router)
    app.include_router(validation_stages_router)
    return TestClient(app)


def _decoded_download_filename(response) -> str:
    header = response.headers["content-disposition"]
    return unquote(header).split("''")[-1].strip('"')


def test_create_then_get_task(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    response = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
            "algorithm": "lgb",
            "target_col": "target",
            "score_col": "score",
            "split_col": "split",
            "time_col": "month",
            "run_mode": "agent",
            "feature_columns": ["x1", "x2"],
            "notebook_path": str(tmp_path / "model.ipynb"),
            "sample_path": str(tmp_path / "sample.csv"),
            "pmml_path": str(tmp_path / "model.pmml"),
            "dictionary_path": str(tmp_path / "dictionary.xlsx"),
            "report_values": {"TEXT:report_title": "自定义标题"},
        },
    )
    assert response.status_code == 200, response.text
    task = response.json()
    assert task["task_type"] == "validation"
    assert task["algorithm"] == "lgb"
    assert task["target_col"] == "target"
    assert task["run_mode"] == "agent"
    assert task["feature_columns"] == ["x1", "x2"]
    assert task["notebook_path"] == str(tmp_path / "model.ipynb")
    assert task["sample_path"] == str(tmp_path / "sample.csv")
    assert task["pmml_path"] == str(tmp_path / "model.pmml")
    assert task["dictionary_path"] == str(tmp_path / "dictionary.xlsx")
    assert task["report_values_revision"] == 0
    assert FakeTaskRepository.report_values[task["id"]] == (
        {"TEXT:report_title": "自定义标题"},
        0,
    )

    got = client.get(f"/api/tasks/{task['id']}")
    assert got.status_code == 200
    assert got.json()["model_name"] == "A卡"
    assert got.json()["task_type"] == "validation"
    assert got.json()["active_job_kind"] is None
    assert got.json()["report_available"] is False


def test_material_upload_preserves_folder_paths_under_workspace(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    response = client.post(
        "/api/material-uploads",
        data={
            "relative_paths": [
                "submitted/model.ipynb",
                "submitted/data/sample.csv",
            ],
        },
        files=[
            ("files", ("model.ipynb", b"{}", "application/x-ipynb+json")),
            ("files", ("sample.csv", b"x,y\n1,0\n", "text/csv")),
        ],
    )

    assert response.status_code == 201, response.text
    payload = response.json()
    source_dir = Path(payload["source_dir"])
    assert source_dir.parent == tmp_path / "material_uploads"
    assert (source_dir / "submitted" / "model.ipynb").read_bytes() == b"{}"
    assert (source_dir / "submitted" / "data" / "sample.csv").read_bytes() == b"x,y\n1,0\n"
    assert payload["files"] == [
        {"relative_path": "submitted/model.ipynb", "size_bytes": 2},
        {"relative_path": "submitted/data/sample.csv", "size_bytes": 8},
    ]


def test_material_upload_rejects_paths_outside_upload_directory(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)

    response = client.post(
        "/api/material-uploads",
        data={"relative_paths": "../escape.csv"},
        files={"files": ("escape.csv", b"x,y\n1,0\n", "text/csv")},
    )

    assert response.status_code == 422
    assert "invalid upload path" in response.json()["detail"]


def test_material_upload_route_is_served_from_dedicated_router():
    routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in materials_router.routes
    }

    assert routes[("/api/material-uploads", ("POST",))] == "marvis.routers.materials"


def test_report_download_routes_are_served_from_dedicated_router():
    routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in reports_router.routes
    }

    assert routes[("/api/tasks/{task_id}/report/download", ("GET",))] == (
        "marvis.routers.reports"
    )
    assert routes[("/api/tasks/{task_id}/report/preview", ("GET",))] == (
        "marvis.routers.reports"
    )
    assert routes[("/api/tasks/{task_id}/analysis/download", ("GET",))] == (
        "marvis.routers.reports"
    )
    assert routes[("/api/tasks/{task_id}/driver-report/download", ("GET",))] == (
        "marvis.routers.reports"
    )


def test_scan_route_is_served_from_dedicated_router():
    routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in scans_router.routes
    }

    assert routes[("/api/tasks/{task_id}/scan", ("POST",))] == "marvis.routers.scans"


def test_evidence_route_is_served_from_dedicated_router():
    routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in evidence_router.routes
    }

    assert routes[("/api/tasks/{task_id}/evidence", ("GET",))] == (
        "marvis.routers.evidence"
    )


def test_report_field_routes_are_served_from_dedicated_router():
    routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in report_fields_router.routes
    }

    assert routes[("/api/tasks/{task_id}/report-fields", ("GET",))] == (
        "marvis.routers.report_fields"
    )
    assert routes[("/api/tasks/{task_id}/report-fields", ("PUT",))] == (
        "marvis.routers.report_fields"
    )


def test_stage_cancel_routes_are_served_from_dedicated_router():
    routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in stage_controls_router.routes
    }

    assert routes[("/api/tasks/{task_id}/notebook/cancel", ("POST",))] == (
        "marvis.routers.stage_controls"
    )
    assert routes[("/api/tasks/{task_id}/metrics/cancel", ("POST",))] == (
        "marvis.routers.stage_controls"
    )
    assert routes[("/api/tasks/{task_id}/report/cancel", ("POST",))] == (
        "marvis.routers.stage_controls"
    )


def test_validation_stage_routes_are_served_from_dedicated_router():
    routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in validation_stages_router.routes
    }

    assert routes[("/api/tasks/{task_id}/notebook", ("POST",))] == (
        "marvis.routers.validation_stages"
    )
    assert routes[("/api/tasks/{task_id}/metrics", ("POST",))] == (
        "marvis.routers.validation_stages"
    )
    assert routes[("/api/tasks/{task_id}/report", ("POST",))] == (
        "marvis.routers.validation_stages"
    )
    assert routes[("/api/tasks/{task_id}/validate", ("POST",))] == (
        "marvis.routers.validation_stages"
    )


def test_validation_agent_routes_are_served_from_dedicated_router():
    routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in validation_agent_router.routes
    }

    assert routes[("/api/tasks/{task_id}/agent/messages", ("GET",))] == (
        "marvis.routers.validation_agent"
    )
    assert routes[("/api/tasks/{task_id}/agent/start", ("POST",))] == (
        "marvis.routers.validation_agent"
    )
    assert routes[("/api/tasks/{task_id}/agent/messages", ("POST",))] == (
        "marvis.routers.validation_agent"
    )
    assert routes[("/api/tasks/{task_id}/agent/stop", ("POST",))] == (
        "marvis.routers.validation_agent"
    )
    assert routes[("/api/tasks/{task_id}/agent/summarize", ("POST",))] == (
        "marvis.routers.validation_agent"
    )
    assert routes[("/api/tasks/{task_id}/agent/report-draft", ("POST",))] == (
        "marvis.routers.validation_agent"
    )
    assert routes[("/api/tasks/{task_id}/agent/report-draft/confirm", ("POST",))] == (
        "marvis.routers.validation_agent"
    )


def test_create_task_dispatches_task_created_hook(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    dispatcher = FakeHookDispatcher()
    client.app.state.hook_dispatcher = dispatcher

    response = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
            "algorithm": "lgb",
            "run_mode": "agent",
        },
    )

    assert response.status_code == 200
    task = response.json()
    assert dispatcher.calls == [
        (
            "task.created",
            {
                "task_id": task["id"],
                "task_type": "validation",
                "status": "created",
                "run_mode": "agent",
                "algorithm": "lgb",
            },
            task["id"],
        )
    ]


def test_scan_task_dispatches_task_scanned_hook(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    dispatcher = FakeHookDispatcher()
    client.app.state.hook_dispatcher = dispatcher
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    dispatcher.calls.clear()

    def fake_scan(repo, task, _settings):
        return {
            "task_id": task.id,
            "status": "scanned",
            "status_message": "材料扫描完成。",
            "checks": [{"code": "notebook_contract", "status": "pass"}],
        }

    monkeypatch.setattr("marvis.routers.scans.perform_scan_task", fake_scan)

    response = client.post(f"/api/tasks/{task_id}/scan")

    assert response.status_code == 200
    assert dispatcher.calls == [
        (
            "task.scanned",
            {
                "task_id": task_id,
                "status": "scanned",
                "status_message": "材料扫描完成。",
                "check_count": 1,
                "failed_check_codes": [],
            },
            task_id,
        )
    ]


def test_task_payload_exposes_active_job_kind_for_reloaded_ui(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.jobs["job-report"] = {
        "id": "job-report",
        "task_id": task_id,
        "kind": "report",
        "status": "queued",
    }

    got = client.get(f"/api/tasks/{task_id}")
    listed = client.get("/api/tasks")

    assert got.status_code == 200
    assert got.json()["active_job_kind"] == "report"
    assert listed.status_code == 200
    assert listed.json()[0]["active_job_kind"] == "report"


def test_latest_task_job_endpoint_exposes_job_error_without_traceback(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.jobs["job-join"] = {
        "id": "job-join",
        "task_id": task_id,
        "kind": "join",
        "status": "failed",
        "error_name": "FanOutError",
        "error_value": "join produced 12 > anchor 10 rows",
        "traceback": "hidden traceback",
    }

    response = client.get(f"/api/tasks/{task_id}/jobs/latest?kind=join")

    assert response.status_code == 200
    assert response.json() == {
        "job": {
            "id": "job-join",
            "task_id": task_id,
            "kind": "join",
            "status": "failed",
            "error_name": "FanOutError",
            "error_value": "join produced 12 > anchor 10 rows",
        }
    }


def test_task_payload_exposes_structured_failure_stage_for_reloaded_ui(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    current = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(current),
            "status": TaskStatus.FAILED,
            "status_message": "unexpected failure wording",
        }
    )
    FakeTaskRepository.jobs["job-metrics"] = {
        "id": "job-metrics",
        "task_id": task_id,
        "kind": "metrics",
        "status": "failed",
    }

    got = client.get(f"/api/tasks/{task_id}")
    listed = client.get("/api/tasks")

    assert got.status_code == 200
    assert got.json()["failure_stage"] == "metrics"
    assert got.json()["stopped"] is False
    assert listed.status_code == 200
    assert listed.json()[0]["failure_stage"] == "metrics"
    assert listed.json()[0]["stopped"] is False


def test_pipeline_failure_with_unknown_message_keeps_failure_stage_unknown(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    current = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(current),
            "status": TaskStatus.FAILED,
            "status_message": "unexpected pipeline failure",
        }
    )
    FakeTaskRepository.jobs["job-pipeline"] = {
        "id": "job-pipeline",
        "task_id": task_id,
        "kind": "pipeline",
        "status": "failed",
    }

    got = client.get(f"/api/tasks/{task_id}")

    assert got.status_code == 200
    assert got.json()["failure_stage"] is None


def test_task_payload_exposes_structured_stopped_state(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    current = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(current),
            "status": TaskStatus.SCANNED,
            "status_message": "普通状态",
            "status_reason_code": TASK_STATUS_REASON_USER_CANCELLED,
        }
    )

    got = client.get(f"/api/tasks/{task_id}")

    assert got.status_code == 200
    assert got.json()["stopped"] is True
    assert got.json()["stop_reason_code"] == TASK_STATUS_REASON_USER_CANCELLED


def test_cancelled_job_history_does_not_mark_rerun_task_stopped(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    current = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(current),
            "status": TaskStatus.RUNNING,
            "status_message": "agent rerun requested: metrics",
            "status_reason_code": "",
        }
    )
    FakeTaskRepository.jobs["job-cancelled"] = {
        "id": "job-cancelled",
        "task_id": task_id,
        "kind": "agent",
        "status": "cancelled",
    }

    got = client.get(f"/api/tasks/{task_id}")

    assert got.status_code == 200
    assert got.json()["stopped"] is False
    assert got.json()["stop_reason_code"] is None


def test_task_payload_normalizes_legacy_stopped_text_server_side(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    current = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(current),
            "status": TaskStatus.SCANNED,
            "status_message": "已停止当前动作",
            "status_reason_code": "",
        }
    )

    got = client.get(f"/api/tasks/{task_id}")

    assert got.status_code == 200
    assert got.json()["stopped"] is True
    assert got.json()["stop_reason_code"] == TASK_STATUS_REASON_USER_CANCELLED


def test_task_payload_exposes_structured_restart_failure_reason(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    current = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(current),
            "status": TaskStatus.FAILED,
            "status_message": "普通失败文案",
            "status_reason_code": TASK_STATUS_REASON_SERVER_RESTART,
        }
    )

    got = client.get(f"/api/tasks/{task_id}")

    assert got.status_code == 200
    assert got.json()["failure_reason_code"] == TASK_STATUS_REASON_SERVER_RESTART


def test_task_payload_exposes_report_availability_for_download_buttons(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )

    before = client.get(f"/api/tasks/{task_id}")
    report_path = tmp_path / "tasks" / task_id / "outputs" / "validation_report.docx"
    report_path.parent.mkdir(parents=True)
    report_path.write_bytes(b"docx")
    after = client.get(f"/api/tasks/{task_id}")

    assert before.json()["report_available"] is False
    assert after.json()["report_available"] is True


def test_create_task_rejects_source_dir_outside_allowed_roots(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    outside_root = Path("/var/empty/rmc-forbidden-source")

    response = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(outside_root),
        },
    )

    assert response.status_code == 422
    assert "allowed material root" in response.json()["detail"]
    assert "RMC_MATERIAL_ROOTS" in response.json()["detail"]


def test_create_task_accepts_source_dir_under_extra_material_root(
    tmp_path: Path,
    monkeypatch,
):
    extra_root = tmp_path.parent / "external-materials"
    source_dir = extra_root / "project-a"
    source_dir.mkdir(parents=True)
    monkeypatch.setenv("RMC_MATERIAL_ROOTS", str(extra_root))
    client = _client(tmp_path, monkeypatch)

    response = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(source_dir),
        },
    )

    assert response.status_code == 200, response.text
    assert response.json()["source_dir"] == str(source_dir.resolve())


def test_create_task_accepts_missing_model_version(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    response = client.post(
        "/api/tasks",
        json={
            "model_name": "贷前评分卡 MOB3 v202604",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    )

    assert response.status_code == 200, response.text
    task = response.json()
    assert task["model_name"] == "贷前评分卡 MOB3 v202604"
    assert task["model_version"] == ""


def test_create_task_rejects_unknown_algorithm(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    response = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
            "algorithm": "random_forest",
        },
    )

    assert response.status_code == 422
    assert "unsupported model algorithm" in response.json()["detail"]


def test_create_task_without_algorithm_keeps_algorithm_pending(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    response = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    )

    assert response.status_code == 200, response.text
    task_id = response.json()["id"]
    assert response.json()["algorithm"] == ""

    fields = client.get(f"/api/tasks/{task_id}/report-fields")

    assert fields.status_code == 200
    description = fields.json()["text_values"]["TEXT:model_training_description"]
    assert "RMC_ALGORITHM" in description
    assert "LightGBM" not in description


def test_report_fields_default_training_description_uses_task_algorithm(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    response = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
            "algorithm": "xgb",
        },
    )
    assert response.status_code == 200, response.text
    task_id = response.json()["id"]

    got = client.get(f"/api/tasks/{task_id}/report-fields")

    assert got.status_code == 200
    text_values = got.json()["text_values"]
    assert text_values["TEXT:model_overview"] == (
        "为了更好的对xx用户进行授信环节风险管控，现开发A卡模型，"
        "对xx客群做前置风险拦截，从授信申请阶段做好风险防范。"
    )
    assert text_values["TEXT:model_scope"] == "本模型适用于xx渠道用户。"
    assert text_values["TEXT:bad_sample_definition"] == "xx逾期 >= xx天"
    assert text_values["TEXT:good_sample_definition"] == "xx未逾期"
    description = text_values["TEXT:model_training_description"]
    assert "XGBoost" in description
    assert "信贷风控" in description


def test_report_fields_get_and_put_round_trip_revision(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
            "report_values": {"TEXT:report_title": "初始标题"},
        },
    ).json()["id"]

    got = client.get(f"/api/tasks/{task_id}/report-fields")

    assert got.status_code == 200
    payload = got.json()
    assert payload["revision"] == 0
    assert payload["text_values"]["TEXT:report_title"] == "初始标题"
    assert payload["text_values"]["TEXT:revision_version"] == "V1"
    assert payload["text_values"]["TEXT:revision_author"] == "qa"
    assert payload["text_values"]["TEXT:revision_description"] == "初稿"
    assert payload["metric_values"] == {}
    assert any(
        field["key"] == "TEXT:final_validation_conclusion"
        for field in payload["fields"]
    )
    assert all(
        field["key"] not in {"TEXT:train_test_period", "TEXT:oot_period"}
        for field in payload["fields"]
    )

    updated = client.put(
        f"/api/tasks/{task_id}/report-fields",
        headers={"If-Match": "0"},
        json={"text_values": {"TEXT:report_title": "更新标题"}},
    )

    assert updated.status_code == 200
    updated_payload = updated.json()
    assert updated_payload["revision"] == 1
    assert updated_payload["text_values"]["TEXT:report_title"] == "更新标题"

    conflict = client.put(
        f"/api/tasks/{task_id}/report-fields",
        headers={"If-Match": "0"},
        json={"text_values": {"TEXT:report_title": "冲突标题"}},
    )
    assert conflict.status_code == 409

    computed_update = client.put(
        f"/api/tasks/{task_id}/report-fields",
        headers={"If-Match": "1"},
        json={"text_values": {"TEXT:train_test_period": "人工覆盖"}},
    )
    assert computed_update.status_code == 422


@pytest.fixture
def complete_validation_results_payload() -> dict:
    return {
        "basic_info": {
            "sample_period": ["20250101", "20250331"],
            "split_summary": [
                {"split": "train", "sample_count": 80, "bad_count": 8, "bad_rate": 0.10,
                 "period_start": "20250101", "period_end": "20250228"},
                {"split": "test",  "sample_count": 20, "bad_count": 2, "bad_rate": 0.10,
                 "period_start": "20250101", "period_end": "20250228"},
                {"split": "oot",   "sample_count": 30, "bad_count": 3, "bad_rate": 0.10,
                 "period_start": "20250301", "period_end": "20250331"},
            ],
            "monthly_distribution": [
                {"month": "202501", "sample_count": 80, "bad_count": 8, "bad_rate": 0.10},
                {"month": "202503", "sample_count": 30, "bad_count": 3, "bad_rate": 0.10},
            ],
            "feature_importance": [
                {"rank": 1, "feature": "income", "category": "征信", "importance": 0.42},
                {"rank": 2, "feature": "behavior", "category": "行为", "importance": 0.31},
            ],
        },
        "effectiveness": {
            "overall": [
                {"split": "train", "sample_count": 80, "bad_count": 8, "bad_rate": 0.10,
                 "ks": 0.3215, "auc": 0.7345, "head_lift_5pct": 2.1, "tail_lift_5pct": 0.4,
                 "psi_vs_train": 0.0},
                {"split": "test", "sample_count": 20, "bad_count": 2, "bad_rate": 0.10,
                 "ks": 0.3308, "auc": 0.7401, "head_lift_5pct": 2.0, "tail_lift_5pct": 0.5,
                 "psi_vs_train": 0.0008},
                {"split": "oot", "sample_count": 30, "bad_count": 3, "bad_rate": 0.10,
                 "ks": 0.321456, "auc": 0.7210, "head_lift_5pct": 1.9, "tail_lift_5pct": 0.4,
                 "psi_vs_train": 0.012345},
            ],
            "monthly_ks": [
                {"month": "202501", "sample_count": 80, "bad_count": 8, "bad_rate": 0.10,
                 "ks": 0.3215, "auc": 0.7345, "head_lift_5pct": 2.1, "tail_lift_5pct": 0.4},
                {"month": "202503", "sample_count": 30, "bad_count": 3, "bad_rate": 0.10,
                 "ks": 0.3050, "auc": 0.7210, "head_lift_5pct": 1.9, "tail_lift_5pct": 0.4},
            ],
            "monthly_psi": [
                {"month": "202501", "psi_first_month": 0.0, "psi_last_month": 0.012, "psi_mom": None},
                {"month": "202503", "psi_first_month": 0.012, "psi_last_month": 0.0, "psi_mom": 0.008},
            ],
            "bin_tables": {
                "train": [
                    {"bin_index": 1, "score_lower": 0.0, "score_upper": 0.5,
                     "sample_count": 80, "bad_count": 8, "bad_rate": 0.10,
                     "cum_sample_pct": 1.0, "cum_bad_pct": 1.0, "lift": 1.0, "ks": 0.0},
                ],
                "test": [],
                "oot": [],
            },
            "roc_ks_curves": {
                "train": {"fpr": [0.0, 0.5, 1.0], "tpr": [0.0, 0.7, 1.0],
                          "ks_curve": [0.0, 0.2, 0.0], "ks": 0.20, "population_at_ks": 0.5},
                "test": {"fpr": [0.0, 0.5, 1.0], "tpr": [0.0, 0.6, 1.0],
                         "ks_curve": [0.0, 0.1, 0.0], "ks": 0.10, "population_at_ks": 0.5},
                "oot": {"fpr": [0.0, 0.4, 1.0], "tpr": [0.0, 0.5, 1.0],
                        "ks_curve": [0.0, 0.1, 0.0], "ks": 0.10, "population_at_ks": 0.4},
            },
        },
        "stress_test": {
            "baseline": {"ks": 0.3215},
            "per_category": [
                {"category": "京东", "ks_after": 0.281, "ks_delta": -0.0404, "psi_vs_baseline": 0.1234},
            ],
        },
    }


def test_report_fields_include_metric_values_from_completed_output(
    tmp_path: Path,
    monkeypatch,
    complete_validation_results_payload: dict,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.WRITING_ARTIFACTS}
    )
    output_dir = tmp_path / "tasks" / task_id / "outputs"
    output_dir.mkdir(parents=True)
    (output_dir / "validation_results.json").write_text(
        json.dumps(complete_validation_results_payload),
        encoding="utf-8",
    )

    got = client.get(f"/api/tasks/{task_id}/report-fields")

    assert got.status_code == 200
    metric_values = got.json()["metric_values"]
    expected_values = {
        "TEXT:sample_period": "20250101-20250331",
        "TEXT:sample_start_month": "20250101",
        "TEXT:sample_end_month": "20250331",
        "TEXT:train_test_period": "20250101-20250228",
        "TEXT:train_test_ratio": "80.00%:20.00%",
        "TEXT:oot_period": "20250301-20250331",
        "TEXT:oot_ks": "0.3215",
        "TEXT:oot_psi": "0.0123",
    }
    assert expected_values.items() <= metric_values.items()


def test_report_fields_include_ordered_metric_table_sections_from_completed_output(
    tmp_path: Path,
    monkeypatch,
    complete_validation_results_payload: dict,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.WRITING_ARTIFACTS}
    )
    output_dir = tmp_path / "tasks" / task_id / "outputs"
    output_dir.mkdir(parents=True)
    (output_dir / "validation_results.json").write_text(
        json.dumps(complete_validation_results_payload),
        encoding="utf-8",
    )

    got = client.get(f"/api/tasks/{task_id}/report-fields")

    assert got.status_code == 200
    sections = got.json()["metric_table_sections"]
    titles = [section["title"] for section in sections]
    assert titles == [
        "样本情况",
        "整体效果&稳定性",
        "分月效果&稳定性",
        "分箱排序性",
        "特征重要性",
        "压力测试",
        "ROC&KS 曲线",
    ]
    assert all(section.get("section_theme") for section in sections)
    assert sections[1]["tables"][0]["layout"] == "kpi_cards"
    assert sections[2]["tables"][0]["layout"] == "trend_table"
    assert sections[5]["tables"][0]["layout"] == "table"
    assert sections[6]["tables"][0]["layout"] == "roc_ks_curve"
    assert len(sections[1]["tables"][0]["column_specs"]) == len(sections[1]["tables"][0]["headers"])


def test_report_preview_endpoint_returns_merged_table_html(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )
    output_dir = tmp_path / "tasks" / task_id / "outputs"
    output_dir.mkdir(parents=True)
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "横向合并"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "纵向合并"
    table.cell(1, 1).text = "明细"
    table.cell(1, 2).text = "值"
    table.cell(2, 1).text = "明细2"
    table.cell(2, 2).text = "值2"
    doc.save(output_dir / "validation_report.docx")

    response = client.get(f"/api/tasks/{task_id}/report/preview")

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/html")
    html = response.text
    assert "横向合并" in html
    assert 'colspan="2"' in html
    assert "纵向合并" in html
    assert 'rowspan="2"' in html
    assert "非 Word 精确排版" in html


def test_list_tasks_returns_array(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    response = client.get("/api/tasks")
    assert response.status_code == 200
    assert response.json() == []


def test_branding_and_task_list_routes_are_served_from_dedicated_routers():
    branding_routes = {
        route.path: route.endpoint.__module__
        for route in branding_router.routes
    }
    task_routes = {
        (route.path, tuple(sorted(route.methods or []))): route.endpoint.__module__
        for route in tasks_router.routes
    }

    assert branding_routes["/api/branding"] == "marvis.routers.branding"
    assert task_routes[("/api/tasks", ("GET",))] == "marvis.routers.tasks"
    assert task_routes[("/api/tasks", ("POST",))] == "marvis.routers.tasks"
    assert task_routes[("/api/tasks/{task_id}", ("GET",))] == "marvis.routers.tasks"
    assert task_routes[("/api/tasks/{task_id}", ("DELETE",))] == "marvis.routers.tasks"


def test_list_tasks_supports_limit_offset_headers(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    first = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    second = client.post(
        "/api/tasks",
        json={
            "model_name": "B卡",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]

    first_page = client.get("/api/tasks", params={"limit": 1})
    assert first_page.status_code == 200
    assert [task["id"] for task in first_page.json()] == [first]
    assert first_page.headers["x-result-limit"] == "1"
    assert first_page.headers["x-result-offset"] == "0"
    assert first_page.headers["x-result-has-more"] == "true"

    second_page = client.get("/api/tasks", params={"limit": 1, "offset": 1})
    assert second_page.status_code == 200
    assert [task["id"] for task in second_page.json()] == [second]
    assert second_page.headers["x-result-limit"] == "1"
    assert second_page.headers["x-result-offset"] == "1"
    assert second_page.headers["x-result-has-more"] == "false"

    capped = client.get("/api/tasks", params={"limit": 9999})
    assert capped.status_code == 200
    assert capped.headers["x-result-limit"] == "500"


def test_execution_environment_settings_round_trip_api(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setattr(
        "marvis.execution_environment.available_kernel_names",
        lambda: ["python3", "marvis-kernel"],
    )

    response = client.put(
        "/api/settings/execution-environment",
        json={
            "execution_mode": "jupyter_kernel",
            "kernel_name": "marvis-kernel",
            "conda_env_name": "",
            "python_executable": "",
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["settings"]["kernel_name"] == "marvis-kernel"
    assert payload["validation"]["ok"] is True

    got = client.get("/api/settings/execution-environment")
    assert got.status_code == 200
    assert got.json()["settings"]["kernel_name"] == "marvis-kernel"


def test_execution_environment_options_api(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setattr(
        "marvis.api_settings.detect_execution_environment_options",
        lambda: [
            ExecutionEnvironmentOption(
                id="kernel:python3",
                label="Kernel · Python 3",
                execution_mode="jupyter_kernel",
                kernel_name="python3",
                source="jupyter",
            )
        ],
    )
    monkeypatch.setattr(
        "marvis.api_settings.available_kernel_names",
        lambda: ["python3"],
        raising=False,
    )
    monkeypatch.setattr(
        "marvis.execution_environment.available_kernel_names",
        lambda: ["python3"],
    )

    response = client.get("/api/settings/execution-environment/options")

    assert response.status_code == 200
    payload = response.json()
    assert payload["options"][0]["id"] == "kernel:python3"
    assert payload["options"][0]["execution_mode"] == "jupyter_kernel"
    assert payload["settings"]["kernel_name"] == "python3"
    assert payload["validation"]["ok"] is True


def test_task_evidence_endpoint_reads_execution_artifacts(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    execution_dir = tmp_path / "tasks" / task_id / "execution"
    output_dir = tmp_path / "tasks" / task_id / "outputs"
    execution_dir.mkdir(parents=True)
    output_dir.mkdir(parents=True)
    (execution_dir / "notebook_steps.json").write_text(
        json.dumps({"steps": [{"title": "模型训练", "status": "succeeded"}], "cells": []}),
        encoding="utf-8",
    )
    (execution_dir / "scan_result.json").write_text(
        json.dumps(
            {
                "artifacts": [{"role": "notebook", "path": "model.ipynb"}],
                "checks": [{"id": "notebook_contract", "status": "success"}],
                "notebook_steps": [{"title": "模型训练", "status": "succeeded"}],
            }
        ),
        encoding="utf-8",
    )
    (execution_dir / "runtime_contract.json").write_text(
        json.dumps({"target_col": "y", "score_decimal_places": 6}),
        encoding="utf-8",
    )
    (output_dir / "validation_results.json").write_text(
        json.dumps({"reproducibility": {"summary": {"status": "pass"}}}),
        encoding="utf-8",
    )

    response = client.get(f"/api/tasks/{task_id}/evidence")

    assert response.status_code == 200
    payload = response.json()
    assert payload["scan"]["checks"][0]["id"] == "notebook_contract"
    assert payload["notebook_steps"][0]["title"] == "模型训练"
    assert payload["contract"]["target_col"] == "y"
    assert payload["reproducibility"]["summary"]["status"] == "pass"


def test_task_evidence_endpoint_reads_notebook_stage_reproducibility_artifact(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    output_dir = tmp_path / "tasks" / task_id / "outputs"
    output_dir.mkdir(parents=True)
    (output_dir / "reproducibility_result.json").write_text(
        json.dumps(
            {
                "sample_size": 2,
                "seed": 42,
                "summary": {
                    "match_count": 2,
                    "mismatch_count": 0,
                    "max_abs_diff": 0.0,
                    "status": "pass",
                },
                "rows": [
                    {
                        "row_index": 0,
                        "score_code_model": 0.1,
                        "score_submitted_pmml": 0.1,
                        "abs_diff": 0.0,
                        "matched": True,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    response = client.get(f"/api/tasks/{task_id}/evidence")

    assert response.status_code == 200
    payload = response.json()
    assert payload["reproducibility"]["summary"]["status"] == "pass"
    assert payload["reproducibility"]["rows"][0]["score_code_model"] == 0.1


def test_notebook_metrics_and_report_endpoints_dispatch_stages(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    calls: list[tuple[str, str, object]] = []

    def fake_notebook_stage(*, task_id, settings, **_kwargs):
        calls.append(("notebook", task_id, settings))

    def fake_metrics_stage(*, task_id, settings, **_kwargs):
        calls.append(("metrics", task_id, settings))

    def fake_report_stage(*, task_id, settings, **_kwargs):
        calls.append(("report", task_id, settings))

    monkeypatch.setattr("marvis.routers.validation_stages.run_notebook_stage", fake_notebook_stage)
    monkeypatch.setattr("marvis.routers.validation_stages.run_metrics_stage", fake_metrics_stage)
    monkeypatch.setattr("marvis.routers.validation_stages.run_report_stage", fake_report_stage)

    create = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
            "algorithm": "lgb",
            "target_col": "y",
            "score_col": "pred",
            "split_col": "split",
            "time_col": "apply_month",
        },
    )
    task_id = create.json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SCANNED}
    )

    notebook_response = client.post(f"/api/tasks/{task_id}/notebook", json={})
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.EXECUTED}
    )
    metrics_response = client.post(f"/api/tasks/{task_id}/metrics")
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.WRITING_ARTIFACTS,
        }
    )
    report_response = client.post(f"/api/tasks/{task_id}/report")

    assert notebook_response.status_code == 202
    assert metrics_response.status_code == 202
    assert report_response.status_code == 202
    assert [call[0] for call in calls] == ["notebook", "metrics", "report"]
    assert calls[0][1] == task_id
    assert calls[0][2].workspace == tmp_path


def test_stage_endpoint_blocks_active_same_task_but_allows_other_tasks(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls: list[str] = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_report_stage",
        lambda **kwargs: calls.append(kwargs["task_id"]),
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    other_task_id = client.post(
        "/api/tasks",
        json={"model_name": "B卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    for current_task_id in (task_id, other_task_id):
        FakeTaskRepository.tasks[current_task_id] = TaskRecord(
            **{
                **asdict(FakeTaskRepository.tasks[current_task_id]),
                "status": TaskStatus.WRITING_ARTIFACTS,
            }
        )
    FakeTaskRepository.jobs["job-existing"] = {
        "id": "job-existing",
        "task_id": task_id,
        "kind": "metrics",
        "status": "running",
    }

    blocked = client.post(f"/api/tasks/{task_id}/report")
    accepted = client.post(f"/api/tasks/{other_task_id}/report")

    assert blocked.status_code == 409
    assert blocked.json()["detail"] == "task already has an active stage"
    assert accepted.status_code == 202
    assert calls == [other_task_id]


def test_completed_task_can_rerun_prior_workflow_stages(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    calls: list[str] = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_notebook_stage",
        lambda **_kwargs: calls.append("notebook"),
    )
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_metrics_stage",
        lambda **_kwargs: calls.append("metrics"),
    )
    monkeypatch.setattr(
        "marvis.routers.validation_stages.get_live_notebook_session",
        lambda _task_id: object(),
    )
    source = tmp_path / "source"
    source.mkdir()
    nbformat.write(
        nbformat.v4.new_notebook(
            cells=[
                nbformat.v4.new_code_cell(
                    "RMC_SAMPLE_DF = sample_df\n"
                    "RMC_TARGET_COL = 'y'\n"
                    "RMC_ALGORITHM = 'lgb'\n"
                    "def RMC_SCORE_FN(df):\n"
                    "    return [0.1] * len(df)\n"
                )
            ]
        ),
        source / "model.ipynb",
    )
    (source / "sample.csv").write_text(
        "x1,pred,y,split,apply_month\n1,0.1,0,train,202501\n",
        encoding="utf-8",
    )
    (source / "model.pmml").write_text("<PMML/>", encoding="utf-8")
    (source / "data_dictionary.csv").write_text("特征名,类别\nx1,基础信息\n", encoding="utf-8")
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(source)},
    ).json()["id"]

    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )
    scan_response = client.post(f"/api/tasks/{task_id}/scan")
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )
    notebook_response = client.post(f"/api/tasks/{task_id}/notebook", json={})
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )
    metrics_response = client.post(f"/api/tasks/{task_id}/metrics")

    assert scan_response.status_code == 200
    assert scan_response.json()["status"] == "scanned"
    assert notebook_response.status_code == 202
    assert metrics_response.status_code == 202
    assert calls == ["notebook", "metrics"]


def test_notebook_endpoint_claims_running_before_dispatching_stage(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_notebook_stage",
        lambda **kwargs: calls.append(kwargs),
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SCANNED}
    )

    first = client.post(f"/api/tasks/{task_id}/notebook", json={})
    second = client.post(f"/api/tasks/{task_id}/notebook", json={})

    assert first.status_code == 202
    assert second.status_code == 409
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.RUNNING
    assert calls[0]["stage_claimed"] is True


def test_cancel_notebook_endpoint_requests_running_notebook_stop(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    requested: list[str] = []
    monkeypatch.setattr(
        "marvis.routers.stage_controls.request_notebook_cancellation",
        lambda task_id: requested.append(task_id) or True,
        raising=False,
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.RUNNING}
    )

    response = client.post(f"/api/tasks/{task_id}/notebook/cancel")

    assert response.status_code == 202
    assert response.json()["status"] == "accepted"
    assert requested == [task_id]


def test_cancel_notebook_endpoint_rejects_non_running_task(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    requested: list[str] = []
    monkeypatch.setattr(
        "marvis.routers.stage_controls.request_notebook_cancellation",
        lambda task_id: requested.append(task_id) or True,
        raising=False,
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SCANNED}
    )

    response = client.post(f"/api/tasks/{task_id}/notebook/cancel")

    assert response.status_code == 409
    assert requested == []


def test_cancel_metrics_endpoint_requests_running_metrics_stop(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    requested: list[str] = []
    monkeypatch.setattr(
        "marvis.routers.stage_controls.request_notebook_cancellation",
        lambda task_id: requested.append(task_id) or True,
        raising=False,
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.COMPUTING_METRICS,
        }
    )

    response = client.post(f"/api/tasks/{task_id}/metrics/cancel")

    assert response.status_code == 202
    assert response.json()["status"] == "accepted"
    assert requested == [task_id]
    marker_path = tmp_path / "tasks" / task_id / "execution" / "metrics_cancel.requested"
    assert marker_path.read_text(encoding="utf-8") == "cancelled\n"


def test_cancel_metrics_endpoint_rejects_non_running_metrics_task(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    requested: list[str] = []
    monkeypatch.setattr(
        "marvis.routers.stage_controls.request_notebook_cancellation",
        lambda task_id: requested.append(task_id) or True,
        raising=False,
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.EXECUTED}
    )

    response = client.post(f"/api/tasks/{task_id}/metrics/cancel")

    assert response.status_code == 409
    assert requested == []


def test_cancel_report_endpoint_requests_report_stop(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    requested: list[str] = []
    monkeypatch.setattr(
        "marvis.routers.stage_controls.request_notebook_cancellation",
        lambda task_id: requested.append(task_id) or True,
        raising=False,
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.WRITING_ARTIFACTS,
        }
    )
    FakeTaskRepository.jobs["job-report"] = {
        "id": "job-report",
        "task_id": task_id,
        "kind": "report",
        "status": "queued",
    }

    response = client.post(f"/api/tasks/{task_id}/report/cancel")

    assert response.status_code == 202
    assert response.json()["status"] == "accepted"
    assert requested == [task_id]


def test_cancel_report_endpoint_rejects_without_active_report_job(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    requested: list[str] = []
    monkeypatch.setattr(
        "marvis.routers.stage_controls.request_notebook_cancellation",
        lambda task_id: requested.append(task_id) or True,
        raising=False,
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.WRITING_ARTIFACTS,
        }
    )

    response = client.post(f"/api/tasks/{task_id}/report/cancel")

    assert response.status_code == 409
    assert response.json()["detail"] == "task has no active report job"
    assert requested == []


def test_stage_job_records_cancelled_when_stage_returns_cancelled_task(
    tmp_path: Path,
    monkeypatch,
):
    from marvis.api import _run_stage_job

    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.COMPUTING_METRICS,
        }
    )
    repo = FakeTaskRepository(tmp_path / "marvis.sqlite")
    job_id = repo.start_job(task_id, "metrics")

    def cancelled_stage(*, task_id: str) -> None:
        repo.update_status(
            task_id,
            TaskStatus.EXECUTED,
            "ordinary resume status",
            expected=TaskStatus.COMPUTING_METRICS,
            reason_code=TASK_STATUS_REASON_USER_CANCELLED,
        )

    _run_stage_job(
        job_id,
        tmp_path / "marvis.sqlite",
        cancelled_stage,
        {"task_id": task_id},
    )

    assert FakeTaskRepository.jobs[job_id]["status"] == "cancelled"


def test_stage_job_dispatches_before_and_after_hooks(tmp_path: Path, monkeypatch):
    from marvis.api import _run_stage_job

    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    repo = FakeTaskRepository(tmp_path / "marvis.sqlite")
    job_id = repo.start_job(task_id, "report")
    dispatcher = FakeHookDispatcher()

    def stage(*, task_id: str) -> None:
        assert task_id

    _run_stage_job(
        job_id,
        tmp_path / "marvis.sqlite",
        stage,
        {"task_id": task_id},
        hook_dispatcher=dispatcher,
        before_hook_event="report.before_generate",
        after_hook_event="report.after_generate",
    )

    assert dispatcher.calls == [
        ("report.before_generate", {"job_id": job_id, "task_id": task_id}, task_id),
        (
            "report.after_generate",
            {"job_id": job_id, "task_id": task_id, "status": "succeeded"},
            task_id,
        ),
    ]


def test_metrics_endpoint_claims_computing_before_dispatching_stage(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_metrics_stage",
        lambda **kwargs: calls.append(kwargs),
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.EXECUTED}
    )

    first = client.post(f"/api/tasks/{task_id}/metrics")
    second = client.post(f"/api/tasks/{task_id}/metrics")

    assert first.status_code == 202
    assert second.status_code == 409
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.COMPUTING_METRICS
    assert calls[0]["stage_claimed"] is True


def test_metrics_endpoint_rejects_terminal_legacy_live_mode_without_explicit_allow(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_metrics_stage",
        lambda **kwargs: calls.append(kwargs),
    )
    monkeypatch.setattr(
        "marvis.routers.validation_stages.pipeline_settings_from_request",
        lambda request, task, feature_columns: PipelineSettings(
            workspace=tmp_path,
            db_path=tmp_path / "marvis.sqlite",
            report_template_path=tmp_path / "template.docx",
            notebook_isolated_execution=False,
        ),
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )

    response = client.post(f"/api/tasks/{task_id}/metrics")

    assert response.status_code == 409
    assert "allow_legacy_live_notebook_execution=True" in response.json()["detail"]
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.SUCCEEDED
    assert calls == []


def test_metrics_endpoint_rejects_terminal_legacy_live_mode_without_process_env(
    tmp_path: Path,
    monkeypatch,
):
    monkeypatch.delenv(LEGACY_LIVE_NOTEBOOK_ENV_VAR, raising=False)
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_metrics_stage",
        lambda **kwargs: calls.append(kwargs),
    )
    monkeypatch.setattr(
        "marvis.routers.validation_stages.pipeline_settings_from_request",
        lambda request, task, feature_columns: PipelineSettings(
            workspace=tmp_path,
            db_path=tmp_path / "marvis.sqlite",
            report_template_path=tmp_path / "template.docx",
            notebook_isolated_execution=False,
            allow_legacy_live_notebook_execution=True,
        ),
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )

    response = client.post(f"/api/tasks/{task_id}/metrics")

    assert response.status_code == 409
    assert f"{LEGACY_LIVE_NOTEBOOK_ENV_VAR}=1" in response.json()["detail"]
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.SUCCEEDED
    assert calls == []


def test_metrics_endpoint_rejects_terminal_rerun_without_live_kernel_when_isolated_disabled(
    tmp_path: Path,
    monkeypatch,
):
    monkeypatch.setenv(LEGACY_LIVE_NOTEBOOK_ENV_VAR, "1")
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_metrics_stage",
        lambda **kwargs: calls.append(kwargs),
    )
    monkeypatch.setattr(
        "marvis.routers.validation_stages.get_live_notebook_session",
        lambda _task_id: None,
    )
    monkeypatch.setattr(
        "marvis.routers.validation_stages.pipeline_settings_from_request",
        lambda request, task, feature_columns: PipelineSettings(
            workspace=tmp_path,
            db_path=tmp_path / "marvis.sqlite",
            report_template_path=tmp_path / "template.docx",
            notebook_isolated_execution=False,
            allow_legacy_live_notebook_execution=True,
        ),
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )

    response = client.post(f"/api/tasks/{task_id}/metrics")

    assert response.status_code == 409
    assert "live notebook kernel is not available" in response.json()["detail"]
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.SUCCEEDED
    assert calls == []


def test_metrics_endpoint_allows_terminal_rerun_without_live_kernel_when_isolated_enabled(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_metrics_stage",
        lambda **kwargs: calls.append(kwargs),
    )
    monkeypatch.setattr(
        "marvis.routers.validation_stages.get_live_notebook_session",
        lambda _task_id: None,
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )

    response = client.post(f"/api/tasks/{task_id}/metrics")

    assert response.status_code == 202
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.COMPUTING_METRICS
    assert calls[0]["settings"].notebook_isolated_execution is True


def test_metrics_endpoint_allows_retry_after_metrics_stage_failure(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_metrics_stage",
        lambda **kwargs: calls.append(kwargs),
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.FAILED,
            "status_message": "模型效果&稳定性验证失败：sample column check failed: split_col='new_flag'",
        }
    )

    response = client.post(f"/api/tasks/{task_id}/metrics")

    assert response.status_code == 202
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.COMPUTING_METRICS
    assert calls[0]["stage_claimed"] is True


def test_metrics_endpoint_allows_retry_after_legacy_sample_column_failure(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_metrics_stage",
        lambda **kwargs: calls.append(kwargs),
    )
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.FAILED,
            "status_message": "sample column check failed: split_col='new_flag'",
        }
    )

    response = client.post(f"/api/tasks/{task_id}/metrics")

    assert response.status_code == 202
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.COMPUTING_METRICS
    assert calls[0]["stage_claimed"] is True


def test_metrics_endpoint_rejects_non_metrics_failure(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    monkeypatch.setattr("marvis.routers.validation_stages.run_metrics_stage", lambda **_kwargs: None)
    task_id = client.post(
        "/api/tasks",
        json={"model_name": "A卡", "validator": "qa", "source_dir": str(tmp_path)},
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.FAILED,
            "status_message": "模型可复现性验证失败：notebook failed",
        }
    )

    response = client.post(f"/api/tasks/{task_id}/metrics")

    assert response.status_code == 409
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.FAILED


def test_legacy_validate_endpoint_runs_staged_pipeline_for_cli_compatibility(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []

    def fake_run_staged_pipeline(*, task_id, settings):
        calls.append((task_id, settings))

    monkeypatch.setattr("marvis.routers.validation_stages.run_staged_pipeline", fake_run_staged_pipeline)
    create = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
            "feature_columns": ["x_task"],
        },
    )
    task_id = create.json()["id"]

    response = client.post(f"/api/tasks/{task_id}/validate", json={})

    assert response.status_code == 202
    assert calls[0][0] == task_id


def test_validate_rejects_terminal_task_without_dispatching_pipeline(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_staged_pipeline",
        lambda **kwargs: calls.append(kwargs),
    )
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
            "feature_columns": ["x1"],
        },
    ).json()["id"]
    task = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(task), "status": TaskStatus.SUCCEEDED}
    )

    response = client.post(f"/api/tasks/{task_id}/validate", json={})

    assert response.status_code == 409
    assert calls == []


def test_validate_accepts_tasks_without_feature_columns(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(
        "marvis.routers.validation_stages.run_staged_pipeline",
        lambda **kwargs: calls.append(kwargs),
    )
    create = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    )
    task_id = create.json()["id"]

    response = client.post(f"/api/tasks/{task_id}/validate", json={})

    assert response.status_code == 202
    assert calls[0]["settings"].feature_columns == []


def test_scan_endpoint_returns_v2_artifacts_and_updates_status(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    source = tmp_path / "source"
    source.mkdir()
    nbformat.write(
        nbformat.v4.new_notebook(
            cells=[
                nbformat.v4.new_code_cell(
                    "RMC_SAMPLE_DF = sample_df\n"
                    "RMC_TARGET_COL = 'y'\n"
                    "RMC_ALGORITHM = 'lgb'\n"
                    "def RMC_SCORE_FN(df):\n"
                    "    return [0.1] * len(df)\n"
                )
            ]
        ),
        source / "model.ipynb",
    )
    (source / "sample.csv").write_text("x1,pred,y,split,apply_month\n1,0.1,0,train,202501\n", encoding="utf-8")
    (source / "model.pmml").write_text("<PMML/>", encoding="utf-8")
    (source / "data_dictionary.csv").write_text("特征名,类别\nx1,基础信息\n", encoding="utf-8")
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(source),
        },
    ).json()["id"]

    response = client.post(f"/api/tasks/{task_id}/scan")

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "scanned"
    assert {artifact["role"] for artifact in payload["artifacts"]} == {
        "notebook",
        "sample",
        "model_pmml",
        "data_dictionary",
    }
    assert [step["title"] for step in payload["notebook_steps"]] == ["Notebook 初始化"]
    saved_scan = json.loads(
        (tmp_path / "tasks" / task_id / "execution" / "scan_result.json").read_text(
            encoding="utf-8"
        )
    )
    assert saved_scan["checks"][-1]["id"] == "notebook_contract"
    assert saved_scan["notebook_steps"][0]["title"] == "Notebook 初始化"
    saved_steps = json.loads(
        (tmp_path / "tasks" / task_id / "execution" / "notebook_steps.json").read_text(
            encoding="utf-8"
        )
    )
    assert saved_steps["steps"][0]["title"] == "Notebook 初始化"
    evidence = client.get(f"/api/tasks/{task_id}/evidence")
    assert evidence.status_code == 200
    assert evidence.json()["notebook_steps"][0]["title"] == "Notebook 初始化"
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.SCANNED


def test_scan_endpoint_reports_notebook_contract_errors_before_execution(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    source = tmp_path / "source"
    source.mkdir()
    nbformat.write(
        nbformat.v4.new_notebook(
            cells=[
                nbformat.v4.new_code_cell(
                    "RMC_SAMPLE_DF = sample_df\n"
                    "RMC_TARGET_COL = 'y'\n"
                    "RMC_ALGORITHM = 'lgb'\n"
                    "print('RMC_SCORE_FN mentioned but not defined')\n"
                )
            ]
        ),
        source / "model.ipynb",
    )
    (source / "sample.csv").write_text("x1,pred,y,split,apply_month\n1,0.1,0,train,202501\n", encoding="utf-8")
    (source / "model.pmml").write_text("<PMML/>", encoding="utf-8")
    (source / "data_dictionary.csv").write_text("特征名,类别\nx1,基础信息\n", encoding="utf-8")
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(source),
        },
    ).json()["id"]

    response = client.post(f"/api/tasks/{task_id}/scan")

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "failed"
    checks = {check["id"]: check for check in payload["checks"]}
    assert checks["notebook_contract"]["status"] == "error"
    assert (
        checks["notebook_contract"]["message"]
        == "Notebook RMC 契约检查失败：缺少 RMC_SCORE_FN（模型打分函数）。请在 Notebook 顶层定义后重新扫描。"
    )
    assert "missing RMC_SCORE_FN" not in checks["notebook_contract"]["message"]
    assert payload["status_message"] == (
        "材料扫描失败：Notebook RMC 契约检查失败：缺少 RMC_SCORE_FN（模型打分函数）。请在 Notebook 顶层定义后重新扫描。"
    )
    assert FakeTaskRepository.tasks[task_id].status == TaskStatus.FAILED
    assert FakeTaskRepository.tasks[task_id].status_message == payload["status_message"]

    notebook_response = client.post(f"/api/tasks/{task_id}/notebook", json={})

    assert notebook_response.status_code == 409
    assert "材料扫描未完整通过" in notebook_response.json()["detail"]


def test_scan_endpoint_translates_multiple_missing_rmc_contract_fields(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    source = tmp_path / "source"
    source.mkdir()
    nbformat.write(
        nbformat.v4.new_notebook(
            cells=[
                nbformat.v4.new_code_cell("print('no rmc contract fields')"),
            ]
        ),
        source / "model.ipynb",
    )
    (source / "sample.csv").write_text("x1,pred,y,split,apply_month\n1,0.1,0,train,202501\n", encoding="utf-8")
    (source / "model.pmml").write_text("<PMML/>", encoding="utf-8")
    (source / "data_dictionary.csv").write_text("特征名,类别\nx1,基础信息\n", encoding="utf-8")
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(source),
        },
    ).json()["id"]

    response = client.post(f"/api/tasks/{task_id}/scan")

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "failed"
    checks = {check["id"]: check for check in payload["checks"]}
    assert checks["notebook_contract"]["message"] == (
        "Notebook RMC 契约检查失败：缺少 "
        "RMC_SAMPLE_DF（样本 DataFrame）、RMC_SCORE_FN（模型打分函数）、"
        "RMC_TARGET_COL（目标列）、RMC_ALGORITHM（模型算法）。请在 Notebook 顶层定义后重新扫描。"
    )
    assert payload["status_message"] == (
        "材料扫描失败：Notebook RMC 契约检查失败：缺少 "
        "RMC_SAMPLE_DF（样本 DataFrame）、RMC_SCORE_FN（模型打分函数）、"
        "RMC_TARGET_COL（目标列）、RMC_ALGORITHM（模型算法）。请在 Notebook 顶层定义后重新扫描。"
    )


def test_scan_endpoint_returns_422_when_source_dir_exceeds_limits(
    tmp_path: Path,
    monkeypatch,
):
    # scan_source_dir raises ValueError when the tree breaches max_files / max_depth.
    # That is a client-side "bad source dir" condition and must surface as 422,
    # not crash into a 500.
    client = _client(tmp_path, monkeypatch)
    source = tmp_path / "source"
    source.mkdir()
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(source),
        },
    ).json()["id"]

    def _raise_limit(*_args, **_kwargs):
        raise ValueError("source_dir has too many files: max_files=2000")

    monkeypatch.setattr("marvis.api_scan_helpers.scan_source_dir", _raise_limit)

    response = client.post(f"/api/tasks/{task_id}/scan")

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail.startswith("source dir invalid")
    assert "too many files" in detail


def test_normalize_task_type_whitelists_known_types():
    from marvis.db import _normalize_task_type
    from marvis.domain import TASK_TYPE_VALIDATION, VALID_TASK_TYPES

    assert _normalize_task_type(None) == TASK_TYPE_VALIDATION
    assert _normalize_task_type("") == TASK_TYPE_VALIDATION
    for task_type in VALID_TASK_TYPES:
        assert _normalize_task_type(task_type) == task_type
    # Unknown / arbitrary strings must not persist as-is.
    assert _normalize_task_type("'; DROP TABLE tasks;--") == TASK_TYPE_VALIDATION


def test_task_stop_reason_code_ignores_legacy_text_for_successful_terminals():
    from marvis.api_task_payloads import task_stop_reason_code, task_stopped
    from marvis.domain import TASK_STATUS_REASON_USER_CANCELLED

    # A SUCCEEDED task whose message incidentally contains "已取消" must NOT be
    # reported as stopped — only the structured status_reason_code can mark it.
    succeeded = SimpleNamespace(
        status=TaskStatus.SUCCEEDED,
        status_reason_code=None,
        status_message="报告包含 3 笔已取消订单的分析",
    )
    assert task_stop_reason_code(None, succeeded) is None
    assert task_stopped(None, succeeded) is False

    # A genuinely cancelled task is still detected via the structured reason code.
    cancelled = SimpleNamespace(
        status=TaskStatus.FAILED,
        status_reason_code=TASK_STATUS_REASON_USER_CANCELLED,
        status_message="已停止",
    )
    assert task_stop_reason_code(None, cancelled) == TASK_STATUS_REASON_USER_CANCELLED
    assert task_stopped(None, cancelled) is True


def test_report_download_endpoint_returns_generated_word(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.SUCCEEDED}
    )
    report_path = tmp_path / "tasks" / task_id / "outputs" / "validation_report.docx"
    report_path.parent.mkdir(parents=True)
    report_path.write_bytes(b"docx-bytes")

    response = client.get(f"/api/tasks/{task_id}/report/download")

    assert response.status_code == 200
    assert response.content == b"docx-bytes"
    assert _decoded_download_filename(response) == "A卡_模型验证报告_20260521.docx"


def test_analysis_download_endpoint_returns_generated_excel(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{**asdict(FakeTaskRepository.tasks[task_id]), "status": TaskStatus.WRITING_ARTIFACTS}
    )
    excel_path = tmp_path / "tasks" / task_id / "outputs" / "validation.xlsx"
    excel_path.parent.mkdir(parents=True)
    excel_path.write_bytes(b"xlsx-bytes")

    response = client.get(f"/api/tasks/{task_id}/analysis/download")

    assert response.status_code == 200
    assert response.content == b"xlsx-bytes"
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert _decoded_download_filename(response) == "A卡_模型验证报告_20260521.xlsx"


def test_analysis_download_allows_failed_report_stage_with_generated_excel(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(FakeTaskRepository.tasks[task_id]),
            "status": TaskStatus.FAILED,
            "status_message": "报告输出失败：RuntimeError: docx failed",
        }
    )
    excel_path = tmp_path / "tasks" / task_id / "outputs" / "validation.xlsx"
    excel_path.parent.mkdir(parents=True)
    excel_path.write_bytes(b"xlsx-after-report-failure")

    response = client.get(f"/api/tasks/{task_id}/analysis/download")

    assert response.status_code == 200
    assert response.content == b"xlsx-after-report-failure"


def test_download_endpoints_ignore_stale_outputs_before_current_stage(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    task_id = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    ).json()["id"]
    output_dir = tmp_path / "tasks" / task_id / "outputs"
    output_dir.mkdir(parents=True)
    (output_dir / "validation_report.docx").write_bytes(b"old-docx")
    (output_dir / "validation.xlsx").write_bytes(b"old-xlsx")

    report_response = client.get(f"/api/tasks/{task_id}/report/download")
    analysis_response = client.get(f"/api/tasks/{task_id}/analysis/download")

    assert report_response.status_code == 404
    assert analysis_response.status_code == 404


def test_delete_task_removes_repo_record_and_task_dir(tmp_path: Path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    create = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    )
    task_id = create.json()["id"]
    task_dir = tmp_path / "tasks" / task_id
    task_dir.mkdir()

    response = client.delete(f"/api/tasks/{task_id}")

    assert response.status_code == 204
    assert FakeTaskRepository.deleted == [task_id]
    assert not task_dir.exists()


def test_delete_task_keeps_record_deleted_when_directory_cleanup_fails(
    tmp_path: Path,
    monkeypatch,
    caplog,
):
    client = _client(tmp_path, monkeypatch)
    create = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    )
    task_id = create.json()["id"]
    task_dir = tmp_path / "tasks" / task_id
    task_dir.mkdir()

    def fail_rmtree(path):
        raise PermissionError(f"locked: {path}")

    monkeypatch.setattr("marvis.routers.tasks.shutil.rmtree", fail_rmtree)

    response = client.delete(f"/api/tasks/{task_id}")

    assert response.status_code == 204
    assert FakeTaskRepository.deleted == [task_id]
    assert task_dir.exists()
    assert "task dir cleanup failed" in caplog.text


def test_agent_stop_ack_detection_scans_past_latest_user_message():
    class Repo:
        def list_agent_messages(self, task_id):
            assert task_id == "task-1"
            return [
                {
                    "role": "assistant",
                    "metadata": {"intent": "stop", "cancel_requested": True},
                },
                {"role": "user", "metadata": {}},
            ]

    assert _agent_has_stop_ack_message(Repo(), "task-1") is True


def test_delete_task_allows_stale_busy_status_when_no_active_job(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    create = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    )
    task_id = create.json()["id"]
    current = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(current),
            "status": TaskStatus.WRITING_ARTIFACTS,
            "status_message": "metrics and excel generated",
        }
    )

    response = client.delete(f"/api/tasks/{task_id}")

    assert response.status_code == 204
    assert FakeTaskRepository.deleted == [task_id]


def test_delete_task_rejects_active_job_even_when_status_is_stale(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path, monkeypatch)
    create = client.post(
        "/api/tasks",
        json={
            "model_name": "A卡",
            "model_version": "v1",
            "validator": "qa",
            "source_dir": str(tmp_path),
        },
    )
    task_id = create.json()["id"]
    current = FakeTaskRepository.tasks[task_id]
    FakeTaskRepository.tasks[task_id] = TaskRecord(
        **{
            **asdict(current),
            "status": TaskStatus.WRITING_ARTIFACTS,
            "status_message": "metrics and excel generated",
        }
    )
    FakeTaskRepository.jobs["job-1"] = {
        "id": "job-1",
        "task_id": task_id,
        "kind": "report",
        "status": "running",
    }

    response = client.delete(f"/api/tasks/{task_id}")

    assert response.status_code == 409
    assert response.json()["detail"] == "task already has an active stage"
    assert FakeTaskRepository.deleted == []


@pytest.mark.parametrize(
    "method,path",
    [
        ("post", "/api/tasks/task-1/run-notebook"),
        ("post", "/api/tasks/task-1/report-template"),
    ],
)
def test_v1_task_endpoints_are_not_registered(
    tmp_path: Path,
    monkeypatch,
    method: str,
    path: str,
):
    client = _client(tmp_path, monkeypatch)
    if method == "get":
        response = client.get(path)
    else:
        response = getattr(client, method)(path, json={})

    assert response.status_code == 404
