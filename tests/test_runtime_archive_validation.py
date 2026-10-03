"""Recompute actual V2 carrier evidence after its application workspace is gone."""
import json
from pathlib import Path
import shutil

import pandas as pd
import pytest
from pydantic import ValidationError

from marvis.orchestrator.eval.runtime_archive_validation import FrozenValidationBinding, revalidate_validation_archive
from marvis.orchestrator.eval.runtime_cases import write_synthetic_suite
from marvis.orchestrator.eval.runtime_contracts import RuntimeCase, digest
from marvis.orchestrator.eval.runtime_runner import run_runtime_suite
from test_runtime_agent_benchmark import fixture_model
from test_runtime_validation_agent import _v2_protocol


@pytest.fixture(scope="module")
def actual_archive(tmp_path_factory):
    root = tmp_path_factory.mktemp("v2-original-custody")
    paths = write_synthetic_suite(root / "suite", normal_validation_agent_only=True)
    case = RuntimeCase.model_validate(json.loads(paths["cases"].read_text())["cases"][0])
    with fixture_model(answer_factory=_v2_protocol) as (model, _):
        report = run_runtime_suite(cases_path=paths["cases"], expected_path=paths["expected"],
            dataset_root=paths["dataset_root"], output_dir=root / "public", evidence_custody_dir=root / "custody",
            model=model, model_source="fixture_model")
    assert report["all_passed"], [(r["runtime_status"], r.get("evidence_custody"), r.get("error_code")) for r in report["cases"]]
    observed = report["cases"][0]
    assert observed["isolated_workspace_removed"]
    pipeline = observed["execution"]["validation_pipeline"]
    binding = FrozenValidationBinding(case=case, run_id=report["run_id"], task_id=observed["execution"]["task_id"],
        cases_sha256=report["cases_sha256"], expected_sha256=report["expected_sha256"], source=report["source"],
        model_connection_sha256=report["model_connection_sha256"], input_contract_sha256=pipeline["input_contract_sha256"],
        confirmed_draft_sha256=pipeline["confirmed_draft_sha256"], report_revision=pipeline["report_revision"],
        report_sha256={f["kind"]: f["sha256"] for f in pipeline["report_files"]}, bin_count=10)
    archive = root / "custody" / report["run_id"] / case.id
    assert not Path(json.loads((archive / "manifest.json").read_text())["original_workspace"]).exists()
    return archive, observed["evidence_custody"]["manifest_sha256"], binding


def _copy(actual_archive, tmp_path):
    archive, expected_digest, binding = actual_archive
    copy = tmp_path / "archive"
    shutil.copytree(archive, copy)
    return copy, expected_digest, binding


def _rewrite_manifest(archive):
    path = archive / "manifest.json"
    manifest = json.loads(path.read_text())
    manifest["files"] = {p.relative_to(archive).as_posix(): {"sha256": digest(p.read_bytes()), "size_bytes": p.stat().st_size}
                         for p in archive.rglob("*") if p.is_file() and p != path}
    manifest["bindings_sha256"] = manifest["files"]["bindings.json"]["sha256"]
    path.write_text(json.dumps(manifest, ensure_ascii=False))
    return digest(path.read_bytes())


def _run(archive, expected_digest, binding):
    return revalidate_validation_archive(archive, expected_manifest_sha256=expected_digest, frozen_binding=binding)


def test_actual_domain_recomputation_reads_retained_materials_and_preserves_originals(actual_archive):
    archive, manifest_digest, binding = actual_archive
    before = {p.relative_to(archive): digest(p.read_bytes()) for p in archive.rglob("*") if p.is_file()}
    result = _run(archive, manifest_digest, binding)
    assert result["supported_checks_verified"], result["checks"]
    assert result["checks"]["pmml_rescoring"]["rows_recomputed"] == 180
    assert result["checks"]["stress_recomputation"]["categories_recomputed"] == 2
    assert result["checks"]["excel_values"]["cells_compared"] > 100
    assert result["archive_unchanged_during_revalidation"]
    assert result["acceptance_claim"] == "not_established"
    assert result["source_authentication"] == "not_established"
    assert result["unsupported"]["native_signature_authentication"] == "original_key_not_available"
    assert result["unsupported"]["narrative_semantic_correctness"] == "requires_independent_review"
    assert before == {p.relative_to(archive): digest(p.read_bytes()) for p in archive.rglob("*") if p.is_file()}


@pytest.mark.parametrize("damage", ["missing_material", "changed_material", "redone_manifest"])
def test_old_external_digest_rejects_missing_or_changed_originals(actual_archive, tmp_path, damage):
    archive, expected_digest, binding = _copy(actual_archive, tmp_path)
    material = next((archive / "inputs").glob("*.csv"))
    if damage == "missing_material":
        material.unlink()
    else:
        material.write_bytes(material.read_bytes() + b"changed")
    if damage == "redone_manifest":
        assert _rewrite_manifest(archive) != expected_digest
    result = _run(archive, expected_digest, binding)
    assert result["checks"]["archive_integrity"]["status"] == "mismatch"
    assert result["checks"]["pmml_rescoring"]["status"] == "unverified"
    assert not result["supported_checks_verified"]


@pytest.mark.parametrize("field", ["source", "run_id", "cases_sha256", "expected_sha256", "model_connection_sha256"])
def test_archive_cannot_supply_its_own_external_binding(actual_archive, tmp_path, field):
    archive, _, binding = _copy(actual_archive, tmp_path)
    path = archive / "bindings.json"
    payload = json.loads(path.read_text())
    payload[field] = {"source_sha256": "b" * 64} if field == "source" else "b" * 64
    path.write_text(json.dumps(payload))
    new_digest = _rewrite_manifest(archive)
    result = _run(archive, new_digest, binding)
    assert result["checks"]["archive_integrity"]["status"] == "verified"
    assert result["checks"]["external_bindings"]["status"] == "mismatch"
    assert result["checks"]["pmml_rescoring"]["status"] == "unverified"


@pytest.mark.parametrize("section", ["effectiveness", "stress_test"])
def test_recomputation_rejects_well_formed_metrics_even_with_redone_inventory_and_success_flags(actual_archive, tmp_path, section):
    archive, _, binding = _copy(actual_archive, tmp_path)
    path = archive / "workspace/tasks" / binding.task_id / "outputs/validation_results.json"
    metrics = json.loads(path.read_text())
    if section == "effectiveness":
        metrics[section]["overall"][0]["ks"] = 0.987654
    else:
        metrics[section]["baseline"]["ks"] = 0.987654
    path.write_text(json.dumps(metrics))
    private_path = archive / "private.json"
    private = json.loads(private_path.read_text())
    private.update(passed=True, verified=True, execution_complete=True)
    private_path.write_text(json.dumps(private))
    result = _run(archive, _rewrite_manifest(archive), binding)
    assert result["checks"]["archive_integrity"]["status"] == "verified"
    assert result["checks"]["pmml_rescoring"]["status"] == "verified"
    target = "metrics_recomputation" if section == "effectiveness" else "stress_recomputation"
    assert result["checks"][target]["status"] == "mismatch", result["checks"]
    assert not result["supported_checks_verified"]


def test_rescoring_rejects_self_consistent_changed_score_artifact(actual_archive, tmp_path):
    archive, _, binding = _copy(actual_archive, tmp_path)
    outputs = archive / "workspace/tasks" / binding.task_id / "outputs"
    path = outputs / "pmml_scores.parquet"
    scores = pd.read_parquet(path)
    scores.loc[0, "pmml_score"] = 0.123456789
    scores.to_parquet(path, index=False)
    scoring_path = outputs / "pmml_scoring_result.json"
    scoring = json.loads(scoring_path.read_text())
    scoring["score_artifact_sha256"] = digest(path.read_bytes())
    scoring_path.write_text(json.dumps(scoring))
    result = _run(archive, _rewrite_manifest(archive), binding)
    assert result["checks"]["archive_integrity"]["status"] == "verified"
    assert result["checks"]["materials"]["status"] == "verified"
    assert result["checks"]["pmml_rescoring"]["status"] == "mismatch", result["checks"]


def test_report_download_bytes_must_match_separately_retained_identity(actual_archive, tmp_path):
    archive, _, binding = _copy(actual_archive, tmp_path)
    path = archive / "workspace/tasks" / binding.task_id / "outputs/validation_report.docx"
    path.write_bytes(path.read_bytes() + b"changed")
    result = _run(archive, _rewrite_manifest(archive), binding)
    assert result["checks"]["metrics_recomputation"]["status"] == "verified", result["checks"]
    assert result["checks"]["report_download_binding"]["status"] == "mismatch"


def test_excel_numeric_cells_are_checked_beyond_file_integrity(actual_archive, tmp_path):
    from openpyxl import load_workbook

    archive, _, binding = _copy(actual_archive, tmp_path)
    path = archive / "workspace/tasks" / binding.task_id / "outputs/validation.xlsx"
    workbook = load_workbook(path)
    workbook["模型效果"]["F2"] = 99.0
    workbook.save(path)
    workbook.close()
    # Even if a new external file identity is supplied, wrong carrier values do
    # not become correctly recomputed business metrics.
    changed = binding.model_copy(update={"report_sha256": {**binding.report_sha256, "excel": digest(path.read_bytes())}})
    result = _run(archive, _rewrite_manifest(archive), changed)
    assert result["checks"]["report_download_binding"]["status"] == "verified", result["checks"]
    assert result["checks"]["excel_values"]["status"] == "mismatch"


def test_absent_confirmation_remains_unverified_with_no_acceptance_promotion(actual_archive):
    archive, expected_digest, binding = actual_archive
    missing = binding.model_copy(update={"confirmed_draft_sha256": None, "report_revision": None,
                                        "report_sha256": {"word": None, "excel": None}})
    result = _run(archive, expected_digest, missing)
    assert result["checks"]["metrics_recomputation"]["status"] == "verified", result["checks"]
    assert result["checks"]["confirmation_record_binding"]["status"] == "unsupported"
    assert result["checks"]["report_download_binding"]["status"] == "unverified"
    assert not result["supported_checks_verified"]


def test_missing_or_untyped_external_binding_is_rejected(actual_archive):
    archive, expected_digest, binding = actual_archive
    with pytest.raises(ValueError, match="explicit FrozenValidationBinding"):
        _run(archive, expected_digest, binding.model_dump())
    incomplete = binding.model_dump()
    del incomplete["source"]
    with pytest.raises(ValidationError):
        FrozenValidationBinding.model_validate(incomplete)


def test_changed_sample_with_new_manifest_still_requires_frozen_input_bytes(actual_archive, tmp_path):
    archive, _, binding = _copy(actual_archive, tmp_path)
    sample_index = next(i for i, m in enumerate(binding.case.materials) if m.role == "sample")
    for path in [archive / f"inputs/{sample_index:03d}.csv", next((archive / "workspace/material_uploads").glob("*/sample.csv"))]:
        path.write_bytes(path.read_bytes() + b"changed")
    result = _run(archive, _rewrite_manifest(archive), binding)
    assert result["checks"]["archive_integrity"]["status"] == "verified"
    assert result["checks"]["materials"]["status"] == "mismatch"


def test_missing_uploaded_copy_is_not_replaced_by_input_copy(actual_archive, tmp_path):
    archive, _, binding = _copy(actual_archive, tmp_path)
    next((archive / "workspace/material_uploads").glob("*/sample.csv")).unlink()
    result = _run(archive, _rewrite_manifest(archive), binding)
    assert result["checks"]["archive_integrity"]["status"] == "verified"
    assert result["checks"]["materials"]["status"] == "mismatch"


@pytest.mark.parametrize("path_case", ["outside_archive", "parent_component", "wrong_selected_material"])
def test_original_absolute_schema_identity_cannot_be_rebound(actual_archive, tmp_path, path_case):
    from contextlib import closing
    import sqlite3
    from marvis.orchestrator.eval.runtime_archive_validation import _read_contract
    from marvis.validation.input_contracts import input_contract_to_dict

    archive, _, binding = _copy(actual_archive, tmp_path)
    task, contract, *_ = _read_contract(archive, binding)
    schema = dict(contract.require_sample_schema().__dict__)
    schema["path"] = {
        "outside_archive": "/another/private/sample.csv",
        "parent_component": str(Path(schema["path"]).parent / ".." / "source" / "sample.csv"),
        "wrong_selected_material": str(Path(task["source_dir"]) / task["dictionary_path"]),
    }[path_case]
    with closing(sqlite3.connect(archive / "workspace/marvis.sqlite")) as conn:
        conn.execute("UPDATE validation_input_contracts SET sample_schema_json=? WHERE task_id=?", (json.dumps(schema), binding.task_id))
        conn.commit()
    # Even a caller deliberately binding this changed contract cannot authorize
    # reading an arbitrary old absolute path or treating another role as sample.
    _, changed_contract, *_ = _read_contract(archive, binding)
    changed = binding.model_copy(update={"input_contract_sha256": digest(input_contract_to_dict(changed_contract))})
    result = _run(archive, _rewrite_manifest(archive), changed)
    assert result["checks"]["task_contract"]["status"] == "verified"
    assert result["checks"]["materials"]["status"] == "mismatch"


def test_duplicate_material_basenames_are_explicitly_unsupported_by_mapping(actual_archive, tmp_path):
    archive, _, binding = _copy(actual_archive, tmp_path)
    case = binding.case.model_copy(deep=True)
    sample = next(m for m in case.materials if m.role == "sample")
    dictionary = next(m for m in case.materials if m.role == "dictionary")
    dictionary.path = "different-folder/" + Path(sample.path).name
    (archive / "case.json").write_text(json.dumps(case.model_dump()))
    changed = binding.model_copy(update={"case": case})
    result = _run(archive, _rewrite_manifest(archive), changed)
    assert result["checks"]["external_bindings"]["status"] == "verified"
    assert result["checks"]["materials"]["status"] == "mismatch"


def test_native_contract_change_cannot_replace_external_contract_identity(actual_archive, tmp_path):
    from contextlib import closing
    import sqlite3

    archive, _, binding = _copy(actual_archive, tmp_path)
    with closing(sqlite3.connect(archive / "workspace/marvis.sqlite")) as conn:
        row = conn.execute("SELECT confirmed_json FROM validation_input_contracts WHERE task_id=?", (binding.task_id,)).fetchone()
        confirmed = json.loads(row[0])
        confirmed["positive_label"], confirmed["negative_label"] = confirmed["negative_label"], confirmed["positive_label"]
        conn.execute("UPDATE validation_input_contracts SET confirmed_json=? WHERE task_id=?", (json.dumps(confirmed), binding.task_id))
        conn.commit()
    result = _run(archive, _rewrite_manifest(archive), binding)
    assert result["checks"]["archive_integrity"]["status"] == "verified"
    assert result["checks"]["task_contract"]["status"] == "mismatch"


def test_word_text_is_checked_beyond_file_integrity(actual_archive, tmp_path):
    import zipfile

    archive, _, binding = _copy(actual_archive, tmp_path)
    path = archive / "workspace/tasks" / binding.task_id / "outputs/validation_report.docx"
    with zipfile.ZipFile(path) as original:
        entries = [(item, original.read(item)) for item in original.infolist()]
    with zipfile.ZipFile(path, "w") as changed_doc:
        for item, content in entries:
            if item.filename == "word/document.xml":
                content = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>forged report</w:t></w:r></w:p></w:body></w:document>'
            changed_doc.writestr(item, content)
    changed = binding.model_copy(update={"report_sha256": {**binding.report_sha256, "word": digest(path.read_bytes())}})
    result = _run(archive, _rewrite_manifest(archive), changed)
    assert result["checks"]["report_download_binding"]["status"] == "verified", result["checks"]
    assert result["checks"]["word_confirmed_text"]["status"] == "mismatch"


def test_digest_is_required_separately_from_archive(actual_archive):
    archive, _, binding = actual_archive
    with pytest.raises(TypeError):
        revalidate_validation_archive(archive, frozen_binding=binding)
    result = _run(archive, "", binding)
    assert result["checks"]["archive_integrity"]["status"] == "mismatch"


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
@pytest.mark.parametrize("element", ["Extension", "TableLocator"])
def test_pmml_resource_filter_is_xml_aware_for_encodings_and_namespaces(encoding, element):
    from marvis.orchestrator.eval.runtime_archive_validation import _validate_pmml_resources

    xml = f'<?xml version="1.0" encoding="{encoding}"?><PMML xmlns:p="urn:foreign"><p:{element}/></PMML>'
    with pytest.raises(ValueError, match="external or extension content"):
        _validate_pmml_resources(xml.encode(encoding))


def test_pmml_resource_filter_rejects_dtd_and_entities():
    from defusedxml.common import DefusedXmlException
    from marvis.orchestrator.eval.runtime_archive_validation import _validate_pmml_resources

    with pytest.raises(DefusedXmlException):
        _validate_pmml_resources(b'<!DOCTYPE PMML [<!ENTITY local "value">]><PMML>&local;</PMML>')


@pytest.mark.parametrize("addition", ["row", "column", "sheet"])
def test_excel_extra_business_values_are_rejected(actual_archive, tmp_path, addition):
    from openpyxl import load_workbook

    archive, _, binding = _copy(actual_archive, tmp_path)
    path = archive / "workspace/tasks" / binding.task_id / "outputs/validation.xlsx"
    workbook = load_workbook(path)
    if addition == "sheet":
        workbook.create_sheet("矛盾的额外结论")["A1"] = "KS=0.999"
    else:
        workbook["模型效果"]["A100" if addition == "row" else "Z2"] = "KS=0.999"
    workbook.save(path)
    workbook.close()
    changed = binding.model_copy(update={"report_sha256": {**binding.report_sha256, "excel": digest(path.read_bytes())}})
    result = _run(archive, _rewrite_manifest(archive), changed)
    assert result["checks"]["report_download_binding"]["status"] == "verified"
    assert result["checks"]["excel_values"]["status"] == "mismatch"


def test_excel_extra_blank_formatting_is_not_business_data(actual_archive, tmp_path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    archive, _, binding = _copy(actual_archive, tmp_path)
    path = archive / "workspace/tasks" / binding.task_id / "outputs/validation.xlsx"
    workbook = load_workbook(path)
    workbook["模型效果"]["Z100"].font = Font(bold=True)
    workbook.save(path)
    workbook.close()
    changed = binding.model_copy(update={"report_sha256": {**binding.report_sha256, "excel": digest(path.read_bytes())}})
    result = _run(archive, _rewrite_manifest(archive), changed)
    assert result["supported_checks_verified"], result["checks"]


def test_recomputation_uses_authenticated_copy_when_original_changes(actual_archive, tmp_path, monkeypatch):
    import marvis.orchestrator.eval.runtime_archive_validation as validation

    archive, expected_digest, binding = _copy(actual_archive, tmp_path)
    original_snapshot = validation._authenticated_snapshot
    destinations = []

    def snapshot_then_replace(source, destination, digest_value):
        original_snapshot(source, destination, digest_value)
        destinations.append(destination)
        next((source / "workspace/material_uploads").glob("*/sample.csv")).write_bytes(b"replaced original")

    monkeypatch.setattr(validation, "_authenticated_snapshot", snapshot_then_replace)
    result = _run(archive, expected_digest, binding)
    assert result["checks"]["recomputation_snapshot"]["status"] == "verified"
    assert result["checks"]["metrics_recomputation"]["status"] == "verified", result["checks"]
    assert result["checks"]["excel_values"]["status"] == "verified"
    assert result["checks"]["archive_integrity"]["status"] == "mismatch"
    assert not result["archive_unchanged_during_revalidation"]
    assert not result["supported_checks_verified"]
    assert destinations and all(not path.exists() for path in destinations)


def test_copy_is_authenticated_before_domain_parsers_run(actual_archive, tmp_path, monkeypatch):
    import marvis.orchestrator.eval.runtime_custody as custody

    archive, expected_digest, binding = _copy(actual_archive, tmp_path)
    original_copy = custody._copy_bounded

    def replace_during_copy(source, destination, expected_size):
        if source.name == "sample.csv":
            raw = source.read_bytes()
            source.write_bytes(bytes([raw[0] ^ 1]) + raw[1:])
        original_copy(source, destination, expected_size)

    monkeypatch.setattr(custody, "_copy_bounded", replace_during_copy)
    result = _run(archive, expected_digest, binding)
    assert result["checks"]["recomputation_snapshot"]["status"] == "mismatch"
    assert result["checks"]["task_contract"]["status"] == "unverified"
    assert result["checks"]["pmml_rescoring"]["status"] == "unverified"
