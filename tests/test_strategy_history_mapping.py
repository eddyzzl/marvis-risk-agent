from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from datetime import date
import hashlib
from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace

from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook, load_workbook
import pytest

from marvis.db_schema import connect
from marvis.packs.strategy.history_mapping import HistoryMappingError, MAPPING_KIND
from marvis.packs.strategy.project_context_tools import load_current_strategy_project_context, run_materialize_project_context
from marvis.packs.strategy.report_bundle_adapters import _history_section
from marvis.repositories.strategy_history_mapping import StrategyHistoryMappingRepository
from marvis.repositories.task_artifacts import TaskArtifactRepository
from marvis.repositories.tasks import TaskRepository
from marvis.routers.plugins import ensure_plugin_admin_token
from marvis.routers.strategy_history_mapping import router
from tests.test_llm_outbound import Wire
from tests.test_strategy_project_context_tool import _request_bound_to_message, _setup


PRIVATE = "PRIVATE_HISTORICAL_REPORT_PARAGRAPH_481"


def _runtime(tmp_path, suffix="csv", *, csv_content=None, json_content=None):
    fx = _setup(tmp_path)
    ensure_plugin_admin_token(fx["settings"].plugin_admin_token_path)
    source = Path(fx["task"].source_dir) / ("历史资料." + suffix)
    if suffix == "csv":
        source.write_text(csv_content if csv_content is not None else "period,version,bad_rate,field,missing,narrative\n2025-01,v1,2.5%,dpd30,," + PRIVATE + "\n", encoding="utf-8")
    elif suffix == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "历史汇总"
        sheet.append(["period", "metric", "version", "formula"])
        sheet.append([date(2025, 1, 31), 0.025, 2, "=B2*100"])
        workbook.save(source)
    elif suffix == "json":
        source.write_text(json_content if json_content is not None else '{"period":"2025-01","version":"v1","metrics":{"bad/rate":0.025},"missing":null}')
    else:
        source.write_bytes(b"opaque " + PRIVATE.encode())
    request = _request_bound_to_message(fx, "请使用 " + source.name + " 作为历史材料。", external_report_filenames=[source.name], explicit_unavailable=[])
    fx["request"] = request
    context = run_materialize_project_context(request, fx["ctx"], fx["runtime"])
    source_id = context["external_artifacts"][0]["artifact_id"]
    app = FastAPI()
    app.state.settings = fx["settings"]
    app.include_router(router)
    @app.middleware("http")
    async def actor(request, call_next):
        identity = request.headers.get("x-test-actor")
        request.state.private_client_authenticated = bool(identity)
        request.state.local_principal = SimpleNamespace(id=identity) if identity else None
        return await call_next(request)
    fx.update(client=TestClient(app, headers={"x-test-actor": "maker"}), source_id=source_id, context=context,
              repo=StrategyHistoryMappingRepository(fx["settings"]))
    return fx


def _base(fx):
    return f"/api/tasks/{fx['task'].id}/strategy/history-mapping"


def _binding(fx, locator, kind="metric", name="bad_rate", unit="percent", period_format="unknown", key="b1"):
    cells = fx["client"].get(_base(fx) + f"/sources/{fx['source_id']}/cells").json()
    cell = next(row for row in cells["cells"] if row["locator"] == locator)
    return {"binding_key": key, "source_id": fx["source_id"], "source_hash": cells["source_hash"], "locator": locator,
            "cell_hash": cell["cell_hash"], "target": {"kind": kind, "name": name, "unit": unit, "period_format": period_format}, "period_key": None, "version_key": None}


def _propose(fx, bindings=None, **kwargs):
    head = fx["client"].get(_base(fx)).json()["head"]
    response = fx["client"].post(_base(fx) + "/proposals", json={"expected_head": head, "bindings": bindings, **kwargs})
    assert response.status_code == 200, response.text
    return response.json()


def _confirm(fx, proposal):
    return fx["client"].post(_base(fx) + "/confirm", json={"proposal_token": proposal["proposal_token"]})


def test_csv_exact_bindings_confirm_unknowns_and_refresh_into_typed_report(tmp_path):
    fx = _runtime(tmp_path)
    period = _binding(fx, {"format": "csv", "row": 2, "column": 1}, kind="period", name="effective_period", unit="unknown", period_format="month", key="period")
    version = _binding(fx, {"format": "csv", "row": 2, "column": 2}, kind="strategy_version", name="strategy_version", unit="unknown", key="version")
    metric = _binding(fx, {"format": "csv", "row": 2, "column": 3})
    metric.update(period_key="period", version_key="version")
    field = _binding(fx, {"format": "csv", "row": 2, "column": 4}, kind="field", name="delinquency_field", unit="unknown", key="field")
    missing = _binding(fx, {"format": "csv", "row": 2, "column": 5}, name="volume", unit="count", key="missing")
    proposal = _propose(fx, [period, version, metric, field, missing])
    assert fx["repo"].current(fx["task"].id) is None
    assert proposal["bindings"][2]["value"] == 2.5
    assert proposal["bindings"][4]["value"] is None and proposal["bindings"][4]["availability"] == "unknown"
    saved = _confirm(fx, proposal)
    assert saved.status_code == 200, saved.text
    mapping = saved.json()["mapping"]
    assert mapping["validation_status"] == "historical_reported_unvalidated"
    assert mapping["bindings"][2]["locator"] == metric["locator"] and mapping["bindings"][2]["cell_hash"] == metric["cell_hash"]
    assert _confirm(fx, proposal).json()["mapping"] == mapping
    before = fx["context"]["revision"]
    refreshed = run_materialize_project_context({**fx["request"], "expected_revision": before["revision"], "expected_revision_id": before["revision_id"], "expected_state_hash": before["state_hash"]}, fx["ctx"], fx["runtime"])
    revision = load_current_strategy_project_context(fx["runtime"], task_id=fx["task"].id)
    assert revision == refreshed["revision"]
    snapshot = revision["state"]["current_project_snapshot"]
    assert snapshot["metric_observation_refs"] == before["state"]["current_project_snapshot"]["metric_observation_refs"]
    assert snapshot["status_fields"] == before["state"]["current_project_snapshot"]["status_fields"]
    section = _history_section(revision, {"kind": "strategy_project_context", "ref_id": "context", "content_hash": "a" * 64})
    table = next(table for table in section["tables"] if table["table_id"] == "confirmed_history_mapping")
    cells = next(row for row in table["rows"] if row["row_id"] == "b1")["cells"]
    assert cells["value"]["value"] == 2.5 and cells["version"]["value"] == "v1"
    assert "未复算验证" in cells["status"]["value"]
    assert next(row for row in table["rows"] if row["row_id"] == "missing")["cells"]["value"]["availability"] == "unavailable"
    assert section["stage_evidence"] == []
    assert PRIVATE not in str(table)


@pytest.mark.parametrize("suffix,locator,kind,name,unit,period_format,expected", [
    ("xlsx", {"format": "xlsx", "sheet": "历史汇总", "cell": "B2"}, "metric", "bad_rate", "ratio", "unknown", 0.025),
    ("xlsx", {"format": "xlsx", "sheet": "历史汇总", "cell": "A2"}, "period", "effective_period", "unknown", "date", {"granularity": "date", "date": "2025-01-31"}),
    ("json", {"format": "json", "pointer": "/metrics/bad~1rate"}, "metric", "bad_rate", "ratio", "unknown", 0.025),
])
def test_xlsx_and_json_pointer_have_exact_native_locations(tmp_path, suffix, locator, kind, name, unit, period_format, expected):
    fx = _runtime(tmp_path, suffix)
    binding = _binding(fx, locator, kind, name, unit, period_format)
    result = _confirm(fx, _propose(fx, [binding]))
    assert result.status_code == 200 and result.json()["mapping"]["bindings"][0]["value"] == expected


def test_natural_and_manual_inputs_share_local_governance_and_make_zero_external_requests(monkeypatch, tmp_path):
    fx = _runtime(tmp_path)
    wire = Wire(monkeypatch)
    proposal = _propose(fx, instruction="将 CSV 第2行第3列映射为指标 bad_rate，单位 percent；将 CSV 第2行第1列映射为期间 effective_period，格式 month", source_id=fx["source_id"])
    assert [item["value"] for item in proposal["bindings"]] == [2.5, {"granularity": "month", "month": "2025-01"}]
    assert _confirm(fx, proposal).status_code == 200 and wire.requests == []
    head = fx["repo"].head(fx["repo"].current(fx["task"].id))
    unsupported = fx["client"].post(_base(fx) + "/proposals", json={"expected_head": head, "source_id": fx["source_id"], "instruction": "将 CSV 第2行第3列映射为指标 bad_rate，单位 percent；并且排除低风险客户"})
    assert unsupported.status_code == 422 and unsupported.json()["detail"]["code"] == "history_instruction_needs_clarification"
    assert wire.requests == [] and fx["repo"].current(fx["task"].id)["revision"] == 1


@pytest.mark.parametrize("mutation", ["actor", "token", "cell", "source", "head"])
def test_stale_unowned_or_modified_confirmation_never_overwrites_mapping(tmp_path, mutation):
    fx = _runtime(tmp_path)
    binding = _binding(fx, {"format": "csv", "row": 2, "column": 3})
    proposal = _propose(fx, [binding])
    if mutation == "actor":
        fx["client"].headers["x-test-actor"] = "checker"
    elif mutation == "token":
        proposal["proposal_token"] = proposal["proposal_token"][:-8] + "aaaaaaaa"
    elif mutation == "cell":
        binding["cell_hash"] = "b" * 64
        response = fx["client"].post(_base(fx) + "/proposals", json={"expected_head": {"revision": 0, "artifact_id": None, "content_hash": None}, "bindings": [binding]})
        assert response.status_code == 409
        assert fx["repo"].current(fx["task"].id) is None
        return
    elif mutation == "source":
        record = fx["runtime"].task_artifacts.get_for_task(fx["task"].id, fx["source_id"])
        Path(record["path"]).write_text("changed source")
    else:
        first = _propose(fx, [deepcopy(binding)])
        assert _confirm(fx, first).status_code == 200
    response = _confirm(fx, proposal)
    assert response.status_code == 409
    with connect(fx["settings"].db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM task_artifacts WHERE kind=?", (MAPPING_KIND,)).fetchone()[0] == (1 if mutation == "head" else 0)


def test_formula_prose_opaque_documents_and_missing_unit_are_not_measured_metrics(tmp_path):
    fx = _runtime(tmp_path / "formula", "xlsx")
    formula = _binding(fx, {"format": "xlsx", "sheet": "历史汇总", "cell": "D2"})
    with pytest.raises(HistoryMappingError, match="公式"):
        fx["repo"].propose(task_id=fx["task"].id, actor_id="maker", expected_head=fx["repo"].head(None), bindings=[formula])
    fx = _runtime(tmp_path / "prose")
    prose = _binding(fx, {"format": "csv", "row": 2, "column": 6})
    response = fx["client"].post(_base(fx) + "/proposals", json={"expected_head": fx["repo"].head(None), "bindings": [prose]})
    assert response.status_code == 422 and response.json()["detail"]["code"] == "history_metric_not_numeric"
    no_unit = _binding(fx, {"format": "csv", "row": 2, "column": 3}, unit="unknown")
    assert _propose(fx, [no_unit])["bindings"][0]["availability"] == "unknown"
    fx = _runtime(tmp_path / "pdf", "pdf")
    sources = fx["client"].get(_base(fx) + "/sources").json()["sources"]
    assert sources[0]["status"] == "unavailable" and sources[0]["code"] == "history_source_unlocatable"
    assert "CSV/XLSX" in sources[0]["next_action"]
    fx["client"].headers.pop("x-test-actor")
    assert fx["client"].get(_base(fx) + "/sources").status_code == 403


def test_mapping_audit_tamper_and_persistence_failure_are_fail_closed(monkeypatch, tmp_path):
    fx = _runtime(tmp_path)
    binding = _binding(fx, {"format": "csv", "row": 2, "column": 3})
    proposal = _propose(fx, [binding])
    import marvis.repositories.strategy_history_mapping as module
    original = module._write_audit_row
    monkeypatch.setattr(module, "_write_audit_row", lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("synthetic failure")))
    assert _confirm(fx, proposal).status_code == 409
    assert fx["repo"].current(fx["task"].id) is None
    assert not list((fx["settings"].tasks_dir / fx["task"].id / "strategy_history_mappings").glob("*.json"))
    monkeypatch.setattr(module, "_write_audit_row", original)
    saved = _confirm(fx, proposal).json()["mapping"]
    with connect(fx["settings"].db_path) as connection:
        connection.execute("DELETE FROM audit WHERE kind='strategy.history_mapping.confirmed' AND target_ref=?", (saved["artifact_id"],))
    assert fx["client"].get(_base(fx)).status_code == 409


def test_historical_declared_metric_renders_in_all_four_real_report_artifacts(tmp_path):
    from tests.test_strategy_report_bundle_tools import _run, _setup as report_setup
    fx = report_setup(tmp_path)
    ensure_plugin_admin_token(fx["settings"].plugin_admin_token_path)
    repository = StrategyHistoryMappingRepository(fx["settings"])
    source = Path(fx["task"].source_dir) / "history.csv"
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_text("bad_rate,narrative\n2.5%," + PRIVATE + "\n")
    message = TaskRepository(fx["settings"].db_path).add_agent_message(fx["task"].id, role="user", stage="chat", content="请使用 history.csv 作为历史材料。")
    ref = fx["request"]["project_context_ref"]
    request = {"expected_revision": ref["expected_revision"], "expected_revision_id": ref["expected_revision_id"], "expected_state_hash": ref["expected_state_hash"],
               "user_message_ref": {"message_id": message["id"], "content_hash": hashlib.sha256(message["content"].encode()).hexdigest()},
               "as_of": "2026-07-23", "scope": "贷前准入策略", "business_context": {}, "explicit_unavailable": [], "external_report_filenames": ["history.csv"]}
    registered = run_materialize_project_context(request, fx["ctx"], fx["runtime"])
    source_id = registered["external_artifacts"][0]["artifact_id"]
    page = repository.cells(fx["task"].id, source_id)
    cell = next(item for item in page["cells"] if item["locator"] == {"format": "csv", "row": 2, "column": 1})
    binding = {"binding_key": "rate", "source_id": source_id, "source_hash": page["source_hash"], "locator": cell["locator"], "cell_hash": cell["cell_hash"],
               "target": {"kind": "metric", "name": "bad_rate", "unit": "percent", "period_format": "unknown"}, "period_key": None, "version_key": None}
    proposal = repository.propose(task_id=fx["task"].id, actor_id="maker", expected_head=repository.head(None), bindings=[binding])
    mapping = repository.confirm(task_id=fx["task"].id, actor_id="maker", token=proposal["proposal_token"])
    prior = registered["revision"]
    refreshed = run_materialize_project_context({**request, "expected_revision": prior["revision"], "expected_revision_id": prior["revision_id"], "expected_state_hash": prior["state_hash"]}, fx["ctx"], fx["runtime"])
    revision = refreshed["revision"]
    fx["request"]["project_context_ref"] = {"artifact_id": refreshed["context_artifact"]["artifact_id"], "expected_artifact_content_hash": refreshed["context_artifact"]["content_hash"],
        "expected_revision": revision["revision"], "expected_revision_id": revision["revision_id"], "expected_state_hash": revision["state_hash"]}
    output = _run(fx)
    registry = TaskArtifactRepository(fx["settings"].db_path)
    raw = {item["format"]: Path(registry.get_for_task(fx["task"].id, item["artifact_id"])["path"]).read_bytes() for item in output["artifacts"]}
    bundle = json.loads(raw["json"])
    history = next(section for section in bundle["sections"] if section["key"] == "historical_versions")
    table = next(table for table in history["tables"] if table["table_id"] == "confirmed_history_mapping")
    assert table["rows"][0]["cells"]["value"]["value"] == 2.5
    assert table["rows"][0]["cells"]["period"]["availability"] == "unavailable"
    assert table["source_refs"][0]["ref_id"] == mapping["artifact_id"]
    assert "oot_validated" not in bundle["effect_stages"] and PRIVATE not in raw["json"].decode()
    assert "未复算验证" in raw["markdown"].decode() and "2.5" in raw["markdown"].decode()
    workbook = load_workbook(BytesIO(raw["xlsx"]), read_only=True, data_only=True)
    try:
        values = [value for row in workbook["02_history"].iter_rows(values_only=True) for value in row]
        assert 2.5 in values and any("未复算验证" in str(value) for value in values)
    finally:
        workbook.close()
    document = Document(BytesIO(raw["docx"]))
    docx_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "2.5" in docx_text and "未复算验证" in docx_text and PRIVATE not in docx_text


def test_known_metric_units_and_complete_chinese_instruction_are_enforced(tmp_path):
    fx = _runtime(tmp_path)
    proposal = _propose(fx, instruction="将 CSV 第2行第3列映射为指标 坏账率，单位 百分比", source_id=fx["source_id"])
    assert proposal["bindings"][0]["target"]["name"] == "bad_rate" and proposal["bindings"][0]["value"] == 2.5
    bad_unit = _binding(fx, {"format": "csv", "row": 2, "column": 3}, unit="count")
    response = fx["client"].post(_base(fx) + "/proposals", json={"expected_head": fx["repo"].head(None), "bindings": [bad_unit]})
    assert response.status_code == 422 and response.json()["detail"]["code"] == "history_metric_unit_invalid"
    fx = _runtime(tmp_path / "json", "json")
    response = fx["client"].post(_base(fx) + "/proposals", json={"expected_head": fx["repo"].head(None), "source_id": fx["source_id"], "instruction": "将 JSON「/metrics」映射为指标 bad_rate，单位 ratio"})
    assert response.status_code == 422 and response.json()["detail"]["code"] == "history_json_non_scalar"


def test_precision_is_not_silently_rounded_and_explicit_clear_preserves_old_revision(tmp_path):
    fx = _runtime(tmp_path, csv_content="large_count,precise_rate,zero\n9007199254740993,0.10000000000000000000001,0\n")
    for column, name, unit in [(1, "volume", "count"), (2, "bad_rate", "ratio")]:
        binding = _binding(fx, {"format": "csv", "row": 2, "column": column}, name=name, unit=unit)
        response = fx["client"].post(_base(fx) + "/proposals", json={"expected_head": fx["repo"].head(None), "bindings": [binding]})
        assert response.status_code == 422 and response.json()["detail"]["code"] == "history_numeric_precision_unsupported"
    zero = _binding(fx, {"format": "csv", "row": 2, "column": 3}, name="volume", unit="count")
    saved = _confirm(fx, _propose(fx, [zero])).json()["mapping"]
    assert saved["bindings"][0]["value"] == 0 and saved["bindings"][0]["availability"] == "present"
    cleared = _confirm(fx, _propose(fx, [])).json()["mapping"]
    assert cleared["bindings"] == [] and cleared["revision"] == 2
    assert cleared["parent_ref"] == fx["repo"].head(saved)
    with connect(fx["settings"].db_path) as connection:
        original = fx["repo"].get_on_connection(connection, fx["task"].id, saved["artifact_id"])
        assert original["bindings"][0]["value"] == 0
    fx = _runtime(tmp_path / "json", "json", json_content='{"metrics":{"bad/rate":0.10000000000000000000001}}')
    precise = _binding(fx, {"format": "json", "pointer": "/metrics/bad~1rate"}, unit="ratio")
    response = fx["client"].post(_base(fx) + "/proposals", json={"expected_head": fx["repo"].head(None), "bindings": [precise]})
    assert response.status_code == 422 and response.json()["detail"]["code"] == "history_numeric_precision_unsupported"


def test_parallel_confirmation_uses_one_cas_revision_and_preserves_other_edit(tmp_path):
    fx = _runtime(tmp_path)
    binding = _binding(fx, {"format": "csv", "row": 2, "column": 3})
    proposals = [_propose(fx, [binding]) for _ in range(2)]
    def confirm(proposal):
        try:
            return fx["repo"].confirm(task_id=fx["task"].id, actor_id="maker", token=proposal["proposal_token"])["revision"]
        except HistoryMappingError as exc:
            return exc.code
    with ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = list(pool.map(confirm, proposals))
    assert set(outcomes) == {1, "history_mapping_revision_conflict"}
    assert fx["repo"].current(fx["task"].id)["revision"] == 1
    with connect(fx["settings"].db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM task_artifacts WHERE kind=?", (MAPPING_KIND,)).fetchone()[0] == 1


def test_user_free_text_cannot_impersonate_authenticated_history_mapping(tmp_path):
    fx = _runtime(tmp_path)
    revision = deepcopy(fx["context"]["revision"])
    revision["state"]["current_project_snapshot"]["user_context_fields"].append({"field_path": "history_mapping.confirmed", "field": {
        "origin": "user", "availability": "present", "value": {"bindings": [{"value": 0.99}]}, "source_refs": [], "as_of": None, "blocking": "none", "note": None,
    }})
    section = _history_section(revision, {"kind": "strategy_project_context", "ref_id": "context", "content_hash": "a" * 64})
    assert all(table["table_id"] != "confirmed_history_mapping" for table in section["tables"])
