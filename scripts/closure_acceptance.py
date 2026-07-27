#!/usr/bin/env python3
"""Generate a machine pre-check for the T4-3 real-material checklist.

This script deliberately does not sign the checklist and never invents external
finance/risk figures.  It inspects an immutable completed plan, its evidence
envelopes, model card and artifacts, then separates:

* machine-verifiable PASS/FAIL/N/A evidence; and
* the external reconciliation/signature work that only a human can complete.

Exit codes:
    0  all machine-verifiable checks pass (manual reconciliation may remain)
    1  at least one machine-verifiable check fails
    2  the requested completed modeling task/evidence cannot be found
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import pickle
import re
import shlex
import sqlite3
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import numpy as np
from joblib.numpy_pickle import (
    NDArrayWrapper,
    NumpyArrayWrapper,
    NumpyUnpickler,
    Unpickler,
)
from openpyxl import load_workbook


_REQUIRED_REGRESSION_TESTS = (
    "tests/test_dirty_shape_regression.py",
    "tests/test_reconcile_reference_numbers.py",
)
_SHA256_RE = re.compile(r"^sha256:[0-9a-f]{64}$")
_PMML_APPLICABLE_RECIPES = frozenset({"lgb", "xgb", "lr", "scorecard"})
_PMML_UNSUPPORTED_RECIPES = frozenset(
    {
        "catboost",
        "mlp",
        "ensemble",
        "lgb_regressor",
        "xgb_regressor",
        "lr_regressor",
        "mlp_regressor",
        "lgb_multiclass",
        "xgb_multiclass",
        "lr_multiclass",
        "mlp_multiclass",
    }
)
_STOCHASTIC_TOOLS = frozenset(
    {
        "reject_inference",
        "prepare_modeling_frame",
        "make_split",
        "resolve_special_values",
        "select_features",
        "tune_hyperparameters",
        "train_model",
        "train_models",
        "calibrate_model",
        "score_dataset",
        "monitor_run",
    }
)
_COMPLETED_STEP_STATUSES = frozenset({"done", "completed", "succeeded"})


@dataclass(frozen=True)
class Check:
    check_id: str
    status: str
    summary: str
    evidence: str
    human_action: str = ""


def _json(value: str | None, default: Any) -> Any:
    if not value:
        return default
    try:
        return json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return default


def _latest_done_modeling_plan(connection: sqlite3.Connection) -> sqlite3.Row | None:
    return connection.execute(
        """
        SELECT p.id AS plan_id, p.task_id, p.template_id, p.status, p.updated_at,
               t.model_name, t.model_version, t.target_col, t.time_col
          FROM plans p
          JOIN tasks t ON t.id = p.task_id
         WHERE t.task_type = 'modeling' AND p.status = 'done'
         ORDER BY p.updated_at DESC, p.id DESC
         LIMIT 1
        """
    ).fetchone()


def _task_plan(connection: sqlite3.Connection, task_id: str) -> sqlite3.Row | None:
    return connection.execute(
        """
        SELECT p.id AS plan_id, p.task_id, p.template_id, p.status, p.updated_at,
               t.model_name, t.model_version, t.target_col, t.time_col
          FROM plans p
          JOIN tasks t ON t.id = p.task_id
         WHERE p.task_id = ?
         ORDER BY (p.status = 'done') DESC, p.updated_at DESC, p.id DESC
         LIMIT 1
        """,
        (task_id,),
    ).fetchone()


def _completed_task_steps(
    connection: sqlite3.Connection,
    task_id: str | None,
) -> tuple[sqlite3.Row | None, list[dict[str, Any]]]:
    if not task_id:
        return None, []
    plan = _task_plan(connection, str(task_id))
    if plan is None or str(plan["status"]) != "done":
        return plan, []
    return plan, _step_record_rows(connection, str(plan["plan_id"]))


def _tool_output(
    rows: list[dict[str, Any]],
    tool_name: str,
) -> dict[str, Any]:
    for row in rows:
        if str(row.get("tool_name") or "") == tool_name:
            output = row.get("output")
            return output if isinstance(output, dict) else {}
    return {}


def _vintage_semantics_check(
    rows: list[dict[str, Any]],
    *,
    task_id: str | None,
) -> Check:
    output = _tool_output(rows, "vintage_curve")
    semantics = str(output.get("label_semantics") or "")
    cohorts = output.get("cohorts")
    curves = output.get("curves")
    valid = (
        semantics in {"incremental", "snapshot"}
        and isinstance(cohorts, list)
        and bool(cohorts)
        and isinstance(curves, dict)
        and bool(curves)
    )
    return Check(
        "A2",
        "PASS" if valid else "FAIL",
        (
            "Vintage 任务记录了标签语义并生成了非空曲线。"
            if valid
            else "缺少可验证的 Vintage 标签语义或曲线证据。"
        ),
        (
            f"vintage_task_id={task_id or '<missing>'}; "
            f"label_semantics={semantics or '<missing>'}; "
            f"cohort_count={len(cohorts) if isinstance(cohorts, list) else 0}"
        ),
        (
            ""
            if valid
            else "传入同一真实材料的已完成 --vintage-task-id，并确认 incremental/snapshot。"
        ),
    )


def _join_acceptance_checks(
    rows: list[dict[str, Any]],
    *,
    task_id: str | None,
) -> list[Check]:
    proposed = _tool_output(rows, "propose_join")
    confirmed = _tool_output(rows, "confirm_join")
    executed = _tool_output(rows, "execute_join")
    joins = proposed.get("joins") if isinstance(proposed.get("joins"), list) else []
    diagnostics = [
        item.get("diagnostics")
        for item in joins
        if isinstance(item, dict) and isinstance(item.get("diagnostics"), dict)
    ]
    per_table = (
        executed.get("per_table")
        if isinstance(executed.get("per_table"), list)
        else []
    )
    execution_complete = bool(
        executed.get("result_dataset_id")
        and per_table
        and confirmed.get("status") == "confirmed"
    )
    precision_checked = bool(diagnostics) and all(
        "precision_loss_columns" in item for item in diagnostics
    )
    precision_flags = sorted({
        str(column)
        for item in diagnostics
        for column in (item.get("precision_loss_columns") or [])
    })
    dtype_checked = bool(diagnostics) and all(
        "key_dtype_divergences" in item for item in diagnostics
    )
    keyed = bool(joins) and all(
        isinstance(item.get("key_pairs"), list) and bool(item.get("key_pairs"))
        for item in joins
        if isinstance(item, dict)
    )
    a4_valid = execution_complete and precision_checked
    a5_valid = execution_complete and dtype_checked and keyed

    finite_match_rates = bool(per_table)
    for item in per_table:
        try:
            rate = float(item["match_rate"])
        except (KeyError, TypeError, ValueError):
            finite_match_rates = False
            break
        if not math.isfinite(rate) or not 0.0 <= rate <= 1.0:
            finite_match_rates = False
            break
    reconciliation = (
        proposed.get("reconcile_summary")
        if isinstance(proposed.get("reconcile_summary"), dict)
        else {}
    )
    join_reconciliations = [
        item.get("reconcile")
        for item in joins
        if isinstance(item, dict) and isinstance(item.get("reconcile"), dict)
    ]
    independently_reconciled = bool(join_reconciliations) and all(
        item.get("consistent") is True for item in join_reconciliations
    )
    anchor_rows = executed.get("anchor_rows")
    joined_rows = executed.get("joined_rows")
    b5_valid = bool(
        execution_complete
        and finite_match_rates
        and independently_reconciled
        and reconciliation.get("blocking") is False
        and executed.get("fan_out") is False
        and isinstance(anchor_rows, int)
        and joined_rows == anchor_rows
    )
    task_evidence = f"join_task_id={task_id or '<missing>'}"
    return [
        Check(
            "A4",
            "PASS" if a4_valid else "FAIL",
            (
                "Join 诊断已执行长 ID 精度检查并完成受控拼接。"
                if a4_valid
                else "缺少长 ID 精度诊断或已完成拼接证据。"
            ),
            (
                f"{task_evidence}; precision_flags={precision_flags}; "
                f"execution_complete={execution_complete}"
            ),
            (
                ""
                if a4_valid
                else "传入同一真实材料的已完成 --join-task-id，并处理全部精度红旗。"
            ),
        ),
        Check(
            "A5",
            "PASS" if a5_valid else "FAIL",
            (
                "Join 键的类型/匹配诊断已记录并完成受控拼接。"
                if a5_valid
                else "缺少空白/零填充键所需的类型与匹配诊断证据。"
            ),
            (
                f"{task_evidence}; diagnostics={len(diagnostics)}; "
                f"keyed={keyed}; execution_complete={execution_complete}"
            ),
            (
                ""
                if a5_valid
                else "补充完成的 join 任务，并核对空白键、零填充和跨文件 dtype。"
            ),
        ),
        Check(
            "B5-INTERNAL",
            "PASS" if b5_valid else "FAIL",
            (
                "Join 匹配率具备独立双路对账，且结果保持 anchor 行数。"
                if b5_valid
                else "Join 匹配率缺少独立双路对账、有限值或行数保持证据。"
            ),
            (
                f"{task_evidence}; match_rates="
                f"{[item.get('match_rate') for item in per_table if isinstance(item, dict)]}; "
                f"anchor_rows={anchor_rows}; joined_rows={joined_rows}; "
                f"independently_reconciled={independently_reconciled}"
            ),
            (
                ""
                if b5_valid
                else "修复 join 红旗并重跑，直至内部双路对账与行数保持通过。"
            ),
        ),
        Check(
            "B5-EXTERNAL",
            "MANUAL",
            "内部 join 证据不能替代真实键抽样核对。",
            f"{task_evidence}; external sampled-key reconciliation not supplied",
            "抽样核对真实键匹配并完成 B5 签字。",
        ),
    ]


def _step_record_rows(
    connection: sqlite3.Connection,
    plan_id: str,
) -> list[dict[str, Any]]:
    rows = connection.execute(
        """
        SELECT s.id AS step_id, s.idx, s.title, s.tool_plugin, s.tool_name,
               s.status, s.output_ref,
               o.output_json, o.evidence_json
          FROM plan_steps s
          LEFT JOIN plan_step_outputs o ON o.step_id = s.id
         WHERE s.plan_id = ?
         ORDER BY s.idx
        """,
        (plan_id,),
    ).fetchall()
    records: list[dict[str, Any]] = []
    for row in rows:
        records.append(
            {
                **dict(row),
                "output": _json(row["output_json"], {}),
                "evidence": _json(row["evidence_json"], {}),
            }
        )
    return records


def _step_records(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {str(row["title"]): row for row in rows}


def _artifact_row(
    connection: sqlite3.Connection, artifact_id: str | None
) -> sqlite3.Row | None:
    if not artifact_id:
        return None
    return connection.execute(
        "SELECT * FROM model_artifacts WHERE id = ?", (artifact_id,)
    ).fetchone()


def _artifact_selected_features(
    artifact: sqlite3.Row | None,
    refine: dict[str, Any],
) -> list[str] | None:
    if artifact is not None and "feature_list_json" in artifact.keys():
        parsed = _json(artifact["feature_list_json"], None)
        if isinstance(parsed, list):
            return [str(feature) for feature in parsed if str(feature)]
    selected = refine.get("selected")
    if isinstance(selected, list):
        return [str(feature) for feature in selected if str(feature)]
    return None


def _numeric_values(values) -> set[float]:
    normalized: set[float] = set()
    for value in values if isinstance(values, (list, tuple)) else []:
        raw = value[0] if isinstance(value, (list, tuple)) and value else value
        try:
            normalized.add(float(raw))
        except (TypeError, ValueError):
            continue
    return normalized


def _sentinel_step_value_sets(
    preprocessing_steps,
    column: str,
) -> list[set[float]]:
    found: list[set[float]] = []
    for step in preprocessing_steps if isinstance(preprocessing_steps, list) else []:
        if not isinstance(step, dict) or str(step.get("kind") or "") != "sentinel":
            continue
        if column not in {str(item) for item in step.get("columns") or []}:
            continue
        params = step.get("params")
        if not isinstance(params, dict) or column not in params:
            continue
        found.append(_numeric_values(params[column]))
    return found


def _a1_check(
    *,
    sentinel_columns: dict,
    artifact_params: dict,
    selected_features: list[str] | None,
) -> Check:
    detected = {
        str(column): _numeric_values(rows)
        for column, rows in sentinel_columns.items()
        if _numeric_values(rows)
    }
    if not detected:
        return Check(
            "A1",
            "PASS",
            "未发现需要治理的哨兵/特殊值。",
            "detected=0; relevant=0",
        )
    if selected_features is None:
        return Check(
            "A1",
            "FAIL",
            "检测到哨兵值，但无法取得冠军模型的最终入模特征证据。",
            f"detected={len(detected)}; selected=unknown",
        )

    selected = set(selected_features)
    relevant = sorted(selected & set(detected))
    governance = artifact_params.get("special_value_governance")
    governance = governance if isinstance(governance, dict) else {}
    preprocessing_steps = artifact_params.get("preprocessing_steps")
    traceable = artifact_params.get("preprocessing_chain_traceable")
    problems: dict[str, str] = {}
    evidence_rows: list[str] = []
    for column in relevant:
        record = governance.get(column)
        if not isinstance(record, dict):
            problems[column] = "missing governance"
            continue
        action = str(record.get("action") or "")
        fingerprint = str(
            record.get("fingerprint") or record.get("decision_fingerprint") or ""
        ).strip()
        recorded_values = _numeric_values(record.get("detected_values") or [])
        if recorded_values != detected[column]:
            problems[column] = "governance values mismatch"
        elif not fingerprint:
            problems[column] = "missing fingerprint"
        elif action == "mask":
            exact = detected[column] in _sentinel_step_value_sets(
                preprocessing_steps,
                column,
            )
            if not exact:
                problems[column] = "missing exact sentinel preprocessing step"
            evidence_rows.append(f"{column}=mask(exact_step={str(exact).lower()})")
        elif action == "retain":
            confirmed = record.get("confirmed") is True
            reason = str(record.get("reason") or "").strip()
            source_dataset_id = str(record.get("source_dataset_id") or "").strip()
            source_content_hash = str(
                record.get("source_dataset_content_hash") or ""
            ).strip()
            if not confirmed:
                problems[column] = "retain not explicitly confirmed"
            elif not reason:
                problems[column] = "retain missing reason"
            elif not source_dataset_id or not source_content_hash:
                problems[column] = "retain missing source dataset fingerprint"
            evidence_rows.append(
                (
                    f"{column}=retain(confirmed={str(confirmed).lower()},"
                    f"reason={str(bool(reason)).lower()},"
                    f"source_fingerprint={str(bool(source_content_hash)).lower()})"
                )
            )
        else:
            # A dropped feature must not survive in the artifact feature list.
            problems[column] = f"selected feature has action {action or 'missing'}"

    base_evidence = (
        f"selected={len(selected)}; detected={len(detected)}; "
        f"relevant={len(relevant)}; preprocessing_chain_traceable={traceable}"
    )
    if problems:
        detail = "; ".join(f"{column}: {reason}" for column, reason in problems.items())
        return Check(
            "A1",
            "FAIL",
            "冠军模型仍包含未被逐列治理的哨兵/特殊值特征。",
            f"{base_evidence}; {detail}",
        )
    return Check(
        "A1",
        "PASS",
        "最终入模的哨兵/特殊值特征均有逐列、可核验的治理证据。",
        "; ".join([base_evidence, *evidence_rows]),
    )


_NATIVE_MODEL_SUFFIXES = frozenset({".pkl", ".joblib", ".txt"})
_REPORT_FILENAME_UNSAFE_RE = re.compile(r"[^0-9A-Za-z_-]+")


def _is_within(candidate: Path, root: Path) -> bool:
    try:
        candidate.relative_to(root)
    except ValueError:
        return False
    return True


def _delivery_path(
    path: str | None,
    workspace: Path,
    task_id: str,
    artifact_name: str,
) -> Path | None:
    if not isinstance(path, str) or not path.strip() or "\x00" in path:
        return None
    try:
        task_root = (workspace / "tasks" / task_id).resolve()
        allowed_root = (
            task_root / "outputs"
            if artifact_name == "report"
            else task_root / "modeling_artifacts"
        ).resolve()
        raw = Path(path.strip()).expanduser()
        candidate = (raw if raw.is_absolute() else allowed_root / raw).resolve()
    except (OSError, RuntimeError):
        return None
    if not _is_within(candidate, allowed_root):
        return None
    return candidate


def _valid_xlsx(path: Path) -> bool:
    if path.suffix.lower() != ".xlsx" or not zipfile.is_zipfile(path):
        return False
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            return bool(workbook.worksheets)
        finally:
            workbook.close()
    # Parser implementations may surface lxml/openpyxl-specific exceptions.
    # Any parse failure is a closed validation failure, never an acceptance crash.
    except Exception:
        return False


class _StructuralPlaceholder:
    """Inert target for pickle globals while validating opcode structure."""

    def __new__(cls, *args, **kwargs):
        del args, kwargs
        return object.__new__(cls)

    def __setstate__(self, state) -> None:
        del state

    # Some estimators persist custom ``dict``/``list``/``set`` subclasses.
    # The restricted unpickler deliberately replaces their globals with this
    # inert object, so the container opcodes must remain parseable without
    # constructing or mutating the original class.
    def __setitem__(self, key, value) -> None:
        del key, value

    def append(self, value) -> None:
        del value

    def extend(self, values) -> None:
        del values

    def add(self, value) -> None:
        del value

    def update(self, *args, **kwargs) -> None:
        del args, kwargs


class _SafePlainUnpickler(Unpickler):
    def find_class(self, module: str, name: str):
        if name == "dtype" and module.startswith("numpy"):
            return np.dtype
        return _StructuralPlaceholder


class _StructuralJoblibUnpickler(NumpyUnpickler):
    """Parse MARVIS' uncompressed joblib stream without constructing its model.

    Unknown globals become inert placeholders, while joblib ndarray payloads
    are bounds-checked and skipped instead of allocated.  This proves that the
    complete stream is structurally readable without executing pickle callables.
    """

    dispatch = Unpickler.dispatch.copy()

    def find_class(self, module: str, name: str):
        if module.startswith("joblib") and name == "NumpyArrayWrapper":
            return NumpyArrayWrapper
        if module.startswith("joblib") and name == "NDArrayWrapper":
            return NDArrayWrapper
        if name == "dtype" and module.startswith("numpy"):
            return np.dtype
        return _StructuralPlaceholder

    def load_build(self) -> None:
        Unpickler.load_build(self)
        wrapper = self.stack[-1]
        if not isinstance(wrapper, (NumpyArrayWrapper, NDArrayWrapper)):
            return
        self.stack.pop()
        if isinstance(wrapper, NDArrayWrapper):
            raise ValueError("legacy joblib ndarray wrapper is not accepted")
        if wrapper.dtype.hasobject:
            _SafePlainUnpickler(self.file_handle).load()
        else:
            alignment = wrapper.safe_get_numpy_array_alignment_bytes()
            if alignment is not None:
                padding = self.file_handle.read(1)
                if len(padding) != 1:
                    raise ValueError("truncated joblib array padding")
                padding_size = int.from_bytes(padding, byteorder="little")
                if len(self.file_handle.read(padding_size)) != padding_size:
                    raise ValueError("truncated joblib array padding")
            shape = tuple(int(value) for value in wrapper.shape)
            if any(value < 0 for value in shape):
                raise ValueError("negative joblib array dimension")
            count = math.prod(shape) if shape else 1
            byte_count = count * int(wrapper.dtype.itemsize)
            current = self.file_handle.tell()
            file_size = os.fstat(self.file_handle.fileno()).st_size
            if byte_count < 0 or current + byte_count > file_size:
                raise ValueError("truncated joblib array payload")
            self.file_handle.seek(byte_count, os.SEEK_CUR)
        self.stack.append(_StructuralPlaceholder())

    dispatch[pickle.BUILD[0]] = load_build


def _valid_native_model(path: Path) -> bool:
    suffix = path.suffix.lower()
    if suffix not in _NATIVE_MODEL_SUFFIXES:
        return False
    if suffix in {".pkl", ".joblib"}:
        try:
            with path.open("rb") as stream:
                _StructuralJoblibUnpickler(
                    str(path),
                    stream,
                    ensure_native_byte_order=True,
                ).load()
                return stream.tell() == path.stat().st_size
        # The restricted parser executes no artifact globals; every malformed
        # opcode/shape/stream error is therefore safely treated as invalid.
        except Exception:
            return False
    try:
        with path.open("rb") as stream:
            prefix = stream.read(4096)
    except OSError:
        return False
    if not prefix:
        return False
    try:
        text = prefix.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return "tree" in text.lower() and "version=" in text.lower()


def _valid_model_card(path: Path, artifact_id: str) -> bool:
    if path.suffix.lower() != ".json":
        return False
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, UnicodeDecodeError):
        return False
    return (
        isinstance(payload, dict)
        and bool(payload)
        and str(payload.get("artifact_id") or "").strip() == artifact_id
        and bool(str(payload.get("recipe") or "").strip())
    )


def _valid_pmml(path: Path) -> bool:
    if path.suffix.lower() != ".pmml":
        return False
    try:
        root = ET.parse(path).getroot()
    except (OSError, ET.ParseError):
        return False
    if root.tag.rsplit("}", 1)[-1] != "PMML" or not root.get("version"):
        return False
    descendants = list(root.iter())
    data_dictionaries = [
        node for node in descendants if node.tag.rsplit("}", 1)[-1] == "DataDictionary"
    ]
    if not any(
        any(child.tag.rsplit("}", 1)[-1] == "DataField" for child in dictionary)
        for dictionary in data_dictionaries
    ):
        return False
    model_names = {
        "AnomalyDetectionModel",
        "AssociationModel",
        "BaselineModel",
        "BayesianNetworkModel",
        "ClusteringModel",
        "GaussianProcessModel",
        "GeneralRegressionModel",
        "MiningModel",
        "NaiveBayesModel",
        "NeuralNetwork",
        "RegressionModel",
        "RuleSetModel",
        "Scorecard",
        "SequenceModel",
        "SupportVectorMachineModel",
        "TextModel",
        "TimeSeriesModel",
        "TreeModel",
    }
    for model in descendants:
        if model.tag.rsplit("}", 1)[-1] not in model_names:
            continue
        mining_schemas = [
            child for child in model if child.tag.rsplit("}", 1)[-1] == "MiningSchema"
        ]
        if any(
            any(child.tag.rsplit("}", 1)[-1] == "MiningField" for child in schema)
            for schema in mining_schemas
        ):
            return True
    return False


def _delivery_artifact(
    path: str | None,
    workspace: Path,
    task_id: str,
    artifact_name: str,
    *,
    artifact_id: str,
) -> tuple[bool, str, str]:
    candidate = _delivery_path(path, workspace, task_id, artifact_name)
    if candidate is None:
        return False, "", "path outside current task governed directory or invalid"
    if not candidate.is_file():
        return False, str(candidate), "file missing"
    try:
        if candidate.stat().st_size <= 0:
            return False, str(candidate), "file empty"
    except OSError:
        return False, str(candidate), "file unreadable"
    validators = {
        "report": _valid_xlsx,
        "native_model": _valid_native_model,
        "model_card": lambda item: _valid_model_card(item, artifact_id),
        "pmml": _valid_pmml,
    }
    if not validators[artifact_name](candidate):
        return False, str(candidate), "file type or suffix invalid"
    return True, str(candidate), "validated"


def _expected_report_filename(recipe: str, experiment_id: str) -> str:
    safe_recipe = _REPORT_FILENAME_UNSAFE_RE.sub("_", recipe).strip("_") or "model"
    digest = hashlib.sha256(experiment_id.encode("utf-8")).hexdigest()[:12]
    return f"model_report_{safe_recipe}_{digest}.xlsx"


def _report_delivery(
    report: dict[str, Any],
    workspace: Path,
    task_id: str,
    connection: sqlite3.Connection,
) -> tuple[bool, str, str]:
    top_level_path = report.get("report_path")
    reports = report.get("reports")
    if reports is None:
        return _delivery_artifact(
            top_level_path,
            workspace,
            task_id,
            "report",
            artifact_id="",
        )
    if not isinstance(reports, list) or not reports:
        return False, "", "reports must be a nonempty list"
    try:
        bindings = {
            str(row["id"]): str(row["recipe_id"])
            for row in connection.execute(
                "SELECT id, recipe_id FROM experiments WHERE task_id = ?",
                (task_id,),
            ).fetchall()
        }
    except sqlite3.Error:
        return False, "", "experiment bindings unavailable"

    validated_paths: list[str] = []
    seen_experiments: set[str] = set()
    seen_paths: set[str] = set()
    problems: list[str] = []
    first_candidate: Path | None = None
    for index, item in enumerate(reports):
        if not isinstance(item, dict):
            problems.append(f"reports[{index}] is not an object")
            continue
        experiment_id = str(item.get("experiment_id") or "").strip()
        recipe = str(item.get("recipe") or "").strip()
        raw_path = item.get("report_path")
        if not experiment_id or not recipe:
            problems.append(f"reports[{index}] missing experiment_id or recipe")
            continue
        if experiment_id in seen_experiments:
            problems.append(f"reports[{index}] duplicate experiment_id")
        seen_experiments.add(experiment_id)
        if bindings.get(experiment_id) != recipe:
            problems.append(
                f"reports[{index}] experiment binding mismatch "
                f"experiment={experiment_id} recipe={recipe}"
            )
        candidate = _delivery_path(raw_path, workspace, task_id, "report")
        if candidate is not None:
            if index == 0:
                first_candidate = candidate
            candidate_text = str(candidate)
            if candidate_text in seen_paths:
                problems.append(f"reports[{index}] duplicate report_path")
            seen_paths.add(candidate_text)
            if candidate.name != _expected_report_filename(recipe, experiment_id):
                problems.append(f"reports[{index}] filename binding mismatch")
        valid, resolved_path, reason = _delivery_artifact(
            raw_path,
            workspace,
            task_id,
            "report",
            artifact_id="",
        )
        if not valid:
            problems.append(f"reports[{index}] {reason}")
        elif resolved_path:
            validated_paths.append(resolved_path)

    top_candidate = _delivery_path(top_level_path, workspace, task_id, "report")
    if (
        top_candidate is None
        or first_candidate is None
        or top_candidate != first_candidate
    ):
        problems.append("top-level report_path must mirror reports[0]")
    if problems:
        return False, ";".join(validated_paths), "; ".join(problems)
    return (
        True,
        ";".join(validated_paths),
        f"validated {len(validated_paths)} experiment-bound reports",
    )


def _pmml_applicability(recipe: str) -> bool | None:
    if recipe in _PMML_APPLICABLE_RECIPES:
        return True
    if recipe in _PMML_UNSUPPORTED_RECIPES:
        return False
    return None


def _current_repo_state() -> tuple[str, bool]:
    repo_root = Path(__file__).resolve().parents[1]
    try:
        commit = subprocess.run(
            ["git", "-C", str(repo_root), "rev-parse", "HEAD"],
            check=True,
            capture_output=True,
            text=True,
        )
        status = subprocess.run(
            [
                "git",
                "-C",
                str(repo_root),
                "status",
                "--porcelain",
                "--untracked-files=normal",
            ],
            check=True,
            capture_output=True,
            text=True,
        )
    except (OSError, subprocess.CalledProcessError):
        return "", False
    return commit.stdout.strip(), not bool(status.stdout.strip())


def _evidence_timestamp(value: Any) -> datetime | None:
    if not isinstance(value, str) or not value.strip():
        return None
    normalized = value.strip().replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(normalized)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        return None
    return parsed.astimezone(UTC)


def _command_text(value: Any) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list) and all(
        isinstance(item, (str, int, float)) for item in value
    ):
        return " ".join(str(item) for item in value).strip()
    return ""


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return "sha256:" + digest.hexdigest()


def _test_evidence_check(
    path: Path,
    *,
    current_commit_sha: str,
    current_worktree_clean: bool,
) -> Check:
    if not path.is_file():
        return Check(
            "C3",
            "FAIL",
            "缺少与当前提交匹配的机器回归测试证据。",
            f"evidence_path={path}; present=false",
        )
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return Check(
            "C3",
            "FAIL",
            "机器回归测试证据不可读取或不是有效 JSON。",
            f"evidence_path={path}; error={exc.__class__.__name__}",
        )
    if not isinstance(payload, dict):
        return Check(
            "C3",
            "FAIL",
            "机器回归测试证据必须是对象。",
            f"evidence_path={path}; payload_type={type(payload).__name__}",
        )

    command = _command_text(payload.get("command"))
    try:
        command_tokens = shlex.split(command) if command else []
    except ValueError:
        command_tokens = []
    exit_code = payload.get("exit_code")
    evidence_commit = str(payload.get("commit_sha") or "").strip()
    evidence_commit_after = str(payload.get("commit_sha_after") or "").strip()
    timestamp = _evidence_timestamp(payload.get("timestamp"))
    log_hash = str(payload.get("log_hash") or "").strip().lower()
    schema_version = str(payload.get("schema_version") or "").strip()
    worktree_clean = payload.get("worktree_clean")
    raw_log_path = str(payload.get("log_path") or "").strip()
    proposed_log_path = (
        (
            Path(raw_log_path)
            if Path(raw_log_path).is_absolute()
            else path.parent / raw_log_path
        )
        if raw_log_path
        else None
    )
    evidence_root = path.parent.resolve()
    log_path_resolution_failed = False
    try:
        log_path = (
            proposed_log_path.resolve() if proposed_log_path is not None else None
        )
    except (OSError, RuntimeError):
        log_path = None
        log_path_resolution_failed = True

    problems: list[str] = []
    if schema_version != "closure-test-evidence.v1":
        problems.append("unsupported or missing schema_version")
    if not command_tokens:
        problems.append("missing command")
    else:
        expected_tail = ["-m", "pytest", "-q", *_REQUIRED_REGRESSION_TESTS]
        executable = Path(command_tokens[0]).name
        if not re.fullmatch(r"python(?:\d+(?:\.\d+)*)?", executable):
            problems.append("command executable is not Python")
        if command_tokens[1:] != expected_tail:
            problems.append(
                "command must run exactly " + " ".join(_REQUIRED_REGRESSION_TESTS)
            )
    if isinstance(exit_code, bool) or not isinstance(exit_code, int):
        problems.append("exit_code is not an integer")
    elif exit_code != 0:
        problems.append(f"exit_code={exit_code}")
    if not current_commit_sha:
        problems.append("current commit unavailable")
    elif evidence_commit != current_commit_sha:
        problems.append(
            f"commit mismatch evidence={evidence_commit or '<missing>'} "
            f"current={current_commit_sha}"
        )
    if evidence_commit_after != evidence_commit:
        problems.append(
            f"commit changed during tests before={evidence_commit or '<missing>'} "
            f"after={evidence_commit_after or '<missing>'}"
        )
    if timestamp is None:
        problems.append("missing or invalid timezone-aware timestamp")
    if worktree_clean is not True:
        problems.append("tests were not run from a clean worktree")
    if current_worktree_clean is not True:
        problems.append("current worktree is not clean")
    if not _SHA256_RE.fullmatch(log_hash):
        problems.append("missing or invalid log_hash")
    if log_path_resolution_failed:
        problems.append("log_path cannot be resolved")
    elif log_path is not None and not _is_within(log_path, evidence_root):
        problems.append("log_path escapes evidence directory")
    elif log_path is None or not log_path.is_file():
        problems.append("log_path missing or unreadable")
    elif _SHA256_RE.fullmatch(log_hash):
        try:
            actual_log_hash = _sha256_file(log_path)
        except OSError:
            problems.append("log_path unreadable")
        else:
            if actual_log_hash != log_hash:
                problems.append(
                    f"log hash mismatch evidence={log_hash} actual={actual_log_hash}"
                )

    evidence = (
        f"evidence_path={path}; command={command!r}; exit_code={exit_code!r}; "
        f"commit_sha={evidence_commit or '<missing>'}; "
        f"commit_sha_after={evidence_commit_after or '<missing>'}; "
        f"timestamp={payload.get('timestamp')!r}; log_hash={log_hash or '<missing>'}"
        f"; worktree_clean={worktree_clean!r}; "
        f"current_worktree_clean={current_worktree_clean!r}"
    )
    if problems:
        return Check(
            "C3",
            "FAIL",
            "机器回归测试证据不完整、失败、过期或与当前提交不匹配。",
            evidence + "; " + "; ".join(problems),
        )
    return Check(
        "C3",
        "PASS",
        "对抗形状与双路对账回归已由当前提交的可核验机器证据证明通过。",
        evidence,
    )


def _lineage_check(step_rows: list[dict[str, Any]]) -> Check:
    executed = [
        row
        for row in step_rows
        if str(row.get("status") or "").lower() in _COMPLETED_STEP_STATUSES
        and str(row.get("tool_plugin") or "").strip()
        and str(row.get("tool_name") or "").strip()
    ]
    if not executed:
        return Check(
            "C2",
            "FAIL",
            "没有可核验的已执行工具步骤，无法证明执行血缘。",
            "required_executed_tool_steps=0",
        )

    failures: list[str] = []
    for row in executed:
        evidence = row.get("evidence")
        evidence = evidence if isinstance(evidence, dict) else {}
        missing: list[str] = []
        if not str(evidence.get("step_run_id") or "").strip():
            missing.append("step_run_id")
        manifest_hash = str(evidence.get("manifest_hash") or "").strip().lower()
        if not _SHA256_RE.fullmatch(manifest_hash):
            missing.append("manifest_hash")
        input_hash = str(evidence.get("input_hash") or "").strip().lower()
        if not _SHA256_RE.fullmatch(input_hash):
            missing.append("input_hash")
        source_refs = evidence.get("source_dataset_refs")
        parent_refs = evidence.get("parent_output_refs")
        if not (
            isinstance(source_refs, list)
            and any(str(item).strip() for item in source_refs)
        ) and not (
            isinstance(parent_refs, list)
            and any(str(item).strip() for item in parent_refs)
        ):
            missing.append("source_refs")

        tool_name = str(row.get("tool_name") or "").strip()
        input_summary = evidence.get("input_summary")
        input_summary = input_summary if isinstance(input_summary, dict) else {}
        stochastic = tool_name in _STOCHASTIC_TOOLS or "seed" in input_summary
        seed = evidence.get("random_seed")
        if stochastic and (isinstance(seed, bool) or not isinstance(seed, int)):
            missing.append("random_seed")
        if missing:
            identity = (
                f"{row.get('step_id') or '<missing-id>'}/"
                f"{row.get('title') or '<untitled>'}"
            )
            failures.append(f"{identity}: {','.join(missing)}")

    if failures:
        return Check(
            "C2",
            "FAIL",
            "逐步执行证据信封不完整。",
            (f"executed_tool_steps={len(executed)}; missing=" + "; ".join(failures)),
        )
    return Check(
        "C2",
        "PASS",
        "每个已执行工具步骤均有运行、代码、输入与来源血缘；随机步骤另有种子。",
        f"executed_tool_steps={len(executed)}; missing=none",
    )


def inspect(
    workspace: Path,
    task_id: str | None = None,
    *,
    test_evidence_path: Path | None = None,
    join_task_id: str | None = None,
    vintage_task_id: str | None = None,
) -> dict[str, Any]:
    db_path = workspace / "marvis.sqlite"
    if not db_path.is_file():
        raise FileNotFoundError(f"workspace database not found: {db_path}")

    connection = sqlite3.connect(db_path)
    connection.row_factory = sqlite3.Row
    plan = (
        _task_plan(connection, task_id)
        if task_id
        else _latest_done_modeling_plan(connection)
    )
    if plan is None:
        raise LookupError("no modeling plan found")
    if plan["status"] != "done":
        raise LookupError(
            f"task {plan['task_id']} latest plan is {plan['status']!r}, not 'done'"
        )

    task_id = str(plan["task_id"])
    step_rows = _step_record_rows(connection, str(plan["plan_id"]))
    primary_has_join = bool(
        _tool_output(step_rows, "propose_join")
        and _tool_output(step_rows, "execute_join")
    )
    resolved_join_task_id = (
        str(join_task_id)
        if join_task_id
        else (task_id if primary_has_join else None)
    )
    if resolved_join_task_id == task_id:
        join_rows = step_rows
    else:
        _join_plan, join_rows = _completed_task_steps(
            connection,
            resolved_join_task_id,
        )
    _vintage_plan, vintage_rows = _completed_task_steps(
        connection,
        str(vintage_task_id) if vintage_task_id else None,
    )
    join_checks = _join_acceptance_checks(
        join_rows,
        task_id=resolved_join_task_id,
    )
    steps = _step_records(step_rows)
    split = (steps.get("切分样本") or {}).get("output") or {}
    screen = (steps.get("特征筛选") or {}).get("output") or {}
    refine = (steps.get("精选特征") or {}).get("output") or {}
    train = (steps.get("训练模型") or {}).get("output") or {}
    select = (steps.get("选择实验") or {}).get("output") or {}
    report = (steps.get("生成模型开发报告") or {}).get("output") or {}
    delivery = (steps.get("模型交付动作") or {}).get("output") or {}
    artifact_id = select.get("artifact_id") or delivery.get("artifact_id")
    artifact = _artifact_row(connection, artifact_id)
    artifact_params = _json(artifact["params_json"], {}) if artifact else {}
    model_card = delivery.get("model_card") or {}
    monitoring = delivery.get("monitoring_policy") or {}
    sentinel_columns = screen.get("sentinel_columns") or {}

    checks: list[Check] = []
    checks.append(
        _a1_check(
            sentinel_columns=sentinel_columns,
            artifact_params=artifact_params,
            selected_features=_artifact_selected_features(artifact, refine),
        )
    )

    checks.append(
        _vintage_semantics_check(
            vintage_rows,
            task_id=str(vintage_task_id) if vintage_task_id else None,
        )
    )

    nan_labels = int(screen.get("nan_labels_dropped") or 0)
    checks.append(
        Check(
            "A3",
            "PASS" if nan_labels == 0 else "FAIL",
            (
                "特征筛选未丢弃标签，未发现静默排除。"
                if nan_labels == 0
                else "特征筛选丢弃了标签；需回到 NaN 标签确认门复核。"
            ),
            f"screen_features.nan_labels_dropped={nan_labels}",
        )
    )

    checks.extend(join_checks[:2])

    split_counts = (split.get("sample_analysis") or {}).get("split_counts") or {}
    holdout = split.get("holdout_values") or []
    oot_count = int(split_counts.get("oot") or 0)
    checks.append(
        Check(
            "A6",
            "PASS" if "oot" in holdout and oot_count > 0 else "FAIL",
            (
                "已生成非空 OOT，切分证据可追溯。"
                if "oot" in holdout and oot_count > 0
                else "没有可核验的非空 OOT。"
            ),
            (
                f"split_col={split.get('split_col')}; holdout_values={holdout}; "
                f"split_counts={split_counts}"
            ),
        )
    )

    for check_id, label in (
        ("B1", "vintage 累计坏账率"),
        ("B2", "组合 EL"),
        ("B3", "分群 bad_rate"),
    ):
        checks.append(
            Check(
                check_id,
                "MANUAL",
                f"{label}必须与外部业务口径对账，仓库内不存在可替代的地面真值。",
                "no external finance/risk ground truth supplied",
                f"填写外部口径来源、实测 vs 口径并由责任人签字（{check_id}）。",
            )
        )

    metrics = select.get("metrics") or {}
    key_metrics = {
        item.get("metric"): item.get("value")
        for item in (model_card.get("key_metrics") or [])
        if isinstance(item, dict)
    }
    ks_names = ("train_ks", "test_ks", "oot_ks")
    ks_problems: list[str] = []
    for name in ks_names:
        if name not in metrics or name not in key_metrics:
            ks_problems.append(f"{name}: missing")
            continue
        try:
            select_value = float(metrics[name])
            card_value = float(key_metrics[name])
        except (TypeError, ValueError):
            ks_problems.append(f"{name}: nonnumeric")
            continue
        if not math.isfinite(select_value) or not math.isfinite(card_value):
            ks_problems.append(f"{name}: nonfinite")
            continue
        delta = abs(select_value - card_value)
        if delta > 1e-12:
            ks_problems.append(f"{name}: delta={delta}")
    ks_consistent = not ks_problems
    checks.append(
        Check(
            "B4-INTERNAL",
            "PASS" if ks_consistent else "FAIL",
            (
                "平台内部 KS 在选择结果与模型卡间逐项一致。"
                if ks_consistent
                else "平台内部 KS 证据缺失、不合法或不一致。"
            ),
            (
                f"internal_consistency={ks_consistent}; "
                + ", ".join(f"{name}={metrics.get(name)}" for name in ks_names)
                + ("; problems=" + ", ".join(ks_problems) if ks_problems else "")
            ),
        )
    )
    checks.append(
        Check(
            "B4-EXTERNAL",
            "MANUAL",
            "内部一致性不能替代独立复算或历史同类模型 KS 对账。",
            "no independent KS recomputation or approved historical benchmark supplied",
            "提供独立复算或历史同类模型 KS，并完成 B4 外部对账签字。",
        )
    )
    checks.extend(join_checks[2:])

    selection_metric = train.get("selection_metric")
    checks.append(
        Check(
            "C1",
            "PASS" if selection_metric else "FAIL",
            (
                "训练输出记录了真实选择指标。"
                if selection_metric
                else "训练输出缺少 selection_metric。"
            ),
            (
                f"train_models.selection_metric={selection_metric!r}; "
                f"select_experiment.selection_reason={select.get('selection_reason')!r}"
            ),
        )
    )

    checks.append(_lineage_check(step_rows))
    test_evidence_path = (
        Path(test_evidence_path)
        if test_evidence_path is not None
        else workspace / "closure_test_evidence.json"
    )
    live_commit_sha, live_worktree_clean = _current_repo_state()
    checks.append(
        _test_evidence_check(
            test_evidence_path,
            current_commit_sha=live_commit_sha,
            current_worktree_clean=live_worktree_clean,
        )
    )

    fit_split = refine.get("fit_split")
    fit_rows = int(refine.get("fit_rows") or 0)
    checks.append(
        Check(
            "D1",
            "PASS" if fit_split == "train" and fit_rows > 0 else "FAIL",
            (
                "精选特征仅在 train 上拟合。"
                if fit_split == "train" and fit_rows > 0
                else "精选特征没有可核验的 train-only 拟合证据。"
            ),
            f"select_features.fit_split={fit_split!r}; fit_rows={fit_rows}",
        )
    )
    checks.append(
        Check(
            "D2",
            "PASS" if nan_labels == 0 else "FAIL",
            (
                "本任务无 NaN 标签；NaN 门另由对抗形状回归覆盖。"
                if nan_labels == 0
                else "本任务出现 NaN 标签丢弃，需核对确认记录。"
            ),
            (
                f"nan_labels_dropped={nan_labels}; "
                "tests/test_dirty_shape_regression.py::test_nan_label_screen_requires_confirmation"
            ),
        )
    )
    refit = select.get("refit") or {}
    refit_enabled = bool(refit.get("enabled") or refit.get("performed"))
    checks.append(
        Check(
            "D3",
            "PASS" if not refit_enabled else "MANUAL",
            (
                "本任务未执行 train+test refit，不存在随机 5% headline 冒充问题。"
                if not refit_enabled
                else "本任务执行了 refit，需要人工核对部署前后指标说明。"
            ),
            f"select_experiment.refit={refit}",
            "若 refit=true，核对模型卡 pre-refit/部署差异说明。",
        )
    )

    report_path = report.get("report_path")
    native_model_path = delivery.get("native_model_path")
    pmml_path = delivery.get("pmml_path")
    model_card_path = delivery.get("model_card_path")
    artifact_id_text = str(artifact_id or "").strip()
    delivery_validation = {
        name: _delivery_artifact(
            raw_path,
            workspace,
            task_id,
            name,
            artifact_id=artifact_id_text,
        )
        for name, raw_path in {
            "native_model": native_model_path,
            "pmml": pmml_path,
            "model_card": model_card_path,
        }.items()
    }
    delivery_validation["report"] = _report_delivery(
        report,
        workspace,
        task_id,
        connection,
    )
    delivery_paths = {
        name: validation[0] for name, validation in delivery_validation.items()
    }
    recipe = str(
        model_card.get("recipe")
        or delivery.get("recipe")
        or select.get("recipe")
        or (
            artifact["algorithm"]
            if artifact is not None and "algorithm" in artifact.keys()
            else ""
        )
        or ""
    ).strip()
    pmml_applicable = _pmml_applicability(recipe)
    artifact_checks = (
        (
            "E1",
            "report",
            True,
            report_path,
            "模型开发报告",
        ),
        (
            "E2",
            "native_model",
            True,
            native_model_path,
            "原生模型",
        ),
        (
            "E3",
            "model_card",
            True,
            model_card_path,
            "模型卡",
        ),
        (
            "E4",
            "pmml",
            pmml_applicable,
            pmml_path,
            "PMML",
        ),
    )
    for check_id, artifact_name, applicable, raw_path, label in artifact_checks:
        if applicable is False:
            checks.append(
                Check(
                    check_id,
                    "N/A",
                    f"{label}不适用于当前 recipe。",
                    f"artifact={artifact_name}; recipe={recipe or '<unknown>'}",
                )
            )
            continue
        if applicable is None:
            checks.append(
                Check(
                    check_id,
                    "FAIL",
                    f"无法判断 {label} 对当前 recipe 是否适用。",
                    f"artifact={artifact_name}; recipe={recipe or '<unknown>'}",
                )
            )
            continue
        present = delivery_paths[artifact_name]
        _, resolved_path, validation_reason = delivery_validation[artifact_name]
        checks.append(
            Check(
                check_id,
                "PASS" if present else "FAIL",
                (
                    f"{label}交付产物存在。"
                    if present
                    else f"缺少强制交付产物：{label}。"
                ),
                (
                    f"artifact={artifact_name}; recipe={recipe or '<unknown>'}; "
                    f"path={raw_path!r}; resolved_path={resolved_path!r}; "
                    f"validation={validation_reason}; "
                    f"present={str(present).lower()}"
                ),
            )
        )

    machine_failures = [item.check_id for item in checks if item.status == "FAIL"]
    manual_items = [item.check_id for item in checks if item.status == "MANUAL"]
    return {
        "schema_version": "closure-real-materials-machine-check.v2",
        "generated_at": datetime.now(UTC).isoformat(),
        "workspace": str(workspace.resolve()),
        "task": {
            "task_id": task_id,
            "model_name": plan["model_name"],
            "model_version": plan["model_version"],
            "plan_id": plan["plan_id"],
            "plan_status": plan["status"],
            "plan_updated_at": plan["updated_at"],
        },
        "companion_tasks": {
            "join_task_id": resolved_join_task_id,
            "vintage_task_id": str(vintage_task_id) if vintage_task_id else None,
        },
        "artifacts": delivery_paths,
        "artifact_applicability": {
            "report": True,
            "native_model": True,
            "model_card": True,
            "pmml": pmml_applicable,
            "recipe": recipe,
        },
        "model_governance": {
            "monitoring_status": monitoring.get("status"),
            "monitoring_recommendation": monitoring.get("recommendation"),
            "limitations": model_card.get("limitations") or [],
        },
        "checks": [asdict(item) for item in checks],
        "machine_failures": machine_failures,
        "manual_items": manual_items,
        "machine_verdict": "FAIL" if machine_failures else "PASS",
        "closure_verdict": (
            "BLOCKED_MACHINE"
            if machine_failures
            else ("BLOCKED_MANUAL" if manual_items else "PASS")
        ),
        "manual_blocker": (
            "B1-B3, B4-EXTERNAL, and B5-EXTERNAL ground-truth reconciliation and accountable "
            "signatures; the script cannot and will not fabricate them."
        ),
    }


def render_markdown(result: dict[str, Any]) -> str:
    task = result["task"]
    lines = [
        "# 真实材料机器预检报告",
        "",
        f"- 生成时间：`{result['generated_at']}`",
        f"- 任务：`{task['task_id']}` · {task['model_name']}",
        f"- 计划：`{task['plan_id']}` · `{task['plan_status']}`",
        f"- 机器判定：**{result['machine_verdict']}**",
        f"- 收口判定：**{result['closure_verdict']}**",
        "",
        "> 本报告只核验仓库/SQLite/产物中可机器验证的证据，不等同于人工签字。",
        "> 外部财务、拨备、风险报表口径及责任人签字绝不自动填充。",
        "",
        "## A/B/C/D 核验",
        "",
        "| 项 | 状态 | 结论 | 证据 | 人工动作 |",
        "|---|---|---|---|---|",
    ]
    for item in result["checks"]:
        cells = [
            item["check_id"],
            item["status"],
            item["summary"],
            item["evidence"],
            item["human_action"] or "—",
        ]
        lines.append(
            "| " + " | ".join(str(cell).replace("|", "\\|") for cell in cells) + " |"
        )
    lines.extend(
        [
            "",
            "## 产物",
            "",
            *[
                (
                    f"- {name}: N/A"
                    if result["artifact_applicability"].get(name, True) is False
                    else f"- {name}: {'PASS' if present else 'MISSING'}"
                )
                for name, present in result["artifacts"].items()
            ],
            "",
            "## 模型治理",
            "",
            (
                f"- 监控状态：`{result['model_governance']['monitoring_status']}` · "
                f"{result['model_governance']['monitoring_recommendation']}"
            ),
            *[f"- 限制：{item}" for item in result["model_governance"]["limitations"]],
            "",
            "## 仍需人工完成",
            "",
            f"- {result['manual_blocker']}",
            "- 完成人应在 `docs/plans/v2-real-materials-reconciliation-checklist.md` "
            "填写外部口径、实测值及签字。",
            "",
        ]
    )
    return "\n".join(lines)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace", default="workspace")
    parser.add_argument("--task-id", default=None)
    parser.add_argument(
        "--join-task-id",
        default=None,
        help=(
            "completed join task for the same governed real-material sample; "
            "optional when the modeling plan itself contains join steps"
        ),
    )
    parser.add_argument(
        "--vintage-task-id",
        default=None,
        help="completed vintage task for the same governed real-material sample",
    )
    parser.add_argument("--output", default=None, help="optional Markdown output path")
    parser.add_argument("--json-output", default=None)
    parser.add_argument(
        "--test-evidence",
        default=None,
        help=(
            "machine-generated regression evidence JSON for the current commit; "
            "defaults to <workspace>/closure_test_evidence.json"
        ),
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        result = inspect(
            Path(args.workspace),
            args.task_id,
            test_evidence_path=(
                Path(args.test_evidence) if args.test_evidence else None
            ),
            join_task_id=args.join_task_id,
            vintage_task_id=args.vintage_task_id,
        )
    except (FileNotFoundError, LookupError, sqlite3.Error) as exc:
        print(json.dumps({"status": "ERROR", "error": str(exc)}, ensure_ascii=False))
        return 2

    markdown = render_markdown(result)
    if args.output:
        output_path = Path(args.output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(markdown, encoding="utf-8")
    if args.json_output:
        json_path = Path(args.json_output)
        json_path.parent.mkdir(parents=True, exist_ok=True)
        json_path.write_text(
            json.dumps(result, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if result["machine_failures"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
