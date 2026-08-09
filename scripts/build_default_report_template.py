#!/usr/bin/env python3
"""Build the public, institution-neutral MARVIS validation report template."""

from __future__ import annotations

import argparse
import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(0xC0, 0x00, 0x00)
ACCENT_DARK = RGBColor(0x7A, 0x1D, 0x1D)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
LIGHT_FILL = "F5F5F5"
FONT_NAME = "Noto Sans CJK SC"
TABLE_WIDTH_DXA = 9360
# Full-width cover metadata table: keep the visible border within the 9360 DXA body.
TABLE_INDENT_DXA = 0
TABLE_CELL_MARGIN_DXA = 120


def _set_run_font(
    run,
    *,
    size: float,
    bold: bool = False,
    color: RGBColor = INK,
    italic: bool = False,
) -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, ACCENT, 16, 8),
        "Heading 2": (13, ACCENT, 12, 6),
        "Heading 3": (12, ACCENT_DARK, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = FONT_NAME
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def _configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def _set_cell_margins(cell) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for side in ("top", "start", "bottom", "end"):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(TABLE_CELL_MARGIN_DXA))
        element.set(qn("w:type"), "dxa")


def _set_table_geometry(table, column_widths: list[int]) -> None:
    if sum(column_widths) != TABLE_WIDTH_DXA:
        raise ValueError("table column widths must total 9360 DXA")
    table.autofit = False
    table_properties = table._tbl.tblPr
    width = table_properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        table_properties.insert(0, width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    indent = table_properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table_properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for column_width in column_widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(column_width))
        grid.append(grid_column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(column_widths[index]))
            _set_cell_margins(cell)


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _add_header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("MARVIS · 模型验证")
    _set_run_font(run, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("MARVIS 自动生成 · 确定性指标以平台输出为准")
    _set_run_font(run, size=8, color=MUTED)


def _add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    _set_run_font(kicker.add_run("MARVIS MODEL VALIDATION"), size=11, bold=True, color=ACCENT)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    _set_run_font(title.add_run("{{TEXT:report_title}}"), size=26, bold=True, color=INK)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    _set_run_font(subtitle.add_run("模型验证报告"), size=15, color=MUTED)

    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    values = (
        ("模型名称", "{{TEXT:model_name}}"),
        ("模型版本", "{{TEXT:model_version}}"),
        ("撰写人", "{{TEXT:drafter}}"),
        ("撰写日期", "{{TEXT:draft_date}}"),
    )
    for row, (label, value) in zip(table.rows, values, strict=True):
        label_paragraph = row.cells[0].paragraphs[0]
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(label_paragraph.add_run(label), size=10.5, bold=True, color=ACCENT_DARK)
        _shade_cell(row.cells[0], LIGHT_FILL)
        value_paragraph = row.cells[1].paragraphs[0]
        _set_run_font(value_paragraph.add_run(value), size=10.5)
    _set_table_geometry(table, [2700, 6660])

    revision = document.add_paragraph()
    revision.alignment = WD_ALIGN_PARAGRAPH.CENTER
    revision.paragraph_format.space_before = Pt(20)
    revision.paragraph_format.space_after = Pt(0)
    _set_run_font(
        revision.add_run(
            "修订 {{TEXT:revision_version}} · {{TEXT:revision_date}} · "
            "{{TEXT:revision_author}} · {{TEXT:revision_description}}"
        ),
        size=9,
        color=MUTED,
    )
    document.add_page_break()


def _add_labeled_text(document: Document, label: str, placeholder: str) -> None:
    paragraph = document.add_paragraph()
    _set_run_font(paragraph.add_run(f"{label}："), size=11, bold=True, color=ACCENT_DARK)
    _set_run_font(paragraph.add_run(placeholder), size=11)


def _add_figure(document: Document, caption: str, placeholder: str) -> None:
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(6)
    caption_paragraph.paragraph_format.space_after = Pt(4)
    _set_run_font(caption_paragraph.add_run(caption), size=9, bold=True, color=MUTED)
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(10)
    _set_run_font(image_paragraph.add_run(placeholder), size=9, color=MUTED)


def _add_body(document: Document) -> None:
    document.add_heading("1. 模型与验证范围", level=1)
    _add_labeled_text(document, "模型概述", "{{TEXT:model_overview}}")
    _add_labeled_text(document, "适用范围", "{{TEXT:model_scope}}")
    _add_labeled_text(document, "坏样本定义", "{{TEXT:bad_sample_definition}}")
    _add_labeled_text(document, "好样本定义", "{{TEXT:good_sample_definition}}")

    document.add_heading("2. 样本与模型开发信息", level=1)
    _add_labeled_text(
        document,
        "样本周期",
        "{{TEXT:sample_start_month}} 至 {{TEXT:sample_end_month}}",
    )
    _add_labeled_text(document, "训练/测试周期", "{{TEXT:train_test_period}}")
    _add_labeled_text(document, "训练/测试比例", "{{TEXT:train_test_ratio}}")
    _add_labeled_text(document, "OOT 周期", "{{TEXT:oot_period}}")
    _add_figure(document, "样本整体分布", "{{IMAGE:sample_overall_distribution}}")
    _add_figure(document, "样本月度分布", "{{IMAGE:sample_month_distribution}}")
    document.add_heading("2.1 模型训练说明", level=2)
    _add_labeled_text(document, "训练方法", "{{TEXT:model_training_description}}")
    _add_figure(document, "模型参数", "{{IMAGE:model_parameters}}")
    _add_figure(document, "特征重要性 Top 20", "{{IMAGE:top20_feature_ranking}}")

    document.add_heading("3. 有效性与稳定性", level=1)
    _add_labeled_text(document, "OOT KS", "{{TEXT:oot_ks}}")
    _add_labeled_text(document, "OOT PSI", "{{TEXT:oot_psi}}")
    _add_figure(document, "总体模型效果", "{{IMAGE:overall_model_effect}}")
    _add_figure(document, "月度模型效果", "{{IMAGE:loan_month_effect}}")
    _add_figure(document, "PSI 稳定性分箱", "{{IMAGE:psi_stability_table}}")
    for section_number, (split, label) in enumerate(
        (("train", "训练集"), ("test", "测试集"), ("oot", "OOT")),
        start=1,
    ):
        document.add_heading(f"3.{section_number} {label}区分能力", level=2)
        _add_figure(document, f"{label} ROC/KS", f"{{{{IMAGE:roc_ks_graph_{split}}}}}")
        _add_figure(document, f"{label}分箱表现", f"{{{{IMAGE:ranking_table_{split}}}}}")

    document.add_heading("4. PMML 打分测试", level=1)
    _add_labeled_text(document, "测试结论", "{{TEXT:reproducibility_summary}}")

    document.add_heading("5. 压力测试", level=1)
    _add_labeled_text(document, "测试说明", "{{TEXT:pressure_test_summary}}")
    _add_figure(document, "压力测试 KS", "{{IMAGE:pressure_ks_table}}")
    _add_figure(document, "压力测试 PSI", "{{IMAGE:pressure_psi_table}}")
    _add_figure(document, "压力测试分数分布", "{{IMAGE:pressure_score_shift}}")
    _add_labeled_text(
        document,
        "风险影响与建议",
        "{{TEXT:pressure_impact_recommendation}}",
    )

    document.add_heading("6. 验证结论", level=1)
    conclusion = document.add_paragraph()
    conclusion.paragraph_format.space_after = Pt(8)
    conclusion.paragraph_format.line_spacing = 1.2
    _set_run_font(conclusion.add_run("{{TEXT:final_validation_conclusion}}"), size=11)


def build_template(output_path: Path) -> Path:
    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "MARVIS 模型验证报告模板"
    document.core_properties.subject = "公开、中性的模型验证报告模板"
    document.core_properties.keywords = "MARVIS, model validation"
    document.core_properties.comments = ""
    document.core_properties.revision = 1
    _configure_styles(document)
    for section in document.sections:
        _configure_section(section)
        _add_header_footer(section)
    _add_cover(document)
    _add_body(document)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(f".{output_path.name}.tmp")
    document.save(temporary_path)
    os.replace(temporary_path, output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("marvis/report_templates/default.docx"),
    )
    args = parser.parse_args()
    print(build_template(args.output.resolve()))


if __name__ == "__main__":
    main()
